"""任务两表纯解析测试；构造载荷不是生产任务或算法验收。"""
from copy import deepcopy
import csv
from datetime import datetime
from io import BytesIO, StringIO
import json
from pathlib import Path
import xml.etree.ElementTree as ET
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from openpyxl import load_workbook
import pytest

from app import task_file_parser as parser
from app.task_file_parser import TaskFileError, build_task_template, parse_task_files


@pytest.fixture
def input_data():
    return {"name": "构造任务", "description": "UNIT_TEST，非生产非实测", "bbox": [113, 36, 115, 38],
            "stationIds": ["N1", "N2"], "presetId": "CURRENT-PRESET", "frequencyId": "CURRENT-FREQUENCY",
            "startTime": None, "durationH": None,
            "constraints": {"connectivityPercent": 100, "maxRelayHops": 2, "survivabilityPercent": 90},
            "demands": [{"id": "D1", "sourceId": "N1", "targetId": "N2", "priority": "P1", "mandatory": False,
                         "dataVolumeMb": None, "requiredRateKbps": 2, "maxDelayMs": 100, "minReliability": 0,
                         "service": "", "preferredClass": None, "endpointMode": "station", "sourceDeviceId": None, "targetDeviceId": None}]}


def csv_files(value):
    return [(file["name"], file["bytes"]) for file in (build_task_template("csv", kind, value) for kind in parser.KINDS)]


def edit_csv(files, kind, edit):
    changed = []
    for name, data in files:
        if name != kind + ".csv":
            changed.append((name, data))
            continue
        rows = list(csv.reader(StringIO(data.decode("utf-8"), newline="")))
        edit(rows)
        output = StringIO(newline="")
        csv.writer(output).writerows(rows)
        changed.append((name, output.getvalue().encode("utf-8")))
    return changed


def set_csv(files, kind, field, value):
    def edit(rows):
        rows[1][rows[0].index(field)] = value
    return edit_csv(files, kind, edit)


def workbook_file(value, edit=None):
    file = build_task_template("xlsx", "task_scenarios", value)
    if edit is None:
        return [(file["name"], file["bytes"])]
    book = load_workbook(BytesIO(file["bytes"]))
    try:
        edit(book)
        output = BytesIO()
        with ZipFile(output, "w", ZIP_DEFLATED) as archive:
            parser._MemoryExcelWriter(book, archive).save()
        return [(file["name"], output.getvalue())]
    finally:
        book.close()


def archive_file(files, edit):
    with ZipFile(BytesIO(files[0][1])) as archive:
        entries = [(item.filename, archive.read(item)) for item in archive.infolist()]
    entries = edit(entries)
    output = BytesIO()
    with ZipFile(output, "w", ZIP_DEFLATED) as archive:
        for name, data in entries:
            info = ZipInfo(name)
            # 在Windows也保留恶意原始ZIP名称，避免构造时被库先行规范化。
            info.filename = name
            info.compress_type = ZIP_DEFLATED
            archive.writestr(info, data)
    return [(files[0][0], output.getvalue())]


def edit_xml(entries, path, edit):
    changed = []
    for name, data in entries:
        if name == path:
            root = ET.fromstring(data)
            edit(root)
            data = ET.tostring(root, encoding="utf-8")
        changed.append((name, data))
    return changed


def result_issue(result, field, sheet=None, row=2):
    assert result["input"] is None
    assert any(issue["field"] == field and issue["row"] == row and (sheet is None or issue["sheet"] == sheet) for issue in result["issues"])
    json.dumps(result, allow_nan=False)


@pytest.mark.parametrize("format", ["csv", "xlsx"])
def test_templates_roundtrip_exact_contract_and_do_not_mutate(input_data, format):
    before = deepcopy(input_data)
    files = csv_files(input_data) if format == "csv" else workbook_file(input_data)
    result = parse_task_files(files, input_data)
    assert result["issues"] == []
    assert result["input"] == before
    assert result["taskId"] == "TASK-NEW"
    assert result["taskRow"] == 2 and result["sourceRows"] == {"D1": 2}
    assert set(result) == {"taskId", "input", "issues", "notes", "differences", "taskRow", "sourceRows"}
    assert input_data == before
    assert set(build_task_template(format, "task_links", input_data)) == {"name", "mime", "bytes"}


@pytest.mark.parametrize("format", ["csv", "xlsx"])
def test_real_extended_columns_zero_false_null_and_device_endpoints(input_data, format):
    input_data.update(startTime="2026-10-23T08:30:00+08:00", durationH=8,
                      deploymentConstraints={"maxSlopeDeg": 0, "elevationRangeM": [-5, 100], "allowedLandcover": ["构造地物"],
                                             "roadMaxDistanceM": 0, "minSiteSpacingM": 0, "losRequiredNodeIds": ["N1"], "gridStepM": 20,
                                             "forcedInclude": [{"lon": 114, "lat": 37}],
                                             "forcedExcludePolygons": [{"type": "Polygon", "coordinates": [[[113, 36], [113.1, 36], [113, 36.1], [113, 36]]]}]},
                      frequencyConstraints={"filters": [{"deviceClass": "HF", "rangeKhz": [3000, 4000], "excludeFreqIds": ["H1"]},
                                                         {"deviceClass": "VUHF", "rangeKhz": None, "excludeFreqIds": ["V1"]}]},
                      deploymentSolve={"mode": "FIXED_COUNT", "fixedCount": 0, "solutionCount": 2})
    input_data["demands"][0].update(endpointMode="device", sourceDeviceId="E1", targetDeviceId="E2", preferredClass="HF",
                                    dataVolumeMb=0, routeStrategy="MIN_DELAY", needBackup=False)
    files = csv_files(input_data) if format == "csv" else workbook_file(input_data)
    result = parse_task_files(files, input_data)
    assert result["issues"] == [] and result["input"] == input_data
    assert result["input"]["deploymentSolve"]["fixedCount"] == 0
    assert result["input"]["demands"][0]["needBackup"] is False
    assert len(parser.HEADERS["task_scenarios"]) == 26 and len(parser.HEADERS["task_links"]) == 17


def test_legacy_headers_use_only_endpoint_union_and_clear_all_optional_conditions(input_data):
    defaults = deepcopy(input_data)
    defaults.update(deploymentConstraints={"unread": "不能继承"}, frequencyConstraints={"unread": "不能继承"},
                    deploymentSolve={"unread": "不能继承"}, stationIds=["N1", "N2", "N3"])
    defaults["demands"][0].update(routeStrategy="MIN_DELAY", needBackup=False)
    files = csv_files(input_data)
    for kind, count in (("task_scenarios", 6), ("task_links", 12)):
        files = edit_csv(files, kind, lambda rows, count=count: [row.__delitem__(slice(count, None)) for row in rows])
    result = parse_task_files(files, defaults)
    assert result["issues"] == []
    assert result["input"]["stationIds"] == ["N1", "N2"]
    assert result["input"]["presetId"] == defaults["presetId"] and result["input"]["frequencyId"] == defaults["frequencyId"]
    for key in parser.OPTION_KEYS:
        assert key not in result["input"]
        assert any(key in note and "清空" in note for note in result["notes"])
        assert any(difference["field"] == key and difference["after"] is None for difference in result["differences"])
    assert "needBackup" not in result["input"]["demands"][0] and "routeStrategy" not in result["input"]["demands"][0]
    assert any("路由条件" in note and "清空" in note for note in result["notes"])
    assert sum("未提供" in note for note in result["notes"]) == 2
    assert any("端ID合并" in note for note in result["notes"])


@pytest.mark.parametrize("field,value", [("required_rate_kbps", ""), ("required_rate_kbps", "NaN"), ("required_rate_kbps", "Infinity"),
                                        ("required_rate_kbps", "1e9999"), ("required_rate_kbps", "true"), ("required_rate_kbps", "=1+1"),
                                        ("data_volume_mb", "-1"), ("min_reliability", "2"), ("max_delay_ms", "0"),
                                        ("is_mandatory", "False"), ("need_backup", "False"), ("strategy", "LOCAL_DEMO")])
def test_illegal_scalar_is_located_and_never_emits_nan(input_data, field, value):
    files = set_csv(csv_files(input_data), "task_links", field, value)
    result_issue(parse_task_files(files, input_data), field, "task_links")


@pytest.mark.parametrize("field,value", [("duration_h", "NaN"), ("area_bbox", "113;36;Infinity;38"),
                                        ("start_time", "2026-02-30T08:00:00Z"), ("start_time", "2026-10-23T08:00:00"),
                                        ("forced_include_json", "["), ("deployment_solve_json", '{"mode":"MIN_COUNT","fixedCount":NaN,"solutionCount":1}'),
                                        ("deployment_solve_json", '{"mode":"MIN_COUNT","mode":"FIXED_COUNT","fixedCount":0,"solutionCount":1}')])
def test_illegal_scenario_scalar_and_json_have_row_and_null_input(input_data, field, value):
    result_issue(parse_task_files(set_csv(csv_files(input_data), "task_scenarios", field, value), input_data), field, "task_scenarios")


@pytest.mark.parametrize("field,value,prefix", [("hf_min_khz", "3000", "frequencyConstraints"),
                                               ("grid_step_m", "0", "deploymentConstraints"),
                                               ("forced_include_json", '[{"lon":114,"lat":37,"unknown":true}]', "deploymentConstraints"),
                                               ("deployment_solve_json", '{"mode":"FIXED_COUNT","fixedCount":false,"solutionCount":1}', "deploymentSolve")])
def test_extension_shape_constraints_are_not_silently_normalized(input_data, field, value, prefix):
    result = parse_task_files(set_csv(csv_files(input_data), "task_scenarios", field, value), input_data)
    assert result["input"] is None and any(issue["field"].startswith(prefix) for issue in result["issues"])
    json.dumps(result, allow_nan=False)


@pytest.mark.parametrize("edit", [lambda rows: rows[0].__setitem__(0, "unknown"),
                                  lambda rows: rows[0].__setitem__(1, rows[0][0]),
                                  lambda rows: [row.pop(0) for row in rows]])
def test_unknown_duplicate_and_missing_headers_block_entire_input(input_data, edit):
    result_issue(parse_task_files(edit_csv(csv_files(input_data), "task_scenarios", edit), input_data), "header", "task_scenarios", 1)


def test_wrong_task_duplicate_ids_extra_columns_and_multiple_tasks(input_data):
    result_issue(parse_task_files(set_csv(csv_files(input_data), "task_links", "task_id", "OTHER"), input_data), "task_id")
    duplicate = edit_csv(csv_files(input_data), "task_links", lambda rows: rows.append(list(rows[1])))
    result = parse_task_files(duplicate, input_data)
    result_issue(result, "demand_id", "task_links", 3)
    assert result["sourceRows"]["D1"] == 2
    extra = edit_csv(csv_files(input_data), "task_links", lambda rows: rows[1].append("unknown"))
    result_issue(parse_task_files(extra, input_data), "columns")
    multi = edit_csv(csv_files(input_data), "task_scenarios", lambda rows: rows.append(list(rows[1])))
    result_issue(parse_task_files(multi, input_data), "task_id", "task_scenarios")


def test_actual_csv_row_locations_include_blank_and_multiline_records(input_data):
    input_data["description"] = "第一行\n第二行"
    files = csv_files(input_data)
    files = edit_csv(files, "task_scenarios", lambda rows: rows.insert(1, []))
    files = edit_csv(files, "task_links", lambda rows: rows.__setitem__(slice(1, 1), [[], []]))
    result = parse_task_files(files, input_data)
    assert result["issues"] == [] and result["taskRow"] == 3 and result["sourceRows"] == {"D1": 4}
    assert result["input"]["description"] == input_data["description"]


@pytest.mark.parametrize("name", ["=SUM(1,2)", "+literal", "-literal", "@literal", "  =literal", '~rp~"literal"'])
def test_csv_formula_prefix_and_escape_prefix_roundtrip_as_text(input_data, name):
    input_data["name"] = name
    files = csv_files(input_data)
    assert "~rp~" in files[0][1].decode("utf-8")
    result = parse_task_files(files, input_data)
    assert result["issues"] == [] and result["input"]["name"] == name


def test_invalid_escape_is_not_silently_rewritten_and_bom_is_accepted(input_data):
    files = csv_files(input_data)
    result = parse_task_files([(name, b"\xef\xbb\xbf" + data) for name, data in files], input_data)
    assert result["issues"] == []
    with pytest.raises(TaskFileError) as error:
        parse_task_files(set_csv(files, "task_scenarios", "task_name", "~rp~false"), input_data)
    assert error.value.code == "IMPORT_FORMAT"


@pytest.mark.parametrize("count", [5000, 5001])
def test_config_exact_demand_boundary(input_data, count):
    # 仅CSV结构解析，不运行算法；5001由文件读取上限明确拒绝。
    files = csv_files(input_data)
    def edit(rows):
        row = rows[1]
        rows[1:] = [["D" + str(index), *row[1:]] for index in range(count)]
    files = edit_csv(files, "task_links", edit)
    assert parser.MAX_DEMANDS == 5000
    if count == 5001:
        with pytest.raises(TaskFileError) as error:
            parse_task_files(files, input_data)
        assert error.value.code == "IMPORT_LIMIT"
    else:
        result = parse_task_files(files, input_data)
        assert result["issues"] == [] and len(result["input"]["demands"]) == count


@pytest.mark.parametrize("kind", ["formula", "error", "date", "hidden-row", "hidden-column"])
def test_xlsx_hidden_formula_errors_and_date_objects_never_become_input(input_data, kind):
    def edit(book):
        sheet = book["task_links"]
        if kind == "formula":
            sheet["G2"] = "=1+1"
        if kind == "error":
            sheet["G2"] = "#DIV/0!"
            sheet["G2"].data_type = "e"
        if kind == "date":
            book["task_scenarios"]["D2"] = datetime(2026, 10, 23, 8)
        if kind == "hidden-row":
            sheet.row_dimensions[2].hidden = True
        if kind == "hidden-column":
            sheet.column_dimensions["G"].hidden = True
    result = parse_task_files(workbook_file(input_data, edit), input_data)
    assert result["input"] is None and result["issues"]
    assert all("sheet" in issue and "row" in issue for issue in result["issues"])


@pytest.mark.parametrize("kind", ["hidden", "hidden-note", "unknown", "missing", "merged"])
def test_workbook_table_set_and_structural_ambiguity_are_rejected(input_data, kind):
    def edit(book):
        if kind == "hidden":
            book["task_links"].sheet_state = "hidden"
        if kind == "hidden-note":
            book["说明"].sheet_state = "hidden"
        if kind == "unknown":
            book.create_sheet("extra")
        if kind == "missing":
            book.remove(book["task_links"])
        if kind == "merged":
            book["task_links"].merge_cells("A2:B2")
    with pytest.raises(TaskFileError) as error:
        parse_task_files(workbook_file(input_data, edit), input_data)
    assert error.value.code == "IMPORT_FORMAT"


def test_xlsx_literal_formula_text_is_not_a_formula_and_no_escape_decoding(input_data):
    for name in ("=literal", '~rp~"literal"'):
        input_data["name"] = name
        files = workbook_file(input_data)
        book = load_workbook(BytesIO(files[0][1]), data_only=False)
        try:
            assert book["task_scenarios"]["B2"].data_type == "s"
        finally:
            book.close()
        assert parse_task_files(files, input_data)["input"]["name"] == name


def test_forged_small_dimension_does_not_hide_data(input_data):
    def change(root):
        next(element for element in root if parser._local(element.tag) == "dimension").set("ref", "A1")
    files = archive_file(workbook_file(input_data), lambda entries: edit_xml(entries, "xl/worksheets/sheet2.xml", change))
    result = parse_task_files(files, input_data)
    assert result["issues"] == [] and result["sourceRows"] == {"D1": 2}
    assert result["input"]["demands"] == input_data["demands"]


@pytest.mark.parametrize("kind", ["duplicate-cell", "wrong-row", "large-row", "large-column"])
def test_actual_xml_coordinates_cannot_be_hidden_or_overwrite_cells(input_data, kind):
    def change(root):
        data = next(element for element in root if parser._local(element.tag) == "sheetData")
        row = list(data)[1]
        if kind == "duplicate-cell":
            row.append(deepcopy(list(row)[0]))
        if kind == "wrong-row":
            list(row)[0].set("r", "A3")
        if kind == "large-row":
            row.set("r", "1048576")
        if kind == "large-column":
            list(row)[0].set("r", "XFD2")
    files = archive_file(workbook_file(input_data), lambda entries: edit_xml(entries, "xl/worksheets/sheet2.xml", change))
    with pytest.raises(TaskFileError):
        parse_task_files(files, input_data)


@pytest.mark.parametrize("kind", ["path", "absolute", "backslash", "duplicate", "case-duplicate", "macro", "entity", "external", "bomb", "members"])
def test_archive_security_is_checked_before_openpyxl(input_data, kind):
    def edit(entries):
        if kind == "path":
            entries.append(("xl/../outside.xml", b"<x/>"))
        if kind == "absolute":
            entries.append(("/outside.xml", b"<x/>"))
        if kind == "backslash":
            entries.append(("xl\\outside.xml", b"<x/>"))
        if kind == "duplicate":
            entries.append(entries[0])
        if kind == "case-duplicate":
            entries.append((entries[0][0].upper(), entries[0][1]))
        if kind == "macro":
            entries.append(("xl/vbaProject.bin", b"UNIT_TEST"))
        if kind == "entity":
            entries.append(("entity.xml", b'<!DOCTYPE x [<!ENTITY secret "unit">]><x>&secret;</x>'))
        if kind == "external":
            entries.append(("external.rels", b'<Relationships><Relationship TargetMode="External" Target="https://invalid.test/x"/></Relationships>'))
        if kind == "bomb":
            entries.append(("padding", b"x" * 2_000_000))
        if kind == "members":
            entries.extend((f"extra/{index}", b"") for index in range(parser.MAX_ZIP_MEMBERS))
        return entries
    with pytest.raises(TaskFileError) as error:
        parse_task_files(archive_file(workbook_file(input_data), edit), input_data)
    assert error.value.code in {"IMPORT_FORMAT", "IMPORT_LIMIT"}


@pytest.mark.parametrize("kind,code", [("one-csv", "IMPORT_FORMAT"), ("bad-encoding", "IMPORT_ENCODING"), ("size", "IMPORT_LIMIT"),
                                      ("empty", "IMPORT_LIMIT"), ("many", "IMPORT_LIMIT"), ("wrong-extension", "IMPORT_FORMAT"), ("duplicate", "IMPORT_FORMAT")])
def test_file_envelope_errors_have_stable_codes(input_data, kind, code):
    files = csv_files(input_data)
    if kind == "one-csv":
        files = files[:1]
    if kind == "bad-encoding":
        files[0] = (files[0][0], b"\xff\xfe")
    if kind == "size":
        files[0] = (files[0][0], b"x" * (parser.MAX_FILE_BYTES + 1))
    if kind == "empty":
        files[0] = (files[0][0], b"")
    if kind == "many":
        files.append(files[0])
    if kind == "wrong-extension":
        files = [("task.xls", files[0][1])]
    if kind == "duplicate":
        files = [files[0], files[0]]
    with pytest.raises(TaskFileError) as error:
        parse_task_files(files, input_data)
    assert error.value.code == code and error.value.message == str(error.value)


def test_template_and_parser_never_open_business_or_temporary_files(input_data, monkeypatch):
    def forbidden(*_args, **_kwargs):
        raise AssertionError("纯parser/template不得打开文件或创建临时XLSX")
    monkeypatch.setattr("builtins.open", forbidden)
    monkeypatch.setattr("tempfile.NamedTemporaryFile", forbidden)
    for files in (csv_files(input_data), workbook_file(input_data)):
        assert parse_task_files(files, input_data)["issues"] == []


def test_source_has_no_database_or_mock_dependency():
    import ast
    tree = ast.parse(Path(parser.__file__).read_text(encoding="utf-8"))
    modules = [node.module for node in ast.walk(tree) if isinstance(node, ast.ImportFrom)]
    assert all(not any(token in (module or "") for token in ("database", "service", "mock", "models")) for module in modules)


def test_default_hidden_rows_are_not_silently_read_as_visible(input_data):
    def change(root):
        element = next(element for element in root if parser._local(element.tag) == "sheetFormatPr")
        element.set("zeroHeight", "1")
    files = archive_file(workbook_file(input_data), lambda entries: edit_xml(entries, "xl/worksheets/sheet2.xml", change))
    result_issue(parse_task_files(files, input_data), "row", "task_links", 1)


def test_template_does_not_allow_openpyxl_to_truncate_long_json_or_text(input_data):
    input_data["deploymentConstraints"] = {"maxSlopeDeg": None, "elevationRangeM": None, "allowedLandcover": ["x" * 33000],
                                           "roadMaxDistanceM": None, "minSiteSpacingM": None, "losRequiredNodeIds": [],
                                           "gridStepM": None, "forcedInclude": [], "forcedExcludePolygons": []}
    with pytest.raises(TaskFileError) as error:
        workbook_file(input_data)
    assert error.value.code == "IMPORT_LIMIT"
    assert parse_task_files(csv_files(input_data), input_data)["input"] == input_data


def test_xlsx_oversized_actual_cell_is_rejected_not_truncated(input_data):
    def change(root):
        cell = next(element for element in root.iter() if parser._local(element.tag) == "c" and element.get("r") == "B2")
        text = next(element for element in cell.iter() if parser._local(element.tag) == "t")
        text.text = "x" * 33000
    files = archive_file(workbook_file(input_data), lambda entries: edit_xml(entries, "xl/worksheets/sheet1.xml", change))
    with pytest.raises(TaskFileError) as error:
        parse_task_files(files, input_data)
    assert error.value.code == "IMPORT_LIMIT"


@pytest.mark.parametrize("value", [float("nan"), float("inf"), datetime(2026, 10, 23), "\ud800"])
def test_invalid_defaults_are_rejected_before_a_non_json_preview_is_produced(input_data, value):
    files = csv_files(input_data)
    input_data["invalid"] = value
    with pytest.raises(TaskFileError) as error:
        parse_task_files(files, input_data)
    assert error.value.code == "IMPORT_FORMAT"


def test_issue_response_has_a_fixed_bound(input_data, monkeypatch):
    monkeypatch.setattr(parser, "MAX_ISSUES", 2)
    files = set_csv(csv_files(input_data), "task_links", "required_rate_kbps", "NaN")
    files = set_csv(files, "task_links", "max_delay_ms", "NaN")
    files = set_csv(files, "task_links", "min_reliability", "NaN")
    with pytest.raises(TaskFileError) as error:
        parse_task_files(files, input_data)
    assert error.value.code == "IMPORT_LIMIT"
