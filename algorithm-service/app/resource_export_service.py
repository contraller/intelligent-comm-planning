"""按不可变资源快照导出，不生成版本、回执或本地业务文件。"""
import csv
from copy import deepcopy
from io import BytesIO, StringIO
import json
import math
import re

from openpyxl import Workbook
from openpyxl.styles import Font

from . import frontend_integration as frontend, resource_dataset_service as datasets
from .resource_dataset_service import ResourceDatasetError
from .resource_read_filter import filter_device_items

HEADERS = {
    "nodes": "node_id,node_name,echelon,node_subtype,parent_node_id,node_role,mobility,lon,lat,elevation_m,region,device_class,relay_capable,status,is_key,remark".split(","),
    "devices": "device_id,node_id,model_id,antenna_id,antenna_height_m,tilt_deg,azimuth_deg,tx_power_dbm,work_freq_khz,status,install_time,availability".split(","),
    "interference_sources": "interference_id,name,lon,lat,elevation_m,antenna_height_m,center_freq_khz,bandwidth_khz,tx_power_dbm,antenna_gain_dbi,interference_type,pattern_type,azimuth_deg,duty_cycle,active".split(","),
}
TABLES = {"nodes": "node.csv", "devices": "device.csv", "interference_sources": "interference_source.csv"}
NUMERIC = set("lon,lat,elevation_m,antenna_height_m,tilt_deg,azimuth_deg,tx_power_dbm,work_freq_khz,center_freq_khz,bandwidth_khz,antenna_gain_dbi,duty_cycle".split(","))
BOOLEANS = {"relay_capable", "is_key", "active"}
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MAX_FILE_BYTES = 20 * 1024 * 1024
OBSERVATION_KEYS = ("snapshot_id", "snapshot_revision", "resource_revision", "data_origin", "source_label")
OBSERVATION_COLUMNS = ["status_mode", *OBSERVATION_KEYS, "node_status_mode"]
OBSERVED_DEVICE_HEADERS = [*HEADERS["devices"], "alert", "observed_at", "record_kind", *OBSERVATION_COLUMNS]


def _project_devices(devices, tables, query):
    try:
        items = [frontend._device_item(row, tables["node.csv"], tables["device_model.csv"], tables["antenna_model.csv"]) for row in devices]
    except (ValueError, TypeError, KeyError, StopIteration) as exc:
        raise ResourceDatasetError("资源关联或筛选参数来源损坏") from exc
    if query.status_mode != "observed":
        return items, None
    try:
        from .resource_observation_projection import project_device_items
        original = deepcopy(items)
        items, metadata = project_device_items(items, query.revision, query.observation_snapshot_id)
    except Exception as exc:
        if isinstance(exc, ResourceDatasetError):
            raise
        raise ResourceDatasetError("观测投影不可用：" + str(exc), getattr(exc, "code", 5003), getattr(exc, "status", 503)) from exc
    if (not isinstance(metadata, dict) or any(not isinstance(metadata.get(key), str) or not metadata[key].strip() for key in OBSERVATION_KEYS)
            or metadata["snapshot_id"] != query.observation_snapshot_id or metadata["resource_revision"] != query.revision
            or metadata["data_origin"] != "SYNTHETIC_INTEGRATION" or not re.fullmatch(r"device-status-[0-9a-f]{64}", metadata["snapshot_revision"])
            or not isinstance(items, list) or len(items) != len(original)):
        raise ResourceDatasetError("观测投影来源或范围不符", 5002, 503)
    for base, row in zip(original, items):
        if (not isinstance(row, dict) or any(key not in row for key in ("status", "alert", "observed_at"))
                or any(row.get(key) != value for key, value in base.items() if key not in ("status", "alert", "observed_at"))
                or row.get("status") not in (None, "NORMAL", "FAULT", "OFFLINE")
                or row.get("alert") is not None and type(row["alert"]) is not bool
                or row.get("observed_at") is not None and (not isinstance(row["observed_at"], str)
                    or not re.fullmatch(r"\d{4}-\d{2}-\d{2}T\d{2}:\d{2}:\d{2}(?:\.\d{1,6})?Z", row["observed_at"]))):
            raise ResourceDatasetError("观测投影修改了冻结资源或状态字段损坏", 5002, 503)
    return items, {key: metadata[key] for key in OBSERVATION_KEYS}


def select_tables(query):
    if not datasets.database_mode():
        raise ResourceDatasetError("资源导出仅支持已初始化的数据库模式", 5003)
    snapshot = datasets.current()
    if snapshot["revision"] != query.revision:
        raise ResourceDatasetError("资源版本已变化，请刷新后重新导出", 1205, 409)
    tables = snapshot["tables"]
    items, observation = (None, None)
    if query.status_mode == "observed" or query.range == "filtered":
        items, observation = _project_devices(tables["device.csv"], tables, query)
    selected = {kind: {row[HEADERS[kind][0]] for row in tables[TABLES[kind]]} for kind in query.kinds}
    if query.range != "all":
        ids = set(query.selected_ids or [])
        devices = tables["device.csv"]
        if query.range == "selected":
            if ids - {row["device_id"] for row in devices}:
                raise ResourceDatasetError("选中设备不存在，未生成部分导出", 1002, 404)
        else:
            try:
                ids = {row["device_id"] for row in filter_device_items(items, query.filter)}
            except (ValueError, TypeError, KeyError, StopIteration) as exc:
                raise ResourceDatasetError("资源关联或筛选参数来源损坏") from exc
        selected = {kind: ids if kind == "devices" else set() for kind in query.kinds}
    added = 0
    if query.format == "xlsx" and "nodes" in selected:
        nodes = {row["node_id"]: row for row in tables["node.csv"]}
        wanted = selected["nodes"]

        def add(node_id):
            nonlocal added
            if node_id not in wanted:
                wanted.add(node_id)
                added += 1

        for device in tables["device.csv"]:
            if device["device_id"] in selected.get("devices", set()):
                add(device["node_id"])
        visited = set()

        def parents(node_id, trail):
            if node_id in trail:
                raise ResourceDatasetError("节点父级存在循环，无法安全补齐导出")
            if node_id in visited:
                return
            for parent in filter(None, nodes[node_id].get("parent_node_id", "").split(";")):
                add(parent)
                parents(parent, trail | {node_id})
            visited.add(node_id)

        for node_id in sorted(wanted):
            parents(node_id, set())
    result = {kind: sorted((row for row in tables[TABLES[kind]] if row[HEADERS[kind][0]] in ids), key=lambda row: row[HEADERS[kind][0]])
              for kind, ids in selected.items()}
    if observation is not None:
        projected = {row["device_id"]: row for row in items}
        metadata = {"status_mode": "observed", **observation, "node_status_mode": "baseline"}
        result["devices"] = [{**row, "status": projected[row["device_id"]]["status"],
            "alert": projected[row["device_id"]]["alert"], "observed_at": projected[row["device_id"]].get("observed_at"),
            "record_kind": "device", **metadata} for row in result["devices"]]
        snapshot = {**snapshot, "observation": observation}
    return snapshot, result, added


def preview(query):
    snapshot, tables, added = select_tables(query)
    result = {"revision": snapshot["revision"], "data_origin": snapshot["data_origin"], "source": snapshot["source"],
              "tables": [{"kind": kind, "count": len(rows)} for kind, rows in tables.items()], "added": added}
    if "observation" in snapshot:
        result.update(status_mode="observed", observation=snapshot["observation"], node_status_mode="baseline")
    return result


def _value(row, header):
    value = row.get(header, "ACTIVE" if header == "availability" else "")
    if value == "":
        return None
    if header in NUMERIC:
        try:
            number = float(value)
            if not math.isfinite(number):
                raise ValueError("nonfinite")
            return number
        except (ValueError, TypeError) as exc:
            raise ResourceDatasetError("资源数值损坏：" + header) from exc
    if header in BOOLEANS:
        if value not in ("true", "false"):
            raise ResourceDatasetError("资源布尔值损坏：" + header)
        return value == "true"
    return value


def _csv_value(value):
    if isinstance(value, str) and (re.match(r"^\s*[=+\-@]", value) or value.startswith("~rp~")):
        return "~rp~" + json.dumps(value, ensure_ascii=False, separators=(",", ":"))
    if isinstance(value, bool):
        return str(value).lower()
    return value


def export_file(query):
    snapshot, tables, added = select_tables(query)
    observation = snapshot.get("observation")
    headers = {**HEADERS, **({"devices": OBSERVED_DEVICE_HEADERS} if observation is not None else {})}
    if query.format == "csv":
        kind = query.kinds[0]
        output = StringIO(newline="")
        writer = csv.writer(output)
        writer.writerow(headers[kind])
        if observation is not None:
            # 独立元数据记录不是设备；零匹配文件也保留完整快照来源。
            metadata = {"record_kind": "metadata", "status_mode": "observed", **observation, "node_status_mode": "baseline"}
            writer.writerow([_csv_value(metadata.get(key)) for key in headers[kind]])
        for row in tables[kind]:
            writer.writerow([_csv_value(_value(row, key)) for key in headers[kind]])
        data = output.getvalue().encode("utf-8")
        mime = "text/csv; charset=utf-8"
    else:
        book = Workbook()
        book.active.title = query.kinds[0]
        for i, (kind, rows) in enumerate(tables.items()):
            sheet = book.active if i == 0 else book.create_sheet(kind)
            sheet.append(headers[kind])
            for row in rows:
                sheet.append([_value(row, key) for key in headers[kind]])
                # 文本显式设置为字符串，避免以 = 开头的来源内容被当作公式执行。
                for cell in sheet[sheet.max_row]:
                    if isinstance(cell.value, str):
                        cell.data_type = "s"
            sheet.freeze_panes = "A2"
            for cell in sheet[1]:
                cell.font = Font(bold=True)
        notes = book.create_sheet("说明")
        for row in [("资源版本", snapshot["revision"]), ("数据来源", snapshot["data_origin"]),
                    ("来源说明", snapshot["source"]), ("内容哈希", snapshot["content_hash"]),
                    ("导出范围", query.range), ("补齐关联节点", added),
                    ("CSV 文本转义", "公式前缀文本采用 ~rp~ 加 JSON 字符串；原始数据未修改"),
                    ("数据性质", "构造联调资源，非生产、非实测；不包含未保存输入")]:
            notes.append(row)
            for cell in notes[notes.max_row]:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
        if observation is not None:
            for row in [("状态口径", "observed"), *observation.items(), ("节点状态口径", "baseline"),
                        ("观测说明", "仅设备状态、告警和观测时间来自所选构造快照；未覆盖设备为空，节点闭包仍为台账基线")]:
                notes.append(row)
                for cell in notes[notes.max_row]:
                    if isinstance(cell.value, str):
                        cell.data_type = "s"
        output = BytesIO()
        book.save(output)
        book.close()
        data, mime = output.getvalue(), XLSX_MIME
    if not data or len(data) > MAX_FILE_BYTES:
        raise ResourceDatasetError("导出文件为空或超过 20 MiB，未生成截断文件", 1001, 413)
    return data, mime, "synthetic-resources-" + snapshot["revision"] + ("-observed" if observation is not None else "") + "." + query.format
