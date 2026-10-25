"""SR-5 service: scenarios, asynchronous diagnosis, impact and reports."""
from __future__ import annotations

import csv
import hashlib
import io
import json
import math
import os
import sys
import threading
import uuid
from copy import deepcopy
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from sqlalchemy import select
from sqlalchemy.exc import IntegrityError

from .database import get_session_factory
from .fault_scenario_models import FaultScenario, FaultScenarioVersion
from .fault_submission_models import FaultTaskSubmission
from .planning_w3 import _spawn, get_task_record, CODE_OK, CODE_TASK_NOT_FOUND
from .fault_plan_source import FaultPlanSourceError, load_plan_source

ROOT = Path(__file__).resolve().parents[2]
if str(ROOT / "scripts") not in sys.path:
    sys.path.insert(0, str(ROOT / "scripts"))

from diagnosis.sr5 import ALLOWED_OBJECTS, FAULT_TYPES, FaultEngine, FaultInputError  # noqa: E402

_ENGINE: FaultEngine | None = None
_SCENARIOS: dict[str, dict[str, Any]] | None = None
_SLOCK = threading.RLock()
_SUBMISSION_LOCK = threading.RLock()
STORE_PATH = Path(os.environ.get("SR5_STORE_PATH", ROOT / "data" / "runtime" / "fault_scenarios.json"))
REPORT_DIR = Path(os.environ.get("SR5_REPORT_DIR", ROOT / "data" / "runtime" / "reports"))


def _engine() -> FaultEngine:
    global _ENGINE
    if _ENGINE is None:
        with _SLOCK:
            if _ENGINE is None:
                _ENGINE = FaultEngine()
    return _ENGINE


def _scenarios() -> dict[str, dict[str, Any]]:
    global _SCENARIOS
    if _SCENARIOS is None:
        with _SLOCK:
            if _SCENARIOS is None:
                if STORE_PATH.exists():
                    with STORE_PATH.open(encoding="utf-8") as handle:
                        seeds = {s["scenario_id"]: s for s in json.load(handle)}
                else:
                    seeds = {s["scenario_id"]: s for s in _engine().default_scenarios}
                _SCENARIOS = seeds
    return _SCENARIOS


def _save_scenarios() -> None:
    STORE_PATH.parent.mkdir(parents=True, exist_ok=True)
    tmp = STORE_PATH.with_suffix(".tmp")
    tmp.write_text(json.dumps(list(_scenarios().values()), ensure_ascii=False, indent=2), encoding="utf-8")
    tmp.replace(STORE_PATH)


def _ok(data):
    return {"code": CODE_OK, "message": "success", "data": data}


def _bad(exc: FaultInputError):
    return {"code": exc.code, "message": str(exc), "data": exc.detail}


def _submission_digest(payload):
    raw = json.dumps(payload, ensure_ascii=False, sort_keys=True,
                     separators=(",", ":")).encode("utf-8")
    return hashlib.sha256(raw).hexdigest()


def _request_key(payload):
    request_id = payload.get("request_id")
    if not isinstance(request_id, str) or not request_id.strip() or len(request_id) > 120:
        raise FaultPlanSourceError(1001, "正式分析须提供长度不超过 120 的 request_id", 422)
    if payload.get("callback_url") not in (None, ""):
        raise FaultPlanSourceError(1001, "callback_url 尚未支持，请轮询任务", 422)
    return request_id, _submission_digest(payload)


def _matching_submission(session, request_id, digest, kind):
    row = session.get(FaultTaskSubmission, request_id)
    if row is not None and (row.request_digest != digest or row.kind != kind):
        raise FaultPlanSourceError(1204, "同一请求编号对应不同内容", 409)
    return row


def _replay_submission(payload, kind):
    request_id, digest = _request_key(payload)
    with get_session_factory()() as session:
        row = _matching_submission(session, request_id, digest, kind)
    if row is None:
        return None
    return _ok({"task_id": row.algorithm_task_id, "estimated_sec": 5,
                "request_id": request_id, "replayed": True})


def _reserve_submission(payload, kind, task_id, plan_id, scenario_id,
                        version_id, diagnose_task_id=None):
    request_id, digest = _request_key(payload)
    with _SUBMISSION_LOCK:
        try:
            with get_session_factory().begin() as session:
                row = _matching_submission(session, request_id, digest, kind)
                if row is not None:
                    return row.algorithm_task_id, False
                algorithm_task_id = "tsk-" + uuid.uuid4().hex[:12]
                session.add(FaultTaskSubmission(
                    request_id=request_id, request_digest=digest, kind=kind,
                    algorithm_task_id=algorithm_task_id, task_id=task_id,
                    plan_id=plan_id, scenario_id=scenario_id,
                    scenario_version_id=version_id,
                    diagnose_task_id=diagnose_task_id,
                    submitted_at=datetime.now(timezone.utc)))
            return algorithm_task_id, True
        except IntegrityError:
            with get_session_factory()() as session:
                row = _matching_submission(session, request_id, digest, kind)
                if row is None:
                    raise
                return row.algorithm_task_id, False


def _spawn_formal(kind, payload, work, *, task_id, plan_id, scenario_id,
                  version_id, diagnose_task_id=None):
    algorithm_task_id, created = _reserve_submission(
        payload, kind, task_id, plan_id, scenario_id, version_id, diagnose_task_id)
    if created:
        callback = None
        if payload.get("frontend_run_id"):
            from .fault_frontend_service import record_runtime
            callback = lambda state: record_runtime(payload["frontend_run_id"], kind, state)
        result = _spawn(kind, payload, work, task_id=algorithm_task_id, on_update=callback)
    else:
        result = _ok({"task_id": algorithm_task_id, "estimated_sec": 5})
    if result["code"] == CODE_OK:
        result["data"].update(request_id=payload["request_id"], replayed=not created)
    return result


def _formal_source(payload):
    if payload.get("context") != "planning":
        raise FaultPlanSourceError(1001, "正式场景须显式设置 context=planning", 422)
    return load_plan_source(payload.get("task_id"), payload.get("plan_id"))


def _formal_scenario_mode(payload):
    context = payload.get("context") if isinstance(payload, dict) else None
    if context not in (None, "demo", "planning"):
        raise FaultPlanSourceError(1001, "context 只能为 planning 或 demo", 422)
    if context != "planning" and isinstance(payload, dict) and any(
        payload.get(key) is not None for key in ("task_id", "plan_id")
    ):
        raise FaultPlanSourceError(1001, "正式场景须显式设置 context=planning", 422)
    return context == "planning"


def _formal_scenario_query(context, task_id):
    if context not in (None, "demo", "planning") or (context != "planning" and task_id is not None):
        raise FaultPlanSourceError(1001, "正式场景查询须显式设置 context=planning", 422)
    return context == "planning"


def _formal_snapshot(payload, source):
    snapshot = {"scenario_name": payload.get("scenario_name"),
                "faults": payload.get("faults")}
    if "description" in payload:
        snapshot["description"] = payload["description"]
    _engine().for_network(source["network"]).validate(snapshot)
    return deepcopy(snapshot)


def _version_dto(scenario, version):
    return {"scenario_id": scenario.scenario_id, "task_id": scenario.task_id,
            "archived": scenario.archived, "version_id": version.version_id,
            "version_no": version.version_no, "plan_id": version.plan_id,
            "plan_revision": version.plan_revision,
            "resource_revision": version.resource_revision,
            "saved_at": version.saved_at.isoformat(), **deepcopy(version.snapshot)}


def _formal_record(session, scenario_id, task_id):
    if not isinstance(task_id, str) or not task_id:
        raise FaultPlanSourceError(1001, "task_id 必填", 422)
    scenario = session.get(FaultScenario, scenario_id)
    if scenario is None or scenario.task_id != task_id:
        raise FaultPlanSourceError(1002, "故障场景不存在或不属于该任务", 404)
    return scenario


def _latest_version(session, scenario_id):
    return session.scalars(
        select(FaultScenarioVersion).where(FaultScenarioVersion.scenario_id == scenario_id)
        .order_by(FaultScenarioVersion.version_no.desc()).limit(1)
    ).one()


def _new_version(session, scenario, source, snapshot, version_no):
    meta = source["meta"]
    version = FaultScenarioVersion(
        version_id="FSV-" + uuid.uuid4().hex[:16].upper(),
        scenario_id=scenario.scenario_id, version_no=version_no,
        plan_id=meta["plan_id"], plan_revision=meta["revision"],
        resource_revision=meta["resource_revision"], snapshot=snapshot,
        saved_at=datetime.now(timezone.utc))
    session.add(version)
    return version


def validate_scenario(payload):
    try:
        if _formal_scenario_mode(payload):
            source = _formal_source(payload)
            snapshot = _formal_snapshot(payload, source)
            return _ok({"valid": True, "fault_count": len(snapshot["faults"]),
                        "plan_id": source["meta"]["plan_id"]})
        return _ok(_engine().validate(payload or {}))
    except FaultInputError as exc:
        return _bad(exc)


def list_scenarios(context=None, task_id=None, include_archived=False):
    if _formal_scenario_query(context, task_id):
        if not isinstance(task_id, str) or not task_id:
            raise FaultPlanSourceError(1001, "task_id 必填", 422)
        with get_session_factory()() as session:
            scenarios = session.scalars(select(FaultScenario).where(
                FaultScenario.task_id == task_id).order_by(FaultScenario.created_at)).all()
            items = [_version_dto(row, _latest_version(session, row.scenario_id))
                     for row in scenarios if include_archived or not row.archived]
        return _ok({"total": len(items), "items": items})
    with _SLOCK:
        rows = sorted(_scenarios().values(), key=lambda s: s["scenario_id"])
        return _ok({"total": len(rows), "items": rows})


def get_scenario(sid, context=None, task_id=None, version_id=None):
    if _formal_scenario_query(context, task_id):
        with get_session_factory()() as session:
            scenario = _formal_record(session, sid, task_id)
            version = session.get(FaultScenarioVersion, version_id) if version_id else _latest_version(session, sid)
            if version is None or version.scenario_id != sid:
                raise FaultPlanSourceError(1002, "场景版本不存在", 404)
            return _ok(_version_dto(scenario, version))
    with _SLOCK:
        row = _scenarios().get(sid)
        return _ok(row) if row else {"code": 1002, "message": "故障场景不存在", "data": None}


def create_scenario(payload):
    if _formal_scenario_mode(payload):
        source = _formal_source(payload)
        try:
            snapshot = _formal_snapshot(payload, source)
        except FaultInputError as exc:
            return _bad(exc)
        now = datetime.now(timezone.utc)
        with get_session_factory().begin() as session:
            scenario = FaultScenario(scenario_id="FS-" + uuid.uuid4().hex[:16].upper(),
                                     task_id=source["meta"]["task_id"], archived=False,
                                     created_at=now, updated_at=now)
            session.add(scenario)
            version = _new_version(session, scenario, source, snapshot, 1)
            data = _version_dto(scenario, version)
        return _ok(data)
    try:
        _engine().validate(payload or {})
    except FaultInputError as exc:
        return _bad(exc)
    with _SLOCK:
        sid = "FS-" + uuid.uuid4().hex[:10].upper()
        row = dict(payload, scenario_id=sid, create_time=datetime.now(timezone.utc).isoformat())
        _scenarios()[sid] = row
        _save_scenarios()
        return _ok(row)


def update_scenario(sid, payload):
    if _formal_scenario_mode(payload):
        source = _formal_source(payload)
        try:
            snapshot = _formal_snapshot(payload, source)
        except FaultInputError as exc:
            return _bad(exc)
        with get_session_factory().begin() as session:
            scenario = _formal_record(session, sid, source["meta"]["task_id"])
            if scenario.archived:
                raise FaultPlanSourceError(1203, "已归档场景不能修改，请先恢复", 409)
            version = _new_version(session, scenario, source, snapshot,
                                   _latest_version(session, sid).version_no + 1)
            scenario.updated_at = version.saved_at
            data = _version_dto(scenario, version)
        return _ok(data)
    with _SLOCK:
        if sid not in _scenarios():
            return {"code": 1002, "message": "故障场景不存在", "data": None}
        try:
            _engine().validate(payload or {})
        except FaultInputError as exc:
            return _bad(exc)
        row = dict(payload, scenario_id=sid, create_time=_scenarios()[sid].get("create_time"),
                   update_time=datetime.now(timezone.utc).isoformat())
        _scenarios()[sid] = row
        _save_scenarios()
        return _ok(row)


def archive_scenario(sid, payload):
    if not isinstance(payload.get("archived"), bool):
        raise FaultPlanSourceError(1001, "archived 必须为布尔值", 422)
    with get_session_factory().begin() as session:
        scenario = _formal_record(session, sid, payload.get("task_id"))
        scenario.archived = payload["archived"]
        scenario.updated_at = datetime.now(timezone.utc)
        data = _version_dto(scenario, _latest_version(session, sid))
    return _ok(data)


def delete_scenario(sid, context=None, task_id=None):
    if _formal_scenario_query(context, task_id):
        data = archive_scenario(sid, {"task_id": task_id, "archived": True})["data"]
        return _ok({"scenario_id": sid, "archived": True,
                    "version_id": data["version_id"]})
    with _SLOCK:
        if sid not in _scenarios():
            return {"code": 1002, "message": "故障场景不存在", "data": None}
        del _scenarios()[sid]
        _save_scenarios()
        return _ok({"scenario_id": sid, "deleted": True})


def import_scenarios(payload):
    """Accept a JSON array or CSV text; reject the entire batch on any bad row."""
    if _formal_scenario_mode(payload):
        source = _formal_source(payload)
        scenarios = payload.get("scenarios")
        if not isinstance(scenarios, list) or not scenarios:
            return _bad(FaultInputError("scenarios 必须是非空列表"))
        try:
            snapshots = [_formal_snapshot(row, source) for row in scenarios]
        except (AttributeError, FaultInputError) as exc:
            return _bad(exc if isinstance(exc, FaultInputError) else FaultInputError("场景行必须是对象"))
        now = datetime.now(timezone.utc)
        with get_session_factory().begin() as session:
            rows = []
            for snapshot in snapshots:
                scenario = FaultScenario(scenario_id="FS-" + uuid.uuid4().hex[:16].upper(),
                                         task_id=source["meta"]["task_id"], archived=False,
                                         created_at=now, updated_at=now)
                session.add(scenario)
                version = _new_version(session, scenario, source, snapshot, 1)
                rows.append(_version_dto(scenario, version))
        return _ok({"imported": len(rows), "items": rows})
    if isinstance(payload, dict) and payload.get("format") == "CSV":
        try:
            rows = list(csv.DictReader(io.StringIO(payload["content"])))
            scenarios = [dict(scenario_name=r["scenario_name"], faults=json.loads(r["faults"]))
                         for r in rows]
        except (KeyError, ValueError, TypeError) as exc:
            return _bad(FaultInputError("CSV 格式错误：%s" % exc))
    else:
        scenarios = payload.get("scenarios") if isinstance(payload, dict) else payload
    if not isinstance(scenarios, list) or not scenarios:
        return _bad(FaultInputError("scenarios 必须是非空列表"))
    try:
        for scenario in scenarios:
            _engine().validate(scenario)
    except FaultInputError as exc:
        return _bad(exc)
    with _SLOCK:
        added = []
        for scenario in scenarios:
            sid = "FS-" + uuid.uuid4().hex[:10].upper()
            row = dict(scenario, scenario_id=sid,
                       create_time=datetime.now(timezone.utc).isoformat())
            _scenarios()[sid] = row
            added.append(sid)
        _save_scenarios()
    return _ok({"imported": len(added), "scenario_ids": added})


def import_scenarios_file(filename, contents, *, context=None, task_id=None, plan_id=None):
    """CSV/XLSX rows contain scenario_name and JSON-encoded faults columns."""
    suffix = Path(filename or "").suffix.lower()
    try:
        if suffix == ".csv":
            rows = list(csv.DictReader(io.StringIO(contents.decode("utf-8-sig"))))
        elif suffix == ".xlsx":
            from openpyxl import load_workbook
            book = load_workbook(io.BytesIO(contents), read_only=True, data_only=True)
            values = book.active.iter_rows(values_only=True)
            header = next(values)
            rows = [dict(zip(header, row)) for row in values]
            book.close()
        else:
            return _bad(FaultInputError("仅支持 CSV/XLSX 场景文件"))
        scenarios = [{"scenario_name": row["scenario_name"],
                      "faults": json.loads(row["faults"])} for row in rows]
    except (KeyError, ValueError, TypeError, UnicodeError, StopIteration) as exc:
        return _bad(FaultInputError("导入文件格式错误：%s" % exc))
    return import_scenarios({"context": context, "task_id": task_id,
                             "plan_id": plan_id, "scenarios": scenarios})


def _diagnosis(tid):
    task = get_task_record(tid)
    if task.get("kind") not in {"FAULT_DIAGNOSE", "FAULT_DIAGNOSE_OBSERVATION"} or task.get("status") != "SUCCESS":
        return None
    return task["result"]


def submit_diagnosis_feedback(task_id, payload, *, session=None):
    """把人工确认结论送入 SR-6 待审核队列，审核发布后才参与后续推理。"""
    diagnosis = _diagnosis(task_id)
    if diagnosis is None:
        return {"code": CODE_TASK_NOT_FOUND, "message": "诊断任务不存在或尚未完成", "data": None}
    payload = payload or {}
    reviewer = str(payload.get("reviewer") or "").strip()
    index = payload.get("result_index", 0)
    if not reviewer or not isinstance(index, int) or not 0 <= index < len(diagnosis["results"]):
        return _bad(FaultInputError("reviewer 必填，result_index 必须指向有效诊断结果"))
    result = diagnosis["results"][index]
    observations = diagnosis.get("input_snapshot", {}).get("faults", [])
    if index >= len(observations):
        return _bad(FaultInputError("该任务由旧版本生成，缺少反馈所需的输入快照，请重新诊断"))
    observation = observations[index]
    fault_type = payload.get("confirmed_fault_type") or result["fused"].get("fault_type")
    object_type = result["fault_object"]["object_type"]
    if fault_type not in FAULT_TYPES or object_type not in ALLOWED_OBJECTS[fault_type]:
        return _bad(FaultInputError("确认的故障类型与对象类型不适配", 3003))
    symptom_ids = observation.get("symptom_ids") or []
    if not symptom_ids:
        return _bad(FaultInputError("至少需要一个已观察现象才能沉淀为可复用案例"))

    kg_branch = next((row for row in result["branches"] if row["source"] == "KG"), {})
    rule_branch = next((row for row in result["branches"] if row["source"] == "RULE"), {})
    root_cause = str(payload.get("root_cause") or result["fused"].get("root_cause") or "").strip()
    component = str(payload.get("fault_component") or result.get("fault_component") or "").strip()
    solution = str(payload.get("solution") or kg_branch.get("action") or rule_branch.get("action") or "").strip()
    steps = payload.get("diagnosis_steps")
    if steps is None:
        steps = kg_branch.get("steps") or []
    if not root_cause or not solution or not isinstance(steps, list) or not all(
            isinstance(step, str) and step.strip() for step in steps):
        return _bad(FaultInputError("root_cause、solution 必填，diagnosis_steps 必须是字符串列表"))

    case_id = "DF-%s-%02d" % (task_id.removeprefix("tsk-").upper(), index + 1)
    provenance = "INTEGRATION_CONFIRMED_CONSTRUCTED" if payload.get("_workflow_request") else "UNAUTHENTICATED_REVIEW"
    source = "diagnosis_feedback#%s by %s [%s]" % (case_id, reviewer, provenance)
    ops = [
        {"op": "add_triple", "head_type": "Fault", "head": case_id,
         "relation": "instanceOf", "tail_type": "FaultType", "tail": fault_type, "source": source},
        {"op": "add_triple", "head_type": "Fault", "head": case_id,
         "relation": "occursOn", "tail_type": "ObjectType", "tail": object_type, "source": source},
        {"op": "add_triple", "head_type": "Fault", "head": case_id,
         "relation": "recordedIn", "tail_type": "Case", "tail": case_id, "source": source},
        {"op": "add_triple", "head_type": "Fault", "head": case_id,
         "relation": "causedBy", "tail_type": "Cause", "tail": root_cause, "source": source},
        {"op": "add_triple", "head_type": "Fault", "head": case_id,
         "relation": "resolvedBy", "tail_type": "Solution", "tail": solution, "source": source},
    ]
    if component:
        ops.append({"op": "add_triple", "head_type": "Fault", "head": case_id,
                    "relation": "locatedAt", "tail_type": "Component", "tail": component,
                    "source": source})
    for symptom_id in symptom_ids:
        ops.append({"op": "add_triple", "head_type": "Fault", "head": case_id,
                    "relation": "hasSymptom", "tail_type": "Symptom",
                    "tail": _engine().symptoms[symptom_id], "source": source})
    for step_index, step in enumerate(steps, 1):
        ops.append({"op": "add_triple", "head_type": "Fault", "head": case_id,
                    "relation": "diagnosedBy", "tail_type": "DiagnosisStep",
                    "tail": "%02d:%s" % (step_index, step.strip()), "source": source})

    from . import kg
    if payload.get("_workflow_request"):
        from . import kg_workflow_service
        if session is not None:
            state = kg_workflow_service._bootstrap(session, kg.STORE_PATH)
            published, _ = kg_workflow_service._current_graph(session, state)
            if published.get("Fault", case_id) is not None:
                raise FaultPlanSourceError(1204, "该诊断结论已发布知识反馈，请勿重复提交", 409)
        request = {**payload["_workflow_request"], "ops": ops, "note": "诊断反馈 " + case_id}
        created = kg_workflow_service.save_draft(kg.STORE_PATH, request, _session=session)
        return _ok({**created, "case_id": case_id, "requires_kg_review": True, "provenance": provenance})
    with kg._LOCK:
        graph, _ = kg._load()
        duplicate_pending = any(change["status"] == "PENDING" and any(
            op.get("op") == "add_triple" and op.get("head_type") == "Fault" and
            op.get("head") == case_id for op in change["ops"]) for change in graph.pending)
        if graph.get("Fault", case_id) is not None or duplicate_pending:
            return _bad(FaultInputError("该诊断结果已经提交知识反馈，请勿重复提交"))
        submitted = kg.submit({"ops": ops, "submitter": reviewer,
                               "note": "SR-5 人工确认诊断反馈 %s" % case_id})
    if submitted["code"] == CODE_OK:
        submitted["data"].update({"case_id": case_id, "diagnose_task_id": task_id,
                                  "requires_kg_review": True,
                                  "provenance": "UNAUTHENTICATED_REVIEW"})
    return submitted


def _impact(tid):
    task = get_task_record(tid)
    if task.get("kind") != "FAULT_IMPACT" or task.get("status") != "SUCCESS":
        return None
    return task["result"]


def submit_diagnosis(payload):
    payload = payload or {}
    if payload.get("context") not in (None, "planning", "demo"):
        return _bad(FaultInputError("context 只能为 planning 或 demo"))
    if payload.get("context") != "planning" and any(
        key in payload for key in ("task_id", "plan_id")
    ):
        return _bad(FaultInputError("正式方案诊断须显式设置 context=planning"))
    if payload.get("context") == "planning" and payload.get("network_solution_id"):
        return _bad(FaultInputError("正式诊断应使用 plan_id，不接受旧方案编号"))
    if payload.get("context") == "planning" and not payload.get("plan_id"):
        raise FaultPlanSourceError(1001, "正式诊断必须指定 plan_id", 422)
    sid = payload.get("scenario_id")
    formal = payload.get("context") == "planning"
    if formal:
        replay = _replay_submission(payload, "FAULT_DIAGNOSE")
        if replay is not None:
            return replay
    source = load_plan_source(payload.get("task_id"), payload.get("plan_id")) if formal else None
    if formal:
        version_id = payload.get("scenario_version_id")
        if not isinstance(version_id, str) or not version_id:
            raise FaultPlanSourceError(1001, "正式诊断必须指定 scenario_version_id", 422)
        with get_session_factory()() as session:
            record = _formal_record(session, sid, source["meta"]["task_id"])
            if record.archived:
                raise FaultPlanSourceError(1203, "已归档场景不能提交新诊断", 409)
            version = session.get(FaultScenarioVersion, version_id)
            if version is None or version.scenario_id != sid:
                raise FaultPlanSourceError(1002, "场景版本不存在或不属于该场景", 404)
            if (version.plan_id != source["meta"]["plan_id"]
                    or version.plan_revision != source["meta"]["revision"]
                    or version.resource_revision != source["meta"]["resource_revision"]):
                raise FaultPlanSourceError(1203, "场景版本绑定的规划方案已变化，请保存新版本", 409)
            scenario = {"scenario_id": sid, **deepcopy(version.snapshot)}
            version_no = version.version_no
    else:
        with _SLOCK:
            stored = _scenarios().get(sid)
            scenario = deepcopy(stored) if stored is not None else None
        if scenario is None:
            return {"code": 1002, "message": "故障场景不存在", "data": None}
    network = source["network"] if source else None
    frozen_graph = None
    if payload.get("frontend_run_id"):
        from . import kg
        with kg._LOCK:
            frozen_graph = deepcopy(kg._load()[0])
    try:
        _diagnose_live(scenario, payload, network=network, frozen_graph=frozen_graph)
    except FaultInputError as exc:
        return _bad(exc)
    if not formal and payload.get("network_solution_id") and _solution(payload["network_solution_id"]) is None:
        return {"code": 1002, "message": "规划方案不存在", "data": None}

    def work(_payload, tick):
        tick(0.4, "规则与图谱推理")
        out = _diagnose_live(scenario, payload, network=network, frozen_graph=frozen_graph)
        if source:
            meta = source["meta"]
            out.update(context="planning", network_solution_id=meta["plan_id"],
                       network_task_id=meta["task_id"], network_revision=meta["revision"],
                       network_resource_revision=meta["resource_revision"],
                       scenario_version_id=version_id, scenario_version_no=version_no)
        else:
            out["network_solution_id"] = payload.get("network_solution_id") or "SYNTHETIC_BASELINE"
        return out
    if formal:
        return _spawn_formal("FAULT_DIAGNOSE", payload, work,
                             task_id=source["meta"]["task_id"],
                             plan_id=source["meta"]["plan_id"],
                             scenario_id=sid, version_id=version_id)
    return _spawn("FAULT_DIAGNOSE", payload, work)


def _diagnose_live(scenario, payload, require_fault_type=True, network=None, frozen_graph=None):
    from . import kg
    with kg._LOCK:
        graph = frozen_graph if frozen_graph is not None else kg._load()[0]
        engine = _engine().for_network(network) if network is not None else _engine()
        result = engine.diagnose(scenario, payload.get("fusion_weights"),
                                 payload.get("confidence_threshold", 0.6), graph=graph,
                                 require_fault_type=require_fault_type)
        result["knowledge_version"] = graph.version
        return result


def submit_observation_diagnosis(payload):
    payload = payload or {}
    scenario = {"scenario_name": payload.get("scenario_name") or "观测诊断",
                "faults": payload.get("faults"), "scenario_id": None}
    try:
        _diagnose_live(scenario, payload, require_fault_type=False)
    except FaultInputError as exc:
        return _bad(exc)

    def work(_payload, tick):
        tick(0.4, "根据观测现象推理")
        return _diagnose_live(scenario, payload, require_fault_type=False)

    return _spawn("FAULT_DIAGNOSE_OBSERVATION", payload, work)


def _solution(sid):
    from .planning_w4 import get_plan
    result = get_plan(sid)
    return result["data"] if result["code"] == CODE_OK else None


def submit_impact(payload):
    payload = payload or {}
    if payload.get("context") == "planning" or (
            payload.get("context") is None and payload.get("task_id") and payload.get("plan_id")):
        replay = _replay_submission(payload, "FAULT_IMPACT")
        if replay is not None:
            return replay
    diagnosis = _diagnosis(payload.get("diagnose_task_id"))
    if diagnosis is None:
        return {"code": CODE_TASK_NOT_FOUND, "message": "诊断任务不存在或尚未完成", "data": None}
    formal = diagnosis.get("context") == "planning"
    if formal:
        if payload.get("context") not in (None, "planning"):
            return _bad(FaultInputError("正式诊断的影响分析不能切换到演示模式"))
        if (payload.get("task_id") != diagnosis["network_task_id"]
                or payload.get("plan_id") != diagnosis["network_solution_id"]):
            return _bad(FaultInputError("影响分析的任务和方案须与诊断结果一致"))
        source = load_plan_source(payload["task_id"], payload["plan_id"],
                                  require_current=False)
        if source["meta"]["revision"] != diagnosis["network_revision"]:
            return _bad(FaultInputError("诊断来源方案版本已经变化"))
        version_id = diagnosis.get("scenario_version_id")
        if not version_id:
            raise FaultPlanSourceError(1203, "原诊断缺少场景版本，请重新提交诊断", 409)
        scenario = {"scenario_id": diagnosis["scenario_id"],
                    **deepcopy(diagnosis["input_snapshot"])}
        solution = source["solution"]
        try:
            _engine().for_network(source["network"]).validate(scenario)
        except FaultInputError as exc:
            return _bad(exc)

        def work(_payload, tick):
            tick(0.4, "沿冻结方案的主备路由计算影响")
            out = _engine().for_network(source["network"]).impact(
                scenario, diagnosis, solution, payload.get("consider_backup_route", True))
            out.update(context="planning", network_task_id=source["meta"]["task_id"],
                       network_revision=source["meta"]["revision"],
                       diagnose_task_id=payload["diagnose_task_id"],
                       scenario_version_id=version_id)
            return out
        return _spawn_formal("FAULT_IMPACT", payload, work,
                             task_id=source["meta"]["task_id"],
                             plan_id=source["meta"]["plan_id"],
                             scenario_id=diagnosis["scenario_id"],
                             version_id=version_id,
                             diagnose_task_id=payload["diagnose_task_id"])

    if payload.get("context") == "planning" or any(
        key in payload for key in ("task_id", "plan_id")
    ):
        return _bad(FaultInputError("正式影响分析须引用正式诊断任务"))

    sid = payload.get("network_solution_id") or diagnosis.get("network_solution_id")
    solution = None if sid == "SYNTHETIC_BASELINE" else _solution(sid)
    if sid != "SYNTHETIC_BASELINE" and solution is None:
        return {"code": 1002, "message": "规划方案不存在", "data": None}
    if solution and solution.get("kind") != "SR4_FLOW":
        return _bad(FaultInputError("影响分析需要 SR4_FLOW 方案"))
    scenario = _scenarios().get(diagnosis["scenario_id"])
    if scenario is None:
        return {"code": 1002, "message": "诊断场景已删除", "data": None}
    if solution:
        plan_links = {r["link_id"] for r in solution["detail"]["links"]}
        unknown = [f["object_id"] for f in scenario["faults"]
                   if f["object_type"] == "LINK" and f["object_id"] not in plan_links]
        if unknown:
            return {"code": 1002, "message": "故障链路不在所选规划方案中：%s" % unknown, "data": None}

    def work(_payload, tick):
        tick(0.4, "沿主备路由计算影响")
        return _engine().impact(scenario, diagnosis, solution,
                                payload.get("consider_backup_route", True))
    return _spawn("FAULT_IMPACT", payload, work)


def get_impact_nodes(task_id, node_role=None, region=None):
    impact = _impact(task_id)
    if impact is None:
        return {"code": CODE_TASK_NOT_FOUND, "message": "影响任务不存在或尚未完成", "data": None}
    rows = [row for row in impact["affected_nodes"]
            if (node_role is None or row.get("node_role") == node_role)
            and (region is None or row.get("region") == region)]
    return _ok({"total": len(rows), "items": rows,
                "filters": {"node_role": node_role, "region": region}})


def get_suggestions(diagnose_task_id, impact_task_id=None):
    diagnosis = _diagnosis(diagnose_task_id)
    if diagnosis is None:
        return {"code": CODE_TASK_NOT_FOUND, "message": "诊断任务不存在或尚未完成", "data": None}
    impact = _impact(impact_task_id) if impact_task_id else None
    if impact_task_id and impact is None:
        return {"code": CODE_TASK_NOT_FOUND, "message": "影响分析任务不存在或尚未完成", "data": None}
    if impact and impact["scenario_id"] != diagnosis["scenario_id"]:
        return _bad(FaultInputError("诊断结果与影响分析结果不属于同一场景"))
    if diagnosis.get("context") == "planning" and impact and (
            impact.get("context") != "planning" or
            impact.get("network_solution_id") != diagnosis.get("network_solution_id")):
        return _bad(FaultInputError("诊断与影响分析的方案来源不一致"))
    return _ok(_engine().suggestions(diagnosis, impact))


def _submission_by_task(session, algorithm_task_id, task_id, kind):
    row = session.scalars(select(FaultTaskSubmission).where(
        FaultTaskSubmission.algorithm_task_id == algorithm_task_id)).one_or_none()
    if row is None or row.task_id != task_id or row.kind != kind:
        raise FaultPlanSourceError(2001, "正式分析任务不存在或不属于该规划任务", 404)
    return row


def _analysis_state(record, phase):
    if not record:
        return "FAILED", 0.0, "%s运行记录不可用，请使用新 request_id 重新提交" % phase
    status = record.get("status")
    if status not in {"PENDING", "RUNNING", "SUCCESS", "FAILED", "CANCELLED"}:
        return "FAILED", 0.0, "%s任务状态无效" % phase
    return status, max(0.0, min(1.0, float(record.get("progress") or 0))), record.get("error")


def get_analysis_view(diagnose_task_id, task_id, impact_task_id=None):
    if not isinstance(task_id, str) or not task_id:
        raise FaultPlanSourceError(1001, "task_id 必填", 422)
    with get_session_factory()() as session:
        diagnosis_receipt = _submission_by_task(session, diagnose_task_id, task_id,
                                                "FAULT_DIAGNOSE")
        impact_receipt = None
        if impact_task_id is not None:
            impact_receipt = _submission_by_task(session, impact_task_id, task_id,
                                                 "FAULT_IMPACT")
            if (impact_receipt.diagnose_task_id != diagnose_task_id
                    or impact_receipt.plan_id != diagnosis_receipt.plan_id
                    or impact_receipt.scenario_version_id != diagnosis_receipt.scenario_version_id):
                raise FaultPlanSourceError(1204, "影响任务与诊断任务不属于同一分析", 409)
        version = session.get(FaultScenarioVersion, diagnosis_receipt.scenario_version_id)
        if version is None:
            raise FaultPlanSourceError(1203, "诊断场景版本已不可用", 409)
        scenario = {"scenario_id": diagnosis_receipt.scenario_id,
                    "version_id": version.version_id, "version_no": version.version_no,
                    **deepcopy(version.snapshot)}
        source = {"task_id": task_id, "plan_id": diagnosis_receipt.plan_id,
                  "plan_revision": version.plan_revision,
                  "resource_revision": version.resource_revision}

    diagnosis_task = get_task_record(diagnose_task_id)
    diagnosis_state, diagnosis_progress, diagnosis_error = _analysis_state(
        diagnosis_task, "诊断")
    if diagnosis_task and diagnosis_task.get("kind") != "FAULT_DIAGNOSE":
        raise FaultPlanSourceError(1203, "诊断任务类型不匹配", 409)
    impact_task = get_task_record(impact_task_id) if impact_receipt else None
    impact_state, impact_progress, impact_error = (
        _analysis_state(impact_task, "影响") if impact_receipt else (None, 0.0, None))
    if impact_task and impact_task.get("kind") != "FAULT_IMPACT":
        raise FaultPlanSourceError(1203, "影响任务类型不匹配", 409)

    diagnosis_ready = diagnosis_state == "SUCCESS"
    impact_ready = impact_state == "SUCCESS"
    if not diagnosis_ready:
        state = diagnosis_state
        progress = diagnosis_progress * 0.5
        error = diagnosis_error
    elif impact_receipt is None:
        state, progress, error = "DIAGNOSIS_READY", 0.5, None
    else:
        state = impact_state
        progress = 0.5 + impact_progress * 0.5
        error = impact_error

    diagnosis = deepcopy(diagnosis_task.get("result")) if diagnosis_ready else None
    impact = deepcopy(impact_task.get("result")) if impact_ready else None
    if diagnosis_ready and (
            diagnosis.get("network_solution_id") != source["plan_id"]
            or diagnosis.get("scenario_version_id") != version.version_id):
        raise FaultPlanSourceError(1203, "诊断结果与提交回执不一致", 409)
    if impact_ready and (
            impact.get("diagnose_task_id") != diagnose_task_id
            or impact.get("scenario_version_id") != version.version_id):
        raise FaultPlanSourceError(1203, "影响结果与诊断结果不一致", 409)
    suggestions = _engine().suggestions(diagnosis, impact) if diagnosis_ready else None
    report_ready = diagnosis_ready and impact_ready
    report = _report_content(diagnosis, impact, suggestions) if report_ready else None
    if report is not None:
        report["generated_at"] = datetime.fromtimestamp(
            impact_task["finished_at"], timezone.utc).isoformat()
    phases = [
        {"name": "故障诊断", "status": diagnosis_state,
         "progress": diagnosis_progress},
        {"name": "影响分析", "status": impact_state or "WAITING",
         "progress": impact_progress},
    ]
    return _ok({"task_id": task_id, "plan_id": source["plan_id"],
                "diagnose_task_id": diagnose_task_id,
                "impact_task_id": impact_task_id,
                "request_id": diagnosis_receipt.request_id,
                "impact_request_id": impact_receipt.request_id if impact_receipt else None,
                "state": state, "progress": round(progress, 3), "phases": phases,
                "diagnosis_ready": diagnosis_ready, "impact_ready": impact_ready,
                "report_ready": report_ready, "error": error,
                "source": source, "scenario": scenario,
                "diagnosis": diagnosis, "impact": impact,
                "suggestions": suggestions, "report": report})


def _report_content(diagnosis, impact, suggestions):
    if diagnosis.get("context") == "planning":
        scenario = {"scenario_id": diagnosis["scenario_id"],
                    **deepcopy(diagnosis["input_snapshot"])}
    else:
        scenario = _scenarios().get(diagnosis["scenario_id"])
    return {"scenario_id": diagnosis["scenario_id"], "data_source": "CONSTRUCTED",
            "scenario": scenario,
            "diagnosis": diagnosis, "impact": impact, "suggestions": suggestions,
            "generated_at": datetime.now(timezone.utc).isoformat()}


def _report_value(value):
    if value is None or value == "":
        return "未提供"
    if isinstance(value, bool):
        return "是" if value else "否"
    if isinstance(value, float) and not math.isfinite(value):
        return "未提供"
    if isinstance(value, (list, tuple)):
        if not value:
            return "未提供"
        shown = "、".join(_report_value(item) for item in value[:12])
        return shown + (f"（另 {len(value) - 12} 项见 JSON）" if len(value) > 12 else "")
    if isinstance(value, dict):
        return "结构化详情见同运行 JSON" if value else "未提供"
    return str(value)


def _report_score(value):
    return f"{value:.3f}" if isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(value) else "未提供"


def _report_sections(content, source=None):
    from .kg_display import FAULT_TYPES as fault_names

    # 只生成展示摘要，不改写原始诊断、证据或 JSON 机器导出。
    source = source or {}
    diagnosis, scenario = content.get("diagnosis") or {}, content.get("scenario") or {}
    sections = [("来源与场景", [
        ("运行编号", source.get("run_id")), ("任务编号", source.get("task_id") or diagnosis.get("network_task_id")),
        ("方案编号", source.get("plan_id") or diagnosis.get("network_solution_id")),
        ("网络快照", source.get("snapshot_id")), ("方案修订", diagnosis.get("network_revision")),
        ("资源版本", diagnosis.get("network_resource_revision")), ("知识版本", diagnosis.get("knowledge_version")),
        ("场景编号", content.get("scenario_id")), ("场景名称", scenario.get("scenario_name")),
        ("场景版本", source.get("scenario_version") or diagnosis.get("scenario_version_id")),
        ("诊断任务", source.get("diagnose_task_id")), ("影响任务", source.get("impact_task_id")),
        ("生成时间", content.get("generated_at")),
    ])]
    object_names = {"NODE": "节点", "DEVICE": "设备", "LINK": "链路"}
    branch_names = {"RULE": "规则推理", "KG": "知识图谱", "ML": "机器学习",
                    "CONSTRUCTED": "构造联调", "MANUAL_REVIEW": "人工研判", "ROUTE_ANALYSIS": "路由分析"}
    state_names = {"INTERRUPTED": "中断", "DEGRADED": "降级"}

    def fault_name(value):
        return f"{fault_names[value]}（{value}）" if value in fault_names else value

    for index, fault in enumerate(scenario.get("faults") or [], 1):
        sections.append((f"故障输入 {index}", [
            ("对象类型", object_names.get(fault.get("object_type"), fault.get("object_type"))),
            ("对象编号", fault.get("object_id")), ("故障类型", fault_name(fault.get("fault_type"))),
            ("现象编号", fault.get("symptom_ids")),
            ("发生时间", fault.get("occur_time")),
        ]))
    for item in diagnosis.get("results") or []:
        fused, obj = item.get("fused") or {}, item.get("fault_object") or {}
        sections.append(("诊断结论 · " + _report_value(obj.get("object_id")), [
            ("故障类型", fault_name(fused.get("fault_type"))), ("根因", fused.get("root_cause")),
            ("故障部件", item.get("fault_component")), ("融合得分", _report_score(fused.get("confidence"))),
            ("需要人工研判", fused.get("need_manual_review")),
        ]))
        for branch in item.get("branches") or []:
            evidence = branch.get("evidence") or {}
            rows = [("分支结论", fault_name(branch.get("fault_type"))), ("分支得分", _report_score(branch.get("confidence"))),
                    ("不可用原因", branch.get("disabled_reason"))]
            if branch.get("source") == "RULE":
                rows.extend([("规则编号", evidence.get("rule_ids")), ("匹配方式", {
                    "EXACT": "精确匹配", "PARTIAL_MISSING_INPUT": "输入缺失下的部分匹配",
                }.get(evidence.get("match_kind"), evidence.get("match_kind"))),
                    ("来源案例编号", evidence.get("source_case_ids")), ("来源标识", evidence.get("source"))])
            elif branch.get("source") == "KG":
                rows.extend([("选用案例", evidence.get("case_id")), ("三元组编号", evidence.get("triple_ids")),
                             ("来源文档", evidence.get("source_docs") or evidence.get("source_doc")),
                             ("同等支持案例", evidence.get("equally_supported_cases"))])
                rankings = evidence.get("candidate_rankings") or []
                rows.append(("候选案例数", len(rankings)))
                for index, candidate in enumerate(rankings[:5], 1):
                    rows.append((f"候选摘要 {index}", "案例 %s；类型 %s；得分 %s；匹配现象数 %s" % (
                        _report_value(candidate.get("case_id")), _report_value(fault_name(candidate.get("fault_type"))),
                        _report_score(candidate.get("score")), _report_value(candidate.get("matched_symptom_count")))))
                if len(rankings) > 5:
                    rows.append(("其余候选", f"另 {len(rankings) - 5} 项见同运行 JSON"))
            sections.append((branch_names.get(branch.get("source"), "其他分支") + "依据", rows))
    impact = content.get("impact") or {}
    summary = impact.get("summary") or {}
    sections.append(("影响统计", [(label, summary.get(key)) for key, label in (
        ("affected_node_count", "受影响节点数"), ("direct_count", "直接影响节点数"),
        ("indirect_count", "间接影响节点数"), ("interrupted_node_count", "中断节点数"),
        ("degraded_node_count", "降级节点数"), ("affected_link_count", "受影响链路数"),
        ("total_bandwidth_loss_kbps", "通联需求速率损失 / kbit/s"),
        ("affected_mandatory_demand_ids", "受影响必要需求编号"),
    )]))
    for key, id_field, title in (("affected_nodes", "node_id", "受影响节点"), ("affected_links", "link_id", "受影响链路")):
        for row in impact.get(key) or []:
            sections.append((title + " · " + _report_value(row.get(id_field)), [
                ("影响状态", state_names.get(row.get("state"), row.get("state"))),
                ("影响程度", {"SEVERE": "严重", "MODERATE": "中等", "MINOR": "轻微"}.get(row.get("level"), row.get("level"))),
                ("原因", row.get("reason")),
            ]))
    for index, row in enumerate((content.get("suggestions") or {}).get("suggestions") or [], 1):
        sections.append((f"处置建议 {index}", [("优先级", row.get("priority")), ("建议行动", row.get("action")),
            ("来源", branch_names.get(row.get("source"), row.get("source"))), ("步骤", row.get("steps")),
            ("所需备件", row.get("required_spares")), ("所需工具", row.get("required_tools")),
            ("预计用时 / min", row.get("estimated_time_min")), ("依据三元组", row.get("kg_source_triple_ids"))]))
    return [(title, [(label, _report_value(value)) for label, value in rows]) for title, rows in sections]


def _report_docx(content, path, *, source=None):
    from docx import Document
    from docx.oxml.ns import qn
    doc = Document()
    for name in ("Normal", "Title", "Heading 1"):
        doc.styles[name].font.name = "宋体"
        doc.styles[name]._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")
    doc.add_heading("故障诊断报告", 0)
    doc.add_paragraph("数据性质：CONSTRUCTED 合成测试数据；结果不能作为真实设备诊断准确率证明。")
    for title, rows in _report_sections(content, source):
        doc.add_heading(title, level=1)
        for label, value in rows:
            doc.add_paragraph(f"{label}：{value}")
    doc.add_paragraph("完整证据、候选排序及机器字段保留在同一运行的 JSON 导出中。")
    doc.save(path)


def _report_pdf(content, path, *, source=None):
    from html import escape
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    style = ParagraphStyle("Chinese", fontName="STSong-Light", fontSize=10, leading=17)
    heading = ParagraphStyle("Heading", parent=style, fontSize=16, leading=24)
    section = ParagraphStyle("Section", parent=style, fontSize=12, leading=20, spaceBefore=10, keepWithNext=True)
    story = [Paragraph("故障诊断报告", heading), Spacer(1, 12),
             Paragraph("数据性质：CONSTRUCTED 合成测试数据，结果须人工研判。", style)]
    for title, rows in _report_sections(content, source):
        story.append(Paragraph(escape(title), section))
        for label, value in rows:
            story.append(Paragraph(escape(f"{label}：{value}"), style))
    story.append(Paragraph("完整证据、候选排序及机器字段保留在同一运行的 JSON 导出中。", style))
    SimpleDocTemplate(str(path), pagesize=A4).build(story)


def create_report(payload):
    payload = payload or {}
    diagnosis = _diagnosis(payload.get("diagnose_task_id"))
    if diagnosis is None:
        return {"code": CODE_TASK_NOT_FOUND, "message": "诊断任务不存在或尚未完成", "data": None}
    impact = _impact(payload.get("impact_task_id")) if payload.get("impact_task_id") else None
    if payload.get("impact_task_id") and impact is None:
        return {"code": CODE_TASK_NOT_FOUND, "message": "影响分析任务不存在或尚未完成", "data": None}
    if impact and impact["scenario_id"] != diagnosis["scenario_id"]:
        return _bad(FaultInputError("诊断结果与影响分析结果不属于同一场景"))
    if diagnosis.get("context") == "planning" and impact and (
            impact.get("context") != "planning" or
            impact.get("network_solution_id") != diagnosis.get("network_solution_id")):
        return _bad(FaultInputError("诊断与影响分析的方案来源不一致"))
    fmt = (payload.get("format") or "json").lower()
    if fmt not in {"json", "docx", "pdf"}:
        return _bad(FaultInputError("报告格式仅支持 json / docx / pdf"))
    content = _report_content(diagnosis, impact, _engine().suggestions(diagnosis, impact))
    REPORT_DIR.mkdir(parents=True, exist_ok=True)
    filename = "report_%s_%s.%s" % (diagnosis["scenario_id"], uuid.uuid4().hex[:8], fmt)
    path = REPORT_DIR / filename
    if fmt == "json":
        path.write_text(json.dumps(content, ensure_ascii=False, indent=2), encoding="utf-8")
    elif fmt == "docx":
        _report_docx(content, path, source={"diagnose_task_id": payload.get("diagnose_task_id"), "impact_task_id": payload.get("impact_task_id")})
    else:
        _report_pdf(content, path, source={"diagnose_task_id": payload.get("diagnose_task_id"), "impact_task_id": payload.get("impact_task_id")})
    return _ok({"report_url": "/api/v1/files/%s" % filename, "format": fmt,
                "data_source": "CONSTRUCTED"})


def report_path(filename):
    if Path(filename).name != filename or not filename.startswith("report_"):
        return None
    path = REPORT_DIR / filename
    return path if path.is_file() else None


def _impact_geometry(impact, network=None):
    if network is not None and network.get("id") != impact.get("network_solution_id"):
        raise ValueError("影响图的冻结方案与影响结果不一致")
    if network is not None and impact.get("network_task_id") and network.get("taskId") != impact["network_task_id"]:
        raise ValueError("影响图的冻结任务与影响结果不一致")
    affected_nodes = {row["node_id"]: row for row in impact.get("affected_nodes") or []}
    affected_links = {row["link_id"]: row for row in impact.get("affected_links") or []}
    nodes, links, missing = {}, [], 0

    def valid(point):
        return (isinstance(point, (list, tuple)) and len(point) == 2
                and all(isinstance(v, (int, float)) and not isinstance(v, bool) and math.isfinite(v) for v in point)
                and abs(point[0]) <= 180 and abs(point[1]) <= 90)

    if network is not None:
        rows = (network.get("inputSnapshot") or {}).get("stations") or []
        for row in rows:
            if not row.get("id") or not valid(row.get("coordinate")):
                missing += 1
                continue
            if row["id"] in nodes:
                raise ValueError("影响图冻结节点编号重复")
            nodes[row["id"]] = {"id": row["id"], "coordinate": tuple(row["coordinate"]),
                                "state": affected_nodes.get(row["id"], {}).get("state")}
        for row in network.get("links") or []:
            if row.get("id") and row.get("a") in nodes and row.get("b") in nodes:
                links.append({"id": row["id"], "a": row["a"], "b": row["b"],
                              "state": affected_links.get(row["id"], {}).get("state")})
    else:
        # 无冻结网络时只使用影响任务已有点，不回填当前资源或构造链路。
        for feature in (impact.get("impact_scope") or {}).get("features") or []:
            geometry, properties = feature.get("geometry") or {}, feature.get("properties") or {}
            if geometry.get("type") != "Point" or not properties.get("node_id") or not valid(geometry.get("coordinates")):
                missing += 1
                continue
            nid = properties["node_id"]
            nodes[nid] = {"id": nid, "coordinate": tuple(geometry["coordinates"]), "state": properties.get("state")}
    return {"nodes": list(nodes.values()), "links": links, "missing_coordinates": missing,
            "unlocated_nodes": len(set(affected_nodes) - nodes.keys()),
            "unlocated_links": len(set(affected_links) - {row["id"] for row in links}),
            "frozen_network": network is not None}


def _impact_font(size):
    from PIL import ImageFont
    # 优先只读本次交付已有共享字体资产，不复制字体或安装后端依赖。
    candidates = [ROOT / "front_design/radio-planning-20261022/public/fonts/NotoSansCJKsc-Regular.otf"]
    candidates += [Path(os.environ.get("WINDIR", "C:/Windows")) / "Fonts" / name for name in ("msyh.ttc", "simsun.ttc")]
    candidates += [Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"),
                   Path("/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc")]
    for candidate in candidates:
        if candidate.is_file():
            try:
                return ImageFont.truetype(str(candidate), size)
            except OSError:
                continue
    raise RuntimeError("影响 PNG 缺少可用中文字体，请配置已有系统中文字体；未生成缺字图片")


def _impact_png(impact, network=None, source=None):
    from PIL import Image, ImageDraw, PngImagePlugin
    diagram, source = _impact_geometry(impact, network), source or {}
    if network and source.get("snapshot_id") and source["snapshot_id"] != network.get("snapshotId"):
        raise ValueError("影响图网络快照与来源编号不一致")
    image = Image.new("RGB", (900, 640), "white")
    draw = ImageDraw.Draw(image)
    title_font, font, small = _impact_font(22), _impact_font(14), _impact_font(12)

    def wrapped(text, y, selected_font=font, color="#26323d"):
        line = ""
        line_height = selected_font.size + 6
        for char in text:
            candidate = line + char
            if line and draw.textbbox((0, 0), candidate, font=selected_font)[2] > 840:
                draw.text((30, y), line, fill=color, font=selected_font)
                y += line_height
                line = char
            else:
                line = candidate
        draw.text((30, y), line, fill=color, font=selected_font)
        return y + line_height

    y = wrapped("故障影响示意（构造联调，非实测）", 12, title_font)
    y = wrapped("运行：%s；场景：%s" % (_report_value(source.get("run_id")), _report_value(impact.get("scenario_id"))), y)
    y = wrapped("任务：%s；方案：%s" % (_report_value(source.get("task_id") or impact.get("network_task_id")),
                                      _report_value(impact.get("network_solution_id"))), y)
    y = wrapped("快照：%s；WGS84 经纬度示意，连线不表示传播路径" % _report_value((network or {}).get("snapshotId")), y)
    top, bottom = max(132, y + 10), 520
    draw.rectangle((30, top, 870, bottom), outline="#cbd1d8")
    colors = {"INTERRUPTED": "#c83832", "DEGRADED": "#e6aa39"}
    nodes = {row["id"]: row for row in diagram["nodes"]}
    positions = {}
    if nodes:
        xs, ys = [row["coordinate"][0] for row in nodes.values()], [row["coordinate"][1] for row in nodes.values()]
        scale = min(780 / max(max(xs) - min(xs), .001), max(1, bottom - top - 50) / max(max(ys) - min(ys), .001))
        cx, cy = (min(xs) + max(xs)) / 2, (min(ys) + max(ys)) / 2
        positions = {nid: (450 + (row["coordinate"][0] - cx) * scale,
                           (top + bottom) / 2 - (row["coordinate"][1] - cy) * scale) for nid, row in nodes.items()}
        priority = {"DEGRADED": 1, "INTERRUPTED": 2}
        for link in sorted(diagram["links"], key=lambda row: priority.get(row["state"], 0)):
            a, b = positions[link["a"]], positions[link["b"]]
            draw.line((*a, *b), fill=colors.get(link["state"], "#a4adb7"), width=3 if link["state"] in colors else 1)
            if link["state"] in colors:
                px, py = (a[0] + b[0]) / 2, (a[1] + b[1]) / 2
                width = draw.textbbox((0, 0), link["id"], font=small)[2]
                draw.text((max(35, min(865 - width, px)), max(top + 3, min(bottom - 18, py))), link["id"], fill="#26323d", font=small)
        # 影响对象后绘制，避免同位置的未知背景覆盖中断或降级标记。
        for nid, row in sorted(nodes.items(), key=lambda item: priority.get(item[1]["state"], 0)):
            px, py = positions[nid]
            radius = 6 if row["state"] in colors else 3
            draw.ellipse((px - radius, py - radius, px + radius, py + radius), fill=colors.get(row["state"], "#697887"))
            if row["state"] in colors or len(nodes) <= 12:
                width = draw.textbbox((0, 0), nid, font=small)[2]
                draw.text((max(35, min(865 - width, px + 8)), max(top + 3, min(bottom - 18, py - 17))), nid, fill="#26323d", font=small)
    else:
        draw.text((55, top + 30), "未提供可绘制的真实节点坐标", font=font, fill="#697887")
    draw.line((35, 550, 61, 550), fill="#a4adb7", width=2)
    draw.text((68, 540), "冻结背景／影响状态未提供", font=font, fill="#26323d")
    draw.ellipse((352, 545, 364, 557), fill=colors["INTERRUPTED"])
    draw.text((373, 540), "中断", font=font, fill="#26323d")
    draw.ellipse((466, 545, 478, 557), fill=colors["DEGRADED"])
    draw.text((487, 540), "降级", font=font, fill="#26323d")
    wrapped("已绘 %s 节点 / %s 链路；缺失坐标 %s；未定位影响节点 %s / 链路 %s" % (
        len(nodes), len(diagram["links"]), diagram["missing_coordinates"], diagram["unlocated_nodes"], diagram["unlocated_links"]), 574, small)
    wrapped("来源：%s；仅示对象位置与已有拓扑，不表示覆盖边界；完整数据见 JSON。" % (
        "本运行冻结网络" if diagram["frozen_network"] else "冻结网络未提供，仅使用原影响点，未构造链路"), 603, small)
    metadata = PngImagePlugin.PngInfo()
    for key, value in {"Title": "故障影响示意", "PlanId": impact.get("network_solution_id"),
                       "SnapshotId": (network or {}).get("snapshotId"), "RunId": source.get("run_id"),
                       "RenderedNodes": len(nodes), "RenderedLinks": len(diagram["links"])}.items():
        metadata.add_itxt(key, _report_value(value))
    buf = io.BytesIO()
    image.save(buf, format="PNG", pnginfo=metadata)
    return buf.getvalue(), "image/png"


def export_impact(tid, fmt, *, network=None, source=None):
    impact = _impact(tid)
    if impact is None:
        return None
    if fmt == "xlsx":
        from openpyxl import Workbook
        buf = io.BytesIO()
        book = Workbook()
        sheet = book.active
        sheet.title = "受影响节点"
        sheet.append(["节点", "影响类型", "程度", "状态", "原因"])
        for row in impact["affected_nodes"]:
            sheet.append([row["node_id"], row["affect_type"], row["level"], row["state"], row["reason"]])
        sheet = book.create_sheet("受影响链路")
        sheet.append(["链路", "状态", "程度", "原因"])
        for row in impact["affected_links"]:
            sheet.append([row["link_id"], row["state"], row["level"], row["reason"]])
        book.save(buf)
        return buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    if fmt == "png":
        return _impact_png(impact, network, source)
    return None
