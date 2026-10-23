"""候选快照纯函数；不访问数据库、HTTP、算法或运行文件。"""
from __future__ import annotations

from copy import deepcopy
import csv
from datetime import date, datetime, time, timedelta
from io import BytesIO, StringIO
import json
import math
import re
import xml.etree.ElementTree as ET
from zipfile import BadZipFile, ZipFile

from openpyxl import Workbook, load_workbook
from pydantic import ValidationError

from . import resource_import_parser as parser
from .resource_candidate_schemas import (
    ASSESSMENT_FIELDS, DATA_FIELDS, EVALUATION_INPUTS,
    CandidateData, CandidateLink, CandidatePatch, CandidateSource,
)
from .resource_dataset_service import ResourceDatasetError

SNAPSHOT_KEY = "frontend_candidate_links"
SHEET = "topology_links"
HEADERS = (
    "link_id", "node_a_id", "node_b_id", "device_a_id", "device_b_id", "device_class",
    "freq_khz", "bandwidth_khz", "distance_m", "path_loss_db", "rx_power_dbm",
    "link_margin_db", "is_available", "link_state", "is_backup",
)
FIELDS = (
    "id", "nodeAId", "nodeBId", "deviceAId", "deviceBId", "deviceClass", "freqKhz",
    "bandwidthKhz", "distanceM", "pathLossDb", "rxPowerDbm", "linkMarginDb",
    "isAvailable", "linkState", "isBackup",
)
HEADER_TO_FIELD = dict(zip(HEADERS, FIELDS))
NUMERIC = {"freqKhz", "bandwidthKhz", "distanceM", "pathLossDb", "rxPowerDbm", "linkMarginDb"}
BOOLEANS = {"isAvailable", "isBackup"}
MAX_FILE_BYTES = parser.MAX_FILE_BYTES
MAX_ROWS = parser.MAX_ROWS
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
NUMBER = re.compile(r"[+-]?(?:[0-9]+(?:\.[0-9]*)?|\.[0-9]+)(?:[eE][+-]?[0-9]+)?\Z")
CONTROL = re.compile(r"[\x00-\x1f\x7f]")


def _issue(row, object_id, field, message):
    return parser._issue(SHEET, row, object_id, field, message)


class _Invalid(ValueError):
    def __init__(self, message, row=0, field="file", object_id="file"):
        super().__init__(message)
        self.detail = _issue(row, object_id, field, message)


def _fail(issues):
    error = ResourceDatasetError("候选校验失败：" + issues[0]["message"], 1001, 422)
    error.issues = deepcopy(issues)
    raise error


def _validation_issues(exc, row, object_id):
    return [_issue(row, object_id, ".".join(map(str, item["loc"])) or "structure", item["msg"])
            for item in exc.errors()]


def _canonical(value, row=0):
    try:
        return CandidateLink.model_validate(value).model_dump(exclude_unset=True)
    except ValidationError as exc:
        _fail(_validation_issues(exc, row, value.get("id", "candidateLinks") if isinstance(value, dict) else "candidateLinks"))


def get_links(snapshot) -> list[dict]:
    """只读取候选扩展；缺省为空，不从历史 link_metric 构造当前网络。"""
    if not isinstance(snapshot, dict):
        _fail([_issue(0, "snapshot", "structure", "快照必须为对象")])
    links = snapshot.get(SNAPSHOT_KEY, [])
    if not isinstance(links, list):
        _fail([_issue(0, "candidateLinks", SNAPSHOT_KEY, "候选集合必须为数组")])
    result, seen = [], set()
    for value in links:
        link = _canonical(value)
        if link["id"] in seen:
            _fail([_issue(0, link["id"], "id", "候选 ID 重复")])
        seen.add(link["id"])
        result.append(link)
    return deepcopy(result)


def _indexes(snapshot):
    indexes = {}
    tables = snapshot.get("tables") if isinstance(snapshot, dict) else None
    for name, field in (("node.csv", "node_id"), ("device.csv", "device_id"), ("device_model.csv", "model_id")):
        if not isinstance(tables, dict) or not isinstance(tables.get(name), list):
            raise _Invalid("快照缺少节点、设备或型号表", field="tables")
        target = indexes[name] = {}
        for value in tables[name]:
            identifier = value.get(field) if isinstance(value, dict) else None
            if not isinstance(identifier, str) or not identifier or identifier in target:
                raise _Invalid("快照资源 ID 缺失或重复", field=field)
            target[identifier] = value
    return indexes


def _model_range(model):
    values = []
    for field in ("freq_min_khz", "freq_max_khz"):
        value = model.get(field)
        if type(value) not in (str, int, float) or isinstance(value, str) and not NUMBER.fullmatch(value):
            raise ValueError("型号频率范围无效")
        number = float(value)
        if not math.isfinite(number) or number <= 0:
            raise ValueError("型号频率范围无效")
        values.append(number)
    if values[0] > values[1]:
        raise ValueError("型号频率范围无效")
    return values


def _reference_issues(link, indexes, row=0):
    issues = []
    for node_field, device_field in (("nodeAId", "deviceAId"), ("nodeBId", "deviceBId")):
        def issue(field, message):
            issues.append(_issue(row, link["id"], field, message))

        if link[node_field] not in indexes["node.csv"]:
            issue(node_field, "节点不存在")
        device = indexes["device.csv"].get(link[device_field])
        if device is None:
            issue(device_field, "设备不存在")
            continue
        if device.get("node_id") != link[node_field]:
            issue(device_field, "设备不属于所选节点")
        model = indexes["device_model.csv"].get(device.get("model_id"))
        if model is None or model.get("device_class") != link["deviceClass"]:
            issue("deviceClass", "两端设备型号类别须与候选类别一致")
        if model is not None and link["freqKhz"] is not None:
            try:
                low, high = _model_range(model)
                if not low <= link["freqKhz"] <= high:
                    issue("freqKhz", "频率须位于两端型号频率范围内")
            except (ValueError, OverflowError):
                issue("freqKhz", "关联型号频率范围无效，不能验证候选频率")
    return issues


def apply_candidate_changes(snapshot, changes: dict) -> list[dict]:
    """只应用候选增量；调用方先更新 tables，再将返回集合写入同一快照。"""
    links = get_links(snapshot)
    if not isinstance(changes, dict):
        _fail([_issue(0, "candidateLinks", "changes", "增量必须为对象")])
    for key in ("newCandidateLinks", "candidateLinks", "importedCandidateLinks"):
        if key in changes and not isinstance(changes[key], list):
            _fail([_issue(0, key, key, "候选增量必须为数组")])
    if changes.get("importedCandidateLinks"):
        _fail([_issue(0, "importedCandidateLinks", "source", "导入候选只能由服务器重新解析原文件产生")])
    existing = {link["id"]: link for link in links}
    seen = set()
    for key in ("newCandidateLinks", "candidateLinks"):
        for value in changes.get(key, []):
            if not isinstance(value, dict):
                _fail([_issue(0, key, "structure", "候选增量记录必须为对象")])
            try:
                model = (CandidateLink if key == "newCandidateLinks" else CandidatePatch).model_validate(value)
            except ValidationError as exc:
                _fail(_validation_issues(exc, 0, value.get("id", key)))
            identifier = model.id
            if identifier in seen:
                _fail([_issue(0, identifier, "id", "同批候选 ID 重复")])
            seen.add(identifier)
            if key == "newCandidateLinks":
                if identifier in existing:
                    _fail([_issue(0, identifier, "id", "新增候选 ID 已存在")])
                if model.source.kind != "MANUAL":
                    _fail([_issue(0, identifier, "source", "新增候选必须为 MANUAL 来源")])
                link = model.model_dump(exclude_unset=True)
                links.append(link)
                existing[identifier] = link
            else:
                if identifier not in existing:
                    _fail([_issue(0, identifier, "id", "待修改候选不存在，不允许替换或删除集合")])
                old = existing[identifier]
                merged = {**deepcopy(old), **model.model_dump(exclude_unset=True)}
                if any(merged[field] != old[field] for field in EVALUATION_INPUTS):
                    merged.update(dict.fromkeys(ASSESSMENT_FIELDS))
                    merged["assessmentState"] = "PENDING"
                link = _canonical(merged)
                old.clear()
                old.update(link)
    if links:
        try:
            indexes = _indexes(snapshot)
        except _Invalid as exc:
            _fail([exc.detail])
        issues = [issue for link in links for issue in _reference_issues(link, indexes)]
        if issues:
            _fail(issues)
    return deepcopy(links)


def _headers(values):
    if (any(not isinstance(value, str) for value in values) or len(values) != len(HEADERS)
            or len(set(values)) != len(values) or set(values) != set(HEADERS)):
        raise _Invalid("表头缺失、重复或包含未知字段", 1, "header")
    return values


def _csv_records(data):
    reader = None
    try:
        reader = csv.reader(StringIO(data.decode("utf-8-sig"), newline=""), strict=True)
        head = _headers(next(reader, []))
        records = []
        while True:
            line = reader.line_num + 1
            values = next(reader, None)
            if values is None:
                break
            if not values:
                continue
            if len(values) != len(head):
                raise _Invalid("列数与表头不一致", line, "row", values[0] if values else "")
            decoded = []
            for field, value in zip(head, values):
                if value.startswith("~rp~"):
                    try:
                        value = json.loads(value[4:])
                        if not isinstance(value, str):
                            raise ValueError("not text")
                    except (ValueError, TypeError, RecursionError) as exc:
                        raise _Invalid("CSV ~rp~ 文本转义无效", line, field, values[0]) from exc
                decoded.append(value)
            records.append((line, dict(zip(head, decoded))))
            if len(records) > MAX_ROWS:
                raise _Invalid("业务记录最多 5000 行", line, "row")
        return records
    except UnicodeError as exc:
        raise _Invalid("CSV 必须为 UTF-8，可带 UTF-8 BOM") from exc
    except csv.Error as exc:
        raise _Invalid("CSV 格式无效：" + str(exc), reader.line_num if reader is not None else 0, "row") from exc


def _xlsx_records(data):
    try:
        with ZipFile(BytesIO(data)) as archive:
            parser._inspect_archive(archive)
        book = load_workbook(BytesIO(data), read_only=True, data_only=False, keep_links=False)
        try:
            if (len(book.sheetnames) != len(set(book.sheetnames))
                    or set(book.sheetnames) not in ({SHEET}, {SHEET, "说明"})):
                raise _Invalid("XLSX 只能包含 topology_links 及可选说明表", 1, "sheet")
            records = []
            for sheet in book:
                if sheet.sheet_state != "visible" or not hasattr(sheet, "reset_dimensions"):
                    raise _Invalid("不接受隐藏表或非业务工作表", 1, "sheet")
                sheet.reset_dimensions()
                head = None
                for line, cells in enumerate(sheet.iter_rows(), 1):
                    values = []
                    for column, cell in enumerate(cells):
                        value = cell.value
                        field = head[column] if head is not None and column < len(head) else "cell"
                        if cell.data_type in {"f", "e"} or isinstance(value, (date, datetime, time, timedelta)):
                            raise _Invalid("不接受公式、错误单元格或日期对象", line, field)
                        if value is not None and type(value) not in (str, int, float, bool):
                            raise _Invalid("单元格标量类型无效", line, field)
                        values.append(value)
                    if sheet.title == "说明":
                        continue
                    if line == 1:
                        head = _headers(values)
                        continue
                    if not any(value is not None and value != "" for value in values):
                        continue
                    if head is None or len(values) > len(head):
                        raise _Invalid("列数与表头不一致", line, "row")
                    values += [None] * (len(head) - len(values))
                    record = {}
                    for header, value in zip(head, values):
                        field = HEADER_TO_FIELD[header]
                        if type(value) is bool:
                            if field not in BOOLEANS:
                                raise _Invalid("布尔值不能代替数字或文本", line, header, str(values[0]))
                            value = str(value).lower()
                        elif type(value) in (int, float):
                            if field not in NUMERIC:
                                raise _Invalid("文本字段必须为字符串", line, header, str(values[0]))
                            value = str(value)
                        record[header] = "" if value is None else value
                    records.append((line, record))
                    if len(records) > MAX_ROWS:
                        raise _Invalid("业务记录最多 5000 行", line, "row")
                if sheet.title == SHEET and head is None:
                    raise _Invalid("缺少表头", 1, "header")
            return records
        finally:
            book.close()
    except _Invalid:
        raise
    except (BadZipFile, ET.ParseError, ValueError, TypeError, KeyError, IndexError, AttributeError,
            OSError, RuntimeError, OverflowError, EOFError, NotImplementedError) as exc:
        detail = getattr(exc, "detail", {})
        raise _Invalid(detail.get("message", "XLSX 文件损坏或格式不受支持"),
                       detail.get("row", 0), detail.get("field", "file"), detail.get("objectId", "file")) from exc


def _import_data(record, line, issues):
    result = {}
    identifier = record["link_id"]
    for header, value in record.items():
        field = HEADER_TO_FIELD[header]
        if CONTROL.search(value):
            issues.append(_issue(line, identifier, header, "不接受控制字符"))
        if field in NUMERIC:
            if value == "":
                result[field] = None
            elif not NUMBER.fullmatch(value):
                issues.append(_issue(line, identifier, header, "须为有限十进制数，不接受布尔、表达式或 NaN"))
                result[field] = None
            else:
                number = float(value)
                if not math.isfinite(number):
                    issues.append(_issue(line, identifier, header, "须为有限数"))
                    number = None
                result[field] = number
        elif field in BOOLEANS:
            if value not in {"", "true", "false"}:
                issues.append(_issue(line, identifier, header, "布尔值仅接受 true、false 或空值"))
            result[field] = None if value == "" else value == "true"
        else:
            result[field] = None if field == "linkState" and value == "" else value
    try:
        return CandidateData.model_validate(result).model_dump()
    except ValidationError as exc:
        for issue in _validation_issues(exc, line, identifier):
            issue["field"] = dict(zip(FIELDS, HEADERS)).get(issue["field"], issue["field"])
            issues.append(issue)
        return None


def _warnings(link, line):
    result = []
    if link["assessmentState"] == "REFERENCE":
        result.append(_issue(line, link["id"], "reference", "外部构造参考，非正式规划网络或当前算法评估"))
    if link["pathLossDb"] is not None and link["pathLossDb"] <= 0:
        result.append(_issue(line, link["id"], "path_loss_db", "来源传播损耗为零或负值，保留原值待核对"))
    if link["linkState"] == "FAULT" and link["isAvailable"] is True:
        result.append(_issue(line, link["id"], "link_state", "来源质量为 FAULT 且可用性为 true，保留原值待核对"))
    return result


def parse_candidate_import(snapshot, filename, data) -> dict:
    """返回全量候选和逐行差异；任一 issue 均退回深拷贝的原集合。"""
    baseline = get_links(snapshot)
    result = {"links": baseline, "rows": [], "issues": [], "warnings": []}
    try:
        CandidateSource.model_validate({"kind": "IMPORT", "file": filename, "row": 2})
        if not isinstance(data, bytes) or not data or len(data) > MAX_FILE_BYTES:
            raise _Invalid("文件不能为空且不得超过 5 MiB")
        if filename.lower().endswith(".csv"):
            records = _csv_records(data)
        elif filename.lower().endswith(".xlsx"):
            records = _xlsx_records(data)
        else:
            raise _Invalid("仅支持 CSV 和 XLSX，不接受宏工作簿")
        if not records:
            raise _Invalid("每批需有 1 至 5000 行业务记录")
        indexes = _indexes(snapshot)
    except ValidationError as exc:
        result["issues"].extend(_validation_issues(exc, 0, "file"))
        return result
    except _Invalid as exc:
        result["issues"].append(exc.detail)
        return result
    links = deepcopy(baseline)
    existing = {link["id"]: link for link in links}
    seen = set()
    for line, record in records:
        identifier = record["link_id"]
        if identifier in seen:
            result["issues"].append(_issue(line, identifier, "link_id", "同批候选 ID 重复"))
            continue
        seen.add(identifier)
        candidate_data = _import_data(record, line, result["issues"])
        if candidate_data is None:
            continue
        source = {"kind": "IMPORT", "file": filename, "row": line}
        present = any(candidate_data[field] is not None for field in ASSESSMENT_FIELDS)
        link = {**candidate_data, "source": source, "assessmentState": "REFERENCE" if present else "PENDING",
                "reference": {"data": deepcopy(candidate_data), "source": deepcopy(source)} if present else None}
        result["issues"].extend(_reference_issues(link, indexes, line))
        result["warnings"].extend(_warnings(link, line))
        old = existing.get(identifier)
        # 重导入相同来源及原始十五列是 no-op，包括已有服务提示。
        unchanged = old is not None and all(old[field] == link[field] for field in (*DATA_FIELDS, "source", "reference", "assessmentState"))
        differences = [{"field": header, "before": old[field] if old is not None else None, "after": link[field]}
                       for header, field in zip(HEADERS, FIELDS) if (old[field] if old is not None else None) != link[field]]
        if old is not None and not unchanged:
            differences.extend({"field": field, "before": deepcopy(old[field]), "after": deepcopy(link[field])}
                               for field in ("source", "reference", "assessmentState") if old[field] != link[field])
        result["rows"].append({"sheet": SHEET, "row": line, "id": identifier,
                               "action": "UNCHANGED" if unchanged else "NEW" if old is None else "UPDATE",
                               "differences": differences})
        if old is None:
            links.append(link)
            existing[identifier] = link
        elif not unchanged:
            old.clear()
            old.update(link)
    for link in links:
        if link["id"] not in seen:
            result["issues"].extend(_reference_issues(link, indexes))
    if not result["issues"]:
        result["links"] = links
    return result


def _template_example(snapshot):
    indexes = _indexes(snapshot)
    first_by_class = {}
    for device in indexes["device.csv"].values():
        node = device.get("node_id")
        model = indexes["device_model.csv"].get(device.get("model_id"))
        if node not in indexes["node.csv"] or model is None or model.get("device_class") not in {"HF", "VUHF"}:
            continue
        device_class = model["device_class"]
        first = first_by_class.get(device_class)
        if first is not None and first["node_id"] != node:
            identifier, suffix = "NEW-CANDIDATE-01", 1
            used = {link["id"] for link in get_links(snapshot)}
            while identifier in used:
                suffix += 1
                identifier = f"NEW-CANDIDATE-{suffix:02d}"
            data = {"id": identifier, "nodeAId": first["node_id"], "nodeBId": node,
                    "deviceAId": first["device_id"], "deviceBId": device["device_id"],
                    "deviceClass": device_class, "freqKhz": None, "bandwidthKhz": None, "isBackup": None,
                    **dict.fromkeys(ASSESSMENT_FIELDS)}
            return CandidateData.model_validate(data).model_dump()
        first_by_class.setdefault(device_class, device)
    raise _Invalid("当前资源缺少位于不同节点的两台同类别设备，无法生成真实引用示例", field="devices")


def make_candidate_template(snapshot, format) -> tuple[bytes, str, str]:
    """只引用当前真实设备；频率、带宽及评估未知留空。"""
    if format not in ("csv", "xlsx"):
        _fail([_issue(0, "file", "format", "模板仅支持 csv 或 xlsx")])
    try:
        example = _template_example(snapshot)
    except _Invalid as exc:
        _fail([exc.detail])
    except ValidationError as exc:
        _fail(_validation_issues(exc, 0, "template"))
    values = [example[field] for field in FIELDS]
    if format == "csv":
        output = StringIO(newline="")
        writer = csv.writer(output)
        writer.writerow(HEADERS)
        writer.writerow(["~rp~" + json.dumps(value, ensure_ascii=False) if isinstance(value, str)
                         and (re.match(r"^\s*[=+\-@]", value) or value.startswith("~rp~")) else value for value in values])
        data, mime = output.getvalue().encode("utf-8"), "text/csv; charset=utf-8"
    else:
        book = Workbook()
        try:
            sheet = book.active
            sheet.title = SHEET
            sheet.append(HEADERS)
            sheet.append(values)
            for cell in sheet[2]:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
            sheet.freeze_panes = "A2"
            notes = book.create_sheet("说明")
            for row in (
                ("数据性质", "SYNTHETIC_INTEGRATION；候选配置及外部构造参考，非生产、非实测、非正式规划网络"),
                ("示例来源", "引用当前资源中不同节点的同类别设备；不编造设备型号或带宽上限"),
                ("空值与单位", "未知频率、带宽及评估留空；单位 kHz、m、dB、dBm；布尔 true、false 或空值"),
                ("导入来源", "服务记录本文件及行号；有任一评估值才冻结为 REFERENCE，不执行传播或质量计算"),
                ("增量保存", "按 ID 新增或更新，未出现记录保留；相同文件相同内容重复导入不变"),
                ("文件限制", "最多 5 MiB、5000 行；不接受公式、宏、隐藏数据、日期对象、外链或额外业务表"),
                ("CSV 转义", "公式前缀文本及 ~rp~ 前缀使用 ~rp~ 加 JSON 字符串"),
            ):
                notes.append(row)
            data, mime = parser._save_memory(book), XLSX_MIME
        finally:
            book.close()
    return data, mime, "SYNTHETIC_INTEGRATION-topology_links-template." + format
