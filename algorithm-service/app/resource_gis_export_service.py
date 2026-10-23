"""GIS 已保存快照只读导出；不写 SQL、资源版本、观测或业务文件。"""
from __future__ import annotations

from copy import deepcopy
import csv
from io import StringIO
import json
import math
import re

from openpyxl import Workbook
from openpyxl.styles import Font
from pydantic import ValidationError

from . import resource_candidate_service as candidates, resource_dataset_service as datasets
from .resource_dataset_service import ResourceDatasetError
from .resource_export_schemas import DeviceFrequencySelection
from .resource_export_service import BOOLEANS, HEADERS, NUMERIC, TABLES, XLSX_MIME
from .resource_gis_export_schemas import DEVICE_FILTER_FIELDS, ResourceGisExportRequest
from .resource_import_parser import _save_memory
from .resource_read_filter import filter_device_items
from .resource_read_projection import _View, _interferences, _nodes
from .resource_read_query_schemas import _NodeFilters

MAX_FILE_BYTES = 20 * 1024 * 1024
BASELINE_NOTICE = "设备运行状态来自已保存台账基线，不包含临时同步观测。"
STATION_NOTICE = "站点筛选导出命中节点下全部已保存设备，包含 HF 节点下的 VUHF 设备；不是逐设备过滤。"
CANDIDATE_NOTICE = "候选文件仅含已保存配置及外部参考，不代表当前算法评估、正式规划网络或算法采用。"
CSV_REFERENCE_NOTICE = "候选 CSV 保留原 15 列，不携带原文件来源和冻结参考历史；重导入记录本次 IMPORT 来源，完整参考请使用 XLSX。"
INVALID_XML_TEXT = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]")
INVALID_ID = re.compile(r"\s|[\x00-\x1f\x7f]")


def _damaged(message="GIS 资源快照或引用损坏，未生成部分文件"):
    return ResourceDatasetError(message, 5002, 503)


def _query(value):
    try:
        return ResourceGisExportRequest.model_validate(value)
    except ValidationError as exc:
        raise ResourceDatasetError("GIS 导出字段、类型或范围无效", 1001, 422) from exc


def _text(value, xlsx=False):
    if not isinstance(value, str) or INVALID_XML_TEXT.search(value):
        raise _damaged("资源文本不能无损导出")
    value.encode("utf-8")
    if xlsx and len(value.encode("utf-16-le")) // 2 > 32767:
        raise ResourceDatasetError("XLSX 单元格超长，未生成截断内容", 1001, 413)
    return value


def _source(snapshot):
    try:
        view = _View(snapshot)
        tables = snapshot["tables"]
        if set(tables) != set(datasets.ID_FIELDS):
            raise _damaged()
        indexes = {}
        for name, field in datasets.ID_FIELDS.items():
            target = indexes[name] = {}
            for row in tables[name]:
                identifier = row[field]
                if not isinstance(identifier, str) or not 0 < len(identifier) <= 120 or INVALID_ID.search(identifier) or identifier in target:
                    raise _damaged("资源编号缺失、重复或无效")
                for key, value in row.items():
                    _text(key)
                    _text(value)
                target[identifier] = row
        nodes = indexes["node.csv"]
        for node in nodes.values():
            parents = node.get("parent_node_id", "").split(";") if node.get("parent_node_id") else []
            if len(parents) != len(set(parents)) or any(parent not in nodes for parent in parents):
                raise _damaged("节点上级引用缺失或重复")
        # 迭代遍历避免深层级触发 Python 递归限制，也拒绝全库范围中的环。
        colors = {}
        for identifier in nodes:
            stack = [(identifier, False)]
            while stack:
                node_id, leaving = stack.pop()
                if leaving:
                    colors[node_id] = 2
                    continue
                if colors.get(node_id) == 2:
                    continue
                if colors.get(node_id) == 1:
                    raise _damaged("节点上级引用存在循环")
                colors[node_id] = 1
                stack.append((node_id, True))
                stack.extend((parent, False) for parent in nodes[node_id].get("parent_node_id", "").split(";") if parent)
        for kind in ("device", "node", "interference"):
            view.items(kind)
        # 只调用无操作增量进行当前引用校验；历史 reference 不按当前资料重新认证。
        links = candidates.apply_candidate_changes(snapshot, {})
        for link in links:
            _text(_json(link))
        return view, indexes, links
    except ResourceDatasetError as exc:
        if exc.code == 5002 and exc.status == 503:
            raise
        raise _damaged() from exc
    except (KeyError, TypeError, ValueError, AttributeError, OverflowError) as exc:
        raise _damaged() from exc


def _node_filter(device_filter):
    values = {field: getattr(device_filter, field) for field in DEVICE_FILTER_FIELDS - {"frequency"}}
    frequency = device_filter.frequency or DeviceFrequencySelection()
    values.update(frequency_mode=frequency.mode, frequency_presence=frequency.presence,
                  min_khz=frequency.min_khz, max_khz=frequency.max_khz)
    return _NodeFilters.model_validate(values)


def _candidate_ids(links, nodes, query):
    condition = query.filter
    if condition.alert == "YES":
        return set()
    term = condition.search.strip().casefold()
    result = set()
    for link in links:
        if condition.device_class is not None and link["deviceClass"] != condition.device_class:
            continue
        labels = [link[field] for field in ("id", "nodeAId", "nodeBId", "deviceAId", "deviceBId")]
        labels.extend(nodes[link[field]]["node_name"] for field in ("nodeAId", "nodeBId"))
        if not term or any(term in value.casefold() for value in labels):
            result.add(link["id"])
    return result


def _selected(query, view, indexes, links):
    collections = {kind: indexes[name] for kind, name in TABLES.items()}
    collections["topology_links"] = {link["id"]: link for link in links}
    if query.range == "all":
        return {kind: set(collections[kind]) for kind in query.kinds}
    if query.range == "selected":
        for kind, identifiers in query.selected_ids.items():
            if set(identifiers) - set(collections[kind]):
                raise ResourceDatasetError("选择编号不存在或类别不匹配，未生成部分文件", 1002, 404)
        return {kind: set(query.selected_ids.get(kind, [])) for kind in query.kinds}
    if query.kinds == ["topology_links"]:
        return {"topology_links": _candidate_ids(links, indexes["node.csv"], query)}
    if query.browse == "devices":
        devices = filter_device_items(view.items("device"), query.filter)
        device_ids = {row["device_id"] for row in devices}
        node_ids = {row["node_id"] for row in devices}
    else:
        node_ids = {row["node_id"] for row in _nodes(view, _node_filter(query.filter))}
        device_ids = {row["device_id"] for row in view.items("device") if row["node_id"] in node_ids}
    matched = {"nodes": node_ids, "devices": device_ids,
               "interference_sources": {row["interference_id"] for row in _interferences(view, query.filter.search)}}
    return {kind: matched[kind] for kind in query.kinds}


def _closure(query, selected, indexes):
    if query.format != "xlsx" or "nodes" not in selected:
        return 0
    wanted = selected["nodes"]
    before = len(wanted)
    for identifier in selected.get("devices", set()):
        wanted.add(indexes["device.csv"][identifier]["node_id"])
    stack, visited = list(wanted), set()
    while stack:
        identifier = stack.pop()
        if identifier in visited:
            continue
        visited.add(identifier)
        for parent in indexes["node.csv"][identifier].get("parent_node_id", "").split(";"):
            if parent:
                wanted.add(parent)
                stack.append(parent)
    return len(wanted) - before


def _warnings(query):
    if query.kinds == ["topology_links"]:
        return [CANDIDATE_NOTICE, *([CSV_REFERENCE_NOTICE] if query.format == "csv" else [])]
    result = [BASELINE_NOTICE]
    if query.range == "filtered" and query.browse == "stations" and "devices" in query.kinds:
        result.append(STATION_NOTICE)
    return result


def _select(value):
    query = _query(value)
    if not datasets.database_mode():
        raise ResourceDatasetError("GIS 导出仅支持已初始化的数据库模式", 5003, 503)
    snapshot = datasets.current()
    if snapshot["revision"] != query.revision:
        raise ResourceDatasetError("资源版本已变化，请刷新后重新导出", 1205, 409)
    view, indexes, links = _source(snapshot)
    selected = _selected(query, view, indexes, links)
    added = _closure(query, selected, indexes)
    by_id = {link["id"]: link for link in links}
    tables = {kind: [deepcopy(by_id[identifier] if kind == "topology_links" else indexes[TABLES[kind]][identifier])
                     for identifier in sorted(identifiers)] for kind, identifiers in selected.items()}
    for kind, rows in tables.items():
        for row in rows:
            _values(kind, row, query.format == "xlsx")
    return query, snapshot, tables, added, _warnings(query)


def _values(kind, row, xlsx=False):
    if kind == "topology_links":
        values = [None if row["assessmentState"] == "PENDING" and field in candidates.ASSESSMENT_FIELDS else row[field]
                  for field in candidates.FIELDS]
    else:
        values = []
        for header in HEADERS[kind]:
            value = row.get(header, "ACTIVE" if header == "availability" else "")
            if value == "":
                value = None
            elif header in NUMERIC:
                try:
                    value = float(value)
                    if not math.isfinite(value):
                        raise ValueError("nonfinite")
                except (ValueError, TypeError, OverflowError) as exc:
                    raise _damaged("资源数值损坏") from exc
            elif header in BOOLEANS:
                if value not in ("true", "false"):
                    raise _damaged("资源布尔值损坏")
                value = value == "true"
            values.append(value)
    return [_text(value, xlsx) if isinstance(value, str) else value for value in values]


def preview(query):
    """与台账相同的 canonical 计数，附加可展示的范围及来源说明。"""
    _, snapshot, tables, added, warnings = _select(query)
    return {"revision": snapshot["revision"], "data_origin": snapshot["data_origin"], "source": snapshot["source"],
            "tables": [{"kind": kind, "count": len(rows)} for kind, rows in tables.items()],
            "added": added, "warnings": warnings}


def _csv_value(value):
    if isinstance(value, str) and (re.match(r"^\s*[=+\-@]", value) or value.startswith("~rp~")):
        return "~rp~" + json.dumps(value, ensure_ascii=False, separators=(",", ":"))
    return str(value).lower() if isinstance(value, bool) else value


def _append(sheet, values):
    sheet.append([_text(value, True) if isinstance(value, str) else value for value in values])
    for cell in sheet[sheet.max_row]:
        if isinstance(cell.value, str):
            cell.data_type = "s"


def _json(value):
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"), allow_nan=False)


def export_file(query):
    """返回 bytes/mime/ASCII 文件名；全部工作表和 ZIP 仅在内存生成。"""
    query, snapshot, tables, added, warnings = _select(query)
    if query.format == "csv":
        kind = query.kinds[0]
        output = StringIO(newline="")
        writer = csv.writer(output)
        writer.writerow(candidates.HEADERS if kind == "topology_links" else HEADERS[kind])
        for row in tables[kind]:
            writer.writerow([_csv_value(value) for value in _values(kind, row)])
        data, mime = output.getvalue().encode("utf-8"), "text/csv; charset=utf-8"
    else:
        book = Workbook()
        try:
            for index, (kind, rows) in enumerate(tables.items()):
                sheet = book.active if index == 0 else book.create_sheet(kind)
                sheet.title = kind
                _append(sheet, candidates.HEADERS if kind == "topology_links" else HEADERS[kind])
                for row in rows:
                    _append(sheet, _values(kind, row, True))
                sheet.freeze_panes = "A2"
                for cell in sheet[1]:
                    cell.font = Font(bold=True)
            notes = book.create_sheet("说明")
            for row in (
                ("资源版本", snapshot["revision"]), ("数据来源", snapshot["data_origin"]),
                ("来源说明", snapshot["source"]), ("内容哈希", snapshot["content_hash"]),
                ("导出范围", query.range), ("浏览方式", query.browse), ("补齐上级节点数", added),
                ("筛选条件", _json(query.filter.model_dump(exclude_unset=True)) if query.filter is not None else ""),
                ("明确选择", _json(query.selected_ids) if query.selected_ids is not None else ""),
                ("数据性质", "已保存构造资源，非生产、非实测；未包含草稿或临时观测"),
                ("CSV 文本转义", "公式前缀及 ~rp~ 文本采用 ~rp~ 加 JSON 字符串，不修改原始内容"),
            ):
                _append(notes, row)
            for warning in warnings:
                _append(notes, ("范围与来源说明", warning))
            if "topology_links" in tables:
                _append(notes, ("参考边界", "reference_json 是来源当时的冻结配置和评估，不能拼接到待评估的新配置；CSV 只含原十五列。"))
                _append(notes, ("link_id", "assessment_state", "source_json", "reference_json"))
                for row in tables["topology_links"]:
                    _append(notes, (row["id"], row["assessmentState"], _json(row["source"]), _json(row["reference"])))
            data, mime = _save_memory(book), XLSX_MIME
        finally:
            book.close()
    if not data or len(data) > MAX_FILE_BYTES:
        raise ResourceDatasetError("文件为空或超过 20 MiB，未生成截断文件", 1001, 413)
    label = "candidate-config-external-reference" if query.kinds == ["topology_links"] else "baseline-resources"
    revision = re.sub(r"[^A-Za-z0-9._-]", "_", snapshot["revision"])
    return data, mime, f"synthetic-gis-{label}-{revision}.{query.format}"
