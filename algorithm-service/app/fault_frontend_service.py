"""新版 SR5 的持久化页面工作区；读取不启动算法任务。"""
from __future__ import annotations

import base64
import csv
import hashlib
import io
import json
import threading
import time
import uuid
from copy import deepcopy
from datetime import datetime, timezone

from sqlalchemy import func, select, update
from sqlalchemy.exc import IntegrityError

from . import fault, frontend_integration, kg, kg_workflow_service, planning_adjustment_service
from .database import get_session_factory
from .fault_frontend_models import FaultReceipt, FaultRun, FaultWorkspace
from .fault_frontend_schemas import FaultDraft, FaultMutation, FaultPage
from .fault_plan_source import FaultPlanSourceError, load_plan_source
from .fault_scenario_models import FaultScenario, FaultScenarioVersion
from .planning_plan_models import PlanningCurrentPlan, PlanningPlan
from .planning_w3 import cancel_task, get_task_record

_LOCK = threading.RLock()
HEADERS = ["scenario_key", "scenario_name", "description", "object_type", "object_id", "fault_type", "symptom_ids", "occur_time", "interference_id"]


def fail(message, code=1001, status=422):
    raise FaultPlanSourceError(code, message, status)


def _iso(value):
    return datetime.fromtimestamp(value, timezone.utc).isoformat() if isinstance(value, (float, int)) else value.isoformat()


def _workspace():
    try:
        with get_session_factory().begin() as session:
            if session.get(FaultWorkspace, 1) is None:
                session.add(FaultWorkspace(id=1, revision=1))
    except IntegrityError:
        pass
    with get_session_factory()() as session:
        return session.get(FaultWorkspace, 1).revision


def _number(value):
    try:
        return float(value) if value not in (None, "") else None
    except (TypeError, ValueError):
        return None


def _terrain(plan):
    # 只有冻结地形指纹仍可核实时才关联当前图片；不能给历史方案套用新 DEM。
    fingerprint = (plan.raw_snapshot or {}).get("terrain_fingerprint")
    try:
        if not fingerprint or fingerprint != planning_adjustment_service.terrain_fingerprint():
            return None, "该方案的 DEM 版本无法与当前地形核对，未显示地形底图。"
        meta = json.loads((frontend_integration.REPO_ROOT / "data/raw/dem/dem_current.json").read_text(encoding="utf-8"))
        return {"sourceSha256": meta["source_sha256"], "bounds": meta["bbox"], "source": meta["source"]}, ""
    except (OSError, ValueError, KeyError, TypeError):
        return None, "当前 DEM 元数据缺失，未显示地形底图。"


def network(plan: PlanningPlan):
    raw = (plan.raw_snapshot or {}).get("detail", {})
    result = plan.result_snapshot
    source_devices = {r["device_id"]: r for r in raw["devices"]}
    result_devices = {r["id"]: r for r in result.get("devices", [])}
    source_nodes = raw["nodes"]
    devices = []
    for did, row in source_devices.items():
        output = result_devices.get(did, {})
        status = row.get("status")
        devices.append({"id": did, "stationId": row["node_id"], "modelId": row.get("model_id") or "型号未提供",
            "antennaId": row.get("antenna_model_id") or row.get("antenna_id") or "天线型号未提供",
            "status": status if status in ("NORMAL", "FAULT", "OFFLINE") else None, "alert": None,
            "antennaHeightM": output.get("antennaHeightM", _number(row.get("antenna_height_m"))),
            "txPowerDbm": output.get("txPowerDbm", _number(row.get("tx_power_dbm"))),
            "workFreqKhz": output.get("workFreqKhz", _number(row.get("work_freq_khz"))), "installedAt": row.get("installed_at", "")})
    stations = []
    for row in source_nodes:
        nid = row["node_id"]
        status = row.get("status")
        classes = sorted({d.get("device_class") for d in source_devices.values() if d["node_id"] == nid and d.get("device_class") in ("HF", "VUHF")})
        stations.append({"id": nid, "name": row.get("node_name") or nid, "echelon": row.get("echelon", ""),
            "parentIds": [], "coordinate": [float(row["lon"]), float(row["lat"])],
            "elevationM": _number(row.get("elevation_m")), "region": row.get("region", ""), "classes": classes,
            "status": status if status in ("NORMAL", "FAULT", "OFFLINE") else None, "alert": None,
            "relayCapable": str(row.get("relay_capable", "false")).lower() == "true", "key": False,
            "mobility": row.get("mobility", ""), "role": row.get("node_role", "")})
    if not result.get("links"):
        fail("保存的方案缺少可展示的完整链路，请重新计算并保存", 1203, 409)
    metrics = result.get("metrics", {})
    terrain, terrain_notice = _terrain(plan)
    return {"id": plan.plan_id, "taskId": plan.task_id, "jobId": plan.job_id, "name": plan.name,
        "version": f"V{plan.version_no}", "snapshotId": result["snapshot_id"], "preset": plan.input_snapshot.get("preset", "SR4"),
        "state": plan.state, "sourceChanged": plan.resource_revision != frontend_integration.resource_revision(),
        "inputSnapshot": {"stations": stations, "devices": devices}, "stationIds": [r["id"] for r in stations],
        "links": deepcopy(result["links"]), "regions": [],
        "terrain": terrain, "terrainNotice": terrain_notice,
        "metrics": {key: metrics.get(key) for key in ("coveragePercent", "connectivityPercent", "survivabilityPercent", "totalRadiatedPowerW", "delayMs")},
        "notes": ["构造联调网络；未提供的状态及指标保持未知，不代表在线监测。"]}


def sources():
    items = []
    with get_session_factory()() as session:
        rows = session.scalars(select(PlanningPlan).join(PlanningCurrentPlan, PlanningCurrentPlan.plan_id == PlanningPlan.plan_id).order_by(PlanningPlan.saved_at.desc())).all()
        for row in rows:
            try:
                load_plan_source(row.task_id, row.plan_id)
                valid, reason = True, None
            except FaultPlanSourceError as exc:
                valid, reason = False, str(exc)
            items.append({"id": row.plan_id, "taskId": row.task_id, "name": row.name,
                "version": f"V{row.version_no}", "preset": row.input_snapshot.get("preset", "SR4"),
                "snapshotId": row.result_snapshot.get("snapshot_id", "缺少快照"), "canCreate": valid, "unavailableReason": reason})
    return {"plans": items, "error": None if items else "暂无当前已确认的完整方案，请先在网络规划中计算、保存并确认。"}


def source_detail(plan_id, *, require_current=True):
    with get_session_factory()() as session:
        row = session.get(PlanningPlan, plan_id)
        if row is None:
            fail("方案不存在", 2001, 404)
        load_plan_source(row.task_id, row.plan_id, require_current=require_current)
        return network(row)


def context():
    revision = _workspace()
    with get_session_factory()() as session:
        latest = session.scalar(select(FaultRun.id).order_by(FaultRun.started_at.desc()).limit(1))
    return {"revision": f"SR5-{revision}", "error": None, "presets": [], "legacy": [], "sources": sources(),
        "dataOrigin": "SAMPLE", "capabilities": {"demoControls": False, "twoStep": True},
        "symptoms": [{"id": key, "label": value} for key, value in fault._engine().symptoms.items()],
        "interferenceIds": [r["interference_id"] for r in frontend_integration._rows("interferences")], "latestRunId": latest}


def _draft(version):
    snap = version.snapshot
    return {"name": snap["scenario_name"], "description": snap.get("description", ""), "planId": version.plan_id,
            "faults": [{"objectType": f["object_type"], "objectId": f["object_id"], "faultType": f["fault_type"],
                        "symptomIds": f.get("symptom_ids", []), "occurTime": f.get("occur_time"), "interferenceId": f.get("interference_id")} for f in snap["faults"]]}


def scenario_version(scenario_id, version_id):
    with get_session_factory()() as session:
        row = session.get(FaultScenario, scenario_id)
        version = session.get(FaultScenarioVersion, version_id)
        if row is None or version is None or version.scenario_id != scenario_id:
            fail("场景版本不存在或归属不符", 1002, 404)
        plan = session.get(PlanningPlan, version.plan_id)
        return {**_draft(version), "id": version.version_id, "scenarioId": scenario_id, "version": version.version_no,
                "savedAt": _iso(version.saved_at), "network": network(plan), "scenarioArchived": row.archived}


def catalog(kind, page: FaultPage):
    with get_session_factory()() as session:
        if kind == "history":
            stmt = select(FaultRun)
            total = session.scalar(select(func.count()).select_from(stmt.subquery()))
            rows = session.scalars(stmt.order_by(FaultRun.started_at.desc(), FaultRun.id).offset((page.pageNum-1)*page.pageSize).limit(page.pageSize)).all()
            items = [{"id": r.id, "startedAt": r.started_at*1000, "name": r.snapshot["name"], "scenarioVersion": r.snapshot["version"], "dataOrigin": "SAMPLE"} for r in rows]
        else:
            stmt = select(FaultScenario).where(FaultScenario.archived == page.archived)
            if page.search:
                stmt = stmt.where(FaultScenario.scenario_id.contains(page.search, autoescape=True) | FaultScenario.scenario_id.in_(select(FaultScenarioVersion.scenario_id).where(FaultScenarioVersion.snapshot["scenario_name"].as_string().contains(page.search, autoescape=True))))
            total = session.scalar(select(func.count()).select_from(stmt.subquery()))
            rows = session.scalars(stmt.order_by(FaultScenario.updated_at.desc(), FaultScenario.scenario_id).offset((page.pageNum-1)*page.pageSize).limit(page.pageSize)).all()
            items = []
            for row in rows:
                versions = session.scalars(select(FaultScenarioVersion).where(FaultScenarioVersion.scenario_id == row.scenario_id).order_by(FaultScenarioVersion.version_no.desc())).all()
                items.append({"id": row.scenario_id, "archived": row.archived, "versions": [{"id": v.version_id, "scenarioId": row.scenario_id,
                    "version": v.version_no, "name": v.snapshot["scenario_name"], "planId": v.plan_id, "savedAt": _iso(v.saved_at), "faultCount": len(v.snapshot["faults"])} for v in versions]})
    return {"items": items, "total": total, "pageNum": page.pageNum, "pageSize": page.pageSize}


def validate(draft: FaultDraft):
    with get_session_factory()() as session:
        plan = session.get(PlanningPlan, draft.planId)
        if plan is None:
            fail("源方案不存在", 2001, 404)
        source = load_plan_source(plan.task_id, plan.plan_id)
    faults = []
    seen = set()
    interferences = {r["interference_id"] for r in frontend_integration._rows("interferences")}
    for f in draft.faults:
        if (f.objectType, f.objectId) in seen or len(set(f.symptomIds)) != len(f.symptomIds):
            fail("故障对象或现象 ID 不能重复")
        seen.add((f.objectType, f.objectId))
        if f.interferenceId is not None and (f.faultType != "EMI" or f.interferenceId not in interferences):
            fail("干扰源不在当前资源中或不是干扰故障")
        faults.append({"object_type": f.objectType, "object_id": f.objectId, "fault_type": f.faultType,
                       "symptom_ids": f.symptomIds, "occur_time": f.occurTime, "interference_id": f.interferenceId})
    snap = {"scenario_name": draft.name.strip(), "description": draft.description, "faults": faults}
    if not snap["scenario_name"]:
        fail("场景名称不能为空")
    fault._engine().for_network(source["network"]).validate(snap)
    return source, snap


def _get_run(session, run_id):
    row = session.get(FaultRun, run_id)
    if row is None:
        fail("诊断运行不存在", 2001, 404)
    return row


def record_runtime(run_id, kind, state):
    with _LOCK, get_session_factory().begin() as session:
        row = _get_run(session, run_id)
        field = "diagnosis" if kind == "FAULT_DIAGNOSE" else "impact"
        setattr(row, field, deepcopy(state))
        setattr(row, "diagnose_task_id" if field == "diagnosis" else "impact_task_id", state["task_id"])


def _result(row):
    diagnosis = (row.diagnosis or {}).get("result") or {}
    impact = (row.impact or {}).get("result") or {}
    conclusions = []
    for index, rec in enumerate(diagnosis.get("results", [])):
        fused = rec["fused"]
        active = [b for b in rec["branches"] if b.get("confidence") is not None]
        weights = {"RULE": .35, "KG": .35, "ML": .30}
        denom = sum(weights[b["source"]] for b in active)
        branches = [{"source": b["source"], "conclusion": b.get("fault_type") or "未形成结论", "confidence": b.get("confidence"),
            "weight": weights[b["source"]]/denom if b in active and denom else None,
            "contribution": weights[b["source"]]*b["confidence"]/denom if b in active and denom else None,
            "disabledReason": b.get("disabled_reason"), "evidence": [{"label": k, "value": json.dumps(v, ensure_ascii=False) if not isinstance(v, str) else v} for k,v in b.get("evidence", {}).items()]} for b in rec["branches"]]
        conclusions.append({"fault": row.snapshot["faults"][index], "label": rec["fault_object"]["object_id"], "faultType": fused.get("fault_type") or "UNKNOWN",
            "confidence": fused["confidence"], "rootCause": fused.get("root_cause") or "未提供根因", "component": rec.get("fault_component"),
            "needManualReview": fused.get("need_manual_review"), "branches": branches, "candidates": [{"faultType": c["fault_type"], "confidence": c["confidence"]} for c in rec["candidates"]], "chain": []})
    suggestions = fault._engine().suggestions(diagnosis, impact)["suggestions"] if diagnosis else []
    summary = impact.get("summary")
    return {"fusionMode": "TWO_WAY_DEGRADED", "disabledReason": "ML 未启用；现有知识与案例未经实测验证。",
        "conclusions": conclusions, "nodes": [{"nodeId": n["node_id"], "affectType": n["affect_type"], "level": n["level"], "state": n["state"], "reason": n["reason"], "detourRoute": []} for n in impact.get("affected_nodes", [])],
        "links": [{"linkId": l["link_id"], "level": l["level"], "state": l["state"], "faultKey": None, "bandwidthLossKbps": None,
                   "isBackbone": l.get("is_backbone"), "reason": l["reason"]} for l in impact.get("affected_links", [])],
        "regions": [], "summary": {"affectedNodeCount": summary["affected_node_count"], "directCount": summary["direct_count"], "indirectCount": summary["indirect_count"],
            "interruptedNodeCount": summary["interrupted_node_count"], "degradedNodeCount": summary["degraded_node_count"], "affectedLinkCount": summary["affected_link_count"],
            "totalBandwidthLossKbps": summary["total_bandwidth_loss_kbps"], "affectedDemandIds": [d["demand_id"] for d in impact.get("affected_demands", [])]} if summary else None,
        "suggestions": [{"priority": s["priority"], "action": s.get("action") or "建议未提供", "applicability": s.get("fault_object_id", "全场景"), "steps": s["steps"],
            "spares": s["required_spares"], "tools": s["required_tools"], "estimatedTimeMin": s["estimated_time_min"], "tripleIds": s["kg_source_triple_ids"],
            "similarCases": [{"id": c["case_id"], "solution": c["solution"], "similarity": None} for c in s["similar_cases"]]} for s in suggestions],
        "notes": ["构造联调数据，非生产、非实测。", "影响区域仅提供对象位置，未提供区域边界与传播路径，不据此构造面。",
                  "诊断使用已发布知识版本：" + str(diagnosis.get("knowledge_version", "未提供"))]}


def view(run_id):
    with get_session_factory()() as session:
        row = _get_run(session, run_id)
        d, i = row.diagnosis or {}, row.impact or {}
        ready, impact_ready = d.get("status") == "SUCCESS", i.get("status") == "SUCCESS"
        ds, ins = d.get("status", "PENDING"), i.get("status", "WAITING")
        state = ins if ready and i else ds
        if ready and not i:
            state = "RUNNING"
        error = i.get("error") or d.get("error")
        if row.cancelled_at:
            state = "CANCELLED"
        # 重启后进程内非终态任务不会继续执行，明确失败，不能永久显示进行中。
        running_id = row.impact_task_id if ready and i else row.diagnose_task_id
        if state in ("PENDING", "RUNNING") and running_id and not get_task_record(running_id) and (not ready or i):
            state, error = "FAILED", "运行进程已中断，请保留历史并重新提交"
        if state not in ("PENDING", "RUNNING", "SUCCESS", "FAILED", "CANCELLED"):
            state, error = "FAILED", "运行状态不可识别"
        try:
            load_plan_source(row.task_id, row.plan_id)
            source = {"state": "CURRENT", "reasons": []}
        except FaultPlanSourceError as exc:
            source = {"state": "HISTORICAL", "reasons": [str(exc)]}
        progress = (float(d.get("progress", 0)) + float(i.get("progress", 0))) * 50
        result_version = hashlib.sha256(json.dumps([row.diagnosis, row.impact], sort_keys=True, ensure_ascii=False).encode()).hexdigest()[:20]
        result = _result(row)
        run = {"id": row.id, "diagnoseTaskId": row.diagnose_task_id, "impactTaskId": row.impact_task_id,
            "scenario": deepcopy(row.snapshot), "presetId": "", "presetVersion": "backend-sr5-v1", "startedAt": row.started_at*1000,
            "cancelledAt": row.cancelled_at*1000 if row.cancelled_at else None}
        return {"run": run, "dataOrigin": "SAMPLE", "state": state, "progress": progress,
            "phases": [{"name": "故障诊断", "state": "done" if ready else "running" if ds == "RUNNING" else "waiting"},
                       {"name": "影响分析", "state": "done" if impact_ready else "running" if ins == "RUNNING" else "waiting"}],
            "diagnosisReady": ready, "impactReady": impact_ready, "error": error, "source": source,
            "id": row.id, "resultVersion": result_version, "result": result}


def report(run_id):
    current = view(run_id)
    if current["state"] != "SUCCESS" or not current["impactReady"]:
        fail("完整诊断与影响分析尚未成功，不生成报告", 2002, 409)
    with get_session_factory()() as session:
        row = _get_run(session, run_id)
        return {"id": row.id, "dataOrigin": "SAMPLE", "legacy": False,
            "scenarioId": row.snapshot["scenarioId"], "scenarioVersion": row.snapshot["id"],
            "name": row.snapshot["name"], "description": row.snapshot["description"], "faults": row.snapshot["faults"], "network": row.snapshot["network"],
            "startedAt": _iso(row.started_at), "completedAt": _iso(row.impact["finished_at"]),
            "diagnoseTaskId": row.diagnose_task_id, "impactTaskId": row.impact_task_id, "result": current["result"], "source": current["source"]}


def mutate(action, request: FaultMutation, identity):
    _workspace()
    if action == "feedback":
        kg_workflow_service.workspace(kg.STORE_PATH)
    normalized = request.model_dump(mode="json")
    digest = hashlib.sha256(json.dumps([identity["account_id"], action, normalized], sort_keys=True, ensure_ascii=False).encode()).hexdigest()
    launch = None
    with _LOCK, get_session_factory().begin() as session:
        old = session.get(FaultReceipt, request.requestId)
        if old:
            if old.digest != digest:
                fail("同一请求编号不能用于不同操作或参数", 1204, 409)
            return deepcopy(old.response)
        revision = int(request.baseRevision[4:])
        changed = session.execute(update(FaultWorkspace).where(FaultWorkspace.id == 1, FaultWorkspace.revision == revision).values(revision=revision+1))
        if changed.rowcount != 1:
            fail("诊断工作区已变化，请保留输入并重新读取", 1203, 409)
        entity_ids = []
        if action in ("save", "import"):
            drafts = request.drafts if action == "import" else [request.draft] if request.draft else []
            if not drafts or len({d.planId for d in drafts}) != 1:
                fail("需提供绑定同一网络的场景")
            validated = [validate(d) for d in drafts]
            now = datetime.now(timezone.utc)
            for source, snap in validated:
                if action == "save" and request.scenarioId:
                    row = fault._formal_record(session, request.scenarioId, source["meta"]["task_id"])
                    if row.archived:
                        fail("已归档场景不能修改", 1203, 409)
                    version_no = fault._latest_version(session, row.scenario_id).version_no + 1
                else:
                    row = FaultScenario(scenario_id="FS-"+uuid.uuid4().hex[:16].upper(), task_id=source["meta"]["task_id"], archived=False, created_at=now, updated_at=now)
                    session.add(row)
                    version_no = 1
                version = fault._new_version(session, row, source, snap, version_no)
                row.updated_at = now
                entity_ids.extend([row.scenario_id, version.version_id])
        elif action == "archive":
            row = session.get(FaultScenario, request.scenarioId)
            if row is None:
                fail("故障场景不存在", 1002, 404)
            if request.archived is None:
                fail("缺少归档状态")
            row.archived = request.archived
            row.updated_at = datetime.now(timezone.utc)
            entity_ids = [row.scenario_id]
        elif action == "submit":
            version = session.get(FaultScenarioVersion, request.scenarioVersionId)
            if version is None:
                fail("场景版本不存在", 1002, 404)
            row = session.get(FaultScenario, version.scenario_id)
            if row.archived:
                fail("已归档场景不能提交新诊断", 1203, 409)
            load_plan_source(row.task_id, version.plan_id)
            plan = session.get(PlanningPlan, version.plan_id)
            frozen = {**_draft(version), "id": version.version_id, "scenarioId": row.scenario_id, "version": version.version_no,
                      "savedAt": _iso(version.saved_at), "network": network(plan)}
            run = FaultRun(id="FR-"+uuid.uuid4().hex[:20], scenario_version_id=version.version_id,
                task_id=row.task_id, plan_id=version.plan_id, started_at=time.time(), snapshot=frozen)
            session.add(run)
            entity_ids = [run.id]
            launch = ("diagnosis", run.id)
        elif action in ("impact", "cancel", "feedback"):
            run = _get_run(session, request.runId)
            entity_ids = [run.id]
            if action == "cancel":
                if (run.impact or {}).get("status") == "SUCCESS":
                    fail("已完成的完整分析不能取消", 1203, 409)
                run.cancelled_at = time.time()
                launch = ("cancel", run.id)
            elif run.cancelled_at:
                fail("已取消的诊断不能继续分析或反馈", 1203, 409)
            elif (run.diagnosis or {}).get("status") != "SUCCESS":
                fail("请先等待诊断成功", 2002, 409)
            elif action == "impact":
                if not run.impact_task_id:
                    launch = ("impact", run.id)
            else:
                from .kg_frontend_service import identity_snapshot
                payload = dict(request.feedback or {})
                allowed = {"result_index", "confirmed_fault_type", "root_cause", "fault_component", "solution", "diagnosis_steps"}
                if set(payload) - allowed:
                    fail("知识反馈包含未知字段")
                payload.update(reviewer=identity["account_id"], _workflow_request={
                    "actor": identity["account_id"], "request_id": "feedback:"+request.requestId,
                    "base_revision": kg_workflow_service._bootstrap(session, kg.STORE_PATH).revision,
                    "_identity": identity_snapshot(identity)})
                response = fault.submit_diagnosis_feedback(run.diagnose_task_id, payload, session=session)
                if response["code"]:
                    fail(response["message"], response["code"])
                entity_ids.append(response["data"]["change_id"])
        else:
            fail("未支持的操作")
        receipt = {"objectId": entity_ids[0], "entityIds": entity_ids, "changed": True, "revision": f"SR5-{revision+1}"}
        session.add(FaultReceipt(request_id=request.requestId, digest=digest, response=deepcopy(receipt)))
    if launch:
        _launch(*launch)
    return receipt


def _launch(stage, run_id):
    with get_session_factory()() as session:
        run = _get_run(session, run_id)
        if stage == "cancel":
            for tid in (run.diagnose_task_id, run.impact_task_id):
                if tid:
                    cancel_task(tid)
            return
        payload = {"context": "planning", "task_id": run.task_id, "plan_id": run.plan_id,
                   "request_id": f"frontend:{run_id}:{stage}", "frontend_run_id": run_id}
        if stage == "diagnosis":
            payload.update(scenario_id=run.snapshot["scenarioId"], scenario_version_id=run.scenario_version_id)
        else:
            payload["diagnose_task_id"] = run.diagnose_task_id
    try:
        result = fault.submit_diagnosis(payload) if stage == "diagnosis" else fault.submit_impact(payload)
        if result["code"]:
            raise ValueError(result["message"])
    except Exception as exc:
        # 已返回受理回执后将启动失败记录为终态，不伪造为提交成功的算法结果。
        with get_session_factory().begin() as session:
            run = _get_run(session, run_id)
            setattr(run, "diagnosis" if stage == "diagnosis" else "impact", {"status": "FAILED", "progress": 0, "error": str(exc)})


def import_preview(payload):
    name, plan_id = payload.get("name", ""), payload.get("planId")
    try:
        contents = base64.b64decode(payload.get("base64", ""), validate=True)
    except (ValueError, TypeError):
        fail("文件编码无效")
    if len(contents) > 5*1024*1024:
        fail("单文件不能超过 5 MiB")
    if name.lower().endswith(".csv"):
        try:
            rows = list(csv.reader(io.StringIO(contents.decode("utf-8-sig"))))
        except (UnicodeError, csv.Error):
            fail("CSV 必须为合法 UTF-8 文件")
    elif name.lower().endswith(".xlsx"):
        from openpyxl import load_workbook
        try:
            book = load_workbook(io.BytesIO(contents), read_only=True, data_only=False)
            if "faults" not in book.sheetnames:
                fail("XLSX 需要 faults 工作表")
            sheet = book["faults"]
            if sheet.max_row > 401 or sheet.max_column != len(HEADERS):
                fail("工作表超出模板行列范围")
            rows = []
            for cells in sheet.iter_rows():
                if any(c.data_type == "f" for c in cells):
                    fail("导入不接受公式")
                rows.append([c.value for c in cells])
            book.close()
        except FaultPlanSourceError:
            raise
        except Exception as exc:
            fail("XLSX 无法解析：" + str(exc))
    else:
        fail("仅支持 CSV 或 XLSX 文件")
    if not rows or rows[0] != HEADERS or len(rows) > 401:
        fail("表头不匹配或超过 400 条故障行")
    groups, issues = {}, []
    for number, row in enumerate(rows[1:], 2):
        if not any(v not in (None, "") for v in row):
            continue
        if len(row) != len(HEADERS):
            issues.append({"row": number, "field": "row", "message": "列数与模板不一致"})
            continue
        data = dict(zip(HEADERS, [str(v).strip() if v is not None else "" for v in row]))
        key = data["scenario_key"]
        if not key:
            issues.append({"row": number, "field": "scenario_key", "message": "场景分组不能为空"})
            continue
        draft = groups.setdefault(key, {"name": data["scenario_name"], "description": data["description"], "planId": plan_id, "faults": []})
        if draft["name"] != data["scenario_name"] or draft["description"] != data["description"]:
            issues.append({"row": number, "field": "scenario_name", "message": "同组名称和备注必须一致"})
        draft["faults"].append({"objectType": data["object_type"], "objectId": data["object_id"], "faultType": data["fault_type"],
            "symptomIds": list(filter(None, data["symptom_ids"].split(";"))), "occurTime": data["occur_time"] or None, "interferenceId": data["interference_id"] or None})
    drafts = list(groups.values())
    if not 1 <= len(drafts) <= 20:
        issues.append({"field": "scenarios", "message": "每批需 1 至 20 个场景"})
    for number, draft in enumerate(drafts, 1):
        try:
            validate(FaultDraft.model_validate(draft))
        except Exception as exc:
            issues.append({"row": number, "field": "draft", "message": str(exc)})
    return {"fileName": name, "drafts": drafts, "issues": issues, "matches": [False]*len(drafts)}


def template(plan_id, fmt):
    frozen = source_detail(plan_id)
    device = frozen["inputSnapshot"]["devices"][0]
    rows = [HEADERS, ["scene-1", "构造联调故障场景", "非实测", "DEVICE", device["id"], "HW_FAILURE", "", "", ""]]
    buf = io.BytesIO()
    if fmt == "csv":
        text = io.StringIO()
        csv.writer(text).writerows(rows)
        return text.getvalue().encode("utf-8"), "text/csv;charset=utf-8"
    if fmt != "xlsx":
        fail("模板格式仅支持 csv/xlsx")
    from openpyxl import Workbook
    book = Workbook()
    book.active.title = "faults"
    for row in rows:
        book.active.append(row)
    book.save(buf)
    return buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def export_file(run_id, kind, fmt, filters=""):
    record = report(run_id)
    with get_session_factory()() as session:
        row = _get_run(session, run_id)
        diagnosis, impact = row.diagnosis["result"], row.impact["result"]
        source = {"run_id": row.id, "task_id": row.task_id, "plan_id": row.plan_id,
                  "snapshot_id": record["network"].get("snapshotId"), "scenario_version": record["scenarioVersion"],
                  "diagnose_task_id": row.diagnose_task_id, "impact_task_id": row.impact_task_id}
        if kind == "impact":
            data = fault.export_impact(row.impact_task_id, fmt, network=record["network"], source=source)
            if data is None:
                fail("影响导出格式不支持或原始任务不可用")
            return data
        if fmt == "json":
            return json.dumps(record, ensure_ascii=False, allow_nan=False).encode("utf-8"), "application/json"
        if fmt not in ("docx", "pdf"):
            fail("报告格式不支持")
        content = fault._report_content(diagnosis, impact, fault._engine().suggestions(diagnosis, impact))
        fault.REPORT_DIR.mkdir(parents=True, exist_ok=True)
        path = fault.REPORT_DIR / f"report_{run_id}_{uuid.uuid4().hex}.{fmt}"
        (fault._report_docx if fmt == "docx" else fault._report_pdf)(content, path, source=source)
        return path.read_bytes(), "application/pdf" if fmt == "pdf" else "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
