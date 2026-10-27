"""版本化 SR4 方案的保存、分页查询和详情读取。"""
from __future__ import annotations

import hashlib
import io
import json
import os
import threading
import uuid
from datetime import datetime, timezone
from typing import TYPE_CHECKING, Any

from sqlalchemy import func, or_, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session, sessionmaker
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

from . import frontend_integration, planning_adjustment_service, planning_job_service
from .database import get_session_factory
from .planning_job_models import PlanningJob
from .planning_result_adapter import project_saved_deployment_display
from .planning_operation_service import _append_owner, _business_context, _require_owner
from .planning_plan_models import (
    PlanningCurrentPlan,
    PlanningPlan,
    PlanningPlanConfirmation,
)
from .planning_plan_schemas import (
    PlanningPlanCompareRequest,
    PlanningPlanConfirmRequest,
    PlanningPlanExportRequest,
    PlanningPlanSaveRequest,
)

if TYPE_CHECKING:
    from .local_business_context import LocalBusinessContext


class PlanningPlanError(ValueError):
    """可直接映射为接口错误的方案管理异常。"""

    def __init__(self, code: int, message: str, status: int,
                 data: Any = None) -> None:
        super().__init__(message)
        self.code = code
        self.status = status
        self.data = data


_SAVE_LOCK = threading.RLock()
_CONFIRM_LOCK = threading.RLock()


def _digest(payload: PlanningPlanSaveRequest) -> str:
    body = payload.model_dump(mode="json", exclude_none=False)
    raw = json.dumps(body, ensure_ascii=False, sort_keys=True,
                     separators=(",", ":")).encode("utf-8")
    return hashlib.sha256(raw).hexdigest()


def _timestamp(value: datetime | None) -> str | None:
    if value is None:
        return None
    if value.tzinfo is None:
        value = value.replace(tzinfo=timezone.utc)
    return value.isoformat().replace("+00:00", "Z")


def _source_changed(plan: PlanningPlan) -> bool:
    return plan.resource_revision != frontend_integration.resource_revision()


def _dto(plan: PlanningPlan, current_plan_id: str | None = None, *,
         source_changed: bool | None = None) -> dict[str, Any]:
    if source_changed is None:
        source_changed = _source_changed(plan)
    data = {
        "plan_id": plan.plan_id,
        "task_id": plan.task_id,
        "job_id": plan.job_id,
        "revision": plan.revision_id,
        "version": f"V{plan.version_no}",
        "name": plan.name,
        "note": plan.note,
        "input_version": plan.revision_id,
        "condition_version": plan.condition_version,
        "resource_revision": plan.resource_revision,
        "input": plan.input_snapshot,
        "parent_plan_id": plan.parent_plan_id,
        "state": plan.state,
        "source_changed": source_changed,
        "current": current_plan_id == plan.plan_id and not source_changed,
        "saved_at": _timestamp(plan.saved_at),
        "confirmed_at": _timestamp(plan.confirmed_at),
        "source": plan.source,
        "result": project_saved_deployment_display(plan.result_snapshot, getattr(plan, "raw_snapshot", None)),
        "outcome": plan.outcome_snapshot,
    }
    if plan.adjustment is not None:
        data["adjustment"] = plan.adjustment
    return data


def _current_plan_id(session: Session, task_id: str) -> str | None:
    pointer = session.get(PlanningCurrentPlan, task_id)
    return pointer.plan_id if pointer is not None else None


def _plan_dto(session: Session, plan: PlanningPlan, *, same_transaction_resource=False) -> dict[str, Any]:
    changed = None
    if same_transaction_resource:
        from .resource_dataset_models import ResourceDatasetCurrent, ResourceDatasetVersion
        dataset = frontend_integration.resource_dataset
        # 大快照flush可能持SQLite排他锁，后台不能另开连接读当前资源而等待自己。
        # 沿原事务完整解码/校验，仍拒绝坏指针、摘要或资源，不用跨请求缓存。
        pointer = session.get(ResourceDatasetCurrent, dataset.SCOPE)
        row = session.get(ResourceDatasetVersion, pointer.revision) if pointer is not None else None
        if pointer is not None and row is None:
            raise dataset.ResourceDatasetError("当前资源指针引用的快照不存在")
        changed = plan.resource_revision != dataset._decode(row)["revision"]
    return _dto(plan, _current_plan_id(session, plan.task_id), source_changed=changed)


def _replay(session: Session, request_id: str,
            digest: str, audit_context=None) -> PlanningPlan | None:
    plan = session.scalar(select(PlanningPlan).where(
        PlanningPlan.request_id == request_id
    ))
    if plan is None:
        return None
    if plan.request_digest != digest:
        raise PlanningPlanError(1204, "同一请求编号对应不同内容", 409)
    if audit_context is not None:
        _require_owner(session, audit_context, operation="plan.save", request_id=request_id,
            object_id=plan.plan_id, page_context="planning", task_id=plan.task_id, revision_id=plan.revision_id)
    return plan


def _job_result(payload: PlanningPlanSaveRequest,
                factory: sessionmaker[Session]) -> dict[str, Any]:
    try:
        response = planning_job_service.query_planning_job(
            payload.job_id, payload.task_id, factory
        )
    except planning_job_service.PlanningJobError as exc:
        raise PlanningPlanError(exc.code, str(exc), exc.status, exc.data) from exc
    if response is None:
        raise PlanningPlanError(2001, "规划作业不存在", 404)
    data = response.get("data") or {}
    if data.get("revision") != payload.expected_revision:
        raise PlanningPlanError(1203, "作业输入版本与保存请求不一致", 409, {
            "expected": payload.expected_revision,
            "actual": data.get("revision"),
            "suggestions": ["重新打开规划结果"],
        })
    if data.get("status") != "SUCCESS" or not data.get("result") or not data.get("outcome"):
        raise PlanningPlanError(2002, "规划作业尚未产生可保存结果", 409)
    if response.get("code") not in {0, 3001, 3002}:
        raise PlanningPlanError(5001, "规划作业结果状态异常", 500)
    if payload.solution_id is not None:
        raise PlanningPlanError(1001, "当前 SR4 作业暂不支持多方案结果保存", 422)
    return data


def _interference_save_proof(session, payload, request):
    """后台只消费数据库已受理的采纳授权，不创建会话或放宽HTTP身份门禁。"""
    from .interference_analysis_engine import digest as canonical_digest
    from .interference_replan_service import InterferenceReplanService
    from .interference_workbench_models import InterferenceWorkbenchAdoption, InterferenceWorkbenchOperation
    from .interference_workbench_service import InterferenceWorkbenchService
    from .system_management_models import LocalAccount, LocalRole

    required = {"adoption_id", "planning_request_id", "planning_job_id", "actor_id", "task_id",
                "source", "measure", "change_set", "interference_origin"}
    if type(request) is not dict or not required.issubset(request):
        raise PlanningPlanError(1203, "后台保存缺少完整采纳授权", 409)
    row = session.get(InterferenceWorkbenchAdoption, request["adoption_id"])
    if row is None or row.task_id != request["task_id"] or row.state not in (
            "PENDING", "RUNNING", "CANCEL_REQUESTED", "SUCCESS"):
        raise PlanningPlanError(1203, "后台保存没有可用的已受理采纳", 409)
    run = InterferenceWorkbenchService._run(session, row.run_id, row.task_id)
    origin = InterferenceReplanService.origin(row)
    measure = next((item for item in (run.result_snapshot or {}).get("measures", [])
                    if item["measure_id"] == row.measure_id), None)
    if (run.state != "SUCCESS" or row.source_hash != run.source_hash
            or row.change_hash != canonical_digest(row.change_set)
            or measure is None or not measure["available"] or measure["change_set"] != row.change_set
            or request["planning_request_id"] != row.planning_request_id
            or request["planning_job_id"] != row.planning_job_id or request["actor_id"] != row.actor_id
            or request["source"] != run.source_snapshot or request["measure"] != measure
            or request["change_set"] != row.change_set or request["interference_origin"] != origin):
        raise PlanningPlanError(1203, "后台保存与已受理来源或措施不一致", 409)
    validator = InterferenceReplanService(None)
    job = validator._job(session, row, run, request["planning_job_id"])
    expected = PlanningPlanSaveRequest(context="planning", request_id="interference-save-" + row.adoption_id,
        task_id=job.task_id, expected_revision=job.revision_id, job_id=job.job_id,
        name="干扰措施重算 " + measure["type"], note="冻结场景算法预测；非实测。来源采纳：" + row.adoption_id)
    if job.status != "SUCCESS" or payload.model_dump() != expected.model_dump():
        raise PlanningPlanError(1203, "后台保存不属于该成功子作业的固定请求", 409)
    # 不可变回执是授权来源；仅接受创建此采纳的原受理人及changed=true原件。
    accepted = []
    operations = session.scalars(select(InterferenceWorkbenchOperation).where(
        InterferenceWorkbenchOperation.operation == "ADOPT_MEASURE",
        InterferenceWorkbenchOperation.object_id == row.adoption_id,
        InterferenceWorkbenchOperation.actor_id == row.actor_id)).all()
    for operation in operations:
        receipt = InterferenceWorkbenchService._operation_dto(operation)
        if receipt.get("changed") is True:
            accepted.append(receipt)
    if (len(accepted) != 1 or accepted[0].get("task_id") != row.task_id
            or accepted[0].get("status") != "ACCEPTED" or accepted[0].get("object_type") != "ADOPTION"
            or accepted[0].get("source_hash") != row.source_hash
            or accepted[0].get("adoption_id") != row.adoption_id
            or accepted[0].get("planning_request_id") != row.planning_request_id):
        raise PlanningPlanError(1203, "后台保存缺少唯一不可变采纳受理证明", 409)
    snapshot = None
    if os.environ.get("AUTH_MODE") == "LOCAL_PASSWORD":
        account = session.get(LocalAccount, row.actor_id, populate_existing=True)
        role = session.get(LocalRole, account.role_id, populate_existing=True) if account is not None else None
        permissions = {"planning.view", "situation.view", "planning.execute", "planning.save"}
        if account is None or not account.enabled or role is None or not permissions.issubset(role.permissions):
            raise PlanningPlanError(1104, "采纳受理人的当前账号或保存权限不可用", 403)
        snapshot = {"id": account.id, "name": account.name, "roleId": role.id, "roleName": role.name}
    return row, run, validator, snapshot, accepted[0]["request_id"]


def _check_interference_saved(session, payload, proof, plan, *, replay=False):
    from .interference_analysis_engine import digest as canonical_digest
    from .system_audit_service import append_system_audit
    from .system_management_models import SystemAuditEvent
    row, run, validator, snapshot, accepted_request_id = proof
    expected_digest = canonical_digest({"operation": "interference.plan.save", "body": _digest(payload),
        "origin": validator.origin(row)})
    if plan.request_id != payload.request_id or plan.request_digest != expected_digest:
        raise PlanningPlanError(1204, "后台保存的原请求摘要不一致", 409)
    validator._plan(session, row, run, plan.plan_id, accepting=True)
    if row.state == "SUCCESS" and row.result_plan_id != plan.plan_id:
        raise PlanningPlanError(1203, "已完成采纳的子方案绑定不一致", 409)
    if snapshot is None:
        return
    changes = {"adoption_id": row.adoption_id, "run_id": row.run_id,
        "accepted_request_id": accepted_request_id, "job_id": plan.job_id,
        "parent_plan_id": run.plan_id, "task_id": plan.task_id, "revision_id": plan.revision_id}
    if replay:
        events = session.scalars(select(SystemAuditEvent).where(
            SystemAuditEvent.action == "interference.plan.save", SystemAuditEvent.request_id == payload.request_id,
            SystemAuditEvent.object_id == plan.plan_id)).all()
        if (len(events) != 1 or events[0].actor != snapshot or events[0].changes != changes
                or events[0].module != "interference" or events[0].outcome != "SUCCESS"):
            raise PlanningPlanError(5002, "后台保存的原子审计证明不一致", 503)
    else:
        # 独立后台动作不冒充当前登录请求的plan.save归属证明；审计与方案同事务。
        append_system_audit(session, None, now=int(datetime.now(timezone.utc).timestamp()),
            action="interference.plan.save", module="interference", snapshot=snapshot,
            object_id=plan.plan_id, request_id=payload.request_id, changes=changes,
            summary="已受理干扰采纳的子方案已保存；未自动转正或下发")


def save_planning_plan(payload: PlanningPlanSaveRequest,
                       factory: sessionmaker[Session] | None = None, *,
                       audit_context: LocalBusinessContext | None = None,
                       _interference_request: dict | None = None, _runtime_root=None
                       ) -> dict[str, Any]:
    factory = factory or get_session_factory()
    worker = _interference_request is not None
    if worker:
        if type(_interference_request) is not dict or audit_context is not None:
            raise PlanningPlanError(1203, "后台采纳授权不能混入交互会话", 409)
        planning_job_service.validate_interference_runtime(factory, _runtime_root)
        owner = None
    else:
        owner = _business_context(audit_context)
    digest = _digest(payload)
    if worker:
        from .interference_analysis_engine import digest as canonical_digest
        digest = canonical_digest({"operation": "interference.plan.save", "body": digest,
            "origin": _interference_request.get("interference_origin")})
    if owner is not None:
        digest = owner.digest(digest, operation="plan.save")
    with factory() as session, session.begin():
        proof = _interference_save_proof(session, payload, _interference_request) if worker else None
        if owner is not None:
            owner.prepare(session, "planning.save")
        replay = _replay(session, payload.request_id, digest, owner)
        if replay is not None:
            if worker:
                _check_interference_saved(session, payload, proof, replay, replay=True)
            return _plan_dto(session, replay, same_transaction_resource=worker)

    job_data = _job_result(payload, factory)
    raw_snapshot = planning_adjustment_service.capture_raw_snapshot(job_data["result"])
    now = datetime.now(timezone.utc)
    try:
        with _SAVE_LOCK, factory() as session, session.begin():
            if worker:
                # 该后台入口已限定SQLite隔离库；冻结授权复核至方案/审计一起提交。
                session.connection().exec_driver_sql("BEGIN IMMEDIATE")
            proof = _interference_save_proof(session, payload, _interference_request) if worker else None
            prepared = owner.prepare(session, "planning.save") if owner is not None else None
            replay = _replay(session, payload.request_id, digest, owner)
            if replay is not None:
                if worker:
                    _check_interference_saved(session, payload, proof, replay, replay=True)
                return _plan_dto(session, replay, same_transaction_resource=worker)
            existing = session.scalar(select(PlanningPlan).where(
                PlanningPlan.job_id == payload.job_id
            ))
            if existing is not None:
                raise PlanningPlanError(1204, "该作业结果已经保存", 409, {
                    "plan_id": existing.plan_id,
                })
            job = session.get(PlanningJob, payload.job_id)
            if job is None or job.task_id != payload.task_id:
                raise PlanningPlanError(2001, "规划作业不存在", 404)
            version_no = (session.scalar(select(func.max(PlanningPlan.version_no)).where(
                PlanningPlan.task_id == payload.task_id
            )) or 0) + 1
            outcome = job_data["outcome"]
            state = ("SAVED" if outcome.get("code") == 0
                     and not job_data.get("source_changed") else "REFERENCE")
            plan = PlanningPlan(
                plan_id="PLAN-" + uuid.uuid4().hex,
                request_id=payload.request_id,
                request_digest=digest,
                task_id=payload.task_id,
                job_id=payload.job_id,
                revision_id=job.revision_id,
                condition_version=job.condition_version,
                resource_revision=job.resource_revision,
                parent_plan_id=job.parent_plan_id,
                adjustment=job.adjustment,
                version_no=version_no,
                name=payload.name,
                note=payload.note,
                state=state,
                input_snapshot=job_data["input"],
                result_snapshot=job_data["result"],
                raw_snapshot=raw_snapshot,
                outcome_snapshot=outcome,
                source="SR4",
                saved_at=now,
                confirmed_at=None,
            )
            session.add(plan)
            session.flush()
            if worker:
                _check_interference_saved(session, payload, proof, plan)
            if owner is not None:
                _append_owner(owner, session, prepared, operation="plan.save", request_id=payload.request_id,
                    object_id=plan.plan_id, page_context="planning", task_id=plan.task_id, revision_id=plan.revision_id)
            return _plan_dto(session, plan, same_transaction_resource=worker)
    except IntegrityError:
        with factory() as session, session.begin():
            proof = _interference_save_proof(session, payload, _interference_request) if worker else None
            if owner is not None:
                owner.prepare(session, "planning.save")
            replay = _replay(session, payload.request_id, digest, owner)
            if replay is not None:
                if worker:
                    _check_interference_saved(session, payload, proof, replay, replay=True)
                return _plan_dto(session, replay, same_transaction_resource=worker)
            if owner is not None:
                raise
            existing = session.scalar(select(PlanningPlan).where(
                PlanningPlan.job_id == payload.job_id
            ))
            if existing is not None:
                raise PlanningPlanError(1204, "该作业结果已经保存", 409, {
                    "plan_id": existing.plan_id,
                }) from None
        raise


def list_saved_plan_history(page: int, page_size: int, search: str | None = None,
                            state: str | None = None,
                            factory: sessionmaker[Session] | None = None
                            ) -> dict[str, Any]:
    """只读已保存方案摘要，不加载完整输入、结果或旧演示文件。"""
    if page < 1 or page_size not in (5, 10, 20, 50):
        raise PlanningPlanError(1001, "分页参数超出范围", 422)
    if state is not None and state not in ("SAVED", "REFERENCE", "CONFIRMED"):
        raise PlanningPlanError(1001, "方案状态取值非法", 422)
    if search is not None and (len(search) > 120 or any(
            ord(char) < 32 or ord(char) == 127 for char in search)):
        raise PlanningPlanError(1001, "搜索条件长度或字符非法", 422)
    filters = []
    if search and search.strip():
        filters.append(or_(*(column.contains(search.strip(), autoescape=True)
                             for column in (PlanningPlan.name, PlanningPlan.plan_id,
                                            PlanningPlan.task_id, PlanningPlan.job_id))))
    if state is not None:
        filters.append(PlanningPlan.state == state)
    factory = factory or get_session_factory()
    with factory() as session:
        total = session.scalar(select(func.count()).select_from(PlanningPlan).where(*filters)) or 0
        # 一次联查当前指针；分页不读取 JSON 快照，也不逐条访问资源文件。
        rows = session.execute(select(
            PlanningPlan.plan_id, PlanningPlan.task_id, PlanningPlan.name,
            PlanningPlan.version_no, PlanningPlan.state, PlanningPlan.resource_revision,
            PlanningPlan.saved_at, PlanningCurrentPlan.plan_id.label("current_plan_id"),
        ).outerjoin(PlanningCurrentPlan, PlanningCurrentPlan.task_id == PlanningPlan.task_id)
          .where(*filters).order_by(PlanningPlan.saved_at.desc(), PlanningPlan.plan_id.desc())
          .offset((page - 1) * page_size).limit(page_size)).all()
        revision = frontend_integration.resource_revision()
        items = []
        for row in rows:
            changed = row.resource_revision != revision
            items.append({
                "plan_id": row.plan_id, "task_id": row.task_id, "name": row.name,
                "version": f"V{row.version_no}", "state": row.state,
                "source_changed": changed, "saved_at": _timestamp(row.saved_at),
                "current": row.state == "CONFIRMED" and row.current_plan_id == row.plan_id and not changed,
            })
        return {"items": items, "total": total, "page": page, "page_size": page_size,
                "data_origin": "SYNTHETIC_INTEGRATION"}


def list_planning_plans(task_id: str, page: int, page_size: int,
                        search: str | None = None,
                        factory: sessionmaker[Session] | None = None
                        ) -> dict[str, Any]:
    if page < 1 or page_size < 1 or page_size > 100:
        raise PlanningPlanError(1001, "分页参数超出范围", 422)
    factory = factory or get_session_factory()
    with factory() as session:
        filters = [PlanningPlan.task_id == task_id]
        if search:
            filters.append(PlanningPlan.name.contains(search.strip(), autoescape=True))
        total = session.scalar(select(func.count()).select_from(PlanningPlan).where(
            *filters
        )) or 0
        rows = session.scalars(select(PlanningPlan).where(*filters).order_by(
            PlanningPlan.saved_at.desc(), PlanningPlan.plan_id.desc()
        ).offset((page - 1) * page_size).limit(page_size)).all()
        current_plan_id = _current_plan_id(session, task_id)
        return {
            "items": [_dto(row, current_plan_id) for row in rows],
            "total": total,
            "page": page,
            "page_size": page_size,
        }


def get_planning_plan(plan_id: str, task_id: str,
                      factory: sessionmaker[Session] | None = None
                      ) -> dict[str, Any]:
    factory = factory or get_session_factory()
    with factory() as session:
        plan = session.get(PlanningPlan, plan_id)
        if plan is None or plan.task_id != task_id:
            raise PlanningPlanError(2001, "规划方案不存在", 404)
        return _plan_dto(session, plan)


def compare_planning_plans(payload: PlanningPlanCompareRequest,
                           factory: sessionmaker[Session] | None = None
                           ) -> list[dict[str, Any]]:
    """按请求顺序返回同一任务、同一冻结条件下的完整方案。"""
    factory = factory or get_session_factory()
    with factory() as session:
        rows = session.scalars(select(PlanningPlan).where(
            PlanningPlan.plan_id.in_(payload.solution_ids),
            PlanningPlan.task_id == payload.task_id,
        )).all()
        by_id = {row.plan_id: row for row in rows}
        missing = [plan_id for plan_id in payload.solution_ids if plan_id not in by_id]
        if missing:
            raise PlanningPlanError(
                2001, "对比方案不存在或不属于当前任务", 404,
                {"solution_ids": missing},
            )

        ordered = [by_id[plan_id] for plan_id in payload.solution_ids]
        baseline = ordered[0]
        incompatible = [
            row.plan_id for row in ordered[1:]
            if row.condition_version != baseline.condition_version
            or row.resource_revision != baseline.resource_revision
        ]
        if incompatible:
            raise PlanningPlanError(
                1001, "方案冻结的任务条件或资源版本不同，不能直接对比", 409,
                {"solution_ids": incompatible,
                 "suggestions": ["请选择同一任务条件和资源版本生成的方案"]},
            )

        current_plan_id = _current_plan_id(session, payload.task_id)
        return [_dto(row, current_plan_id) for row in ordered]


def _excel_value(value: Any) -> str | int | float | bool | None:
    if value is None or isinstance(value, (int, float, bool)):
        return value
    if not isinstance(value, str):
        value = json.dumps(value, ensure_ascii=False, default=str)
    return ILLEGAL_CHARACTERS_RE.sub("", value)


def _append_records(book: Workbook, title: str,
                    records: list[dict[str, Any]]) -> None:
    sheet = book.create_sheet(title)
    if not records:
        sheet.append(["无数据"])
        return
    headers = list(dict.fromkeys(
        key for record in records for key in record.keys()
    ))
    sheet.append(headers)
    for record in records:
        sheet.append([_excel_value(record.get(header)) for header in headers])


def _plan_workbook(plan: dict[str, Any]) -> bytes:
    """把数据库冻结方案展开为可复核的工作簿。"""
    book = Workbook()
    summary = book.active
    summary.title = "方案概要"
    summary.append(["字段", "值"])
    for key in ("plan_id", "task_id", "job_id", "version", "name", "note",
                "state", "current", "source_changed", "source", "saved_at",
                "confirmed_at", "condition_version", "resource_revision"):
        summary.append([key, _excel_value(plan.get(key))])

    task_input = plan["input"]
    summary_input = {key: value for key, value in task_input.items()
                     if key not in {"demands", "deploymentConstraints"}}
    _append_records(book, "任务输入", [summary_input])
    _append_records(book, "通联需求", task_input.get("demands") or [])
    deployment_constraints = task_input.get("deploymentConstraints")
    _append_records(book, "部署约束", [deployment_constraints]
                    if isinstance(deployment_constraints, dict) else [])

    result = plan["result"]
    configuration = result.get("configuration", {})
    _append_records(book, "节点", result.get("nodes") or [])
    _append_records(book, "设备", result.get("devices") or [])
    _append_records(book, "链路", result.get("links") or [])
    _append_records(book, "部署结果", configuration.get("deployment") or [])
    _append_records(book, "设备参数", configuration.get("devices") or [])
    _append_records(book, "路由", configuration.get("routes") or [])
    _append_records(book, "频率", configuration.get("frequencies") or [])
    _append_records(book, "指标", [{
        **result.get("metrics", {}),
        "deployment_cost": result.get("deployment_cost"),
        "cost_unit": result.get("cost_unit"),
    }])
    _append_records(book, "检查项", result.get("checks") or [])

    snapshot = book.create_sheet("完整快照")
    snapshot.append(["顺序", "JSON片段"])
    raw = json.dumps(plan, ensure_ascii=False, default=str)
    for index, start in enumerate(range(0, len(raw), 16000), start=1):
        snapshot.append([index, _excel_value(raw[start:start + 16000])])

    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


def export_planning_plan(plan_id: str, payload: PlanningPlanExportRequest,
                         factory: sessionmaker[Session] | None = None
                         ) -> dict[str, Any] | bytes:
    """读取数据库冻结方案并返回完整 DTO 或 XLSX。"""
    plan = get_planning_plan(plan_id, payload.task_id, factory)
    if payload.format == "json":
        return plan
    return _plan_workbook(plan)


def _confirmation_digest(plan_id: str,
                         payload: PlanningPlanConfirmRequest) -> str:
    body = payload.model_dump(mode="json")
    body["plan_id"] = plan_id
    raw = json.dumps(body, ensure_ascii=False, sort_keys=True,
                     separators=(",", ":")).encode("utf-8")
    return hashlib.sha256(raw).hexdigest()


def _confirmation_replay(session: Session, request_id: str,
                         digest: str, audit_context=None) -> PlanningPlan | None:
    receipt = session.get(PlanningPlanConfirmation, request_id)
    if receipt is None:
        return None
    if receipt.request_digest != digest:
        raise PlanningPlanError(1204, "同一请求编号对应不同内容", 409)
    plan = session.get(PlanningPlan, receipt.plan_id)
    if plan is None:
        raise PlanningPlanError(5001, "确认记录引用的方案不存在", 500)
    if audit_context is not None:
        _require_owner(session, audit_context, operation="plan.confirm", request_id=request_id,
            object_id=receipt.plan_id, page_context="planning", task_id=receipt.task_id, revision_id=receipt.revision_id)
    return plan


def confirm_planning_plan(plan_id: str, payload: PlanningPlanConfirmRequest,
                          factory: sessionmaker[Session] | None = None, *,
                          audit_context: LocalBusinessContext | None = None
                          ) -> dict[str, Any]:
    """确认达标且来源未变化的方案，并切换任务的当前方案指针。"""
    owner = _business_context(audit_context)
    factory = factory or get_session_factory()
    digest = _confirmation_digest(plan_id, payload)
    if owner is not None:
        digest = owner.digest(digest, operation="plan.confirm")
    try:
        with _CONFIRM_LOCK, factory() as session, session.begin():
            prepared = owner.prepare(session, "planning.confirm") if owner is not None else None
            replay = _confirmation_replay(session, payload.request_id, digest, owner)
            if replay is not None:
                return _plan_dto(session, replay)
            plan = session.get(PlanningPlan, plan_id)
            if plan is None or plan.task_id != payload.task_id:
                raise PlanningPlanError(2001, "规划方案不存在", 404)
            if plan.revision_id != payload.expected_revision:
                raise PlanningPlanError(1203, "方案输入版本与确认请求不一致", 409, {
                    "expected": payload.expected_revision,
                    "actual": plan.revision_id,
                    "suggestions": ["重新打开方案"],
                })
            if _source_changed(plan):
                raise PlanningPlanError(1205, "方案资源版本已经变化，不能确认", 409, {
                    "expected": plan.resource_revision,
                    "actual": frontend_integration.resource_revision(),
                    "suggestions": ["按当前资源重新运行规划"],
                })
            current_plan_id = _current_plan_id(session, payload.task_id)
            if plan.state == "REFERENCE":
                raise PlanningPlanError(3001, "参考方案未满足全部硬约束，不能确认", 409, {
                    "plan_id": plan.plan_id,
                    "gap": plan.outcome_snapshot.get("gap", []),
                })
            if plan.state == "CONFIRMED":
                if current_plan_id != plan.plan_id:
                    raise PlanningPlanError(3001, "历史确认方案不能重新设为当前方案", 409)
                session.add(PlanningPlanConfirmation(
                    request_id=payload.request_id,
                    request_digest=digest,
                    plan_id=plan.plan_id,
                    task_id=plan.task_id,
                    revision_id=plan.revision_id,
                    created_at=datetime.now(timezone.utc),
                ))
                session.flush()
                if owner is not None:
                    _append_owner(owner, session, prepared, operation="plan.confirm", request_id=payload.request_id,
                        object_id=plan.plan_id, page_context="planning", task_id=plan.task_id,
                        revision_id=plan.revision_id, changed=False)
                return _plan_dto(session, plan)
            if plan.state != "SAVED":
                raise PlanningPlanError(1001, "方案状态不允许确认", 409)

            now = datetime.now(timezone.utc)
            plan.state = "CONFIRMED"
            plan.confirmed_at = now
            pointer = session.get(PlanningCurrentPlan, plan.task_id)
            if pointer is None:
                session.add(PlanningCurrentPlan(
                    task_id=plan.task_id, plan_id=plan.plan_id, updated_at=now
                ))
            else:
                pointer.plan_id = plan.plan_id
                pointer.updated_at = now
            session.add(PlanningPlanConfirmation(
                request_id=payload.request_id,
                request_digest=digest,
                plan_id=plan.plan_id,
                task_id=plan.task_id,
                revision_id=plan.revision_id,
                created_at=now,
            ))
            session.flush()
            if owner is not None:
                _append_owner(owner, session, prepared, operation="plan.confirm", request_id=payload.request_id,
                    object_id=plan.plan_id, page_context="planning", task_id=plan.task_id, revision_id=plan.revision_id)
            return _plan_dto(session, plan)
    except IntegrityError:
        with factory() as session, session.begin():
            if owner is not None:
                owner.prepare(session, "planning.confirm")
            replay = _confirmation_replay(session, payload.request_id, digest, owner)
            if replay is not None:
                return _plan_dto(session, replay)
        raise
