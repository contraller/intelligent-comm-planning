"""完整任务文件的纯解析与模板；不访问数据库、资源目录或业务文件。"""
from __future__ import annotations

from copy import deepcopy
import csv
from datetime import date, datetime, time, timedelta
from io import BytesIO, StringIO
import json
import math
import posixpath
import re
import stat
import xml.etree.ElementTree as ET
from zipfile import BadZipFile, ZIP_DEFLATED, ZIP_STORED, ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.drawing.spreadsheet_drawing import SpreadsheetDrawing
from openpyxl.worksheet._writer import WorksheetWriter
from openpyxl.writer.excel import ExcelWriter
from pydantic import ValidationError

from .task_input_schemas import DeploymentConstraints, DeploymentSolve, FrequencyConstraints

MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_DEMANDS = 5000  # 与新版 config/demo.ts 的 TASK_MAX_DEMANDS 一致。
MAX_ZIP_BYTES = 32 * 1024 * 1024
MAX_ZIP_MEMBER_BYTES = 16 * 1024 * 1024
MAX_ZIP_MEMBERS = 256
MAX_ISSUES = 1000
MAX_XLSX_CELL_CHARS = 32767
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
KINDS = ("task_scenarios", "task_links")
DEPLOYMENT_HEADERS = (
    "max_slope_deg", "min_elevation_m", "max_elevation_m", "allowed_landcover",
    "road_max_distance_m", "min_site_spacing_m", "los_required_node_ids", "grid_step_m",
    "forced_include_json", "forced_exclude_polygons_json",
)
OPTION_HEADERS = (
    "hf_min_khz", "hf_max_khz", "hf_exclude_freq_ids", "vuhf_min_khz", "vuhf_max_khz",
    "vuhf_exclude_freq_ids", "deployment_solve_json",
)
HEADERS = {
    "task_scenarios": (
        "task_id", "task_name", "area_bbox", "start_time", "duration_h", "description",
        "station_ids", "preset_id", "frequency_id", *DEPLOYMENT_HEADERS, *OPTION_HEADERS,
    ),
    "task_links": (
        "demand_id", "task_id", "src_node_id", "dst_node_id", "priority", "is_mandatory",
        "data_volume_mb", "required_rate_kbps", "max_delay_ms", "min_reliability", "service_type",
        "preferred_class", "endpoint_mode", "src_device_id", "dst_device_id", "strategy", "need_backup",
    ),
}
REQUIRED = {"task_scenarios": HEADERS["task_scenarios"][:6], "task_links": HEADERS["task_links"][:12]}
OPTION_KEYS = ("deploymentConstraints", "frequencyConstraints", "deploymentSolve")
NUMBER = re.compile(r"[+-]?(?:[0-9]+(?:\.[0-9]*)?|\.[0-9]+)(?:[eE][+-]?[0-9]+)?\Z")
CELL_REF = re.compile(r"([A-Z]+)([1-9][0-9]*)\Z")
CONTROL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")


class TaskFileError(ValueError):
    def __init__(self, message: str, code: str = "IMPORT_FORMAT"):
        super().__init__(message)
        self.message = message
        self.code = code


def _json_copy(value):
    try:
        encoded = json.dumps(value, ensure_ascii=False, allow_nan=False)
        encoded.encode("utf-8")
        return json.loads(encoded)
    except (TypeError, ValueError, OverflowError, RecursionError) as exc:
        raise TaskFileError("输入须为有限数值和JSON结构，不接受日期对象或NaN") from exc


def _json_read(value):
    def unique(pairs):
        result = {}
        for key, item in pairs:
            if key in result:
                raise ValueError("JSON字段重复")
            result[key] = item
        return result

    def nonfinite(_value):
        raise ValueError("JSON不接受NaN或Infinity")

    return _json_copy(json.loads(value, object_pairs_hook=unique, parse_constant=nonfinite))


def _identifier(value):
    return isinstance(value, str) and 0 < len(value) <= 120 and not re.search(r"\s|[\x00-\x1f\x7f]", value)


def _add(issues, sheet, row, field, message, object_id=None):
    if len(issues) >= MAX_ISSUES:
        raise TaskFileError("文件错误超过1000项，请先修正文件结构", "IMPORT_LIMIT")
    issues.append({"objectId": object_id if object_id is not None else sheet, "field": field,
                   "message": message, "sheet": sheet, "row": row})


def _scalar(value):
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    return str(value).strip()


def _unescape(value):
    if not value.startswith("~rp~"):
        return value
    try:
        decoded = _json_read(value[4:])
        if not isinstance(decoded, str):
            raise ValueError("转义不是字符串")
        return decoded
    except (ValueError, TypeError, RecursionError) as exc:
        raise TaskFileError("CSV ~rp~ 文本转义无效") from exc


def _add_table(tables, kind, headers, rows, issues):
    if kind in tables:
        raise TaskFileError("业务工作表重复：" + kind)
    tables[kind] = (headers, rows)
    if (len(set(headers)) != len(headers) or set(headers) - set(HEADERS[kind])
            or set(REQUIRED[kind]) - set(headers)):
        _add(issues, kind, 1, "header", "表头缺失、未知或重复，请使用模板")
    for line, values in rows:
        if len(values) > len(headers):
            _add(issues, kind, line, "columns", "数据列多于表头")


def _csv_table(data):
    try:
        reader = csv.reader(StringIO(data.decode("utf-8-sig"), newline=""), strict=True)
        headers = [_scalar(value) for value in next(reader, [])]
        rows = []
        while True:
            line = reader.line_num + 1
            raw = next(reader, None)
            if raw is None:
                break
            values = [_unescape(_scalar(value)) for value in raw]
            if any(value != "" for value in values):
                rows.append((line, values))
                if len(rows) > MAX_DEMANDS:
                    raise TaskFileError("业务记录最多5000行", "IMPORT_LIMIT")
        return headers, rows
    except UnicodeError as exc:
        raise TaskFileError("CSV必须是UTF-8编码，可带BOM", "IMPORT_ENCODING") from exc
    except csv.Error as exc:
        code = "IMPORT_LIMIT" if "field larger than field limit" in str(exc) else "IMPORT_FORMAT"
        raise TaskFileError("CSV结构损坏或字段超过解析上限", code) from exc


def _local(tag):
    return tag.rsplit("}", 1)[-1]


def _inspect_archive(archive, issues):
    members = archive.infolist()
    if len(members) > MAX_ZIP_MEMBERS or sum(item.file_size for item in members) > MAX_ZIP_BYTES:
        raise TaskFileError("XLSX ZIP成员数或展开总量超限", "IMPORT_LIMIT")
    seen, roots = set(), {}
    for item in members:
        # Windows会规范化反斜杠，ZipInfo也可能截断NUL；先核对原始成员名。
        name = item.orig_filename
        normalized = posixpath.normpath(name).casefold()
        if (name != item.filename or not name or normalized in seen or name.startswith(("/", "\\")) or "\\" in name or ":" in name
                or ".." in name.split("/") or re.search(r"[\x00-\x1f\x7f]", name) or normalized in {".", ".."}
                or stat.S_ISLNK(item.external_attr >> 16) or item.flag_bits & 1
                or item.compress_type not in {ZIP_STORED, ZIP_DEFLATED}):
            raise TaskFileError("XLSX ZIP路径、重复成员、加密或成员类型不安全")
        seen.add(normalized)
        if item.file_size > MAX_ZIP_MEMBER_BYTES or item.file_size > max(1, item.compress_size) * 1000:
            raise TaskFileError("XLSX ZIP成员展开体积或比例超限", "IMPORT_LIMIT")
        if (any(part in normalized for part in ("vbaproject", "macrosheets/", "embeddings/", "activex/", "externallinks/"))
                or normalized.endswith(".bin")):
            raise TaskFileError("不接受宏、嵌入对象或外部链接")
    for item in members:
        name = item.filename.casefold()
        if not name.endswith((".xml", ".rels")):
            continue
        raw = archive.read(item)
        if re.search(br"<!\s*(?:DOCTYPE|ENTITY)", raw, re.I) or b"\x00" in raw:
            raise TaskFileError("不接受XML实体声明或非UTF-8 XML")
        roots[name] = ET.fromstring(raw.decode("utf-8-sig"))
        for element in roots[name].iter():
            tag = _local(element.tag)
            if tag == "Relationship" and (element.get("TargetMode", "").casefold() == "external"
                                          or re.match(r"^[a-zA-Z][a-zA-Z0-9+.-]*:", element.get("Target", ""))):
                raise TaskFileError("不接受外部链接")
            if name == "[content_types].xml" and any(token in str(element.attrib).casefold() for token in ("macroenabled", "vba")):
                raise TaskFileError("不接受含宏工作簿")
    if not {"xl/workbook.xml", "xl/_rels/workbook.xml.rels", "[content_types].xml"} <= roots.keys():
        raise TaskFileError("XLSX缺少必要结构")
    targets = {}
    for element in roots["xl/_rels/workbook.xml.rels"]:
        target = element.get("Target", "")
        targets[element.get("Id")] = (target.lstrip("/") if target.startswith("/")
                                       else posixpath.normpath(posixpath.join("xl", target))).casefold()
    sheet_names = {}
    for element in roots["xl/workbook.xml"].iter():
        if _local(element.tag) != "sheet":
            continue
        title = element.get("name", "")
        if title not in {*KINDS, "说明"} or title in sheet_names.values() or element.get("state", "visible") != "visible":
            raise TaskFileError("未知、重复或隐藏工作表：" + title)
        path = targets.get(element.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"))
        if not path or not path.startswith("xl/worksheets/") or path not in roots or path in sheet_names:
            raise TaskFileError("业务工作表引用无效或重复")
        sheet_names[path] = title
    if not set(KINDS) <= set(sheet_names.values()):
        raise TaskFileError("必须同时包含task_scenarios和task_links")
    for path, root in roots.items():
        if not path.startswith("xl/worksheets/"):
            continue
        if path not in sheet_names:
            raise TaskFileError("不接受未登记的工作表数据")
        kind = sheet_names[path]
        last_row, refs = 0, set()
        for element in root.iter():
            tag = _local(element.tag)
            if tag == "row":
                line = int(element.get("r", "0"))
                if not 1 <= line <= MAX_DEMANDS + 1:
                    raise TaskFileError("XLSX实际行号超限", "IMPORT_LIMIT")
                if line <= last_row:
                    raise TaskFileError("XLSX行号重复或次序无效")
                last_row = line
                if element.get("hidden", "0").casefold() in {"1", "true"}:
                    _add(issues, kind, line, "row", "不支持隐藏数据行")
            if tag == "col" and element.get("hidden", "0").casefold() in {"1", "true"}:
                _add(issues, kind, 1, "column", "不支持隐藏业务列")
            if tag == "sheetFormatPr" and element.get("zeroHeight", "0").casefold() in {"1", "true"}:
                _add(issues, kind, 1, "row", "不支持默认隐藏数据行")
            if tag in {"mergeCell", "oleObject"}:
                raise TaskFileError("不支持合并单元格或嵌入对象")
            if tag != "c":
                continue
            ref = element.get("r", "")
            match = CELL_REF.fullmatch(ref)
            if not match or ref in refs or int(match[2]) != last_row:
                raise TaskFileError("XLSX单元格坐标重复或与所在行不符")
            refs.add(ref)
            column = 0
            for char in match[1]:
                column = column * 26 + ord(char) - ord("A") + 1
            if column > len(HEADERS.get(kind, HEADERS["task_scenarios"])):
                raise TaskFileError("XLSX实际列超出任务模板限制", "IMPORT_LIMIT")


def _xlsx_tables(data, tables, issues):
    try:
        with ZipFile(BytesIO(data)) as archive:
            _inspect_archive(archive, issues)
        book = load_workbook(BytesIO(data), read_only=True, data_only=False, keep_links=False)
        try:
            for sheet in book:
                kind = sheet.title
                if not hasattr(sheet, "reset_dimensions"):
                    raise TaskFileError("不支持非业务工作表")
                # dimension可能被伪造；真实单元格已经过有界XML扫描，再按实际行读取。
                sheet.reset_dimensions()
                headers, rows = [], []
                for line, cells in enumerate(sheet.iter_rows(), 1):
                    values = []
                    for column, cell in enumerate(cells):
                        raw = cell.value
                        field = headers[column] if column < len(headers) else str(column + 1)
                        if isinstance(raw, str) and len(raw) > MAX_XLSX_CELL_CHARS:
                            raise TaskFileError("XLSX单元格文本超过32767字符，请使用CSV", "IMPORT_LIMIT")
                        if (cell.data_type in {"f", "e"} or isinstance(raw, (date, datetime, time, timedelta))
                                or raw is not None and not isinstance(raw, (str, int, float, bool))):
                            _add(issues, kind, line, field, "只接受文本、有限数值和布尔；拒绝公式、错误及日期对象")
                            raw = None
                        if isinstance(raw, float) and not math.isfinite(raw):
                            _add(issues, kind, line, field, "数值必须有限")
                            raw = None
                        values.append(_scalar(raw))
                    if kind == "说明":
                        continue
                    if line == 1:
                        headers = values
                    elif any(value != "" for value in values):
                        rows.append((line, values))
                if kind != "说明":
                    _add_table(tables, kind, headers, rows, issues)
        finally:
            book.close()
    except TaskFileError:
        raise
    except (BadZipFile, ET.ParseError, UnicodeError, ValueError, TypeError, KeyError, IndexError,
            AttributeError, OSError, RuntimeError, OverflowError, EOFError, NotImplementedError) as exc:
        raise TaskFileError("XLSX文件损坏或格式不受支持") from exc


def _number(value, field, add, *, nullable=False, minimum=None, maximum=None, positive=False):
    if value == "" and nullable:
        return None
    if not NUMBER.fullmatch(value):
        add(field, "须为有限十进制数，空值不能替代必填数值")
        return None
    result = float(value)
    if not math.isfinite(result):
        add(field, "数值必须有限，不接受NaN或Infinity")
        return None
    if (minimum is not None and result < minimum or maximum is not None and result > maximum
            or positive and result <= 0):
        add(field, "数值超出允许范围")
    return result


def _validate_model(model, value, prefix, add):
    if value is None:
        return
    try:
        model.model_validate(value)
    except ValidationError as exc:
        for error in exc.errors(include_url=False):
            location = ".".join(str(part) for part in error["loc"])
            add(prefix + ("." + location if location else ""), error["msg"])


def _extensions(row, add):
    def number(key):
        return _number(row.get(key, ""), key, add, nullable=True)

    def values(key):
        return [value.strip() for value in row.get(key, "").split(";")] if row.get(key, "").strip() else []

    def structured(key, empty):
        if not row.get(key, "").strip():
            return empty
        try:
            return _json_read(row[key])
        except (ValueError, TypeError, RecursionError):
            add(key, "JSON格式无效、字段重复或含非有限数值")
            return empty

    result = {}
    if any(row.get(key, "").strip() for key in DEPLOYMENT_HEADERS):
        low, high = number("min_elevation_m"), number("max_elevation_m")
        deployment = {"maxSlopeDeg": number("max_slope_deg"), "elevationRangeM": None if low is None and high is None else [low, high],
                      "allowedLandcover": values("allowed_landcover") if row.get("allowed_landcover", "").strip() else None,
                      "roadMaxDistanceM": number("road_max_distance_m"), "minSiteSpacingM": number("min_site_spacing_m"),
                      "losRequiredNodeIds": values("los_required_node_ids"), "gridStepM": number("grid_step_m"),
                      "forcedInclude": structured("forced_include_json", []),
                      "forcedExcludePolygons": structured("forced_exclude_polygons_json", [])}
        _validate_model(DeploymentConstraints, deployment, "deploymentConstraints", add)
        result["deploymentConstraints"] = deployment
    filters = []
    for device_class in ("HF", "VUHF"):
        prefix = device_class.lower()
        low, high, excluded = (row.get(prefix + suffix, "") for suffix in ("_min_khz", "_max_khz", "_exclude_freq_ids"))
        if low.strip() or high.strip() or excluded.strip():
            filters.append({"deviceClass": device_class,
                            "rangeKhz": [number(prefix + "_min_khz"), number(prefix + "_max_khz")] if low or high else None,
                            "excludeFreqIds": values(prefix + "_exclude_freq_ids")})
    if filters:
        result["frequencyConstraints"] = {"filters": filters}
        _validate_model(FrequencyConstraints, result["frequencyConstraints"], "frequencyConstraints", add)
    if row.get("deployment_solve_json", "").strip():
        result["deploymentSolve"] = structured("deployment_solve_json", None)
        _validate_model(DeploymentSolve, result["deploymentSolve"], "deploymentSolve", add)
    return result


def _routing(row, add):
    result = {}
    if row.get("strategy"):
        if row["strategy"] not in {"SHORTEST_PATH", "MIN_DELAY", "MAX_RELIABILITY"}:
            add("strategy", "路由策略须为SHORTEST_PATH、MIN_DELAY或MAX_RELIABILITY")
        else:
            result["routeStrategy"] = row["strategy"]
    if row.get("need_backup"):
        if row["need_backup"] not in {"true", "false"}:
            add("need_backup", "备份请求只接受空白、true或false")
        else:
            result["needBackup"] = row["need_backup"] == "true"
    return result


def parse_task_files(files: list[tuple[str, bytes]], defaults: dict) -> dict:
    """解析两表并返回预览；资源引用、版本、签名和持久化交外层服务。"""
    if not isinstance(defaults, dict):
        raise TaskFileError("defaults须为当前任务输入对象")
    before = _json_copy(defaults)
    if "demands" in before and not isinstance(before["demands"], list):
        raise TaskFileError("defaults.demands须为任务需求列表")
    if (not isinstance(files, list) or not 1 <= len(files) <= 2
            or any(not isinstance(file, tuple) or len(file) != 2 or not isinstance(file[0], str)
                   or not isinstance(file[1], bytes) or not file[1] or len(file[1]) > MAX_FILE_BYTES for file in files)):
        raise TaskFileError("请选择一个XLSX或配套两个CSV，单文件最多5MiB", "IMPORT_LIMIT")
    if any(not name or re.search(r"[/\\\x00-\x1f]", name) for name, _data in files):
        raise TaskFileError("文件名不能包含目录或控制字符")
    tables, issues, notes = {}, [], []
    if len(files) == 1 and files[0][0].lower().endswith(".xlsx"):
        _xlsx_tables(files[0][1], tables, issues)
    else:
        if len(files) != 2 or any(not name.lower().endswith(".csv") for name, _data in files):
            raise TaskFileError("CSV需同时选择任务信息和通联需求两个文件")
        for _name, data in files:
            headers, rows = _csv_table(data)
            kind = "task_links" if "demand_id" in headers else "task_scenarios"
            _add_table(tables, kind, headers, rows, issues)
    if not set(KINDS) <= tables.keys():
        raise TaskFileError("必须同时包含task_scenarios和task_links")
    scene_headers, scenes = tables["task_scenarios"]
    link_headers, links = tables["task_links"]
    task_row = scenes[0][0] if scenes else 2
    if len(scenes) != 1:
        _add(issues, "task_scenarios", task_row, "task_id", "每次必须且只能导入一个任务")
    if not 1 <= len(links) <= MAX_DEMANDS:
        _add(issues, "task_links", 2, "rows", "通联需求须为1至5000条")
    row = dict(zip(scene_headers, scenes[0][1] if scenes else []))
    task_id = row.get("task_id", "")

    def add(field, message):
        _add(issues, "task_scenarios", task_row, field, message, task_id or "TASK")

    if not _identifier(task_id):
        add("task_id", "任务编号为空、超长或含空白/控制字符")
    candidate = deepcopy(before)
    for key in OPTION_KEYS:
        candidate.pop(key, None)
    candidate.update(name=row.get("task_name", ""), description=row.get("description", ""), demands=[])
    for key, limit, required in (("name", 60, True), ("description", 500, False)):
        value = candidate[key]
        if len(value) > limit or CONTROL.search(value) or required and not value.strip():
            add("task_name" if key == "name" else key, "文本为空、超长或包含不支持的控制字符")
    bbox = row.get("area_bbox", "").split(";")
    candidate["bbox"] = [_number(part.strip(), "area_bbox", add) for part in bbox]
    if len(bbox) != 4:
        add("area_bbox", "范围须为分号分隔的四个有限经纬度")
    elif all(value is not None for value in candidate["bbox"]):
        west, south, east, north = candidate["bbox"]
        if not (-180 <= west < east <= 180 and -90 <= south < north <= 90):
            add("area_bbox", "范围须为有效WGS84经纬度矩形")
    candidate["startTime"] = row.get("start_time") or None
    if candidate["startTime"] is not None:
        value = candidate["startTime"]
        try:
            if not re.fullmatch(r"\d{4}-\d{2}-\d{2}T\d{2}:\d{2}:\d{2}(?:\.\d+)?(?:Z|[+-]\d{2}:\d{2})", value):
                raise ValueError("缺少时区")
            datetime.fromisoformat(value.replace("Z", "+00:00"))
        except ValueError:
            add("start_time", "时间须为有效日历日期及带时区ISO8601文本")
    candidate["durationH"] = _number(row.get("duration_h", ""), "duration_h", add, nullable=True, positive=True)
    candidate.update(_extensions(row, add))
    for field, key, label in (("preset_id", "presetId", "参数预设"), ("frequency_id", "frequencyId", "频率资源")):
        if row.get(field):
            candidate[key] = row[field]
        else:
            notes.append("未提供" + field + "，沿用当前" + label + "：" + str(candidate.get(key, "")))
        if not _identifier(candidate.get(key)):
            add(field, "缺少有效文件引用或defaults引用")
    source_rows = {}
    for line, values in links:
        demand = dict(zip(link_headers, values))
        demand_id = demand.get("demand_id", "")

        def demand_issue(field, message):
            _add(issues, "task_links", line, field, message, demand_id or "TASK")

        if demand.get("task_id") != task_id:
            demand_issue("task_id", "需求必须属于本次唯一任务")
        for field in ("demand_id", "src_node_id", "dst_node_id"):
            if not _identifier(demand.get(field)):
                demand_issue(field, "编号为空、超长或含空白/控制字符")
        if demand_id in source_rows:
            demand_issue("demand_id", "通联需求编号重复")
        source_rows.setdefault(demand_id, line)
        mandatory = demand.get("is_mandatory", "")
        if mandatory not in {"true", "false"}:
            demand_issue("is_mandatory", "必需标志只接受true或false")
        mode = demand.get("endpoint_mode") or ("device" if demand.get("src_device_id") or demand.get("dst_device_id") else "station")
        for field, allowed in (("priority", {"P1", "P2", "P3"}), ("service_type", {"", "话音", "数据", "短消息"}), ("preferred_class", {"", "HF", "VUHF"})):
            if demand.get(field, "") not in allowed:
                demand_issue(field, "字段值不受支持")
        if mode not in {"station", "device"}:
            demand_issue("endpoint_mode", "端点模式须为station或device")
        parsed = {"id": demand_id, "sourceId": demand.get("src_node_id", ""), "targetId": demand.get("dst_node_id", ""),
                  "priority": demand.get("priority", ""), "mandatory": mandatory == "true" if mandatory in {"true", "false"} else None,
                  "service": demand.get("service_type", ""), "preferredClass": demand.get("preferred_class") or None,
                  "endpointMode": mode, "sourceDeviceId": demand.get("src_device_id") or None, "targetDeviceId": demand.get("dst_device_id") or None,
                  **_routing(demand, demand_issue)}
        for field, key, nullable, low, high, positive in (
            ("data_volume_mb", "dataVolumeMb", True, 0, None, False),
            ("required_rate_kbps", "requiredRateKbps", False, None, None, True),
            ("max_delay_ms", "maxDelayMs", False, None, None, True),
            ("min_reliability", "minReliability", False, 0, 1, False),
        ):
            parsed[key] = _number(demand.get(field, ""), field, demand_issue, nullable=nullable, minimum=low, maximum=high, positive=positive)
        candidate["demands"].append(parsed)
    candidate["stationIds"] = ([value.strip() for value in row["station_ids"].split(";")] if row.get("station_ids")
                               else list(dict.fromkeys(node for demand in candidate["demands"] for node in (demand["sourceId"], demand["targetId"]))))
    if not row.get("station_ids"):
        notes.append("参与节点由通联需求两端ID合并得到；没有额外补入当前草稿节点。")
    if len(set(candidate["stationIds"])) != len(candidate["stationIds"]) or any(not _identifier(value) for value in candidate["stationIds"]):
        add("station_ids", "参与节点编号无效或重复")
    for key in OPTION_KEYS:
        if before.get(key) is not None and candidate.get(key) is None:
            notes.append("文件未设置" + key + "，确认应用后将清空原设置；尚未保存。")
    old_demands = {item.get("id"): item for item in before.get("demands", []) if isinstance(item, dict)}
    for demand in candidate["demands"]:
        old = old_demands.get(demand["id"], {})
        if any(old.get(key) is not None and demand.get(key) is None for key in ("routeStrategy", "needBackup")):
            notes.append("需求" + demand["id"] + "的路由条件有未设置项，确认应用后将清空对应原设置；尚未保存。")
    differences = []
    for key in dict.fromkeys([*candidate, *before]):
        old, new = before.get(key), candidate.get(key)
        if key not in before or key not in candidate or json.dumps(old, sort_keys=True, ensure_ascii=False) != json.dumps(new, sort_keys=True, ensure_ascii=False):
            differences.append({"field": key, "before": old, "after": new})
    return {"taskId": task_id, "input": None if issues else candidate, "issues": issues, "notes": notes,
            "differences": differences, "taskRow": task_row, "sourceRows": source_rows}


class _MemoryExcelWriter(ExcelWriter):
    def write_worksheet(self, sheet):
        # openpyxl默认保存会创建临时文件；沿用仓库模式，仅替换XML输出流为内存。
        sheet._drawing = SpreadsheetDrawing()
        output = BytesIO()
        writer = WorksheetWriter(sheet, out=output)
        writer.write()
        sheet._rels = writer._rels
        self._archive.writestr(sheet.path[1:], output.getvalue())
        self.manifest.append(sheet)


def build_task_template(format: str, kind: str, input: dict) -> dict:
    """只从调用方当前输入生成模板；不选择预设、不提供示例资源或写业务文件。"""
    if format not in {"csv", "xlsx"} or kind not in KINDS or not isinstance(input, dict):
        raise TaskFileError("模板只支持task_scenarios/task_links的CSV或XLSX")
    value = _json_copy(input)
    problems = []
    add = lambda field, message: problems.append((field, message))
    for key, model in (("deploymentConstraints", DeploymentConstraints), ("frequencyConstraints", FrequencyConstraints), ("deploymentSolve", DeploymentSolve)):
        _validate_model(model, value.get(key), key, add)
    if not isinstance(value.get("demands"), list) or not 1 <= len(value["demands"]) <= MAX_DEMANDS:
        raise TaskFileError("模板通联需求须为1至5000条", "IMPORT_LIMIT")
    if (not isinstance(value.get("bbox"), list) or len(value["bbox"]) != 4
            or any(type(number) not in {int, float} for number in value["bbox"])
            or not isinstance(value.get("stationIds"), list) or not all(_identifier(item) for item in value["stationIds"])
            or any(not isinstance(value.get(key), str) for key in ("name", "description", "presetId", "frequencyId"))):
        raise TaskFileError("模板须由当前完整任务输入生成，不补造缺失引用")
    if (not value["name"].strip() or len(value["name"]) > 60 or len(value["description"]) > 500
            or CONTROL.search(value["name"]) or CONTROL.search(value["description"])):
        raise TaskFileError("模板名称或说明为空、超长或包含不支持的控制字符")
    for demand in value["demands"]:
        if not isinstance(demand, dict) or not _identifier(demand.get("id")) or type(demand.get("mandatory")) is not bool:
            raise TaskFileError("模板通联需求结构无效")
        _routing({"strategy": demand.get("routeStrategy") or "",
                  "need_backup": "" if demand.get("needBackup") is None else _scalar(demand["needBackup"])}, add)
        if demand.get("needBackup") is not None and type(demand["needBackup"]) is not bool:
            problems.append(("needBackup", "备份请求须为布尔或未指定"))
    if problems:
        raise TaskFileError("请先修正规划条件再下载模板：" + problems[0][0] + "：" + problems[0][1])
    deployment = value.get("deploymentConstraints")
    deployment_values = [None] * len(DEPLOYMENT_HEADERS) if deployment is None else [
        deployment["maxSlopeDeg"], (deployment["elevationRangeM"] or [None, None])[0], (deployment["elevationRangeM"] or [None, None])[1],
        ";".join(deployment["allowedLandcover"]) if deployment["allowedLandcover"] is not None else None,
        deployment["roadMaxDistanceM"], deployment["minSiteSpacingM"], ";".join(deployment["losRequiredNodeIds"]), deployment["gridStepM"],
        json.dumps(deployment["forcedInclude"], ensure_ascii=False), json.dumps(deployment["forcedExcludePolygons"], ensure_ascii=False),
    ]
    option_values = []
    filters = (value.get("frequencyConstraints") or {}).get("filters", [])
    for device_class in ("HF", "VUHF"):
        selected = next((row for row in filters if row["deviceClass"] == device_class), None)
        option_values.extend([*(selected["rangeKhz"] or [None, None]), ";".join(selected["excludeFreqIds"])] if selected else [None, None, None])
    option_values.append(json.dumps(value["deploymentSolve"], ensure_ascii=False) if value.get("deploymentSolve") is not None else None)
    rows = {"task_scenarios": [["TASK-NEW", value["name"], ";".join(map(str, value["bbox"])), value.get("startTime"), value.get("durationH"),
                                value["description"], ";".join(value["stationIds"]), value["presetId"], value["frequencyId"], *deployment_values, *option_values]],
            "task_links": [[demand["id"], "TASK-NEW", demand.get("sourceId"), demand.get("targetId"), demand.get("priority"), _scalar(demand["mandatory"]),
                            demand.get("dataVolumeMb"), demand.get("requiredRateKbps"), demand.get("maxDelayMs"), demand.get("minReliability"), demand.get("service"),
                            demand.get("preferredClass"), demand.get("endpointMode"), demand.get("sourceDeviceId"), demand.get("targetDeviceId"), demand.get("routeStrategy"),
                            "" if demand.get("needBackup") is None else _scalar(demand["needBackup"])] for demand in value["demands"]]}
    if format == "csv":
        output = StringIO(newline="")
        writer = csv.writer(output)
        writer.writerow(HEADERS[kind])
        for row in rows[kind]:
            writer.writerow(["~rp~" + json.dumps(item, ensure_ascii=False) if isinstance(item, str)
                             and (re.match(r"^\s*[=+\-@]", item) or item.startswith("~rp~")) else item for item in row])
        data, mime, name = output.getvalue().encode("utf-8"), "text/csv;charset=utf-8", kind + ".csv"
    else:
        book = Workbook()
        try:
            for index, table in enumerate(KINDS):
                sheet = book.active if index == 0 else book.create_sheet(table)
                sheet.title = table
                sheet.append(HEADERS[table])
                for row in rows[table]:
                    if any(isinstance(item, str) and len(item) > MAX_XLSX_CELL_CHARS for item in row):
                        raise TaskFileError("XLSX单元格文本超过32767字符，请使用CSV", "IMPORT_LIMIT")
                    sheet.append(row)
                for cells in sheet:
                    for cell in cells:
                        if isinstance(cell.value, str):
                            cell.data_type = "s"
                sheet.freeze_panes = "A2"
            notes = book.create_sheet("说明")
            for row in [("范围", "每批一个任务，预检不保存；任务保存和规划执行分别操作。"),
                        ("来源", "模板仅使用传入任务输入；TASK-NEW不代表已创建任务。"),
                        ("空值", "未知数值留空，不补零；完整导入清空文件未设置的部署、频率及路由条件。"),
                        ("CSV", "UTF-8配套两表；公式前缀和~rp~文本使用~rp~加JSON字符串可逆转义。")]:
                notes.append(row)
            output = BytesIO()
            with ZipFile(output, "w", ZIP_DEFLATED, allowZip64=True) as archive:
                _MemoryExcelWriter(book, archive).save()
            data, mime, name = output.getvalue(), XLSX_MIME, "任务导入模板.xlsx"
        finally:
            book.close()
    if len(data) > MAX_FILE_BYTES:
        raise TaskFileError("生成的单文件超过5MiB", "IMPORT_LIMIT")
    return {"name": name, "mime": mime, "bytes": data}
