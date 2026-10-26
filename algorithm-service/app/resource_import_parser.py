"""资源文件纯解析与模板；不读取原件，不访问数据库，不产生业务文件。"""
from __future__ import annotations

from copy import deepcopy
import csv
from datetime import date, datetime, time, timedelta
from io import BytesIO, StringIO
import json
import math
import posixpath
import re
import unicodedata
import xml.etree.ElementTree as ET
from zipfile import BadZipFile, ZIP_DEFLATED, ZIP_STORED, ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.drawing.spreadsheet_drawing import SpreadsheetDrawing
from openpyxl.worksheet._writer import WorksheetWriter
from openpyxl.writer.excel import ExcelWriter

from .resource_export_service import BOOLEANS, HEADERS, NUMERIC, TABLES, XLSX_MIME
from .resource_geography import valid_bounds

MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_ROWS = 5000
MAX_ZIP_BYTES = 32 * 1024 * 1024
MAX_ZIP_MEMBER_BYTES = 16 * 1024 * 1024
MAX_ZIP_MEMBERS = 256
TABLE_IDS = {
    "node.csv": "node_id", "device.csv": "device_id",
    "device_model.csv": "model_id", "antenna_model.csv": "antenna_id",
    "frequency_resource.csv": "freq_id", "interference_source.csv": "interference_id",
    "task_scenario.csv": "task_id", "comm_demand.csv": "demand_id",
    "link_metric.csv": "record_id", "candidate_site.csv": "site_id",
}
OPTIONAL = {
    "nodes": {"parent_node_id", "elevation_m", "region", "remark"},
    "devices": {"tilt_deg", "azimuth_deg", "work_freq_khz", "install_time", "availability"},
    "interference_sources": {"elevation_m", "antenna_gain_dbi", "pattern_type", "azimuth_deg", "duty_cycle"},
}
NODE_EDITABLE = {"node_name", "lon", "lat", "elevation_m", "region", "remark"}
DEVICE_READONLY = {"node_id", "model_id", "antenna_id", "status", "install_time", "availability"}
STATION_TYPES = {
    "I_FIXED": ("I", "FIXED", {"HF"}),
    "I_MOBILE": ("I", "VEHICLE", {"HF"}),
    "II_FIXED": ("II", "FIXED", {"HF", "VUHF"}),
    "II_MOBILE": ("II", "VEHICLE", {"HF", "VUHF"}),
    "II_NODE": ("II", "VEHICLE", {"HF", "VUHF"}),
    "III_MOBILE": ("III", "VEHICLE", {"HF", "VUHF"}),
    "IV_MOBILE": ("IV", "VEHICLE", {"VUHF"}),
    "VEHICLE_A": ("I", "VEHICLE", {"HF"}),
    "VEHICLE_B": ("I", "VEHICLE", {"HF"}),
    "MANPACK": ("I", "MANPACK", {"HF"}),
}
NUMBER = re.compile(r"[+-]?(?:[0-9]+(?:\.[0-9]*)?|\.[0-9]+)(?:[eE][+-]?[0-9]+)?\Z")
CONTROL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")
CELL_REF = re.compile(r"([A-Z]+)([1-9][0-9]*)\Z")
ISO_TIME = re.compile(r"[0-9]{4}-[0-9]{2}-[0-9]{2}T[0-9]{2}:[0-9]{2}(?::[0-9]{2}(?:\.[0-9]+)?)?(?:Z|[+-][0-9]{2}:[0-9]{2})\Z")


class _FileError(ValueError):
    def __init__(self, message, sheet="file", row=0, field="file", object_id="file"):
        super().__init__(message)
        self.detail = _issue(sheet, row, object_id, field, message)


def _diagnostic_label(value, limit):
    raw = str(value).encode("utf-16-le", errors="surrogatepass")
    original_length = len(raw) // 2
    tokens, length = [], 0
    # 只清理显示标签；完整码点和可见转义作为整体截断，不改变来源行。
    for char in raw.decode("utf-16-le", errors="surrogatepass"):
        if unicodedata.category(char) in {"Cc", "Cf", "Cs"} or char in {"\u2028", "\u2029"}:
            units = char.encode("utf-16-le", errors="surrogatepass")
            token = "".join("\\u" + format(int.from_bytes(units[index:index + 2], "little"), "04x")
                            for index in range(0, len(units), 2))
        else:
            token = char
        width = len(token.encode("utf-16-le")) // 2
        tokens.append((token, width))
        length += width
        if length > limit:
            break
    if length <= limit:
        return "".join(token for token, _ in tokens)
    suffix = "... [已截断，原长" + str(original_length) + " UTF-16]"
    available = limit - len(suffix.encode("utf-16-le")) // 2
    prefix, length = [], 0
    for token, width in tokens:
        if length + width > available:
            break
        prefix.append(token)
        length += width
    return "".join(prefix) + suffix


def _issue(sheet, row, object_id, field, message):
    return {"sheet": _diagnostic_label(sheet, 128), "row": row,
            "objectId": _diagnostic_label(object_id, 512), "field": _diagnostic_label(field, 128),
            "message": _diagnostic_label(message, 4096)}


def _number(value):
    if not isinstance(value, str) or not NUMBER.fullmatch(value):
        raise ValueError("须为有限十进制数，不接受布尔、表达式或 NaN")
    result = float(value)
    if not math.isfinite(result):
        raise ValueError("须为有限数")
    return result


def _valid_id(value):
    return isinstance(value, str) and 0 < len(value) <= 120 and not re.search(r"\s|[\x00-\x1f\x7f]", value)


def _headers(kind, values):
    if (any(not isinstance(value, str) for value in values)
            or len(values) != len(set(values))
            or set(values) - set(HEADERS[kind])
            or set(HEADERS[kind]) - OPTIONAL[kind] - set(values)):
        raise _FileError("表头缺失、重复或包含未知字段", kind, 1, "header")
    return values


def _csv_records(data, kind):
    try:
        reader = csv.reader(StringIO(data.decode("utf-8-sig"), newline=""), strict=True)
        head = _headers(kind, next(reader, []))
        records = []
        while True:
            line = reader.line_num + 1
            values = next(reader, None)
            if values is None:
                break
            if not values:
                continue
            if len(values) != len(head):
                raise _FileError("列数与表头不一致", kind, line, "row", values[0] if values else "")
            decoded = []
            for field, value in zip(head, values):
                if value.startswith("~rp~"):
                    try:
                        value = json.loads(value[4:])
                        if not isinstance(value, str):
                            raise ValueError("not a JSON string")
                    except (ValueError, TypeError) as exc:
                        raise _FileError("CSV ~rp~ 文本转义无效", kind, line, field, values[0]) from exc
                decoded.append(value)
            records.append((kind, line, dict(zip(head, decoded))))
            if len(records) > MAX_ROWS:
                raise _FileError("业务记录最多 5000 行", kind, line, "row")
        return records
    except UnicodeError as exc:
        raise _FileError("CSV 必须为 UTF-8，可带 UTF-8 BOM", kind) from exc
    except csv.Error as exc:
        raise _FileError("CSV 格式无效：" + str(exc), kind, reader.line_num, "row") from exc


def _inspect_archive(archive):
    members = archive.infolist()
    if (len(members) > MAX_ZIP_MEMBERS
            or len({posixpath.normpath(item.filename).lower() for item in members}) != len(members)):
        raise _FileError("XLSX ZIP 成员过多或名称重复")
    if sum(item.file_size for item in members) > MAX_ZIP_BYTES:
        raise _FileError("XLSX ZIP 展开总量超过 32 MiB")
    roots = {}
    for item in members:
        name = item.filename.lower()
        if (item.flag_bits & 1 or item.compress_type not in {ZIP_STORED, ZIP_DEFLATED}
                or item.file_size > MAX_ZIP_MEMBER_BYTES
                or item.file_size > max(item.compress_size, 1) * 1000
                or name.startswith(("/", "\\")) or ".." in name.replace("\\", "/").split("/")):
            raise _FileError("XLSX ZIP 成员不安全或展开体积超限")
        if any(token in name for token in ("vbaproject", "macrosheets/", "embeddings/", "activex/", "externallinks/")) or name.endswith(".bin"):
            raise _FileError("不接受宏、嵌入对象或外部链接")
    for item in members:
        name = item.filename.lower()
        if not name.endswith((".xml", ".rels")):
            continue
        raw = archive.read(item)
        # 在 openpyxl 之前拒绝实体声明，所有 XML 都受 ZIP 展开上限约束。
        if re.search(br"<!\s*(?:DOCTYPE|ENTITY)", raw, re.I) or b"\x00" in raw:
            raise _FileError("不接受 XML 实体声明或非 UTF-8 XML")
        roots[name] = ET.fromstring(raw)
    workbook = roots.get("xl/workbook.xml")
    relationships = roots.get("xl/_rels/workbook.xml.rels")
    targets = {} if relationships is None else {
        el.get("Id"): posixpath.normpath(posixpath.join("xl", el.get("Target", "")))
        if not el.get("Target", "").startswith("/") else el.get("Target").lstrip("/")
        for el in relationships
    }
    sheet_names = {} if workbook is None else {
        targets.get(el.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"), "").lower(): el.get("name", "file")
        for el in workbook.iter() if el.tag.rsplit("}", 1)[-1] == "sheet"
    }
    for name, root in roots.items():
        if name == "[content_types].xml" and any("macroenabled" in str(el.attrib).lower() or "vba" in str(el.attrib).lower() for el in root):
            raise _FileError("不接受含宏工作簿")
        sheet_xml = name.startswith("xl/worksheets/")
        sheet_name = sheet_names.get(name, "file")
        last_row, cell_refs = 0, set()
        for el in root.iter():
            tag = el.tag.rsplit("}", 1)[-1]
            if tag == "Relationship" and el.get("TargetMode", "").lower() == "external":
                raise _FileError("不接受外部链接")
            if tag == "sheet" and el.get("state", "visible") != "visible":
                raise _FileError("不接受隐藏工作表", el.get("name", "file"), 1, "sheet")
            if not sheet_xml:
                continue
            if tag == "row" and not 1 <= int(el.get("r", "0")) <= MAX_ROWS + 1:
                raise _FileError("XLSX 实际行号超限", sheet_name, int(el.get("r", "0")), "row")
            if tag == "row":
                row_number = int(el.get("r", "0"))
                if row_number <= last_row:
                    raise _FileError("XLSX 行号重复或顺序无效", sheet_name, row_number, "row")
                last_row = row_number
            if tag in {"row", "col"} and el.get("hidden", "0").lower() in {"1", "true"}:
                raise _FileError("不接受隐藏行或列", sheet_name, int(el.get("r", "1")), "hidden")
            if tag in {"mergeCell", "oleObject"}:
                raise _FileError("不接受合并单元格或嵌入对象", sheet_name, 1, "structure")
            if tag == "c":
                ref = el.get("r", "")
                match = CELL_REF.fullmatch(ref)
                if not match:
                    raise _FileError("XLSX 单元格坐标无效")
                if ref in cell_refs or int(match[2]) != last_row:
                    raise _FileError("XLSX 单元格坐标重复或与所在行不符", sheet_name, int(match[2]), ref)
                cell_refs.add(ref)
                col = 0
                for char in match[1]:
                    col = col * 26 + ord(char) - ord("A") + 1
                if int(match[2]) > MAX_ROWS + 1 or col > max(map(len, HEADERS.values())):
                    raise _FileError("XLSX 实际行列超限", sheet_name, int(match[2]), ref)


def _xlsx_records(data):
    try:
        with ZipFile(BytesIO(data)) as archive:
            _inspect_archive(archive)
        book = load_workbook(BytesIO(data), read_only=True, data_only=False, keep_links=False)
        try:
            if (len(book.sheetnames) != len(set(book.sheetnames))
                    or set(book.sheetnames) not in (set(HEADERS), set(HEADERS) | {"说明"})):
                raise _FileError("XLSX 必须包含三类业务表，可附加说明表；不接受额外工作表")
            records = []
            for sheet in book:
                kind = sheet.title
                if sheet.sheet_state != "visible" or not hasattr(sheet, "reset_dimensions"):
                    raise _FileError("不接受隐藏表或非业务工作表", kind, 1, "sheet")
                # 不相信来源 dimension，防止将范围外数据静默截断。
                sheet.reset_dimensions()
                head = None
                for line, cells in enumerate(sheet.iter_rows(), 1):
                    values = []
                    for column, cell in enumerate(cells):
                        value = cell.value
                        field = head[column] if head is not None and column < len(head) else getattr(cell, "column_letter", "cell")
                        object_id = str(cells[0].value or "")
                        if cell.data_type in {"f", "e"} or isinstance(value, (date, datetime, time, timedelta)):
                            raise _FileError("不接受公式、错误单元格或日期对象；日期请使用 ISO 文本", kind, line, field, object_id)
                        if value is not None and not isinstance(value, (str, int, float, bool)):
                            raise _FileError("单元格标量类型无效", kind, line, "cell")
                        values.append(value)
                    if kind == "说明":
                        continue
                    if line == 1:
                        head = _headers(kind, values)
                        continue
                    if not any(value is not None and value != "" for value in values):
                        continue
                    if head is None or len(values) > len(head):
                        raise _FileError("列数与表头不一致", kind, line, "row")
                    values += [None] * (len(head) - len(values))
                    row = {}
                    for field, value in zip(head, values):
                        if isinstance(value, bool):
                            if field not in BOOLEANS:
                                raise _FileError("布尔值不能代替数字或文本", kind, line, field, str(values[0]))
                            value = str(value).lower()
                        elif isinstance(value, (int, float)):
                            if field not in NUMERIC:
                                raise _FileError("文本字段必须为字符串", kind, line, field, str(values[0]))
                            value = str(value)
                        row[field] = "" if value is None else value
                    records.append((kind, line, row))
                    if len(records) > MAX_ROWS:
                        raise _FileError("业务记录合计最多 5000 行", kind, line, "row")
                if kind != "说明" and head is None:
                    raise _FileError("缺少表头", kind, 1, "header")
            return records
        finally:
            book.close()
    except _FileError:
        raise
    except (BadZipFile, ET.ParseError, ValueError, TypeError, KeyError, IndexError, AttributeError,
            OSError, RuntimeError, OverflowError, EOFError, NotImplementedError) as exc:
        raise _FileError("XLSX 文件损坏或格式不受支持") from exc


def _validate_record(kind, line, row, old, changed, indexes, issues, warnings, bounds):
    object_id = row[HEADERS[kind][0]]

    def issue(field, message):
        issues.append(_issue(kind, line, object_id, field, message))

    def number(field, nullable=False, low=None, high=None, exclusive_high=False):
        value = row.get(field, "")
        if value == "" and nullable:
            return None
        try:
            result = _number(value)
        except ValueError as exc:
            issue(field, str(exc))
            return None
        if (low is not None and result < low or high is not None
                and (result >= high if exclusive_high else result > high)):
            issue(field, "数值超出允许范围")
        return result

    def text(field, limit, required=False):
        value = row.get(field, "")
        if len(value) > limit or required and (not value or value != value.strip()):
            issue(field, "文本长度或首尾空白无效")

    def coords():
        try:
            coordinate = [_number(row.get(field, "")) for field in ("lon", "lat")]
        except ValueError:
            for field in ("lon", "lat"):
                number(field)
            return
        if old is not None:
            try:
                if coordinate == [_number(old.get(field, "")) for field in ("lon", "lat")]:
                    return
            except ValueError:
                pass
        if not valid_bounds(bounds):
            issue("coordinate", "当前 DEM 范围缺失或无效，不能新增或修改坐标")
            return
        for field, low, high in (("lon", bounds[0], bounds[2]), ("lat", bounds[1], bounds[3])):
            number(field, low=low, high=high)

    if kind == "nodes":
        coords()
        for field, limit, required in (("node_name", 60, True), ("region", 60, False), ("remark", 500, False)):
            if old is None or field in changed:
                text(field, limit, required)
        if old is None or "elevation_m" in changed:
            number("elevation_m", nullable=True)
        if old is not None:
            return
        classes = row.get("device_class", "").split(";")
        definition = STATION_TYPES.get(row.get("node_subtype"))
        if (not definition or row.get("echelon") != definition[0] or row.get("mobility") != definition[1]
                or len(classes) != len(set(classes)) or set(classes) != definition[2]):
            issue("node_subtype", "节点子类型、层级、移动类型及 HF/VUHF 类别组合无效")
        if row.get("node_role") not in {"TASK", "FIXED_STATION", "RELAY"}:
            issue("node_role", "节点角色无效")
        if row.get("status", "") not in {"", "UNKNOWN", "NORMAL", "FAULT", "OFFLINE"}:
            issue("status", "运行状态无效，未知值须留空或为 UNKNOWN")
        parents = row.get("parent_node_id", "").split(";") if row.get("parent_node_id") else []
        if len(parents) != len(set(parents)) or any(not _valid_id(parent) or parent not in indexes["node.csv"] or parent == object_id for parent in parents):
            issue("parent_node_id", "上级节点不存在、重复或为自身")
        # 仅沿本批新增节点的祖先路径检查，旧库无关环不阻断本次操作。
        visiting, visited = set(), set()
        stack = [(object_id, False)]
        while stack:
            current, leaving = stack.pop()
            if leaving:
                visiting.discard(current)
                visited.add(current)
                continue
            if current in visiting:
                issue("parent_node_id", "上级节点形成循环引用")
                break
            if current in visited or current not in indexes["node.csv"]:
                continue
            visiting.add(current)
            stack.append((current, True))
            stack.extend((parent, False) for parent in indexes["node.csv"][current].get("parent_node_id", "").split(";") if parent)
    elif kind == "devices":
        model = indexes["device_model.csv"].get(row.get("model_id"))
        node = indexes["node.csv"].get(row.get("node_id"))
        antenna = indexes["antenna_model.csv"].get(row.get("antenna_id"))
        for field, related in (("model_id", model), ("node_id", node), ("antenna_id", antenna)):
            if related is None:
                issue(field, "引用对象不存在")
        if model is not None and node is not None and antenna is not None:
            device_class = model.get("device_class")
            if device_class not in {"HF", "VUHF"} or device_class not in node.get("device_class", "").split(";"):
                issue("node_id", "设备型号类别与节点 HF/VUHF 能力不一致")
            if device_class not in antenna.get("device_class", "").split(";"):
                issue("antenna_id", "天线类别与设备型号不一致")
        for field, options in (("antenna_height_m", {"low": 0}), ("tilt_deg", {"nullable": True, "low": -90, "high": 90}),
                               ("azimuth_deg", {"nullable": True, "low": 0, "high": 360, "exclusive_high": True})):
            if old is None or field in changed:
                number(field, **options)
        if old is None or "tx_power_dbm" in changed:
            number("tx_power_dbm")
        if model is not None:
            try:
                power, maximum = _number(row.get("tx_power_dbm", "")), _number(model.get("tx_power_max_dbm", ""))
                if power > maximum:
                    if old is not None and _number(old.get("tx_power_dbm", "")) == power and power - maximum <= 0.040001:
                        warnings.append(_issue(kind, line, object_id, "tx_power_dbm", "来源功率舍入超限，本次未改动予以保留"))
                    else:
                        issue("tx_power_dbm", "发射功率不得超过型号上限")
                freq = row.get("work_freq_khz", "")
                if freq and not 0 < _number(model.get("freq_min_khz", "")) <= _number(freq) <= _number(model.get("freq_max_khz", "")):
                    issue("work_freq_khz", "工作频率须为正数且位于型号频率范围内")
            except ValueError:
                issue("model_id", "设备参数或关联型号数值无效")
        if old is None:
            if row.get("status", "") not in {"", "UNKNOWN", "NORMAL", "FAULT", "OFFLINE"}:
                issue("status", "运行状态无效，未知值须留空或为 UNKNOWN")
            stamp = row.get("install_time", "")
            if stamp:
                try:
                    if not ISO_TIME.fullmatch(stamp) or datetime.fromisoformat(stamp.replace("Z", "+00:00")).tzinfo is None:
                        raise ValueError("invalid timestamp")
                except ValueError:
                    issue("install_time", "安装时间须为有效带时区 ISO 文本或留空")
    else:
        coords()
        if old is None or "name" in changed:
            text("name", 60, True)
        for field, options in (
            ("elevation_m", {"nullable": True}), ("antenna_height_m", {"low": 0}),
            ("center_freq_khz", {"low": 0}), ("bandwidth_khz", {"low": 0}),
            ("tx_power_dbm", {}), ("antenna_gain_dbi", {"nullable": True}),
            ("azimuth_deg", {"nullable": True, "low": 0, "high": 360, "exclusive_high": True}),
            ("duty_cycle", {"nullable": True, "low": 0, "high": 1}),
        ):
            if old is None or field in changed:
                value = number(field, **options)
                if field in {"center_freq_khz", "bandwidth_khz"} and value == 0:
                    issue(field, "频率及带宽必须大于 0")
        if (old is None or "interference_type" in changed) and row.get("interference_type") not in {"NARROWBAND", "BROADBAND", "SWEEP", "PULSE", "NOISE"}:
            issue("interference_type", "干扰类型无效")
        if old is None or {"pattern_type", "azimuth_deg"} & changed:
            if row.get("pattern_type", "") not in {"", "OMNI", "DIRECTIONAL"}:
                issue("pattern_type", "方向图无效")
            if row.get("pattern_type") == "DIRECTIONAL" and not row.get("azimuth_deg"):
                issue("azimuth_deg", "定向干扰源须填写方位角")


def parse_import(tables: dict, file_name: str, data: bytes, kind: str, bounds=None) -> dict:
    """返回完整十表深拷贝、文件行差异及逐行问题；任一问题均撤销整批候选。"""
    baseline = deepcopy(tables)
    result = {"tables": baseline, "rows": [], "issues": [], "warnings": []}
    issues, warnings = result["issues"], result["warnings"]
    try:
        if not isinstance(tables, dict) or set(tables) != set(TABLE_IDS) or any(not isinstance(rows, list) for rows in tables.values()):
            raise _FileError("输入必须为完整十张资源表", field="tables")
        if kind not in HEADERS or not isinstance(file_name, str):
            raise _FileError("导入类型或文件名无效")
        if not isinstance(data, bytes) or not data or len(data) > MAX_FILE_BYTES:
            raise _FileError("文件不能为空且不得超过 5 MiB")
        if file_name.lower().endswith(".csv"):
            records = _csv_records(data, kind)
        elif file_name.lower().endswith(".xlsx"):
            records = _xlsx_records(data)
        else:
            raise _FileError("仅支持 CSV 和 XLSX，不接受宏工作簿")
        if not records:
            raise _FileError("每批需有 1–5000 行业务记录")
    except _FileError as exc:
        issues.append(exc.detail)
        return result

    candidate = deepcopy(tables)
    indexes = {name: {row.get(field): row for row in rows if isinstance(row, dict)}
               for name, rows in candidate.items() for field in [TABLE_IDS[name]]}
    seen, pending = set(), []
    for sheet, line, provided in records:
        id_field = HEADERS[sheet][0]
        object_id = provided[id_field]
        if not _valid_id(object_id):
            issues.append(_issue(sheet, line, object_id, id_field, "ID 须为 1–120 字符且不含空白或控制字符"))
            continue
        if (sheet, object_id) in seen:
            issues.append(_issue(sheet, line, object_id, id_field, "同类资源 ID 重复"))
            continue
        seen.add((sheet, object_id))
        target = indexes[TABLES[sheet]]
        old = deepcopy(target.get(object_id))
        row = dict(old) if old is not None else {field: "" for field in HEADERS[sheet]}
        patch = dict(provided)
        if sheet == "devices":
            # 空的只读生命周期列沿用当前值；旧基线缺列时仅使用既有 ACTIVE 语义。
            availability = old.get("availability", "ACTIVE") if old is not None else "ACTIVE"
            if patch.get("availability", "") not in {"", availability}:
                issues.append(_issue(sheet, line, object_id, "availability", "导入不能改变生命周期；新增设备必须为 ACTIVE"))
            if old is not None:
                patch.pop("availability", None)
            else:
                patch["availability"] = "ACTIVE"
        for field, value in patch.items():
            if CONTROL.search(value):
                issues.append(_issue(sheet, line, object_id, field, "不接受控制字符"))
            if field in NUMERIC and value:
                try:
                    _number(value)
                except ValueError as exc:
                    issues.append(_issue(sheet, line, object_id, field, str(exc)))
            if field in BOOLEANS and value not in {"true", "false"}:
                issues.append(_issue(sheet, line, object_id, field, "布尔字段仅接受 true 或 false"))
            if old is not None and field not in old and value == "":
                continue
            if old is not None and field != id_field and value != old.get(field, ""):
                if sheet == "nodes" and field not in NODE_EDITABLE or sheet == "devices" and field in DEVICE_READONLY:
                    issues.append(_issue(sheet, line, object_id, field, "既有对象的只读字段或关联不可修改"))
            row[field] = value
        differences = [{"field": field, "before": old.get(field) if old is not None else None, "after": row.get(field)}
                       for field in HEADERS[sheet] if (old.get(field) if old is not None else None) != row.get(field)]
        changed = {diff["field"] for diff in differences}
        if old is None or changed:
            row["origin"] = "IMPORT"
        result["rows"].append({"sheet": sheet, "row": line, "id": object_id,
                               "action": "NEW" if old is None else "UPDATE" if changed else "UNCHANGED", "differences": differences})
        if old is None:
            candidate[TABLES[sheet]].append(row)
        else:
            target[object_id].clear()
            target[object_id].update(row)
        target[object_id] = row if old is None else target[object_id]
        pending.append((sheet, line, row, old, changed))
    # 先完成三表合并，再校验引用，允许设备引用本批次任意顺序新增的站点。
    for sheet, line, row, old, changed in pending:
        if old is None or changed:
            _validate_record(sheet, line, row, old, changed, indexes, issues, warnings, bounds)
    if not issues:
        result["tables"] = candidate
    return result


def _template_rows(tables):
    selected = {kind: [] for kind in HEADERS}
    devices = tables["device.csv"]
    if devices:
        selected["devices"] = [devices[0]]
    nodes = {row["node_id"]: row for row in tables["node.csv"]}
    first = devices[0].get("node_id") if devices else next(iter(nodes), None)
    wanted, stack = set(), [first] if first else []
    while stack:
        node_id = stack.pop()
        if node_id in wanted or node_id not in nodes:
            continue
        wanted.add(node_id)
        stack.extend(parent for parent in nodes[node_id].get("parent_node_id", "").split(";") if parent)
    selected["nodes"] = [row for node_id, row in nodes.items() if node_id in wanted]
    selected["interference_sources"] = tables["interference_source.csv"][:1]
    return selected


def _template_value(row, field):
    return row.get(field, "ACTIVE" if field == "availability" else "")


class _MemoryExcelWriter(ExcelWriter):
    def write_worksheet(self, sheet):
        # 默认 openpyxl 保存会落临时文件；复用 XML writer，仅将输出改为内存。
        sheet._drawing = SpreadsheetDrawing()
        output = BytesIO()
        writer = WorksheetWriter(sheet, out=output)
        writer.write()
        sheet._rels = writer._rels
        self._archive.writestr(sheet.path[1:], output.getvalue())
        self.manifest.append(sheet)


def _save_memory(book):
    output = BytesIO()
    with ZipFile(output, "w", ZIP_DEFLATED, allowZip64=True) as archive:
        _MemoryExcelWriter(book, archive).save()
    return output.getvalue()


def make_template(tables, format, kind):
    """使用当前真实记录及关联站点示例；CSV 来源性质在文件名中明确标注。"""
    if format not in {"csv", "xlsx"} or kind not in HEADERS:
        raise ValueError("模板仅支持三类资源的 CSV 或 XLSX")
    selected = _template_rows(tables)
    filename = "SYNTHETIC_INTEGRATION-非生产非实测-资源模板-" + (kind if format == "csv" else "三表") + "." + format
    if format == "csv":
        output = StringIO(newline="")
        writer = csv.writer(output)
        writer.writerow(HEADERS[kind])
        for row in selected[kind]:
            values = [_template_value(row, field) for field in HEADERS[kind]]
            writer.writerow(["~rp~" + json.dumps(value, ensure_ascii=False) if isinstance(value, str) and
                             (re.match(r"^\s*[=+\-@]", value) or value.startswith("~rp~")) else value for value in values])
        data, mime = output.getvalue().encode("utf-8"), "text/csv; charset=utf-8"
    else:
        book = Workbook()
        book.remove(book.active)
        try:
            for sheet_kind, rows in selected.items():
                sheet = book.create_sheet(sheet_kind)
                sheet.append(HEADERS[sheet_kind])
                for row in rows:
                    sheet.append([_template_value(row, field) for field in HEADERS[sheet_kind]])
                    for cell in sheet[sheet.max_row]:
                        cell.data_type = "s"
                sheet.freeze_panes = "A2"
            notes = book.create_sheet("说明")
            for row in [
                ("数据来源", "SYNTHETIC_INTEGRATION"), ("数据性质", "构造联调，非生产、非实测"),
                ("模板示例", "来自当前资源真实行及关联行；不新造设备型号或状态；按 ID 新增或更新，未出现记录保留"),
                ("只读字段", "既有节点结构、设备关联及状态不可修改；availability 只读，新增设备必须 ACTIVE"),
                ("单位与空值", "WGS84 lon/lat；kHz、dBm、m；布尔值仅 true/false；未知数值留空，不补零或正常"),
                ("文件限制", "最多 5 MiB、5000 业务行；不接受公式、宏、隐藏数据、额外工作表或 Excel 日期对象"),
                ("CSV 文本转义", "公式前缀及 ~rp~ 前缀文本采用 ~rp~ 加 JSON 字符串；安装时间使用带时区 ISO 文本"),
            ]:
                notes.append(row)
                for cell in notes[notes.max_row]:
                    cell.data_type = "s"
            data, mime = _save_memory(book), XLSX_MIME
        finally:
            book.close()
    row_count = len(selected[kind]) if format == "csv" else sum(len(rows) for rows in selected.values())
    if len(data) > MAX_FILE_BYTES or row_count > MAX_ROWS:
        raise ValueError("模板超过导入上限，未返回截断内容")
    return data, mime, filename
