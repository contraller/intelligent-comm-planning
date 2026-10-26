"""纯内存资源文件测试；不初始化数据库，不读写项目原件。"""
from copy import deepcopy
import csv
from datetime import datetime
from io import BytesIO, StringIO
import json
from pathlib import Path
import re
import sys
import xml.etree.ElementTree as ET
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "algorithm-service"))
from app import resource_import_parser as parser

# 仅供旧解析用例显式注入，不是生产默认范围。
TEST_BOUNDS = (114.0465, 36.911, 115.4045, 37.989)


@pytest.fixture
def tables():
    result = {name: [] for name in parser.TABLE_IDS}
    result["node.csv"] = [{
        "node_id": "N1", "node_name": "当前站点", "echelon": "II", "node_subtype": "II_FIXED",
        "parent_node_id": "", "node_role": "FIXED_STATION", "mobility": "FIXED",
        "lon": "114.5", "lat": "37.5", "elevation_m": "", "region": "区域", "device_class": "HF;VUHF",
        "relay_capable": "true", "status": "UNKNOWN", "is_key": "false", "remark": "原备注",
    }]
    result["device.csv"] = [{
        "device_id": "D1", "node_id": "N1", "model_id": "M1", "antenna_id": "A1",
        "antenna_height_m": "10", "tilt_deg": "", "azimuth_deg": "", "tx_power_dbm": "43.05",
        "work_freq_khz": "", "status": "", "install_time": "", "availability": "OFF_SHELF",
    }]
    result["device_model.csv"] = [{"model_id": "M1", "device_class": "HF", "freq_min_khz": "1500",
                                    "freq_max_khz": "30000", "tx_power_max_dbm": "43.01"}]
    result["antenna_model.csv"] = [{"antenna_id": "A1", "device_class": "HF"}]
    result["interference_source.csv"] = [{
        "interference_id": "I1", "name": "当前干扰源", "lon": "114.5", "lat": "37.5", "elevation_m": "",
        "antenna_height_m": "0", "center_freq_khz": "1500", "bandwidth_khz": "10", "tx_power_dbm": "-10",
        "antenna_gain_dbi": "", "interference_type": "NOISE", "pattern_type": "", "azimuth_deg": "",
        "duty_cycle": "", "active": "false",
    }]
    # 其余七表的内容也必须完整保留，不由导入器投影或补造。
    for name, id_field in parser.TABLE_IDS.items():
        if not result[name]:
            result[name] = [{id_field: name, "source": "保留来源"}]
    return result


def csv_bytes(kind, rows, headers=None):
    head = parser.HEADERS[kind] if headers is None else headers
    output = StringIO(newline="")
    writer = csv.writer(output)
    writer.writerow(head)
    for row in rows:
        writer.writerow([row.get(field, "") for field in head])
    return output.getvalue().encode("utf-8")


def parse(tables, kind, rows, headers=None):
    return parser.parse_import(tables, "input.csv", csv_bytes(kind, rows, headers), kind, TEST_BOUNDS)


def workbook_bytes(tables, rows=None, mutate=None, order=None):
    book = Workbook()
    book.remove(book.active)
    for kind in order or parser.HEADERS:
        sheet = book.create_sheet(kind)
        sheet.append(parser.HEADERS[kind])
        source = rows.get(kind, []) if rows is not None else tables[parser.TABLES[kind]]
        for row in source:
            sheet.append([row.get(field, "") for field in parser.HEADERS[kind]])
            for cell in sheet[sheet.max_row]:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
    if mutate:
        mutate(book)
    data = parser._save_memory(book)
    book.close()
    return data


def parse_xlsx(tables, **kwargs):
    return parser.parse_import(tables, "input.xlsx", workbook_bytes(tables, **kwargs), "nodes", TEST_BOUNDS)


def assert_atomic_failure(tables, result):
    assert result["issues"]
    assert result["tables"] == tables
    assert result["tables"] is not tables
    assert all(set(issue) == {"sheet", "row", "objectId", "field", "message"} for issue in result["issues"])


def utf16_length(value):
    return len(value.encode("utf-16-le")) // 2


def assert_displayable_issue(issue):
    for field, limit in (("sheet", 128), ("field", 128), ("objectId", 512), ("message", 4096)):
        label = issue[field]
        assert isinstance(label, str) and utf16_length(label) <= limit
        assert not re.search(r"[\x00-\x1f\x7f-\x9f]", label)
        label.encode("utf-8")


@pytest.mark.parametrize("field,limit", [("sheet", 128), ("field", 128), ("objectId", 512), ("message", 4096)])
def test_issue_labels_escape_controls_and_bound_js_utf16_length(field, limit):
    value = "L\x00\t\n\r\x7f\x85\u202e" + "\U00020000" * limit
    labels = {"sheet": "nodes", "field": "node_id", "objectId": "N1", "message": "原始诊断"}
    labels[field] = value
    issue = parser._issue(labels["sheet"], 5002, labels["objectId"], labels["field"], labels["message"])
    assert_displayable_issue(issue)
    assert issue["row"] == 5002
    assert issue[field].startswith(r"L\u0000\u0009\u000a\u000d\u007f\u0085\u202e")
    assert "已截断" in issue[field] and "原长" + str(utf16_length(value)) in issue[field]
    assert labels[field] == value
    assert all(issue[key] == label for key, label in labels.items() if key != field)


@pytest.mark.parametrize("field,limit", [("sheet", 128), ("field", 128), ("objectId", 512), ("message", 4096)])
def test_issue_utf16_boundary_preserves_complete_supplementary_characters(field, limit):
    labels = {"sheet": "nodes", "field": "node_id", "objectId": "N1", "message": "诊断"}
    exact = "\U00020000" * (limit // 2)
    labels[field] = exact
    issue = parser._issue(labels["sheet"], 2, labels["objectId"], labels["field"], labels["message"])
    assert issue[field] == exact and utf16_length(issue[field]) == limit
    labels[field] += "\U00020000"
    issue = parser._issue(labels["sheet"], 2, labels["objectId"], labels["field"], labels["message"])
    assert_displayable_issue(issue)
    prefix, marker = issue[field].split("... ", 1)
    assert prefix and all(ord(char) == 0x20000 for char in prefix)
    assert "原长" + str(limit + 2) in marker


def test_issue_controls_and_surrogates_are_visible_and_json_safe():
    controls = "".join(chr(code) for code in (*range(32), *range(127, 160), 0x2028, 0x2029, 0x202e))
    issue = parser._issue("file", 0, "\ud840\udc00\ud800X\udfff", "field", controls)
    assert_displayable_issue(issue)
    assert issue["objectId"] == "\U00020000" + r"\ud800X\udfff"
    assert issue["message"] == "".join("\\u" + format(ord(char), "04x") for char in controls)
    assert json.loads(json.dumps(issue, ensure_ascii=False).encode("utf-8")) == issue


def test_issue_does_not_split_visible_control_escapes_when_truncated():
    issue = parser._issue("\n" * 100, 0, "N1", "field", "message")
    assert_displayable_issue(issue)
    prefix, marker = issue["sheet"].split("... ", 1)
    assert prefix == r"\u000a" * (len(prefix) // 6)
    assert "原长100" in marker


@pytest.mark.parametrize("identifier", ["I" * 513, "\U00020000" * 300, "I" * 510 + "\x00\t"])
def test_invalid_long_ids_have_bounded_diagnostics_without_mutating_input(tables, identifier):
    row = {**tables["node.csv"][0], "node_id": identifier}
    original = deepcopy(tables)
    data = csv_bytes("nodes", [row])
    result = parser.parse_import(tables, "input.csv", data, "nodes", TEST_BOUNDS)
    assert_atomic_failure(tables, result)
    assert tables == original and row["node_id"] == identifier
    assert result["rows"] == []
    issue = result["issues"][0]
    assert_displayable_issue(issue)
    assert issue["row"] == 2 and issue["field"] == "node_id"
    assert "已截断" in issue["objectId"] and "原长" + str(utf16_length(identifier)) in issue["objectId"]


def test_sanitized_issues_do_not_rewrite_raw_row_differences_or_candidate(tables):
    remark = "\x00" + "R" * 700
    row = {**tables["node.csv"][0], "remark": remark}
    result = parse(tables, "nodes", [row])
    assert_atomic_failure(tables, result)
    assert result["rows"][0]["id"] == "N1"
    assert result["rows"][0]["differences"] == [{"field": "remark", "before": "原备注", "after": remark}]
    assert row["remark"] == remark
    for issue in result["issues"]:
        assert_displayable_issue(issue)


@pytest.mark.parametrize("filename,data", [("bad\x00.xlsm", b"unsupported"), ("bad\n.xlsx", b"bad zip"), ("input.csv", b"")])
def test_file_level_diagnostics_are_displayable_without_relabeling_business_rows(tables, filename, data):
    result = parser.parse_import(tables, filename, data, "nodes", TEST_BOUNDS)
    assert_atomic_failure(tables, result)
    assert result["rows"] == []
    assert result["issues"][0]["sheet"] == "file" and result["issues"][0]["row"] == 0
    assert_displayable_issue(result["issues"][0])


def test_update_retains_ten_tables_and_input_is_deeply_immutable(tables):
    original = deepcopy(tables)
    updated = {**tables["node.csv"][0], "node_name": "更名,含逗号\n第二行"}
    result = parse(tables, "nodes", [updated])
    assert not result["issues"]
    assert result["tables"]["node.csv"][0] == {**updated, "origin": "IMPORT"}
    assert tables == original
    assert result["rows"] == [{"sheet": "nodes", "row": 2, "id": "N1", "action": "UPDATE",
                               "differences": [{"field": "node_name", "before": "当前站点", "after": updated["node_name"]}]}]
    assert set(result["tables"]) == set(parser.TABLE_IDS)
    result["tables"]["device_model.csv"][0]["device_class"] = "VUHF"
    assert tables == original


def test_optional_columns_preserve_values_and_unmentioned_rows(tables):
    tables["node.csv"].append({**tables["node.csv"][0], "node_id": "N2"})
    head = [field for field in parser.HEADERS["nodes"] if field not in parser.OPTIONAL["nodes"]]
    result = parse(tables, "nodes", [{**tables["node.csv"][0], "node_name": "改名"}], head)
    assert not result["issues"]
    assert result["tables"]["node.csv"][0]["remark"] == "原备注"
    assert result["tables"]["node.csv"][1] == tables["node.csv"][1]
    device = tables["device.csv"][0]
    head = [field for field in parser.HEADERS["devices"] if field not in parser.OPTIONAL["devices"]]
    result = parse(tables, "devices", [{**device, "antenna_height_m": "11"}], head)
    assert not result["issues"]
    assert result["tables"]["device.csv"][0]["availability"] == "OFF_SHELF"
    assert result["tables"]["device.csv"][0]["work_freq_khz"] == ""


@pytest.mark.parametrize("bom", [b"", b"\xef\xbb\xbf"])
def test_bom_formula_text_and_multiline_csv_roundtrip(tables, bom):
    raw = '~rp~=SUM(1,2)\n"原文"'
    row = {**tables["node.csv"][0], "remark": "~rp~" + json.dumps(raw, ensure_ascii=False)}
    result = parser.parse_import(tables, "input.csv", bom + csv_bytes("nodes", [row]), "nodes", TEST_BOUNDS)
    assert not result["issues"]
    assert result["tables"]["node.csv"][0]["remark"] == raw


@pytest.mark.parametrize("value", ["~rp~x", "~rp~null", "~rp~true", "~rp~[1]"])
def test_bad_formula_escape_is_located_and_atomic(tables, value):
    result = parse(tables, "nodes", [{**tables["node.csv"][0], "remark": value}])
    assert_atomic_failure(tables, result)
    assert result["issues"][0]["row"] == 2
    assert result["issues"][0]["field"] == "remark"


@pytest.mark.parametrize("mutation", ["duplicate", "unknown", "missing", "short", "long", "quote"])
def test_csv_headers_and_column_count_are_strict(tables, mutation):
    headers = list(parser.HEADERS["nodes"])
    if mutation == "duplicate":
        headers.append("node_id")
    elif mutation == "unknown":
        headers.append("source")
    elif mutation == "missing":
        headers.remove("node_name")
    data = csv_bytes("nodes", [tables["node.csv"][0]], headers)
    if mutation == "short":
        data = data.replace(b",\r\n", b"\r\n") if data.endswith(b",\r\n") else data.rsplit(b",", 1)[0] + b"\r\n"
    elif mutation == "long":
        data = data[:-2] + b",extra\r\n"
    elif mutation == "quote":
        data = data.split(b"\r\n")[0] + b'\r\n"unclosed'
    assert_atomic_failure(tables, parser.parse_import(tables, "input.csv", data, "nodes", TEST_BOUNDS))


def test_duplicate_id_and_mixed_valid_invalid_batch_are_atomic(tables):
    row = {**tables["node.csv"][0], "node_name": "可改"}
    result = parse(tables, "nodes", [row, row])
    assert_atomic_failure(tables, result)
    assert result["issues"][0]["row"] == 3
    bad = {**tables["interference_source.csv"][0], "interference_id": "I2", "active": "1"}
    rows = {"nodes": [row], "interference_sources": [bad]}
    result = parse_xlsx(tables, rows=rows)
    assert_atomic_failure(tables, result)
    assert any(issue["objectId"] == "I2" and issue["field"] == "active" for issue in result["issues"])


@pytest.mark.parametrize("kind", list(parser.HEADERS))
@pytest.mark.parametrize("identifier", ["BAD ID", "BAD\tID", "BAD\nID", "BAD\rID", "BAD\u00a0ID", "BAD\u3000ID", "BAD\u2003ID"])
def test_ids_reject_all_whitespace_for_identifier_contract(tables, kind, identifier):
    id_field = parser.HEADERS[kind][0]
    row = {**tables[parser.TABLES[kind]][0], id_field: identifier}
    result = parse(tables, kind, [row])
    assert_atomic_failure(tables, result)
    assert result["rows"] == []
    assert len(result["issues"]) == 1
    assert result["issues"][0]["field"] == id_field
    expected_label = identifier.replace("\t", r"\u0009").replace("\n", r"\u000a").replace("\r", r"\u000d")
    assert result["issues"][0]["objectId"] == expected_label


@pytest.mark.parametrize("value", ["true", "false", "NaN", "nan", "Infinity", "-inf", "1e999", "1_000", "1+2", " 10", "0x10"])
def test_numeric_grammar_is_finite_and_not_boolean(tables, value):
    result = parse(tables, "devices", [{**tables["device.csv"][0], "antenna_height_m": value}])
    assert_atomic_failure(tables, result)
    assert any(issue["field"] == "antenna_height_m" and issue["row"] == 2 for issue in result["issues"])


@pytest.mark.parametrize("value", ["TRUE", "False", "1", "0", "", " true"])
def test_booleans_are_exact(tables, value):
    result = parse(tables, "interference_sources", [{**tables["interference_source.csv"][0], "active": value}])
    assert_atomic_failure(tables, result)


def test_unknown_values_remain_unknown(tables):
    device = {**tables["device.csv"][0], "antenna_height_m": "11", "availability": ""}
    result = parse(tables, "devices", [device])
    assert not result["issues"]
    saved = result["tables"]["device.csv"][0]
    assert saved["status"] == saved["work_freq_khz"] == saved["tilt_deg"] == saved["install_time"] == ""
    assert saved["availability"] == "OFF_SHELF"
    assert result["warnings"][0]["field"] == "tx_power_dbm"
    assert result["warnings"][0]["objectId"] == "D1"


@pytest.mark.parametrize("kind", ["nodes", "devices"])
def test_new_empty_status_is_unknown_and_actual_changes_mark_import(tables, kind):
    row = {**tables[parser.TABLES[kind]][0], "status": ""}
    if kind == "nodes":
        row["node_id"] = "N2"
    else:
        row.update(device_id="D2", availability="ACTIVE", tx_power_dbm="40")
    result = parse(tables, kind, [row])
    assert not result["issues"], result["issues"]
    saved = result["tables"][parser.TABLES[kind]][-1]
    assert saved["status"] == ""
    assert saved["origin"] == "IMPORT"
    if kind == "devices":
        assert saved["work_freq_khz"] == ""


@pytest.mark.parametrize("origin", [None, "SOURCE", "MANUAL", "IMPORT"])
def test_noop_template_preserves_origin_and_empty_status(tables, origin):
    for name in parser.TABLES.values():
        if name != "interference_source.csv":
            tables[name][0]["status"] = ""
        if origin is not None:
            tables[name][0]["origin"] = origin
    data, _, filename = parser.make_template(tables, "xlsx", "nodes")
    result = parser.parse_import(tables, filename, data, "nodes", TEST_BOUNDS)
    assert not result["issues"]
    assert all(row["action"] == "UNCHANGED" for row in result["rows"])
    assert result["tables"] == tables
    updated = parse(tables, "devices", [{**tables["device.csv"][0], "antenna_height_m": "12"}])
    assert not updated["issues"]
    assert updated["tables"]["device.csv"][0]["origin"] == "IMPORT"


def test_empty_missing_readonly_fields_do_not_create_noop_changes(tables):
    for field in ("status", "install_time", "availability", "tilt_deg", "work_freq_khz"):
        tables["device.csv"][0].pop(field)
    data, _, filename = parser.make_template(tables, "csv", "devices")
    result = parser.parse_import(tables, filename, data, "devices", TEST_BOUNDS)
    assert not result["issues"]
    assert result["rows"][0]["action"] == "UNCHANGED"
    assert result["rows"][0]["differences"] == []
    assert result["tables"] == tables


@pytest.mark.parametrize("field,value", [
    ("model_id", "M2"), ("node_id", "N2"), ("antenna_id", "A2"), ("status", "NORMAL"),
    ("install_time", "2026-10-23T00:00:00Z"), ("availability", "ACTIVE"),
])
def test_existing_device_identity_and_lifecycle_are_read_only(tables, field, value):
    result = parse(tables, "devices", [{**tables["device.csv"][0], field: value}])
    assert_atomic_failure(tables, result)
    assert any(issue["field"] == field for issue in result["issues"])


@pytest.mark.parametrize("field,value", [
    ("echelon", "I"), ("node_subtype", "I_FIXED"), ("parent_node_id", "N2"),
    ("device_class", "HF"), ("mobility", "VEHICLE"), ("status", "NORMAL"),
    ("is_key", "true"), ("relay_capable", "false"), ("node_role", "RELAY"),
])
def test_existing_node_structure_is_read_only(tables, field, value):
    result = parse(tables, "nodes", [{**tables["node.csv"][0], field: value}])
    assert_atomic_failure(tables, result)
    assert any(issue["field"] == field for issue in result["issues"])


def test_cross_table_insert_resolves_after_all_sheets_are_merged(tables):
    node = {**tables["node.csv"][0], "node_id": "N2", "parent_node_id": "N3"}
    parent = {**node, "node_id": "N3", "parent_node_id": "N1"}
    device = {**tables["device.csv"][0], "device_id": "D2", "node_id": "N2", "availability": "ACTIVE", "tx_power_dbm": "40"}
    result = parse_xlsx(tables, rows={"nodes": [node, parent], "devices": [device]}, order=["devices", "nodes", "interference_sources"])
    assert not result["issues"], result["issues"]
    assert [row["action"] for row in result["rows"]] == ["NEW", "NEW", "NEW"]
    assert len(result["tables"]["device.csv"]) == 2
    assert result["tables"]["device.csv"][1]["availability"] == "ACTIVE"
    assert all(isinstance(value, str) for rows in result["tables"].values() for row in rows for value in row.values())
    assert all(diff["before"] is None for diff in result["rows"][0]["differences"])


@pytest.mark.parametrize("parent", ["N2", "missing", "N1;N1", ";N1"])
def test_new_node_parent_constraints(tables, parent):
    node = {**tables["node.csv"][0], "node_id": "N2", "parent_node_id": parent}
    assert_atomic_failure(tables, parse(tables, "nodes", [node]))


def test_parent_cycle_and_unrelated_legacy_cycle(tables):
    nodes = [{**tables["node.csv"][0], "node_id": "N2", "parent_node_id": "N3"},
             {**tables["node.csv"][0], "node_id": "N3", "parent_node_id": "N2"}]
    result = parse(tables, "nodes", nodes)
    assert_atomic_failure(tables, result)
    assert {issue["row"] for issue in result["issues"] if "循环" in issue["message"]} == {2, 3}
    tables["node.csv"].extend(nodes)
    result = parse(tables, "nodes", [{**tables["node.csv"][0], "node_name": "更名"}])
    assert not result["issues"]


@pytest.mark.parametrize("field,value", [("device_class", "UHF"), ("mobility", "MANPACK"), ("echelon", "I"), ("node_role", "X")])
def test_new_node_structure_is_fully_validated(tables, field, value):
    row = {**tables["node.csv"][0], "node_id": "N2", field: value}
    assert_atomic_failure(tables, parse(tables, "nodes", [row]))


@pytest.mark.parametrize("field,value", [
    ("lon", "114.046499999"), ("lat", "37.989000001"), ("antenna_height_m", "-1"),
    ("tilt_deg", "91"), ("azimuth_deg", "360"), ("work_freq_khz", "0"),
    ("work_freq_khz", "30000.001"), ("tx_power_dbm", "43.06"),
])
def test_changed_numeric_bounds(tables, field, value):
    kind = "nodes" if field in {"lon", "lat"} else "devices"
    row = {**tables[parser.TABLES[kind]][0], field: value}
    assert_atomic_failure(tables, parse(tables, kind, [row]))


def test_power_rounding_is_only_preserved_for_existing_numeric_value(tables):
    result = parse(tables, "devices", [{**tables["device.csv"][0], "tx_power_dbm": "43.050"}])
    assert not result["issues"] and result["warnings"]
    assert result["rows"][0]["differences"] == [{"field": "tx_power_dbm", "before": "43.05", "after": "43.050"}]
    result = parse(tables, "devices", [{**tables["device.csv"][0], "tx_power_dbm": "43.04"}])
    assert_atomic_failure(tables, result)
    new = {**tables["device.csv"][0], "device_id": "D2", "availability": ""}
    assert_atomic_failure(tables, parse(tables, "devices", [new]))
    new["tx_power_dbm"] = "43.01"
    result = parse(tables, "devices", [new])
    assert not result["issues"]
    assert result["tables"]["device.csv"][1]["availability"] == "ACTIVE"


@pytest.mark.parametrize("target", ["node", "model", "antenna", "class", "antenna_class"])
def test_device_references_and_class_compatibility(tables, target):
    row = {**tables["device.csv"][0], "device_id": "D2", "tx_power_dbm": "40", "availability": "ACTIVE"}
    if target in {"node", "model", "antenna"}:
        row[target + "_id"] = "missing"
    elif target == "class":
        tables["node.csv"][0]["device_class"] = "VUHF"
    else:
        tables["antenna_model.csv"][0]["device_class"] = "VUHF"
    assert_atomic_failure(tables, parse(tables, "devices", [row]))


@pytest.mark.parametrize("field,value", [
    ("center_freq_khz", "0"), ("bandwidth_khz", "-1"), ("duty_cycle", "1.0001"),
    ("azimuth_deg", "360"), ("interference_type", "UNKNOWN"), ("pattern_type", "X"),
])
def test_interference_changes_are_validated(tables, field, value):
    result = parse(tables, "interference_sources", [{**tables["interference_source.csv"][0], field: value}])
    assert_atomic_failure(tables, result)


def test_directional_interference_requires_azimuth(tables):
    row = {**tables["interference_source.csv"][0], "pattern_type": "DIRECTIONAL"}
    assert_atomic_failure(tables, parse(tables, "interference_sources", [row]))
    row["azimuth_deg"] = "0"
    assert not parse(tables, "interference_sources", [row])["issues"]


def test_unchanged_legacy_doubts_do_not_reject_whole_database(tables):
    tables["interference_source.csv"][0]["lon"] = "113.7"
    tables["device.csv"][0]["node_id"] = "missing"
    result = parse(tables, "nodes", [{**tables["node.csv"][0], "remark": "修改备注"}])
    assert not result["issues"]
    result = parse_xlsx(tables)
    assert not result["issues"]
    assert all(row["action"] == "UNCHANGED" for row in result["rows"])


@pytest.mark.parametrize("kind", list(parser.HEADERS))
@pytest.mark.parametrize("format", ["csv", "xlsx"])
def test_templates_use_current_rows_and_roundtrip_without_invention(tables, kind, format):
    tables["node.csv"][0]["remark"] = '=HYPERLINK("example")'
    original = deepcopy(tables)
    data, mime, filename = parser.make_template(tables, format, kind)
    assert "SYNTHETIC_INTEGRATION" in filename and "非生产非实测" in filename
    assert not data.startswith(b"\xef\xbb\xbf")
    result = parser.parse_import(tables, filename, data, kind, TEST_BOUNDS)
    assert not result["issues"], result["issues"]
    assert result["tables"] == tables == original
    assert all(row["action"] == "UNCHANGED" for row in result["rows"])
    if format == "xlsx":
        assert mime == parser.XLSX_MIME
        book = load_workbook(BytesIO(data), read_only=True)
        assert book.sheetnames == ["nodes", "devices", "interference_sources", "说明"]
        assert "SYNTHETIC_INTEGRATION" in dict(book["说明"].values)["数据来源"]
        assert book["nodes"]["P2"].data_type == "s"
        book.close()
    else:
        assert mime.startswith("text/csv")


def test_template_includes_current_device_ancestors(tables):
    tables["node.csv"].append({**tables["node.csv"][0], "node_id": "parent"})
    tables["node.csv"][0]["parent_node_id"] = "parent"
    data, _, filename = parser.make_template(tables, "xlsx", "devices")
    result = parser.parse_import(tables, filename, data, "devices", TEST_BOUNDS)
    assert not result["issues"]
    assert {row["id"] for row in result["rows"] if row["sheet"] == "nodes"} == {"N1", "parent"}


@pytest.mark.parametrize("attack", ["formula", "hidden_sheet", "hidden_row", "hidden_column", "extra", "missing", "date", "error", "boolean_numeric", "merged", "notes_formula"])
def test_xlsx_rejects_unsafe_or_ambiguous_content(tables, attack):
    def mutate(book):
        if attack == "formula":
            book["devices"]["E2"] = "=1+2"
        elif attack == "hidden_sheet":
            book["devices"].sheet_state = "hidden"
        elif attack == "hidden_row":
            book["devices"].row_dimensions[2].hidden = True
        elif attack == "hidden_column":
            book["devices"].column_dimensions["E"].hidden = True
        elif attack == "extra":
            book.create_sheet("other")
        elif attack == "missing":
            book.remove(book["devices"])
        elif attack == "date":
            book["devices"]["K2"] = datetime(2026, 10, 23)
        elif attack == "error":
            book["devices"]["E2"] = "#DIV/0!"
        elif attack == "boolean_numeric":
            book["devices"]["E2"] = True
        elif attack == "merged":
            book["devices"].merge_cells("A2:B2")
        elif attack == "notes_formula":
            book.create_sheet("说明")["A1"] = "=1+2"
    assert_atomic_failure(tables, parse_xlsx(tables, mutate=mutate))


def test_xlsx_accepts_native_numbers_and_booleans_and_keeps_strings(tables):
    def mutate(book):
        book["devices"]["E2"] = 12.5
        book["interference_sources"]["O2"] = True
    result = parse_xlsx(tables, mutate=mutate)
    assert not result["issues"]
    assert result["tables"]["device.csv"][0]["antenna_height_m"] == "12.5"
    assert result["tables"]["interference_source.csv"][0]["active"] == "true"


def rewrite_archive(data, transform, extra=None):
    output = BytesIO()
    with ZipFile(BytesIO(data)) as source, ZipFile(output, "w", ZIP_DEFLATED) as target:
        for item in source.infolist():
            target.writestr(item.filename, transform(item.filename, source.read(item)))
        if extra:
            for name, content in extra.items():
                target.writestr(name, content)
    return output.getvalue()


@pytest.mark.parametrize("title", ["bad\n\x7f\x85\u202esheet", "\U00020000" * 100 + "\x85"])
def test_hidden_invalid_sheet_names_are_sanitized_only_in_diagnostics(tables, title):
    def transform(name, raw):
        if name != "xl/workbook.xml":
            return raw
        root = ET.fromstring(raw)
        sheet = next(element for element in root.iter() if element.tag.rsplit("}", 1)[-1] == "sheet")
        sheet.set("name", title)
        sheet.set("state", "hidden")
        return ET.tostring(root, encoding="utf-8")

    data = rewrite_archive(workbook_bytes(tables), transform)
    result = parser.parse_import(tables, "input.xlsx", data, "nodes", TEST_BOUNDS)
    assert_atomic_failure(tables, result)
    assert result["rows"] == []
    issue = result["issues"][0]
    assert_displayable_issue(issue)
    assert issue["row"] == 1 and issue["field"] == "sheet"
    if "\n" in title:
        assert issue["sheet"] == r"bad\u000a\u007f\u0085\u202esheet"
    else:
        assert "已截断" in issue["sheet"] and "原长201" in issue["sheet"]


def test_xlsx_row_5002_is_preserved_as_diagnostic_position(tables):
    def transform(name, raw):
        if name == "xl/worksheets/sheet1.xml":
            return raw.replace(b"</sheetData>", b'<row r="5002"><c r="A5002" t="inlineStr"><is><t>overflow</t></is></c></row></sheetData>')
        return raw

    data = rewrite_archive(workbook_bytes(tables), transform)
    result = parser.parse_import(tables, "input.xlsx", data, "nodes", TEST_BOUNDS)
    assert_atomic_failure(tables, result)
    assert result["rows"] == []
    assert result["issues"][0]["sheet"] == "nodes" and result["issues"][0]["row"] == 5002
    assert_displayable_issue(result["issues"][0])


@pytest.mark.parametrize("attack", ["vba", "entity", "sparse_row", "empty_sparse_row", "bad_dimension", "external", "duplicate_cell", "duplicate_row", "bad_shared_string"])
def test_xlsx_zip_and_xml_security(tables, attack):
    def transform(name, raw):
        if attack == "entity" and name == "xl/worksheets/sheet1.xml":
            return b'<!DOCTYPE x [<!ENTITY a "expanded">]>' + raw
        if attack == "sparse_row" and name == "xl/worksheets/sheet1.xml":
            return raw.replace(b"</sheetData>", b'<row r="1048576"><c r="A1048576" t="inlineStr"><is><t>hidden</t></is></c></row></sheetData>')
        if attack == "empty_sparse_row" and name == "xl/worksheets/sheet1.xml":
            return raw.replace(b"</sheetData>", b'<row r="1048576"/></sheetData>')
        if attack == "duplicate_cell" and name == "xl/worksheets/sheet1.xml":
            return raw.replace(b'</row></sheetData>', b'<c r="A2" t="inlineStr"><is><t>other</t></is></c></row></sheetData>')
        if attack == "duplicate_row" and name == "xl/worksheets/sheet1.xml":
            return raw.replace(b"</sheetData>", b'<row r="2"/></sheetData>')
        if attack == "bad_shared_string" and name == "xl/worksheets/sheet1.xml":
            return raw.replace(b'<c r="A2" t="inlineStr"><is><t>N1</t></is></c>', b'<c r="A2" t="s"><v>999999</v></c>')
        if attack == "bad_dimension" and name == "xl/worksheets/sheet1.xml":
            return raw.replace(b'ref="A1:P2"', b'ref="A1:A1"')
        if attack == "external" and name == "_rels/.rels":
            return raw.replace(b"</Relationships>", b'<Relationship Type="x" Id="unsafe" Target="https://example.invalid" TargetMode="External"/></Relationships>')
        return raw
    data = rewrite_archive(workbook_bytes(tables), transform, {"xl/vbaProject.bin": b"macro"} if attack == "vba" else None)
    result = parser.parse_import(tables, "input.xlsx", data, "nodes", TEST_BOUNDS)
    if attack == "bad_dimension":
        assert not result["issues"]
        assert len(result["rows"]) == 3
    else:
        assert_atomic_failure(tables, result)


def test_zip_expansion_limit_is_checked_before_openpyxl(tables, monkeypatch):
    data = rewrite_archive(workbook_bytes(tables), lambda name, raw: raw, {"padding.bin": b"A" * (parser.MAX_ZIP_BYTES + 1)})
    monkeypatch.setattr(parser, "load_workbook", lambda *args, **kwargs: pytest.fail("must reject before openpyxl"))
    assert_atomic_failure(tables, parser.parse_import(tables, "input.xlsx", data, "nodes", TEST_BOUNDS))


def test_zip_rejects_case_aliases_before_xml_can_be_shadowed(tables):
    data = rewrite_archive(workbook_bytes(tables), lambda name, raw: raw,
                           {"xl/worksheets/SHEET1.xml": b"<worksheet/>"})
    assert_atomic_failure(tables, parser.parse_import(tables, "input.xlsx", data, "nodes", TEST_BOUNDS))


def test_file_encoding_extension_size_and_record_limits(tables):
    for filename, data, kind in [("input.csv", b"\xff", "nodes"), ("input.xlsm", b"data", "nodes"),
                                  ("input.xlsx", b"broken zip", "nodes"), ("input.csv", b"", "nodes"),
                                  ("input.csv", b"x" * (parser.MAX_FILE_BYTES + 1), "nodes"),
                                  ("input.csv", b"x", "topology_links")]:
        assert_atomic_failure(tables, parser.parse_import(tables, filename, data, kind, TEST_BOUNDS))
    row = tables["node.csv"][0]
    data = csv_bytes("nodes", [{**row, "node_id": "N" + str(i)} for i in range(parser.MAX_ROWS + 1)])
    invalid = parser.parse_import(tables, "input.csv", data, "nodes", TEST_BOUNDS)
    assert_atomic_failure(tables, invalid)
    assert invalid["rows"] == []
    assert invalid["issues"][0]["sheet"] == "nodes" and invalid["issues"][0]["row"] == 5002
    assert_displayable_issue(invalid["issues"][0])
    data = csv_bytes("nodes", [{**row, "node_id": "N" + str(i)} for i in range(parser.MAX_ROWS)])
    result = parser.parse_import(tables, "input.csv", data, "nodes", TEST_BOUNDS)
    assert not result["issues"], result["issues"][:2]
    assert len(result["rows"]) == parser.MAX_ROWS


def test_xlsx_total_limit_is_across_all_three_tables(tables, monkeypatch):
    monkeypatch.setattr(parser, "MAX_ROWS", 2)
    result = parse_xlsx(tables)
    assert_atomic_failure(tables, result)
    assert any("合计" in issue["message"] for issue in result["issues"])


def test_invalid_timestamp_and_new_device_lifecycle(tables):
    row = {**tables["device.csv"][0], "device_id": "D2", "tx_power_dbm": "40", "availability": "ACTIVE"}
    for stamp in ["2026-02-30T00:00:00Z", "2026-10-23", "2026-10-23T00:00:00"]:
        assert_atomic_failure(tables, parse(tables, "devices", [{**row, "install_time": stamp}]))
    assert_atomic_failure(tables, parse(tables, "devices", [{**row, "availability": "OFF_SHELF"}]))
    assert not parse(tables, "devices", [{**row, "install_time": "2026-10-23T00:00:00+08:00"}])["issues"]


def test_parser_has_no_runtime_file_or_network_io(tables, monkeypatch):
    import builtins
    import socket
    from openpyxl.worksheet import _writer

    data = csv_bytes("nodes", [{**tables["node.csv"][0], "remark": "纯函数"}])
    xlsx = workbook_bytes(tables)
    monkeypatch.setattr(builtins, "open", lambda *args, **kwargs: pytest.fail("unexpected file IO"))
    monkeypatch.setattr(socket, "socket", lambda *args, **kwargs: pytest.fail("unexpected network IO"))
    monkeypatch.setattr(_writer, "create_temporary_file", lambda *args, **kwargs: pytest.fail("unexpected temporary file IO"))
    assert not parser.parse_import(tables, "input.csv", data, "nodes", TEST_BOUNDS)["issues"]
    assert not parser.parse_import(tables, "input.xlsx", xlsx, "nodes", TEST_BOUNDS)["issues"]
    for format in ("csv", "xlsx"):
        data, _, filename = parser.make_template(tables, format, "nodes")
        assert not parser.parse_import(tables, filename, data, "nodes", TEST_BOUNDS)["issues"]


@pytest.mark.parametrize("attack", ["formula", "date", "hidden_row"])
def test_xlsx_unsafe_data_has_sheet_and_row_location(tables, attack):
    def mutate(book):
        if attack == "formula":
            book["devices"]["E2"] = "=1+2"
        elif attack == "date":
            book["devices"]["K2"] = datetime(2026, 10, 23)
        else:
            book["devices"].row_dimensions[2].hidden = True
    result = parse_xlsx(tables, mutate=mutate)
    assert_atomic_failure(tables, result)
    issue = result["issues"][0]
    assert issue["sheet"] == "devices" and issue["row"] == 2
    if attack in {"formula", "date"}:
        assert issue["objectId"] == "D1"
        assert issue["field"] == ("antenna_height_m" if attack == "formula" else "install_time")
