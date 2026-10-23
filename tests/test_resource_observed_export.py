"""同版构造状态导出；只使用隔离测试库，不更新资源历史。"""
from copy import deepcopy
import csv
from io import BytesIO, StringIO
from pathlib import Path
import sys

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pydantic import ValidationError
from sqlalchemy import create_engine, func, select
from sqlalchemy.orm import sessionmaker

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "algorithm-service"))
from app import database, device_status_service as statuses, integration_identity_service as identity
from app import resource_dataset_service as datasets, resource_export_service as exports
from app import resource_observation_projection as projection
from app.database import Base
from app.device_status_models import DeviceStatusImportOperation, DeviceStatusRecord, DeviceStatusSnapshot
from app.device_status_schemas import DeviceStatusImport
from app.resource_dataset_models import ResourceDatasetCurrent, ResourceDatasetOperation, ResourceDatasetVersion
from app.resource_export_router import router
from app.resource_export_schemas import ResourceExportRequest


def request(revision="resource-0123456789abcdef", **patch):
    return {"context": "ledger", "revision": revision, "format": "csv", "kinds": ["devices"], "range": "all",
            "status_mode": "observed", "observation_snapshot_id": "selected-snapshot", **patch}


@pytest.fixture
def api(tmp_path, monkeypatch):
    path = tmp_path / "observed-export.sqlite"
    engine = create_engine("sqlite:///" + str(path), connect_args={"check_same_thread": False})
    models = (ResourceDatasetVersion, ResourceDatasetCurrent, ResourceDatasetOperation,
              DeviceStatusSnapshot, DeviceStatusRecord, DeviceStatusImportOperation)
    Base.metadata.create_all(engine, tables=[model.__table__ for model in models])
    factory = sessionmaker(bind=engine, expire_on_commit=False)
    for module in (database, datasets, statuses, projection):
        monkeypatch.setattr(module, "get_session_factory", lambda: factory)
    monkeypatch.setenv("RESOURCE_DATASET_MODE", "database")
    monkeypatch.setenv("DATABASE_URL", "sqlite:///" + str(path))
    datasets.initialize()
    baseline = datasets.current()
    records = []
    for index in (0, 1, 40):
        device = baseline["tables"]["device.csv"][index]
        records.append({"device_id": device["device_id"], "node_id": device["node_id"], "model_id": device["model_id"],
            "status": None if index == 1 else "FAULT", "alert": None if index == 1 else index == 0,
            "observed_at": None if index == 1 else "2026-10-23T10:00:00+08:00",
            "unknown_reasons": ["来源未提供"] if index == 1 else []})
    receipt = statuses.import_snapshot(DeviceStatusImport.model_validate({
        "schema_version": "device-status-snapshot-a1", "context": "ledger", "request_id": "observed-export-source",
        "base_revision": baseline["revision"], "data_origin": "SYNTHETIC_INTEGRATION", "source_label": '=构造,"来源"',
        "provenance_note": "只验证构造状态，不代表实测或在线", "records": records,
    }), "export-test-actor")

    def require(req, permission, writing=False):
        if req.headers.get("x-test-csrf") != "valid":
            raise identity.IdentityError("CSRF rejected", 1105, 403)
        assert permission == "resources.view" and writing
        return {"account_id": "export-test-actor"}

    monkeypatch.setattr(identity, "require", require)
    app = FastAPI()
    app.include_router(router)
    with TestClient(app) as client:
        yield client, factory, baseline, receipt, projection
    engine.dispose()


def post(api, download=False, **patch):
    client, _, baseline, receipt, _ = api
    payload = request(**{"revision": baseline["revision"], "observation_snapshot_id": receipt["snapshot_id"], **patch})
    return client.post("/api/v1/resource/export" + ("" if download else "/preview"), json=payload, headers={"x-test-csrf": "valid"})


def csv_rows(response):
    assert response.status_code == 200, response.text
    return list(csv.DictReader(StringIO(response.content.decode("utf-8"))))


@pytest.mark.parametrize("patch", [
    {"observation_snapshot_id": None}, {"status_mode": "baseline"}, {"status_mode": "latest"},
    {"kinds": ["nodes"]}, {"kinds": ["interference_sources"]}, {"filter": {}},
    {"range": "selected", "selected_ids": ["D1"], "filter": {"status": "FAULT"}},
    {"range": "filtered", "filter": {"status_mode": "observed"}},
    {"range": "filtered", "filter": {"observation_snapshot_id": "snapshot"}},
    {"range": "filtered", "filter": {"status": "FAULT", "online": True}},
])
def test_observation_is_explicit_root_source_not_business_filter(patch):
    with pytest.raises(ValidationError):
        ResourceExportRequest.model_validate(request(**patch))


@pytest.mark.parametrize("scope", ["all", "selected", "filtered"])
def test_every_range_projects_same_snapshot_and_does_not_fill_unknowns(api, scope):
    _, _, baseline, receipt, _ = api
    chosen = [row["device_id"] for row in baseline["tables"]["device.csv"][:3]]
    patch = {"range": scope}
    if scope == "selected":
        patch["selected_ids"] = chosen
    elif scope == "filtered":
        patch["filter"] = {}
    preview = post(api, **patch).json()["data"]
    assert preview["status_mode"] == "observed" and preview["node_status_mode"] == "baseline"
    assert preview["observation"] == {key: receipt[key] if key != "source_label" else '=构造,"来源"'
        for key in exports.OBSERVATION_KEYS}
    rows = csv_rows(post(api, True, **patch))
    assert rows[0]["record_kind"] == "metadata" and rows[0]["device_id"] == ""
    assert rows[0]["snapshot_id"] == receipt["snapshot_id"]
    devices = {row["device_id"]: row for row in rows[1:]}
    assert len(devices) == preview["tables"][0]["count"] == (3 if scope == "selected" else 414)
    assert devices[chosen[0]]["status"] == "FAULT" and devices[chosen[0]]["alert"] == "true"
    assert devices[chosen[0]]["observed_at"] == "2026-10-23T02:00:00.000000Z"
    assert devices[chosen[1]]["status"] == devices[chosen[2]]["status"] == ""
    assert devices[chosen[2]]["alert"] == devices[chosen[2]]["observed_at"] == ""
    assert devices[chosen[2]]["record_kind"] == "device"
    assert devices[chosen[0]]["source_label"].startswith("~rp~")


def test_filtered_export_projects_entire_catalog_before_filtering(api):
    _, _, baseline, _, _ = api
    rows = csv_rows(post(api, True, range="filtered", filter={"status": "FAULT", "alert": "NO"}))
    assert [row["device_id"] for row in rows[1:]] == [baseline["tables"]["device.csv"][40]["device_id"]]
    unknown = csv_rows(post(api, True, range="filtered", filter={"status": "UNKNOWN", "alert": "UNKNOWN"}))
    assert len(unknown) - 1 == 412
    assert all(row["status"] == row["alert"] == "" for row in unknown[1:])


def test_empty_export_still_contains_snapshot_metadata(api):
    _, _, _, receipt, _ = api
    assert post(api, range="filtered", filter={"status": "NORMAL"}).json()["data"]["tables"][0]["count"] == 0
    rows = csv_rows(post(api, True, range="filtered", filter={"status": "NORMAL"}))
    assert len(rows) == 1 and rows[0]["record_kind"] == "metadata"
    assert rows[0]["snapshot_revision"] == receipt["snapshot_revision"]


def test_observed_xlsx_marks_node_closure_as_baseline(api):
    _, _, baseline, receipt, _ = api
    selected = baseline["tables"]["device.csv"][40]["device_id"]
    patch = {"format": "xlsx", "kinds": ["devices", "nodes"], "range": "selected", "selected_ids": [selected]}
    preview = post(api, **patch).json()["data"]
    response = post(api, True, **patch)
    assert response.status_code == 200, response.text
    assert response.headers["x-status-mode"] == "observed" and response.headers["x-node-status-mode"] == "baseline"
    book = load_workbook(BytesIO(response.content))
    assert list(next(book["devices"].values)) == exports.OBSERVED_DEVICE_HEADERS
    rows = list(book["devices"].values)
    device = dict(zip(rows[0], rows[1]))
    assert device["alert"] is False and device["status"] == "FAULT"
    assert book["nodes"].max_row - 1 == preview["added"] >= 1
    original = {row["node_id"]: row for row in baseline["tables"]["node.csv"]}
    nodes = list(book["nodes"].values)
    for values in nodes[1:]:
        node = dict(zip(nodes[0], values))
        assert node["status"] == original[node["node_id"]]["status"]
    notes = dict(book["说明"].values)
    assert notes["节点状态口径"] == "baseline" and notes["snapshot_revision"] == receipt["snapshot_revision"]
    assert notes["source_label"] == '=构造,"来源"'
    assert next(row[1] for row in book["说明"].iter_rows() if row[0].value == "source_label").data_type == "s"
    book.close()


def test_baseline_headers_response_and_resource_history_are_unchanged(api):
    client, factory, baseline, _, _ = api
    with factory() as session:
        before = tuple(session.scalar(select(func.count()).select_from(model)) for model in
            (ResourceDatasetVersion, ResourceDatasetOperation, DeviceStatusSnapshot, DeviceStatusRecord, DeviceStatusImportOperation))
    assert post(api).status_code == post(api, True).status_code == 200
    body = {key: value for key, value in request(baseline["revision"]).items() if key not in ("status_mode", "observation_snapshot_id")}
    preview = client.post("/api/v1/resource/export/preview", json=body, headers={"x-test-csrf": "valid"}).json()["data"]
    assert set(preview) == {"revision", "data_origin", "source", "tables", "added"}
    response = client.post("/api/v1/resource/export", json=body, headers={"x-test-csrf": "valid"})
    assert response.content.decode("utf-8").splitlines()[0] == ",".join(exports.HEADERS["devices"])
    assert "x-status-mode" not in response.headers
    assert csv_rows(response)[0]["status"] == baseline["tables"]["device.csv"][0]["status"]
    assert datasets.current() == baseline
    with factory() as session:
        assert tuple(session.scalar(select(func.count()).select_from(model)) for model in
            (ResourceDatasetVersion, ResourceDatasetOperation, DeviceStatusSnapshot, DeviceStatusRecord, DeviceStatusImportOperation)) == before


def test_missing_stale_and_failed_projection_never_export_baseline(api, monkeypatch):
    client, _, baseline, receipt, projection = api
    unknown = request(baseline["revision"], observation_snapshot_id="missing-snapshot")
    response = client.post("/api/v1/resource/export", json=unknown, headers={"x-test-csrf": "valid"})
    assert response.status_code == 404 and response.json()["code"] == 1002
    assert post(api, revision="resource-ffffffffffffffff").status_code == 409
    fake = deepcopy(baseline)
    fake["revision"] = "resource-ffffffffffffffff"
    monkeypatch.setattr(datasets, "current", lambda: fake)
    mismatch = request(fake["revision"], observation_snapshot_id=receipt["snapshot_id"])
    response = client.post("/api/v1/resource/export", json=mismatch, headers={"x-test-csrf": "valid"})
    assert response.status_code == 409 and response.json()["code"] == 1205
    monkeypatch.setattr(datasets, "current", lambda: baseline)
    monkeypatch.setattr(projection, "project_device_items", lambda *_: ([], {}))
    response = post(api, True)
    assert response.status_code == 503 and "content-disposition" not in response.headers


def test_selected_unknown_device_is_not_partial_observed_export(api):
    response = post(api, True, range="selected", selected_ids=["DV-0001", "missing"])
    assert response.status_code == 404 and response.json()["code"] == 1002
