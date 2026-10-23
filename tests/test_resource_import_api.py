"""资源导入 API：独立 SQLite、真实服务编排，不启动监听服务。"""
import base64
from copy import deepcopy
import csv
from email.message import Message
from email.utils import collapse_rfc2231_value
from hashlib import sha256
from io import BytesIO, StringIO
import json
from pathlib import Path
import re
import sys
from uuid import uuid4

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine, func, select
from sqlalchemy.orm import sessionmaker

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "algorithm-service"))
from app import database, integration_identity_service as identity
from app import resource_dataset_service as datasets, resource_import_service as imports
from app import resource_import_parser as parser
from app.database import Base
from app.main import app
from app.integration_identity_models import IntegrationSession
from app.planning_plan_models import PlanningPlan
from app.resource_dataset_models import ResourceDatasetCurrent, ResourceDatasetOperation, ResourceDatasetVersion
from app.resource_export_service import HEADERS, TABLES

PREFIX = "/api/v1/resource/import"


@pytest.fixture
def import_api(monkeypatch, request):
    directory = ROOT / "data/runtime/tmp" / ("import-api-" + uuid4().hex)
    directory.mkdir(parents=True, exist_ok=False)
    path = directory / "import.sqlite"
    request.node.user_properties.append(("database", str(path)))
    engine = create_engine("sqlite:///" + str(path), connect_args={"check_same_thread": False, "timeout": 30})
    Base.metadata.create_all(engine, tables=[model.__table__ for model in (
        ResourceDatasetVersion, ResourceDatasetCurrent, ResourceDatasetOperation, IntegrationSession, PlanningPlan)])
    factory = sessionmaker(bind=engine, expire_on_commit=False)
    for name, module in tuple(sys.modules.items()):
        if name.startswith("app.") and hasattr(module, "get_session_factory"):
            monkeypatch.setattr(module, "get_session_factory", lambda: factory)
    monkeypatch.setattr(database, "get_engine", lambda: engine)
    monkeypatch.setenv("DATABASE_URL", "sqlite:///" + str(path))
    monkeypatch.setenv("FRONTEND_INTEGRATION_DEMO", "1")
    monkeypatch.setenv("RESOURCE_DATASET_MODE", "database")
    try:
        baseline = datasets.initialize()
        with TestClient(app) as client:
            yield client, factory, baseline, path
    finally:
        engine.dispose()


def headers(client):
    response = client.get("/api/v1/auth/csrf")
    assert response.status_code == 200, response.text
    data = response.json()["data"]
    return {"X-Session-Id": data["session_id"], "X-CSRF-Token": data["csrf_token"]}


def switch(client, account="integration-editor"):
    response = client.post("/api/v1/auth/integration/switch", headers=headers(client), json={"accountId": account})
    assert response.status_code == 200, response.text
    return response.json()["data"]


def counts(factory):
    with factory() as session:
        return tuple(session.scalar(select(func.count()).select_from(model))
                     for model in (ResourceDatasetVersion, ResourceDatasetOperation))


def csv_data(kind, rows):
    output = StringIO(newline="")
    writer = csv.writer(output)
    writer.writerow(HEADERS[kind])
    for row in rows:
        values = [row.get(field, "ACTIVE" if field == "availability" else "") for field in HEADERS[kind]]
        writer.writerow(["~rp~" + json.dumps(value, ensure_ascii=False) if isinstance(value, str) and
                         (re.match(r"^\s*[=+\-@]", value) or value.startswith("~rp~")) else value for value in values])
    return output.getvalue().encode("utf-8")


def body(baseline, kind="devices", rows=None, data=None, filename="resources.csv", context="ledger"):
    if data is None:
        data = csv_data(kind, rows if rows is not None else baseline["tables"][TABLES[kind]][:1])
    return {"context": context, "base_revision": baseline["revision"], "kind": kind,
            "file": {"name": filename, "base64": base64.b64encode(data).decode("ascii")}}


def preview(client, payload):
    return client.post(PREFIX + "/preview", headers=headers(client), json=payload)


def commit_body(client, payload, request_id="import-save"):
    response = preview(client, payload)
    assert response.status_code == 200, response.text
    return {**deepcopy(payload), "request_id": request_id, "signature": response.json()["data"]["signature"]}


def commit(client, payload):
    return client.post(PREFIX + "/commit", headers=headers(client), json=payload)


def assert_no_resource_write(factory, baseline, expected=(1, 0)):
    assert counts(factory) == expected
    assert datasets.current()["revision"] == baseline["revision"]
    assert datasets.current()["tables"] == baseline["tables"]


@pytest.mark.parametrize("format,kind", [("csv", "nodes"), ("csv", "devices"), ("csv", "interference_sources"), ("xlsx", "devices")])
def test_template_preview_and_noop_commit_are_read_only(import_api, format, kind):
    client, factory, baseline, _ = import_api
    switch(client)
    raw_hashes = {meta["path"]: sha256((ROOT / meta["path"]).read_bytes()).hexdigest()
                  for meta in baseline["source_files"].values()}
    response = client.get(PREFIX + "/template", params={"context": "ledger", "revision": baseline["revision"], "format": format, "kind": kind})
    assert response.status_code == 200, response.text
    assert response.headers["x-resource-revision"] == baseline["revision"]
    assert response.headers["x-data-origin"] == "SYNTHETIC_INTEGRATION"
    assert response.headers["cache-control"] == "no-store"
    assert response.headers["x-content-type-options"] == "nosniff"
    for name, value in response.headers.raw:
        assert name.isascii() and value.isascii()
        assert not re.search(rb"[\x00-\x1f\x7f]", name + value)
    disposition = Message()
    disposition["Content-Disposition"] = response.headers["content-disposition"]
    assert disposition.get_content_disposition() == "attachment"
    names = [value for key, value in disposition.get_params(header="content-disposition") if key == "filename"]
    assert any(isinstance(value, str) and value.isascii() and value.endswith("." + format) for value in names)
    extended_names = [collapse_rfc2231_value(value) for value in names if isinstance(value, tuple)]
    assert len(extended_names) == 1
    assert "SYNTHETIC_INTEGRATION" in extended_names[0] and "非生产非实测" in extended_names[0]
    assert extended_names[0].endswith("." + format)
    expected_mime = "text/csv" if format == "csv" else parser.XLSX_MIME
    assert response.headers["content-type"].split(";", 1)[0] == expected_mime
    if format == "csv":
        assert "charset=utf-8" in response.headers["content-type"].lower()
    if format == "xlsx":
        book = load_workbook(BytesIO(response.content), read_only=True)
        assert book.sheetnames == ["nodes", "devices", "interference_sources", "说明"]
        assert dict(book["说明"].values)["数据来源"] == "SYNTHETIC_INTEGRATION"
        book.close()
    payload = body(baseline, kind, data=response.content, filename="template." + format)
    response = preview(client, payload)
    assert response.status_code == 200, response.text
    data = response.json()["data"]
    assert data["valid"] and data["issues"] == []
    assert re.fullmatch("[0-9a-f]{64}", data["signature"])
    assert all(row["action"] == "UNCHANGED" and row["differences"] == [] for row in data["rows"])
    assert data["candidate"]["version"] == baseline["revision"]
    assert data["changes"] == {"stations": [], "devices": []}
    assert_no_resource_write(factory, baseline)
    saved = commit(client, {**payload, "signature": data["signature"], "request_id": "template-noop"})
    assert saved.status_code == 200, saved.text
    receipt = saved.json()["data"]["receipt"]
    assert receipt["changed"] is False and receipt["affectedIds"] == []
    assert receipt["revision"] == receipt["readRevision"] == baseline["revision"]
    assert_no_resource_write(factory, baseline, (1, 1))
    assert all(sha256((ROOT / name).read_bytes()).hexdigest() == digest for name, digest in raw_hashes.items())


@pytest.mark.parametrize("kind,field,value,collection,canonical", [
    ("devices", "antenna_height_m", "19.75", "devices", "antennaHeightM"),
    ("nodes", "node_name", "API导入更名", "stations", "name"),
    ("interference_sources", "name", "API导入干扰源", "interferences", "name"),
])
def test_preview_then_commit_persists_only_file_changes_and_history(import_api, kind, field, value, collection, canonical):
    client, factory, baseline, _ = import_api
    switch(client)
    row = {**baseline["tables"][TABLES[kind]][0], field: value}
    identifier = row[HEADERS[kind][0]]
    payload = body(baseline, kind, [row])
    response = preview(client, payload)
    assert response.status_code == 200, response.text
    data = response.json()["data"]
    assert data["valid"] and data["rows"][0]["action"] == "UPDATE"
    if kind == "devices":
        assert data["warnings"]
        assert all(set(warning) == {"sheet", "row", "objectId", "field", "message"} for warning in data["warnings"])
        assert data["warnings"][0]["objectId"] == identifier
    assert any(diff["field"] == field for diff in data["rows"][0]["differences"])
    changed = data["changes"][collection][0]
    assert changed == {"id": identifier, canonical: float(value) if kind == "devices" else value}
    assert next(item for item in data["candidate"][collection] if item["id"] == identifier)["origin"] == "IMPORT"
    assert_no_resource_write(factory, baseline)
    response = commit(client, {**payload, "signature": data["signature"], "request_id": "actual-import"})
    assert response.status_code == 200, response.text
    saved = response.json()["data"]
    receipt = saved["receipt"]
    assert receipt["changed"] and receipt["revision"] != baseline["revision"]
    assert receipt["readRevision"] == receipt["revision"] == saved["result"]["resources"]["version"]
    assert receipt["affectedIds"] == [identifier]
    assert receipt["warnings"] == data["warnings"]
    expected = deepcopy(baseline["tables"])
    expected[TABLES[kind]][0].update({field: value, "origin": "IMPORT"})
    assert datasets.current()["tables"] == expected
    with factory() as session:
        assert datasets._decode(session.get(ResourceDatasetVersion, baseline["revision"]))["tables"] == baseline["tables"]
        operation = session.get(ResourceDatasetOperation, "actual-import")
        assert operation.actor_id == "integration-editor" and operation.revision == receipt["revision"]
    assert counts(factory) == (2, 1)


def test_xlsx_cross_table_new_records_preserve_unknown_status_and_import_origin(import_api):
    client, factory, baseline, _ = import_api
    switch(client)
    template, _, _ = parser.make_template(baseline["tables"], "xlsx", "devices")
    book = load_workbook(BytesIO(template))
    node_id, device_id = "IMPORT-NEW-NODE", "IMPORT-NEW-DEVICE"
    book["nodes"]["A2"] = node_id
    book["nodes"]["E2"] = ""
    book["nodes"]["H2"] = (parser.BOUNDS[0] + parser.BOUNDS[2]) / 2
    book["nodes"]["I2"] = (parser.BOUNDS[1] + parser.BOUNDS[3]) / 2
    book["nodes"]["N2"] = ""
    device = baseline["tables"]["device.csv"][0]
    model = next(row for row in baseline["tables"]["device_model.csv"] if row["model_id"] == device["model_id"])
    book["devices"]["A2"] = device_id
    book["devices"]["B2"] = node_id
    book["devices"]["H2"] = float(model["tx_power_max_dbm"]) - 1
    book["devices"]["J2"] = ""
    data = parser._save_memory(book)
    book.close()
    payload = body(baseline, data=data, filename="new-records.xlsx")
    response = preview(client, payload)
    assert response.status_code == 200, response.text
    previewed = response.json()["data"]
    assert previewed["valid"], previewed["issues"]
    assert {row["id"] for row in previewed["rows"] if row["action"] == "NEW"} == {node_id, device_id}
    assert_no_resource_write(factory, baseline)
    response = commit(client, {**payload, "signature": previewed["signature"], "request_id": "xlsx-new-records"})
    assert response.status_code == 200, response.text
    resources = response.json()["data"]["result"]["resources"]
    node = next(row for row in resources["stations"] if row["id"] == node_id)
    saved_device = next(row for row in resources["devices"] if row["id"] == device_id)
    assert node["status"] is node["alert"] is saved_device["status"] is saved_device["alert"] is None
    assert node["origin"] == saved_device["origin"] == "IMPORT"
    assert saved_device["availability"] == "ACTIVE" and saved_device["stationId"] == node_id
    assert counts(factory) == (2, 1)


def test_invalid_reference_batch_preview_is_not_savable(import_api):
    client, factory, baseline, _ = import_api
    switch(client)
    first, second = baseline["tables"]["device.csv"][:2]
    rows = [{**first, "antenna_height_m": "17.25"}, {**second, "model_id": "missing-model"}]
    payload = body(baseline, "devices", rows)
    response = preview(client, payload)
    assert response.status_code == 200, response.text
    data = response.json()["data"]
    assert data["valid"] is False
    assert any(issue["objectId"] == second["device_id"] and issue["row"] == 3 and issue["sheet"] == "devices" for issue in data["issues"])
    assert data["changes"] == {"stations": [], "devices": []}
    canonical = next(row for row in data["candidate"]["devices"] if row["id"] == first["device_id"])
    assert canonical["antennaHeightM"] == float(first["antenna_height_m"])
    assert_no_resource_write(factory, baseline)
    response = commit(client, {**payload, "signature": data["signature"], "request_id": "invalid-batch"})
    assert response.status_code == 422 and response.json()["code"] == 1001
    assert_no_resource_write(factory, baseline)


@pytest.mark.parametrize("kind,identifier", [("nodes", "BAD NODE"), ("devices", "BAD\tDEVICE"), ("interference_sources", "BAD\u3000SOURCE")])
def test_import_ids_with_whitespace_cannot_be_previewed_as_valid_or_committed(import_api, kind, identifier):
    client, factory, baseline, _ = import_api
    switch(client)
    id_field = HEADERS[kind][0]
    row = {**baseline["tables"][TABLES[kind]][0], id_field: identifier}
    payload = body(baseline, kind, [row])
    response = preview(client, payload)
    assert response.status_code == 200, response.text
    data = response.json()["data"]
    assert data["valid"] is False and data["rows"] == []
    assert len(data["issues"]) == 1
    assert data["issues"][0]["objectId"] == identifier.replace("\t", r"\u0009")
    assert data["issues"][0]["field"] == id_field
    assert data["issues"][0]["sheet"] == kind and data["issues"][0]["row"] == 2
    response = commit(client, {**payload, "signature": data["signature"], "request_id": "invalid-whitespace-id"})
    assert response.status_code == 422 and response.json()["code"] == 1001
    assert_no_resource_write(factory, baseline)


def test_commit_reparses_original_bytes_and_ignores_no_client_candidate(import_api, monkeypatch):
    client, factory, baseline, _ = import_api
    switch(client)
    row = {**baseline["tables"]["device.csv"][0], "antenna_height_m": "16.125"}
    payload = body(baseline, rows=[row])
    original_parser, calls = imports.parse_import, []

    def tracked(*args):
        calls.append(args[2])
        return original_parser(*args)

    monkeypatch.setattr(imports, "parse_import", tracked)
    request = commit_body(client, payload)
    assert len(calls) == 1
    assert commit(client, {**request, "candidate": {"devices": []}}).status_code == 422
    assert counts(factory) == (1, 0)
    response = commit(client, request)
    assert response.status_code == 200, response.text
    assert len(calls) == 2 and calls[0] == calls[1] == base64.b64decode(payload["file"]["base64"])


def test_replay_after_later_commit_returns_original_resources_even_after_reopen(import_api, monkeypatch):
    client, factory, baseline, path = import_api
    switch(client)
    device = baseline["tables"]["device.csv"][0]
    first_request = commit_body(client, body(baseline, rows=[{**device, "antenna_height_m": "20.125"}]), "first-import")
    first = commit(client, first_request).json()["data"]
    current = datasets.current()
    second_request = commit_body(client, body(current, rows=[{**current["tables"]["device.csv"][0], "antenna_height_m": "21.125"}]), "second-import")
    second = commit(client, second_request).json()["data"]
    assert first["receipt"]["revision"] != second["receipt"]["revision"]
    assert commit(client, first_request).json()["data"] == first
    reopened = create_engine("sqlite:///" + str(path), connect_args={"check_same_thread": False})
    replacement = sessionmaker(bind=reopened, expire_on_commit=False)
    try:
        for module in (database, datasets, identity):
            monkeypatch.setattr(module, "get_session_factory", lambda: replacement)
        assert commit(client, first_request).json()["data"] == first
        assert datasets.current()["revision"] == second["receipt"]["revision"]
    finally:
        reopened.dispose()
    original = next(row for row in first["result"]["resources"]["devices"] if row["id"] == device["device_id"])
    assert original["antennaHeightM"] == 20.125
    assert counts(factory) == (3, 2)


def test_signature_conflicts_stale_versions_and_actor_bound_idempotency(import_api, monkeypatch):
    client, factory, baseline, _ = import_api
    switch(client)
    row = baseline["tables"]["device.csv"][0]
    request = commit_body(client, body(baseline, rows=[{**row, "antenna_height_m": "14.125"}]))
    for field in ("context", "base_revision", "kind", "name", "bytes"):
        altered = deepcopy(request)
        if field == "context":
            altered[field] = "gis"
        elif field == "base_revision":
            altered[field] = "resource-0000000000000000"
        elif field == "kind":
            altered[field] = "nodes"
        elif field == "name":
            altered["file"]["name"] = "changed.csv"
        else:
            altered["file"]["base64"] = base64.b64encode(csv_data("devices", [{**row, "antenna_height_m": "15.125"}])).decode("ascii")
        response = commit(client, altered)
        assert response.status_code == 409 and response.json()["code"] == 1206
    first = commit(client, request)
    assert first.status_code == 200, first.text
    assert commit(client, request).json() == first.json()
    assert preview(client, body(baseline)).json()["code"] == 1205
    stale = commit(client, {**request, "request_id": "new-stale-request"})
    assert stale.status_code == 409 and stale.json()["code"] == 1205
    current = datasets.current()
    reused = commit_body(client, body(current, rows=[{**current["tables"]["device.csv"][0], "antenna_height_m": "16.125"}]), request["request_id"])
    assert commit(client, reused).json()["code"] == 1206
    monkeypatch.setitem(identity.CATALOG, "second-import-editor", ("另一联调编辑者", "EDITOR", list(identity.CATALOG["integration-editor"][2])))
    switch(client, "second-import-editor")
    response = commit(client, request)
    assert response.status_code == 409 and response.json()["code"] == 1206
    assert counts(factory) == (2, 1)


@pytest.mark.parametrize("name", ["../input.csv", "sub/input.csv", "sub\\input.csv", "C:\\input.csv", "C:input.csv", "input\x00.csv", "input\r\n.csv", "input\x7f.csv"])
def test_malicious_names_do_not_write(import_api, name):
    client, factory, baseline, _ = import_api
    switch(client)
    payload = body(baseline, filename=name)
    response = preview(client, payload)
    assert response.status_code == 422, response.text[:1000]
    assert response.json()["code"] == 1001
    assert_no_resource_write(factory, baseline)


@pytest.mark.parametrize("encoded", ["%%%", "YQ", "YQ==\n", "_w==", "!!!!", "", True, 123, "é"])
def test_invalid_base64_is_rejected_without_writes(import_api, encoded):
    client, factory, baseline, _ = import_api
    switch(client)
    payload = body(baseline)
    payload["file"]["base64"] = encoded
    response = preview(client, payload)
    assert response.status_code == 422, response.text
    assert_no_resource_write(factory, baseline)


@pytest.mark.parametrize("mutation", ["extra", "file_extra", "kind", "context", "name_type", "name_empty", "name_long", "signature"])
def test_strict_import_schema(import_api, mutation):
    client, factory, baseline, _ = import_api
    switch(client)
    payload = body(baseline)
    if mutation == "extra":
        payload["actor_id"] = "admin"
    elif mutation == "file_extra":
        payload["file"]["mime"] = "text/csv"
    elif mutation == "kind":
        payload["kind"] = "topology_links"
    elif mutation == "context":
        payload["context"] = "other"
    elif mutation == "name_type":
        payload["file"]["name"] = True
    elif mutation == "name_empty":
        payload["file"]["name"] = ""
    elif mutation == "name_long":
        payload["file"]["name"] = "x" * 201
    else:
        payload.update(request_id="bad-signature", signature="0" * 63)
    response = commit(client, payload) if mutation == "signature" else preview(client, payload)
    assert response.status_code == 422, response.text
    assert_no_resource_write(factory, baseline)


def test_permissions_csrf_expired_identity_and_template_query_validation(import_api, monkeypatch):
    client, factory, baseline, _ = import_api
    query = {"context": "ledger", "revision": baseline["revision"], "format": "csv", "kind": "nodes"}
    assert client.get(PREFIX + "/template", params=query).status_code == 403
    switch(client, "integration-reviewer")
    assert preview(client, body(baseline)).status_code == 403
    editor = switch(client)
    assert "resources.import" in editor["permissions"] and "resources.edit" not in editor["permissions"]
    request = commit_body(client, body(baseline))
    for endpoint, payload in (("preview", body(baseline)), ("commit", request)):
        assert client.post(PREFIX + "/" + endpoint, json=payload).status_code == 409
        assert client.post(PREFIX + "/" + endpoint, headers={**headers(client), "X-CSRF-Token": "bad"}, json=payload).status_code == 403
    for params in ({**query, "format": "xlsm"}, {**query, "extra": "x"}, {**query, "kind": "other"}):
        assert client.get(PREFIX + "/template", params=params).status_code == 422
    assert client.get(PREFIX + "/template", params=[*query.items(), ("revision", "another")]).status_code == 422
    assert client.get(PREFIX + "/template", params={**query, "revision": "resource-0000000000000000"}).json()["code"] == 1205
    last_headers = headers(client)
    monkeypatch.setenv("FRONTEND_INTEGRATION_DEMO", "0")
    assert client.post(PREFIX + "/preview", headers=last_headers, json=body(baseline)).status_code == 401
    assert_no_resource_write(factory, baseline)


def test_special_permission_and_resources_view_are_both_required(import_api, monkeypatch):
    client, factory, baseline, _ = import_api
    name, role, permissions = identity.CATALOG["integration-editor"]
    for excluded in ("resources.import", "resources.view"):
        monkeypatch.setitem(identity.CATALOG, "integration-editor", (name, role, [item for item in permissions if item != excluded]))
        switch(client)
        assert preview(client, body(baseline)).status_code == 403
    assert_no_resource_write(factory, baseline)


def test_file_size_limit_and_transaction_failure_have_no_partial_write(import_api, monkeypatch):
    client, factory, baseline, _ = import_api
    switch(client)
    assert preview(client, body(baseline, data=b"x" * (imports.MAX_BYTES + 1))).status_code == 422
    row = {**baseline["tables"]["device.csv"][0], "antenna_height_m": "18.125"}
    request = commit_body(client, body(baseline, rows=[row]))
    original = datasets._validate_snapshot

    def reject_changed(snapshot):
        original(snapshot)
        if "change" in snapshot:
            raise datasets.ResourceDatasetError("隔离测试模拟提交校验失败")

    monkeypatch.setattr(datasets, "_validate_snapshot", reject_changed)
    assert commit(client, request).status_code == 503
    assert_no_resource_write(factory, baseline)
