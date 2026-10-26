"""构造资源导出契约；测试数据库与原件均和主运行环境隔离。"""
from copy import deepcopy
import csv
from hashlib import sha256
from io import BytesIO, StringIO
from pathlib import Path
import sys

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine, func, select
from sqlalchemy.orm import sessionmaker

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "algorithm-service"))
from app import database, integration_identity_service as identity
from app import resource_dataset_service as datasets, resource_export_service as exports
from app.database import Base
from app.main import app
from app.resource_dataset_models import ResourceDatasetOperation, ResourceDatasetVersion


@pytest.fixture
def api(tmp_path, monkeypatch):
    engine = create_engine("sqlite:///" + str(tmp_path / "export.sqlite"), connect_args={"check_same_thread": False})
    Base.metadata.create_all(engine)
    factory = sessionmaker(bind=engine, expire_on_commit=False)
    for module in (database, datasets, identity):
        monkeypatch.setattr(module, "get_session_factory", lambda: factory)
    monkeypatch.setenv("FRONTEND_INTEGRATION_DEMO", "1")
    monkeypatch.setenv("RESOURCE_DATASET_MODE", "database")
    snapshot = datasets.initialize()
    with TestClient(app) as client:
        yield client, factory, snapshot
    engine.dispose()


def headers(client):
    data = client.get("/api/v1/auth/csrf").json()["data"]
    return {"X-Session-Id": data["session_id"], "X-CSRF-Token": data["csrf_token"]}


def query(snapshot, **patch):
    return {"context": "ledger", "revision": snapshot["revision"], "format": "csv",
            "kinds": ["devices"], "range": "all", **patch}


def post(client, body, download=False):
    return client.post("/api/v1/resource/export" + ("" if download else "/preview"), json=body, headers=headers(client))


def counts(factory):
    with factory() as session:
        return tuple(session.scalar(select(func.count()).select_from(model))
                     for model in (ResourceDatasetVersion, ResourceDatasetOperation))


def test_all_export_and_preview_are_read_only(api):
    client, factory, snapshot = api
    source = {name: sha256((ROOT / meta["path"]).read_bytes()).hexdigest() for name, meta in snapshot["source_files"].items()}
    expected = counts(factory)
    preview = post(client, query(snapshot))
    assert preview.status_code == 200, preview.text
    assert preview.json()["data"]["tables"] == [{"kind": "devices", "count": 414}]
    assert preview.json()["data"]["data_origin"] == "SYNTHETIC_INTEGRATION"
    result = post(client, query(snapshot), True)
    assert result.status_code == 200
    assert result.headers["content-disposition"].startswith("attachment;")
    assert snapshot["revision"] in result.headers["content-disposition"]
    assert result.headers["x-resource-revision"] == snapshot["revision"]
    rows = list(csv.DictReader(StringIO(result.content.decode("utf-8"))))
    assert len(rows) == 414 and rows[0]["device_id"] == "DV-0001"
    assert rows[0]["work_freq_khz"] == "" and rows[0]["availability"] == "ACTIVE"
    assert counts(factory) == expected
    assert datasets.current()["revision"] == snapshot["revision"]
    assert all(sha256((ROOT / snapshot["source_files"][name]["path"]).read_bytes()).hexdigest() == value for name, value in source.items())


def test_selected_xlsx_closes_nodes_and_keeps_sheet_order(api):
    client, _, snapshot = api
    body = query(snapshot, format="xlsx", kinds=["devices", "nodes", "interference_sources"], range="selected", selected_ids=["DV-0001", "DV-0002"])
    preview = post(client, body).json()["data"]
    assert preview["tables"][0] == {"kind": "devices", "count": 2}
    assert preview["tables"][1]["count"] == preview["added"] >= 1
    assert preview["tables"][2]["count"] == 0
    response = post(client, body, True)
    assert response.status_code == 200, response.text
    book = load_workbook(BytesIO(response.content))
    assert book.sheetnames == ["devices", "nodes", "interference_sources", "说明"]
    assert book["devices"].max_row == 3
    assert book["nodes"].max_row - 1 == preview["added"]
    assert dict(book["说明"].values)["资源版本"] == snapshot["revision"]
    assert book["devices"]["I2"].value is None
    assert book["nodes"]["M2"].value in (True, False)
    book.close()


@pytest.mark.parametrize("filter,params", [
    ({"device_class": "HF"}, {"device_class": "HF"}),
    ({"device_class": "VUHF", "status": "NORMAL", "alert": "NO"}, {"device_class": "VUHF", "status": "NORMAL", "alert": "NO"}),
    ({"mobility": "FIXED", "relay_capable": True}, {"mobility": "FIXED", "relay_capable": "true"}),
    ({"node_role": "FIXED_STATION"}, {"node_role": "FIXED_STATION"}),
    ({"availability": "OFF_SHELF"}, {"availability": "OFF_SHELF"}),
    ({"status": "UNKNOWN", "alert": "UNKNOWN"}, {"status": "UNKNOWN", "alert": "UNKNOWN"}),
    ({"search": "DV-0001"}, {"search": "DV-0001"}),
    ({"frequency": {"mode": "MODEL", "min_khz": 3000, "max_khz": 3000}}, {"frequency_mode": "MODEL", "min_khz": 3000, "max_khz": 3000}),
    ({"frequency": {"mode": "DEVICE", "presence": "PRESENT"}}, {"frequency_mode": "DEVICE", "frequency_presence": "PRESENT"}),
    ({"frequency": {"mode": "DEVICE", "presence": "MISSING"}}, {"frequency_mode": "DEVICE", "frequency_presence": "MISSING"}),
])
def test_list_and_export_filter_the_whole_set_identically(api, filter, params):
    client, _, snapshot = api
    body = query(snapshot, range="filtered", filter=filter)
    preview = post(client, body)
    assert preview.status_code == 200, preview.text
    page = client.get("/api/v1/resource/devices", params={"context": "ledger", "revision": snapshot["revision"], "page_size": 5, **params})
    assert page.status_code == 200, page.text
    assert page.json()["data"]["total"] == preview.json()["data"]["tables"][0]["count"]
    csv_rows = list(csv.DictReader(StringIO(post(client, body, True).content.decode("utf-8"))))
    assert len(csv_rows) == page.json()["data"]["total"]
    assert [r["device_id"] for r in csv_rows[:5]] == [r["device_id"] for r in page.json()["data"]["items"]]
    if filter == {"device_class": "HF"}:
        assert len(csv_rows) == 174


@pytest.mark.parametrize("patch", [
    {"format": "pdf"}, {"kinds": []}, {"kinds": ["devices", "devices"]}, {"kinds": ["nodes", "devices"]},
    {"kinds": ["topology_links"]}, {"context": "gis"}, {"actor": "admin"}, {"filter": {}}, {"selected_ids": []},
    {"range": "selected", "selected_ids": ["DV-0001", "DV-0001"]},
    {"range": "selected", "selected_ids": [True]},
    {"range": "selected", "selected_ids": ["D" + str(i) for i in range(51)]},
    {"range": "filtered", "filter": {"status_mode": "observed"}},
    {"range": "filtered", "filter": {"frequency": {"mode": "ALL", "min_khz": 0}}},
    {"range": "filtered", "filter": {"frequency": {"mode": "MODEL", "min_khz": 10, "max_khz": 1}}},
    {"range": "filtered", "filter": {"frequency": {"mode": "MODEL", "min_khz": True}}},
])
def test_invalid_export_rejected_without_writes(api, patch):
    client, factory, snapshot = api
    before = counts(factory)
    for download in (False, True):
        response = post(client, query(snapshot, **patch), download)
        assert response.status_code == 422 and response.json()["code"] == 1001
    assert counts(factory) == before


def test_security_stale_and_unknown_selection(api, monkeypatch):
    client, factory, snapshot = api
    body = query(snapshot)
    assert client.post("/api/v1/resource/export/preview", json=body).status_code == 409
    assert client.post("/api/v1/resource/export", json=body, headers={**headers(client), "X-CSRF-Token": "wrong"}).status_code == 403
    unknown = post(client, query(snapshot, range="selected", selected_ids=["DV-0001", "missing"]))
    assert unknown.status_code == 404 and unknown.json()["code"] == 1002
    stale = post(client, query(snapshot, revision="old"))
    assert stale.status_code == 409 and stale.json()["code"] == 1205
    stale = client.get("/api/v1/resource/devices", params={"revision": "old"})
    assert stale.status_code == 409 and stale.json()["code"] == 1205
    assert counts(factory) == (1, 0)
    last_headers = headers(client)
    monkeypatch.setenv("FRONTEND_INTEGRATION_DEMO", "0")
    assert client.post("/api/v1/resource/export/preview", json=body, headers=last_headers).status_code == 401


def test_formula_safety_and_missing_values(api, monkeypatch):
    client, _, snapshot = api
    fake = deepcopy(snapshot)
    fake["tables"]["node.csv"][0]["node_name"] = '=HYPERLINK("unsafe")'
    monkeypatch.setattr(datasets, "current", lambda *args, **kwargs: fake)
    text = post(client, query(snapshot, kinds=["nodes"]), True).content.decode("utf-8")
    first = list(csv.DictReader(StringIO(text)))[0]
    assert first["node_name"].startswith("~rp~")
    assert __import__("json").loads(first["node_name"][4:]) == fake["tables"]["node.csv"][0]["node_name"]
    book = load_workbook(BytesIO(post(client, query(snapshot, format="xlsx", kinds=["nodes"]), True).content))
    assert book["nodes"]["B2"].data_type == "s"
    assert book["nodes"]["B2"].value.startswith("=HYPERLINK")
    book.close()


def test_resource_change_between_preview_and_download_is_not_mixed(api):
    client, _, snapshot = api
    body = query(snapshot)
    assert post(client, body).status_code == 200
    switch = client.post("/api/v1/auth/integration/switch", json={"accountId": "integration-editor"}, headers=headers(client))
    assert switch.status_code == 200
    saved = client.post("/api/v1/resource/changes", headers=headers(client), json={"context": "ledger", "base_revision": snapshot["revision"], "request_id": "export-change",
        "stations": [], "devices": [{"id": "DV-0001", "antenna_height_m": 17.25}]})
    assert saved.status_code == 200, saved.text
    assert post(client, body, True).status_code == 409
    fresh = query(snapshot, revision=saved.json()["data"]["revision"])
    first = list(csv.DictReader(StringIO(post(client, fresh, True).content.decode("utf-8"))))[0]
    assert float(first["antenna_height_m"]) == 17.25


@pytest.mark.parametrize("params", [{"extra": "ignored"}, {"status_mode": "observed"}, {"page": "1.0"}, {"page": "true"},
    {"page_size": "11"}, {"relay_capable": "1"}, {"min_khz": "nan"}, {"frequency_mode": "MODEL", "frequency_presence": "PRESENT"}])
def test_device_query_does_not_ignore_bad_fields(api, params):
    client, _, snapshot = api
    response = client.get("/api/v1/resource/devices", params={"revision": snapshot["revision"], **params})
    assert response.status_code == 422 and response.json()["code"] == 1001


def test_duplicate_query_and_stable_sort(api):
    client, _, snapshot = api
    duplicate = client.get("/api/v1/resource/devices", params=[("revision", snapshot["revision"]), ("page", "1"), ("page", "2")])
    assert duplicate.status_code == 422
    response = client.get("/api/v1/resource/devices", params={"revision": snapshot["revision"], "sort": "id", "order": "desc", "page_size": 5})
    assert response.status_code == 200, response.text
    ids = [row["device_id"] for row in response.json()["data"]["items"]]
    assert ids == sorted((row["device_id"] for row in snapshot["tables"]["device.csv"]), reverse=True)[:5]


def test_file_limit_and_database_mode_fail_closed(api, monkeypatch):
    client, factory, snapshot = api
    monkeypatch.setattr(exports, "MAX_FILE_BYTES", 4)
    response = post(client, query(snapshot), True)
    assert response.status_code == 413 and response.json()["code"] == 1001
    assert "content-disposition" not in response.headers
    assert counts(factory) == (1, 0)
    monkeypatch.setenv("RESOURCE_DATASET_MODE", "file")
    assert post(client, query(snapshot)).status_code == 503
