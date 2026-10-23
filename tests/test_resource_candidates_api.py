"""GIS 候选真实 TestClient/SQLite 闭环；只使用每用例独立的临时资源库。"""
import base64
from copy import deepcopy
import csv
from email.message import Message
from email.utils import collapse_rfc2231_value
from io import BytesIO, StringIO
import json
from pathlib import Path
import re
import sys
import unicodedata
import xml.etree.ElementTree as ET
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "algorithm-service"))
from app import integration_identity_service as identity
from app import resource_candidate_service as candidates, resource_dataset_service as datasets
from app import resource_import_parser as parser
from app.resource_candidate_schemas import ASSESSMENT_FIELDS, CONFIG_FIELDS, CandidateLink
from app.resource_dataset_models import ResourceDatasetCurrent, ResourceDatasetOperation, ResourceDatasetVersion
from test_resource_workspace_api import counts, headers, switch, workspace_api

PREFIX = "/api/v1/resource/import"
SAVE = "/api/v1/resource/workspace/changes"
CONTEXT = "/api/v1/resource/edit-context"


def post(client, endpoint, payload):
    return client.post(endpoint, headers=headers(client), json=payload)


def ok(response):
    assert response.status_code == 200, response.text[:2000]
    assert response.json()["code"] == 0, response.text[:2000]
    return response.json()["data"]


def assert_error(response, status=422, code=1001):
    assert response.status_code == status, response.text[:1500]
    assert response.json()["code"] == code, response.text[:1500]


def stored(factory, revision=None):
    with factory() as session:
        if revision is None:
            revision = session.get(ResourceDatasetCurrent, datasets.SCOPE).revision
        row = session.get(ResourceDatasetVersion, revision)
        assert row is not None
        return deepcopy(row.snapshot)


def state(factory):
    return {"counts": counts(factory), "current": deepcopy(datasets.current()), "stored": stored(factory)}


def assert_no_write(factory, before):
    assert state(factory) == before


def assert_resource_files_and_tables(factory, baseline, revision=None):
    saved = stored(factory, revision)
    assert saved["source_files"] == baseline["source_files"]
    assert len(saved["tables"]) == 10
    assert saved["tables"] == baseline["tables"]
    return saved


def template(client, snapshot, format="csv"):
    response = client.get(PREFIX + "/template", params={"context": "gis", "kind": "topology_links",
                          "revision": snapshot["revision"], "format": format})
    assert response.status_code == 200, response.text[:1500]
    assert response.headers["x-resource-revision"] == snapshot["revision"]
    assert response.headers["x-data-origin"] == "SYNTHETIC_INTEGRATION"
    assert response.headers["cache-control"] == "no-store"
    assert response.headers["x-content-type-options"] == "nosniff"
    assert response.headers["content-type"].startswith("text/csv" if format == "csv" else candidates.XLSX_MIME)
    message = Message()
    message["Content-Disposition"] = response.headers["content-disposition"]
    filename = message.get_param("filename", header="Content-Disposition")
    if isinstance(filename, tuple):
        filename = collapse_rfc2231_value(filename)
    assert filename and filename.isascii() and filename.endswith("." + format)
    return response.content, filename


def csv_data(rows, head=candidates.HEADERS):
    output = StringIO(newline="")
    writer = csv.writer(output)
    writer.writerow(head)
    for row in rows:
        values = [row.get(field, "") for field in head]
        writer.writerow(["~rp~" + json.dumps(value, ensure_ascii=False) if isinstance(value, str)
                         and (re.match(r"^\s*[=+\-@]", value) or value.startswith("~rp~")) else value for value in values])
    return output.getvalue().encode("utf-8")


def import_body(snapshot, data, filename, context="gis"):
    return {"context": context, "kind": "topology_links", "base_revision": snapshot["revision"],
            "file": {"name": filename, "base64": base64.b64encode(data).decode("ascii")}}


def workspace_body(snapshot, changes, request_id="candidate-workspace", context="gis"):
    return {"context": context, "base_revision": snapshot["revision"], "request_id": request_id, "changes": changes}


def prepare(client, snapshot, format="csv", reference=False):
    data, filename = template(client, snapshot, format)
    if reference:
        if format == "csv":
            rows = list(csv.DictReader(StringIO(data.decode("utf-8-sig"))))
            rows[0].update(path_loss_db="0", rx_power_dbm="-100", is_available="false", link_state="FAULT")
            data = csv_data(rows)
        else:
            book = load_workbook(BytesIO(data))
            try:
                sheet = book["topology_links"]
                for field, value in (("path_loss_db", 0), ("rx_power_dbm", -100), ("is_available", False), ("link_state", "FAULT")):
                    sheet.cell(2, candidates.HEADERS.index(field) + 1, value)
                data = parser._save_memory(book)
            finally:
                book.close()
    payload = import_body(snapshot, data, filename)
    preview = ok(post(client, PREFIX + "/preview", payload))
    assert preview["valid"], preview["issues"]
    assert len(preview["candidate"]["candidateLinks"]) >= 1
    return payload, preview


def commit_body(payload, preview, request_id="candidate-import"):
    return {**deepcopy(payload), "signature": preview["signature"], "request_id": request_id}


def seed_reference(client, baseline, format="csv"):
    payload, preview = prepare(client, baseline, format, reference=True)
    request = commit_body(payload, preview)
    result = ok(post(client, PREFIX + "/commit", request))
    return datasets.current(), result["result"]["resources"]["candidateLinks"][0], request, result


def manual_from(link, identifier="MANUAL-CANDIDATE"):
    return {**deepcopy(link), "id": identifier, "source": {"kind": "MANUAL", "file": None, "row": None},
            "assessmentState": "PENDING", "reference": None, **dict.fromkeys(ASSESSMENT_FIELDS)}


def candidate_patch(snapshot, link, field):
    tables = snapshot["tables"]
    models = {row["model_id"]: row for row in tables["device_model.csv"]}
    devices = tables["device.csv"]
    if field in {"nodeAId", "deviceAId", "nodeBId", "deviceBId"}:
        alternate = next(row for row in devices if row["node_id"] not in {link["nodeAId"], link["nodeBId"]}
                         and models[row["model_id"]]["device_class"] == link["deviceClass"])
        end = "A" if field in {"nodeAId", "deviceAId"} else "B"
        return {"id": link["id"], f"node{end}Id": alternate["node_id"], f"device{end}Id": alternate["device_id"]}
    if field == "deviceClass":
        other = "VUHF" if link["deviceClass"] == "HF" else "HF"
        first = next(row for row in devices if models[row["model_id"]]["device_class"] == other)
        second = next(row for row in devices if row["node_id"] != first["node_id"] and models[row["model_id"]]["device_class"] == other)
        return {"id": link["id"], "nodeAId": first["node_id"], "nodeBId": second["node_id"],
                "deviceAId": first["device_id"], "deviceBId": second["device_id"], "deviceClass": other}
    if field == "freqKhz":
        endpoints = [next(row for row in devices if row["device_id"] == link[key]) for key in ("deviceAId", "deviceBId")]
        minimum = max(float(models[row["model_id"]]["freq_min_khz"]) for row in endpoints)
        maximum = min(float(models[row["model_id"]]["freq_max_khz"]) for row in endpoints)
        assert minimum <= maximum
        return {"id": link["id"], field: minimum}
    return {"id": link["id"], field: True if field == "isBackup" else 1.25}


@pytest.mark.parametrize("format", ["csv", "xlsx"])
def test_template_preview_commit_persists_only_sidecar_in_real_version(workspace_api, format):
    client, factory, baseline, _ = workspace_api
    switch(client)
    before = state(factory)
    initial = ok(client.get(CONTEXT, params={"context": "gis"}))
    assert initial["resources"]["candidateLinks"] == []
    assert baseline["tables"]["link_metric.csv"]
    payload, preview = prepare(client, baseline, format)
    assert_no_write(factory, before)
    assert preview["base_revision"] == baseline["revision"] and preview["context"] == "gis"
    assert preview["kind"] == "topology_links" and preview["candidate"]["source"] == datasets.SOURCE
    link = preview["candidate"]["candidateLinks"][0]
    assert CandidateLink.model_validate(link).assessmentState == "PENDING"
    assert link["source"] == {"kind": "IMPORT", "file": payload["file"]["name"], "row": 2}
    assert link["reference"] is None and all(link[field] is None for field in ASSESSMENT_FIELDS)
    assert preview["changes"] == {"stations": [], "devices": [], "newCandidateLinks": [link]}
    assert preview["rows"][0]["action"] == "NEW" and preview["rows"][0]["sheet"] == "topology_links"
    request = commit_body(payload, preview)
    result = ok(post(client, PREFIX + "/commit", request))
    receipt, resources = result["receipt"], result["result"]["resources"]
    assert set(result) == {"receipt", "result"}
    assert receipt["changed"] is result["result"]["changed"] is True
    assert receipt["revision"] == receipt["readRevision"] == resources["version"] != baseline["revision"]
    assert receipt["affectedIds"] == [link["id"]] and receipt["objectId"] == "resource-import"
    assert resources["source"] == datasets.SOURCE and resources["candidateLinks"] == [link]
    saved = assert_resource_files_and_tables(factory, baseline)
    assert saved[candidates.SNAPSHOT_KEY] == [link]
    assert candidates.SNAPSHOT_KEY not in stored(factory, baseline["revision"])
    with factory() as session:
        row = session.get(ResourceDatasetVersion, receipt["revision"])
        assert row.source == datasets.SOURCE and row.data_origin == "SYNTHETIC_INTEGRATION"
        assert session.get(ResourceDatasetOperation, request["request_id"]).revision == receipt["revision"]
    context = ok(client.get(CONTEXT, params={"context": "gis", "revision": receipt["revision"]}))
    assert context["resources"] == resources and context["validationSource"]["candidateLinks"] == [link]
    assert context["situation"]["links"] == []
    assert counts(factory) == (2, 1)


@pytest.mark.parametrize("format", ["csv", "xlsx"])
def test_import_reference_is_external_frozen_and_zero_false_preserved(workspace_api, format):
    client, factory, baseline, _ = workspace_api
    switch(client)
    current, link, request, result = seed_reference(client, baseline, format)
    assert link["assessmentState"] == "REFERENCE" and link["pathLossDb"] == 0
    assert link["isAvailable"] is False and link["rxPowerDbm"] == -100
    assert link["reference"]["source"] == link["source"]
    assert link["reference"]["data"] == {field: link[field] for field in candidates.DATA_FIELDS}
    assert {warning["field"] for warning in result["receipt"]["warnings"]} == {"reference", "path_loss_db"}
    assert_resource_files_and_tables(factory, baseline)
    before = state(factory)
    assert ok(post(client, PREFIX + "/commit", request)) == result
    assert_no_write(factory, before)


def test_update_preview_uses_imported_changes_and_original_file_is_reparsed(workspace_api, monkeypatch):
    client, factory, baseline, _ = workspace_api
    switch(client)
    current, link, _, _ = seed_reference(client, baseline)
    payload, preview = prepare(client, current)
    rows = list(csv.DictReader(StringIO(base64.b64decode(payload["file"]["base64"]).decode("utf-8"))))
    rows[0].update(link_id=link["id"], path_loss_db="-5", is_available="true", link_state="FAULT")
    payload["file"]["base64"] = base64.b64encode(csv_data(rows)).decode("ascii")
    calls, original = [], candidates.parse_candidate_import

    def tracked(snapshot, filename, data):
        calls.append((filename, data))
        return original(snapshot, filename, data)

    monkeypatch.setattr(candidates, "parse_candidate_import", tracked)
    before = state(factory)
    preview = ok(post(client, PREFIX + "/preview", payload))
    changed = preview["candidate"]["candidateLinks"][0]
    assert preview["valid"] and preview["rows"][0]["action"] == "UPDATE"
    assert preview["changes"] == {"stations": [], "devices": [], "importedCandidateLinks": [changed]}
    assert changed["reference"]["data"]["pathLossDb"] == -5 and changed["isAvailable"] is True
    assert {warning["field"] for warning in preview["warnings"]} == {"reference", "path_loss_db", "link_state"}
    assert_no_write(factory, before)
    request = commit_body(payload, preview, "update-import")
    assert_error(post(client, PREFIX + "/commit", {**request, "candidate": {"candidateLinks": []}}))
    assert_no_write(factory, before)
    result = ok(post(client, PREFIX + "/commit", request))
    assert len(calls) == 2 and calls[0] == calls[1]
    assert result["result"]["resources"]["candidateLinks"] == [changed]
    assert_resource_files_and_tables(factory, baseline)


def test_original_file_reimport_is_noop_and_creates_only_idempotent_receipt(workspace_api):
    client, factory, baseline, _ = workspace_api
    switch(client)
    current, link, original, first = seed_reference(client, baseline)
    payload = {key: deepcopy(value) for key, value in original.items() if key not in {"request_id", "signature"}}
    payload["base_revision"] = current["revision"]
    preview = ok(post(client, PREFIX + "/preview", payload))
    assert preview["valid"] and preview["rows"][0]["action"] == "UNCHANGED" and preview["rows"][0]["differences"] == []
    assert preview["changes"] == {"stations": [], "devices": []}
    assert preview["candidate"]["candidateLinks"] == [link]
    request = commit_body(payload, preview, "candidate-noop")
    result = ok(post(client, PREFIX + "/commit", request))
    assert result["receipt"]["changed"] is result["result"]["changed"] is False
    assert result["receipt"]["revision"] == first["receipt"]["revision"]
    assert result["receipt"]["affectedIds"] == [] and counts(factory) == (2, 2)
    before = state(factory)
    assert ok(post(client, PREFIX + "/commit", request)) == result
    assert_no_write(factory, before)
    assert_resource_files_and_tables(factory, baseline)


@pytest.mark.parametrize("field", CONFIG_FIELDS)
def test_eight_configuration_patch_fields_invalidate_only_when_required(workspace_api, field):
    client, factory, baseline, _ = workspace_api
    switch(client)
    current, original, _, _ = seed_reference(client, baseline)
    patch = candidate_patch(current, original, field)
    response = ok(post(client, SAVE, workspace_body(current, {"candidateLinks": [patch]})))
    link = response["result"]["resources"]["candidateLinks"][0]
    assert all(link[key] == value for key, value in patch.items())
    assert link["source"] == original["source"] and link["reference"] == original["reference"]
    if field == "isBackup":
        assert link["assessmentState"] == "REFERENCE" and link["pathLossDb"] == 0
    else:
        assert link["assessmentState"] == "PENDING" and all(link[key] is None for key in ASSESSMENT_FIELDS)
    assert response["receipt"]["affectedIds"] == [original["id"]]
    assert counts(factory) == (3, 2)
    saved = assert_resource_files_and_tables(factory, baseline)
    assert saved[candidates.SNAPSHOT_KEY] == [link]
    assert stored(factory, current["revision"])[candidates.SNAPSHOT_KEY] == [original]


def test_manual_candidate_is_pending_and_unchanged_patch_is_noop(workspace_api):
    client, factory, baseline, _ = workspace_api
    switch(client)
    _, preview = prepare(client, baseline)
    manual = manual_from(preview["candidate"]["candidateLinks"][0])
    result = ok(post(client, SAVE, workspace_body(baseline, {"newCandidateLinks": [manual]}, "manual-create")))
    assert result["result"]["resources"]["candidateLinks"] == [manual]
    assert result["result"]["resources"]["source"] == datasets.SOURCE
    assert result["receipt"]["changed"] is True and result["receipt"]["affectedIds"] == [manual["id"]]
    current = datasets.current()
    patch = {"id": manual["id"], **{field: manual[field] for field in CONFIG_FIELDS}}
    noop = ok(post(client, SAVE, workspace_body(current, {"candidateLinks": [patch]}, "manual-noop")))
    assert noop["receipt"]["changed"] is False and noop["receipt"]["revision"] == current["revision"]
    assert counts(factory) == (2, 2)
    assert assert_resource_files_and_tables(factory, baseline)[candidates.SNAPSHOT_KEY] == [manual]


@pytest.mark.parametrize("mutation", ["import", "source", "assessment", "state", "reference", "readonly-source", "readonly-assessment", "readonly-reference", "delete"])
def test_forged_import_assessment_readonly_source_and_deletion_are_atomic_rejections(workspace_api, mutation):
    client, factory, baseline, _ = workspace_api
    switch(client)
    current, link, _, _ = seed_reference(client, baseline)
    row = manual_from(link)
    changes = {"newCandidateLinks": [row]}
    if mutation == "import":
        changes = {"importedCandidateLinks": [link]}
    elif mutation == "source":
        row["source"] = deepcopy(link["source"])
    elif mutation == "assessment":
        row["isAvailable"] = False
    elif mutation == "state":
        row["assessmentState"] = "REFERENCE"
    elif mutation == "reference":
        row["reference"] = deepcopy(link["reference"])
    else:
        field, value = {"readonly-source": ("source", link["source"]), "readonly-assessment": ("pathLossDb", 0),
                        "readonly-reference": ("reference", link["reference"]), "delete": ("deleted", True)}[mutation]
        changes = {"candidateLinks": [{"id": link["id"], field: value}]}
    device = current["tables"]["device.csv"][0]
    changes["devices"] = [{"id": device["device_id"], "antennaHeightM": float(device["antenna_height_m"]) + 0.125}]
    before = state(factory)
    assert_error(post(client, SAVE, workspace_body(current, changes)))
    assert_no_write(factory, before)


def test_mixed_valid_invalid_import_batch_never_saves_partial_sidecar(workspace_api):
    client, factory, baseline, _ = workspace_api
    switch(client)
    payload, _ = prepare(client, baseline)
    row = next(csv.DictReader(StringIO(base64.b64decode(payload["file"]["base64"]).decode("utf-8"))))
    payload["file"]["base64"] = base64.b64encode(csv_data([row, {**row, "link_id": "BAD", "device_a_id": "MISSING"}])).decode("ascii")
    before = state(factory)
    preview = ok(post(client, PREFIX + "/preview", payload))
    assert preview["valid"] is False and preview["candidate"]["candidateLinks"] == []
    assert preview["changes"] == {"stations": [], "devices": []}
    assert any(issue["row"] == 3 and issue["objectId"] == "BAD" for issue in preview["issues"])
    assert_error(post(client, PREFIX + "/commit", commit_body(payload, preview)))
    assert_no_write(factory, before)


def test_historical_import_and_workspace_receipts_read_frozen_sql_versions_after_reopen(workspace_api, monkeypatch):
    client, factory, baseline, path = workspace_api
    switch(client)
    first_snapshot, original, import_request, imported = seed_reference(client, baseline)
    patch = candidate_patch(first_snapshot, original, "freqKhz")
    patch_request = workspace_body(first_snapshot, {"candidateLinks": [patch]}, "historic-patch")
    patched = ok(post(client, SAVE, patch_request))
    current = datasets.current()
    later = ok(post(client, SAVE, workspace_body(current, {"newCandidateLinks": [manual_from(original)]}, "later-manual")))
    latest = datasets.current()
    assert latest["revision"] == later["receipt"]["revision"]
    before = state(factory)
    reopened = create_engine("sqlite:///" + str(path), connect_args={"check_same_thread": False})
    replacement = sessionmaker(bind=reopened, expire_on_commit=False)
    try:
        for name, module in tuple(sys.modules.items()):
            if name.startswith("app.") and hasattr(module, "get_session_factory"):
                monkeypatch.setattr(module, "get_session_factory", lambda: replacement)
        assert ok(post(client, PREFIX + "/commit", import_request)) == imported
        assert ok(post(client, SAVE, patch_request)) == patched
        assert_no_write(factory, before)
        assert stored(replacement, first_snapshot["revision"])[candidates.SNAPSHOT_KEY] == [original]
        assert len(stored(replacement)[candidates.SNAPSHOT_KEY]) == 2
        assert_resource_files_and_tables(replacement, baseline)
    finally:
        reopened.dispose()


@pytest.mark.parametrize("endpoint", ["template", "preview", "commit", "create", "patch"])
def test_ledger_candidate_operations_are_422_without_resource_writes(workspace_api, endpoint):
    client, factory, baseline, _ = workspace_api
    switch(client)
    payload, preview = prepare(client, baseline)
    link = preview["candidate"]["candidateLinks"][0]
    if endpoint == "patch":
        created = ok(post(client, SAVE, workspace_body(baseline, {"newCandidateLinks": [manual_from(link)]}, "ledger-seed")))
        baseline = datasets.current()
        link = created["result"]["resources"]["candidateLinks"][0]
    before = state(factory)
    if endpoint == "template":
        response = client.get(PREFIX + "/template", params={"context": "ledger", "kind": "topology_links", "format": "csv", "revision": baseline["revision"]})
    elif endpoint in {"preview", "commit"}:
        request = {**payload, "context": "ledger"}
        if endpoint == "commit":
            request = commit_body(request, preview)
        response = post(client, PREFIX + "/" + endpoint, request)
    else:
        changes = {"newCandidateLinks": [manual_from(link)]} if endpoint == "create" else {"candidateLinks": [{"id": link["id"], "isBackup": True}]}
        response = post(client, SAVE, workspace_body(baseline, changes, context="ledger"))
    assert_error(response)
    assert_no_write(factory, before)


@pytest.mark.parametrize("permission,action", [("resources.import", "import"), ("resources.create", "create"), ("resources.workspace.edit", "patch")])
def test_import_create_and_edit_permissions_are_independent(workspace_api, monkeypatch, permission, action):
    client, factory, baseline, _ = workspace_api
    switch(client)
    current, original, _, _ = seed_reference(client, baseline)
    payload, preview = prepare(client, current)
    import_request = commit_body(payload, preview, "isolated-permission-import")
    manual_request = workspace_body(current, {"newCandidateLinks": [manual_from(original)]}, "isolated-create")
    patch_request = workspace_body(current, {"candidateLinks": [{"id": original["id"], "isBackup": True}]}, "isolated-edit")
    allowed = ["resources.view", "situation.view", permission]
    monkeypatch.setitem(identity.CATALOG, "candidate-isolated", ("候选专项测试", "EDITOR", allowed))
    account = switch(client, "candidate-isolated")
    assert set(account["permissions"]) == set(allowed) and "resources.edit" not in account["permissions"]
    operations = {"import": (PREFIX + "/commit", import_request), "create": (SAVE, manual_request), "patch": (SAVE, patch_request)}
    before = state(factory)
    for name, (endpoint, request) in operations.items():
        if name != action:
            assert_error(post(client, endpoint, request), 403, 1104)
            assert_no_write(factory, before)
    result = ok(post(client, *operations[action]))
    assert result["receipt"]["changed"] is True and counts(factory) == (3, 2)
    assert_resource_files_and_tables(factory, baseline)


@pytest.mark.parametrize("missing", ["situation.view", "resources.view"])
@pytest.mark.parametrize("action", ["template", "preview", "commit", "context", "create", "patch"])
def test_gis_and_resource_page_permissions_required_for_all_candidate_paths(workspace_api, monkeypatch, missing, action):
    client, factory, baseline, _ = workspace_api
    switch(client)
    current, original, _, _ = seed_reference(client, baseline)
    payload, preview = prepare(client, current)
    permissions = [item for item in identity.CATALOG["integration-editor"][2] if item != missing]
    monkeypatch.setitem(identity.CATALOG, "candidate-no-page", ("缺少页面权限", "EDITOR", permissions))
    switch(client, "candidate-no-page")
    before = state(factory)
    if action == "template":
        response = client.get(PREFIX + "/template", params={"context": "gis", "kind": "topology_links", "format": "csv", "revision": current["revision"]})
    elif action == "context":
        response = client.get(CONTEXT, params={"context": "gis"})
    elif action == "preview":
        response = post(client, PREFIX + "/preview", payload)
    elif action == "commit":
        response = post(client, PREFIX + "/commit", commit_body(payload, preview, "no-page-import"))
    else:
        changes = {"newCandidateLinks": [manual_from(original)]} if action == "create" else {"candidateLinks": [{"id": original["id"], "isBackup": True}]}
        response = post(client, SAVE, workspace_body(current, changes, "no-page-workspace"))
    assert_error(response, 403, 1104)
    assert_no_write(factory, before)


@pytest.mark.parametrize("permission", ["resources.create", "resources.workspace.edit"])
def test_mixed_manual_add_and_existing_patch_requires_both_permissions(workspace_api, monkeypatch, permission):
    client, factory, baseline, _ = workspace_api
    switch(client)
    current, original, _, _ = seed_reference(client, baseline)
    monkeypatch.setitem(identity.CATALOG, "candidate-mixed", ("单一权限", "EDITOR", ["resources.view", "situation.view", permission]))
    switch(client, "candidate-mixed")
    before = state(factory)
    changes = {"newCandidateLinks": [manual_from(original)], "candidateLinks": [{"id": original["id"], "isBackup": True}]}
    assert_error(post(client, SAVE, workspace_body(current, changes)), 403, 1104)
    assert_no_write(factory, before)


@pytest.mark.parametrize("action", ["preview", "commit", "create", "patch"])
@pytest.mark.parametrize("security", ["missing-session", "wrong-csrf", "old-session"])
def test_writes_bind_csrf_and_active_session(workspace_api, action, security):
    client, factory, baseline, _ = workspace_api
    switch(client)
    current, original, _, _ = seed_reference(client, baseline)
    payload, preview = prepare(client, current)
    auth = headers(client)
    if security == "missing-session":
        auth.pop("X-Session-Id")
    elif security == "wrong-csrf":
        auth["X-CSRF-Token"] = "bad"
    else:
        switch(client)
    if action in {"preview", "commit"}:
        endpoint = PREFIX + "/" + action
        request = payload if action == "preview" else commit_body(payload, preview, "csrf-import")
    else:
        endpoint = SAVE
        changes = {"newCandidateLinks": [manual_from(original)]} if action == "create" else {"candidateLinks": [{"id": original["id"], "isBackup": True}]}
        request = workspace_body(current, changes, "csrf-workspace")
    before = state(factory)
    assert_error(client.post(endpoint, headers=auth, json=request), 403 if security == "wrong-csrf" else 409,
                 1105 if security == "wrong-csrf" else 1106)
    assert_no_write(factory, before)


def test_stale_template_preview_import_and_workspace_cannot_overwrite_candidate_version(workspace_api):
    client, factory, baseline, _ = workspace_api
    switch(client)
    payload, preview = prepare(client, baseline)
    manual = manual_from(preview["candidate"]["candidateLinks"][0])
    ok(post(client, SAVE, workspace_body(baseline, {"newCandidateLinks": [manual]}, "advance-version")))
    before = state(factory)
    response = client.get(PREFIX + "/template", params={"context": "gis", "kind": "topology_links", "format": "csv", "revision": baseline["revision"]})
    assert_error(response, 409, 1205)
    assert_error(post(client, PREFIX + "/preview", payload), 409, 1205)
    assert_error(post(client, PREFIX + "/commit", commit_body(payload, preview, "stale-import")), 409, 1205)
    assert_error(post(client, SAVE, workspace_body(baseline, {"newCandidateLinks": [manual_from(manual, "OTHER")]}, "stale-workspace")), 409, 1205)
    assert_no_write(factory, before)


@pytest.mark.parametrize("mutation", ["signature", "name", "bytes"])
def test_import_signature_binds_original_candidate_file(workspace_api, mutation):
    client, factory, baseline, _ = workspace_api
    switch(client)
    payload, preview = prepare(client, baseline)
    request = commit_body(payload, preview)
    if mutation == "signature":
        request["signature"] = "0" * 64
    elif mutation == "name":
        request["file"]["name"] = "different.csv"
    else:
        request["file"]["base64"] = base64.b64encode(base64.b64decode(request["file"]["base64"]) + b"\n").decode("ascii")
    before = state(factory)
    assert_error(post(client, PREFIX + "/commit", request), 409, 1206)
    assert_no_write(factory, before)


@pytest.mark.parametrize("action", ["import", "workspace"])
def test_request_id_is_bound_to_actor_and_request_content(workspace_api, monkeypatch, action):
    client, factory, baseline, _ = workspace_api
    switch(client)
    payload, preview = prepare(client, baseline)
    endpoint = PREFIX + "/commit" if action == "import" else SAVE
    request = commit_body(payload, preview) if action == "import" else workspace_body(baseline, {"newCandidateLinks": [manual_from(preview["candidate"]["candidateLinks"][0])]})
    first = ok(post(client, endpoint, request))
    before = state(factory)
    assert ok(post(client, endpoint, request)) == first
    changed = deepcopy(request)
    if action == "import":
        changed_payload, changed_preview = prepare(client, datasets.current())
        changed = commit_body(changed_payload, changed_preview, request["request_id"])
    else:
        changed["changes"]["newCandidateLinks"][0]["isBackup"] = True
    assert_error(post(client, endpoint, changed), 409, 1206)
    monkeypatch.setitem(identity.CATALOG, "candidate-second-actor", ("第二测试编辑者", "EDITOR", list(identity.CATALOG["integration-editor"][2])))
    switch(client, "candidate-second-actor")
    assert_error(post(client, endpoint, request), 409, 1206)
    assert_no_write(factory, before)


@pytest.mark.parametrize("action", ["import", "workspace"])
def test_transaction_validation_failure_rolls_back_sidecar_pointer_and_receipt(workspace_api, monkeypatch, action):
    client, factory, baseline, _ = workspace_api
    switch(client)
    payload, preview = prepare(client, baseline)
    before = state(factory)
    original, staged = datasets._validate_snapshot, []

    def reject(snapshot):
        original(snapshot)
        if "change" in snapshot:
            staged.append(deepcopy(snapshot[candidates.SNAPSHOT_KEY]))
            raise datasets.ResourceDatasetError("候选原子保存故障注入")

    monkeypatch.setattr(datasets, "_validate_snapshot", reject)
    if action == "import":
        response = post(client, PREFIX + "/commit", commit_body(payload, preview))
    else:
        response = post(client, SAVE, workspace_body(baseline, {"newCandidateLinks": [manual_from(preview["candidate"]["candidateLinks"][0])]}))
    assert_error(response, 503, 5002)
    assert staged and len(staged[0]) == 1
    assert_no_write(factory, before)


def assert_displayable(issue):
    assert set(issue) == {"sheet", "row", "objectId", "field", "message"}
    for field, limit in (("sheet", 128), ("field", 128), ("objectId", 512), ("message", 4096)):
        value = issue[field]
        assert len(value.encode("utf-16-le")) // 2 <= limit
        assert not any(unicodedata.category(char) in {"Cc", "Cf", "Cs"} or char in {"\u2028", "\u2029"} for char in value)
        value.encode("utf-8")


@pytest.mark.parametrize("identifier,truncated", [
    pytest.param("BAD\x00\t\n\x7f\x85\u202eID", False, id="controls"),
    pytest.param("\U00020000" * 300, True, id="supplementary-overflow"),
    pytest.param("I" * 510 + "\x00\t", True, id="escaped-overflow"),
])
def test_malicious_import_ids_are_display_only_sanitized_and_never_persisted(workspace_api, identifier, truncated):
    client, factory, baseline, _ = workspace_api
    switch(client)
    data, filename = template(client, baseline)
    row = next(csv.DictReader(StringIO(data.decode("utf-8"))))
    row["link_id"] = identifier
    raw = csv_data([row])
    payload = import_body(baseline, raw, filename)
    original_payload, before = deepcopy(payload), state(factory)
    preview = ok(post(client, PREFIX + "/preview", payload))
    assert preview["valid"] is False and preview["candidate"]["candidateLinks"] == []
    for issue in preview["issues"]:
        assert_displayable(issue)
        assert issue["sheet"] == "topology_links" and issue["row"] == 2
    if truncated:
        assert "已截断" in preview["issues"][0]["objectId"]
        assert "原长" + str(len(identifier.encode("utf-16-le")) // 2) in preview["issues"][0]["objectId"]
    else:
        assert r"\u0000" in preview["issues"][0]["objectId"]
    assert_error(post(client, PREFIX + "/commit", commit_body(payload, preview)))
    assert payload == original_payload and base64.b64decode(payload["file"]["base64"]) == raw
    assert_no_write(factory, before)


def test_unknown_control_field_is_sanitized_in_error_without_rewriting_request(workspace_api):
    client, factory, baseline, _ = workspace_api
    switch(client)
    _, preview = prepare(client, baseline)
    row = manual_from(preview["candidate"]["candidateLinks"][0])
    field = "BAD\x00\t\u202e" + "\U00020000" * 100
    row[field] = "raw\x00value"
    request = workspace_body(baseline, {"newCandidateLinks": [row]})
    before, original = state(factory), deepcopy(request)
    response = post(client, SAVE, request)
    assert_error(response)
    assert_displayable(candidates._issue(0, row["id"], field, response.json()["message"]))
    with pytest.raises(datasets.ResourceDatasetError) as caught:
        candidates.apply_candidate_changes(baseline, request["changes"])
    assert_displayable(caught.value.issues[0])
    assert "已截断" in caught.value.issues[0]["field"] and r"\u0000" in caught.value.issues[0]["field"]
    assert request == original
    assert_no_write(factory, before)


@pytest.mark.parametrize("title", ["bad\n\x7f\x85\u202esheet", "\U00020000" * 100 + "\x85"])
def test_hidden_malicious_sheet_is_bounded_diagnostic_without_changing_source(workspace_api, title):
    client, factory, baseline, _ = workspace_api
    switch(client)
    data, filename = template(client, baseline, "xlsx")
    output = BytesIO()
    with ZipFile(BytesIO(data)) as original, ZipFile(output, "w", ZIP_DEFLATED) as target:
        for member in original.infolist():
            value = original.read(member)
            if member.filename == "xl/workbook.xml":
                root = ET.fromstring(value)
                sheet = next(element for element in root.iter() if element.tag.rsplit("}", 1)[-1] == "sheet")
                sheet.set("name", title)
                sheet.set("state", "hidden")
                value = ET.tostring(root, encoding="utf-8")
            target.writestr(member.filename, value)
    raw = output.getvalue()
    payload, before = import_body(baseline, raw, filename), state(factory)
    preview = ok(post(client, PREFIX + "/preview", payload))
    assert preview["valid"] is False and preview["rows"] == []
    assert preview["issues"][0]["sheet"] == "topology_links" and preview["issues"][0]["field"] == "sheet"
    for issue in preview["issues"]:
        assert_displayable(issue)
    assert_error(post(client, PREFIX + "/commit", commit_body(payload, preview)))
    assert base64.b64decode(payload["file"]["base64"]) == raw
    assert_no_write(factory, before)


def test_candidate_issue_delegates_all_display_labels_without_mutating_raw_values():
    values = ["I\x00\t\n\x85\u202e" + "\U00020000" * 400,
              "F\x00\t" + "\U00020000" * 100, "M\x00\ud800" + "\U00020000" * 2200]
    before = deepcopy(values)
    issue = candidates._issue(5002, *values)
    assert issue == parser._issue("topology_links", 5002, *values)
    assert issue["row"] == 5002
    assert_displayable(issue)
    assert all("已截断" in issue[field] for field in ("objectId", "field", "message"))
    assert values == before
