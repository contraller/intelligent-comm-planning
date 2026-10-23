"""版本化 SR4 作业提交接口测试。"""
from __future__ import annotations

import json
import sys
import threading
from copy import deepcopy
from io import BytesIO
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine, func, select
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool


ROOT = Path(__file__).resolve().parents[1]
SERVICE_ROOT = ROOT / "algorithm-service"
if str(SERVICE_ROOT) not in sys.path:
    sys.path.insert(0, str(SERVICE_ROOT))

from app import frontend_integration, planning_job_models  # noqa: E402
from app import planning_job_service, planning_plan_service, task_input_service  # noqa: E402
from app.database import Base  # noqa: E402
from app.main import app  # noqa: E402


@pytest.fixture
def api(monkeypatch, tmp_path):
    monkeypatch.setenv("FRONTEND_INTEGRATION_DEMO", "1")
    engine = create_engine(
        "sqlite+pysqlite:///:memory:",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(engine)
    factory = sessionmaker(bind=engine, expire_on_commit=False)
    monkeypatch.setattr(task_input_service, "get_session_factory", lambda: factory)
    monkeypatch.setattr(planning_job_service, "get_session_factory", lambda: factory)
    monkeypatch.setattr(planning_plan_service, "get_session_factory", lambda: factory)
    monkeypatch.setattr(planning_job_service.planning_w3, "TASK_DIR", tmp_path)
    calls = []

    def fake_submit(payload, *, job_id, on_update):
        calls.append((payload, job_id))
        on_update({"status": "PENDING", "progress": 0.0,
                   "stage": "排队中", "error": None})
        return {"code": 0, "message": "success",
                "data": {"task_id": job_id, "estimated_sec": 5}}

    monkeypatch.setattr(planning_job_service.planning_w4, "submit_sr4", fake_submit)
    yield TestClient(app), factory, calls
    engine.dispose()


def _save_task(client: TestClient, task_id: str = "frontend-task-1",
               demand_id: str | None = None) -> dict:
    workspace = client.get(
        "/api/v1/resource/context", params={"context": "planning"}
    ).json()["data"]
    preset = client.get(
        "/api/v1/resource/param-presets", params={"context": "planning"}
    ).json()["data"]["items"][0]
    if demand_id is not None:
        preset["defaults"]["demands"][0]["id"] = demand_id
    response = client.put(f"/api/v1/task-inputs/{task_id}", json={
        "context": "planning",
        "task_id": task_id,
        "expected_revision": None,
        "resource_revision": workspace["revision"],
        "request_id": f"save-{task_id}",
        "input": preset["defaults"],
    })
    assert response.status_code == 200
    return response.json()["data"]


def _submit(saved: dict, request_id: str = "submit-1") -> dict:
    return {
        "context": "planning",
        "request_id": request_id,
        "task_id": saved["task_id"],
        "expected_revision": saved["revision"],
        "input_version": saved["input_version"],
        "condition_version": saved["condition_version"],
        "resource_revision": saved["resource_revision"],
        "parent_plan_id": None,
    }


def _cancel(saved: dict, request_id: str = "cancel-1") -> dict:
    return {
        "context": "planning", "request_id": request_id,
        "task_id": saved["task_id"], "expected_revision": saved["revision"],
    }


def _save_plan(saved: dict, job: dict, request_id: str = "save-plan-1") -> dict:
    return {
        "context": "planning",
        "request_id": request_id,
        "task_id": saved["task_id"],
        "expected_revision": saved["revision"],
        "job_id": job["job_id"],
        "name": "太行山区规划方案 V1",
        "note": "接口测试方案",
    }


def _confirm_plan(saved: dict, plan: dict,
                  request_id: str = "confirm-plan-1") -> dict:
    return {
        "context": "planning",
        "request_id": request_id,
        "task_id": saved["task_id"],
        "expected_revision": plan["revision"],
    }


def _complete_job(monkeypatch, saved: dict, *, conflict: bool = False) -> None:
    monkeypatch.setattr(planning_job_service.planning_w3, "get_task_record", lambda _: {
        "status": "SUCCESS", "progress": 1.0, "stage": "完成", "error": None,
        "result": {"solution_id": "PL-TEST0001"},
    })
    monkeypatch.setattr(planning_job_service.planning_w4, "get_plan", lambda _: {
        "code": 0, "message": "success", "data": _raw_plan(saved, conflict=conflict),
    })


def _raw_plan(saved: dict, *, conflict: bool = False) -> dict:
    station_ids = saved["input"]["stationIds"]
    source, target = station_ids[:2]
    demand_id = saved["input"]["demands"][0]["id"]
    constraints = {
        "mandatory_e2e_connectivity": {"value": 1.0, "required": 1.0, "pass": True},
        "survivability": {"value": 0.92, "required": 0.9, "pass": True},
        "max_relay_hops": {"value": 0, "required": 2, "pass": True},
    }
    conflicts = ([{"kind": "FREQ_CONFLICT", "severity": "SEVERE",
                   "detail": "测试频率冲突", "link_ids": ["LK-TEST"]}]
                 if conflict else [])
    return {
        "solution_id": "PL-TEST0001",
        "summary": {"objectives": {"e2e_connectivity": 1.0,
                                      "total_radiated_power_w": 20.0,
                                      "survivability": 0.92}},
        "detail": {
            "nodes": [
                {"node_id": node_id, "node_name": f"参与节点{index + 1}",
                 "lon": 114.0 + index * 0.1, "lat": 37.0 + index * 0.1,
                 "status": "NORMAL", "node_role": "任务节点"}
                for index, node_id in enumerate(station_ids)
            ],
            "links": [{"link_id": "LK-TEST", "node_a_id": source,
                       "node_b_id": target, "device_class": "HF",
                       "device_a_id": "", "device_b_id": "",
                       "link_margin_db": 15.0, "is_available": True,
                       "link_state": "GOOD", "is_backup": False}],
            "routes": {"routes": [{
                "demand_id": demand_id, "reachable": True,
                "primary": {"node_path": [source, target],
                            "link_path": ["LK-TEST"], "total_delay_ms": 12.0},
                "backup": None,
            }]},
            "frequency": {
                "link_assignments": [{"link_id": "LK-TEST",
                                      "member_link_ids": ["LK-TEST"],
                                      "freq_khz": 7000.0}],
                "conflicts": conflicts,
            },
            "params": {"params": [
                {"node_id": node_id, "band": "HF", "model_id": "DM-0004",
                 "tx_power_dbm": 40.0, "antenna_height_m": 10.0}
                for node_id in station_ids
            ]},
            "optimization": {
                "best": {"constraint_check": constraints},
                "infeasible": ({"conflicting_constraints": ["存在频率冲突"],
                                "suggestions": ["调整频率"]} if conflict else None),
            },
            "check": {"issues": conflicts},
        },
    }


def test_submit_saved_task_returns_frontend_job_and_starts_real_entry(api):
    client, factory, calls = api
    saved = _save_task(client)
    response = client.post("/api/v1/plan/sr4-flow", json=_submit(saved))
    assert response.status_code == 202
    body = response.json()
    assert body["code"] == 0
    assert body["request_id"] == "submit-1"
    job = body["data"]
    assert job["task_id"] == saved["task_id"]
    assert job["revision"] == saved["revision"]
    assert job["input_version"] == saved["input_version"]
    assert job["condition_version"] == saved["condition_version"]
    assert job["resource_revision"] == saved["resource_revision"]
    assert job["input"] == saved["input"]
    assert job["status"] == "PENDING"
    assert job["progress"] == 0
    assert job["result"] is None and job["outcome"] is None
    assert len(job["phases"]) == 5
    assert len(calls) == 1
    algorithm_payload, job_id = calls[0]
    assert job_id == job["job_id"]
    assert algorithm_payload["task_id"] == saved["task_id"]
    assert algorithm_payload["preset"] == "TERRAIN"
    assert algorithm_payload["hard_constraints"] == {
        "mandatory_e2e_connectivity": 1.0,
        "survivability_threshold": 0.9,
        "max_relay_hops": 2,
    }
    assert algorithm_payload["demand_rows"][0]["demand_id"] == "DEMAND-01"
    assert algorithm_payload["demand_rows"][0]["task_id"] == saved["task_id"]
    with factory() as session:
        assert session.scalar(select(func.count()).select_from(
            planning_job_models.PlanningJob)) == 1


def test_submit_is_idempotent_and_blocks_second_active_job(api):
    client, _, calls = api
    saved = _save_task(client)
    request = _submit(saved)
    first = client.post("/api/v1/plan/sr4-flow", json=request)
    replay = client.post("/api/v1/plan/sr4-flow", json=request)
    assert replay.status_code == 202
    assert replay.json()["data"]["job_id"] == first.json()["data"]["job_id"]
    assert len(calls) == 1

    conflict = dict(request, condition_version="TCOND-different")
    reused = client.post("/api/v1/plan/sr4-flow", json=conflict)
    assert reused.status_code == 409
    assert reused.json()["code"] == 1204

    another = client.post(
        "/api/v1/plan/sr4-flow", json=_submit(saved, "submit-2")
    )
    assert another.status_code == 409
    assert another.json()["code"] == 2002
    assert another.json()["data"]["job_id"] == first.json()["data"]["job_id"]


def test_submit_rejects_stale_task_and_resource_versions(api, monkeypatch):
    client, _, calls = api
    saved = _save_task(client)
    stale = _submit(saved)
    stale["expected_revision"] = "TREV-stale"
    response = client.post("/api/v1/plan/sr4-flow", json=stale)
    assert response.status_code == 409
    assert response.json()["code"] == 1203

    monkeypatch.setattr(frontend_integration, "resource_revision",
                        lambda: "resource-new-version")
    changed = client.post(
        "/api/v1/plan/sr4-flow", json=_submit(saved, "submit-resource-changed")
    )
    assert changed.status_code == 409
    assert changed.json()["code"] == 1205
    assert calls == []


def test_submit_rejects_adjustment_and_unknown_fields(api):
    client, _, calls = api
    saved = _save_task(client)
    adjusted = _submit(saved)
    adjusted["adjustment"] = {"manual_override": []}
    response = client.post("/api/v1/plan/sr4-flow", json=adjusted)
    assert response.status_code == 422
    assert response.json()["code"] == 1001

    unknown = _submit(saved, "submit-extra")
    unknown["unexpected"] = True
    invalid = client.post("/api/v1/plan/sr4-flow", json=unknown)
    assert invalid.status_code == 422
    assert invalid.json()["code"] == 1001
    assert invalid.json()["data"][0]["field"] == "unexpected"
    assert calls == []


def test_submit_rejects_demand_id_owned_by_another_task(api):
    client, _, calls = api
    saved = _save_task(client, demand_id="CD-0001")
    response = client.post("/api/v1/plan/sr4-flow", json=_submit(saved))
    assert response.status_code == 422
    assert response.json()["code"] == 1001
    assert "CD-0001" in response.json()["message"]
    assert calls == []


def test_query_versioned_job_running_and_checks_owner(api, monkeypatch):
    client, _, _ = api
    saved = _save_task(client)
    job = client.post("/api/v1/plan/sr4-flow", json=_submit(saved)).json()["data"]
    monkeypatch.setattr(planning_job_service.planning_w3, "get_task_record", lambda _: {
        "status": "RUNNING", "progress": 0.55, "stage": "路由与频率规划",
        "error": None, "result": None,
    })
    response = client.get(f"/api/v1/plan/tasks/{job['job_id']}", params={
        "context": "planning", "task_id": saved["task_id"],
    })
    assert response.status_code == 200
    body = response.json()
    assert body["code"] == 2002
    assert body["request_id"] == "submit-1"
    assert body["data"]["status"] == "RUNNING"
    assert body["data"]["progress"] == 0.55
    assert body["data"]["input"] == saved["input"]

    wrong_owner = client.get(f"/api/v1/plan/tasks/{job['job_id']}", params={
        "context": "planning", "task_id": "another-task",
    })
    assert wrong_owner.status_code == 404
    assert wrong_owner.json()["code"] == 2001


def test_query_success_returns_frontend_result_contract(api, monkeypatch):
    client, _, _ = api
    saved = _save_task(client)
    job = client.post("/api/v1/plan/sr4-flow", json=_submit(saved)).json()["data"]
    monkeypatch.setattr(planning_job_service.planning_w3, "get_task_record", lambda _: {
        "status": "SUCCESS", "progress": 1.0, "stage": "完成", "error": None,
        "result": {"solution_id": "PL-TEST0001"},
    })
    monkeypatch.setattr(planning_job_service.planning_w4, "get_plan", lambda _: {
        "code": 0, "message": "success", "data": _raw_plan(saved),
    })
    response = client.get(f"/api/v1/plan/tasks/{job['job_id']}", params={
        "context": "planning", "task_id": saved["task_id"],
    })
    assert response.status_code == 200
    body = response.json()
    assert body["code"] == 0
    result = body["data"]["result"]
    assert body["data"]["status"] == "SUCCESS"
    assert body["data"]["outcome"]["code"] == 0
    assert {row["id"] for row in result["nodes"]} == set(saved["input"]["stationIds"])
    assert result["configuration"]["routes"][0]["demandId"] == "DEMAND-01"
    link = result["links"][0]
    devices = {row["id"]: row for row in result["devices"]}
    assert devices[link["deviceA"]]["stationId"] == link["a"]
    assert devices[link["deviceB"]]["stationId"] == link["b"]
    assert result["coverage"] == {"contour": None, "raster_png": None}


def test_query_preserves_infeasible_result_and_reports_gap(api, monkeypatch):
    client, _, _ = api
    saved = _save_task(client)
    job = client.post("/api/v1/plan/sr4-flow", json=_submit(saved)).json()["data"]
    monkeypatch.setattr(planning_job_service.planning_w3, "get_task_record", lambda _: {
        "status": "SUCCESS", "progress": 1.0, "stage": "完成", "error": None,
        "result": {"solution_id": "PL-TEST0001"},
    })
    monkeypatch.setattr(planning_job_service.planning_w4, "get_plan", lambda _: {
        "code": 0, "message": "success", "data": _raw_plan(saved, conflict=True),
    })
    response = client.get(f"/api/v1/plan/tasks/{job['job_id']}", params={
        "context": "planning", "task_id": saved["task_id"],
    })
    body = response.json()
    assert response.status_code == 200
    assert body["code"] == 3001
    assert body["data"]["status"] == "SUCCESS"
    assert body["data"]["result"] is not None
    assert body["data"]["outcome"]["gap"]
    frequency = next(row for row in body["data"]["result"]["checks"]
                     if row["id"] == "frequency")
    assert frequency["passed"] is False


def test_query_rejects_malformed_algorithm_result(api, monkeypatch):
    client, _, _ = api
    saved = _save_task(client)
    job = client.post("/api/v1/plan/sr4-flow", json=_submit(saved)).json()["data"]
    raw_plan = _raw_plan(saved)
    raw_plan["detail"]["nodes"][0]["lon"] = "invalid"
    monkeypatch.setattr(planning_job_service.planning_w3, "get_task_record", lambda _: {
        "status": "SUCCESS", "progress": 1.0, "stage": "完成", "error": None,
        "result": {"solution_id": "PL-TEST0001"},
    })
    monkeypatch.setattr(planning_job_service.planning_w4, "get_plan", lambda _: {
        "code": 0, "message": "success", "data": raw_plan,
    })
    response = client.get(f"/api/v1/plan/tasks/{job['job_id']}", params={
        "context": "planning", "task_id": saved["task_id"],
    })
    assert response.status_code == 500
    assert response.json()["code"] == 5001
    assert response.json()["message"] == "SR4 方案结果不符合前端数据契约"


def test_query_marks_unrecoverable_job_failed_after_restart(api):
    client, _, _ = api
    saved = _save_task(client)
    job = client.post("/api/v1/plan/sr4-flow", json=_submit(saved)).json()["data"]
    response = client.get(f"/api/v1/plan/tasks/{job['job_id']}", params={
        "context": "planning", "task_id": saved["task_id"],
    })
    body = response.json()
    assert body["code"] == 0
    assert body["data"]["status"] == "FAILED"
    assert "无法恢复" in body["data"]["error"]


def test_query_without_task_id_keeps_legacy_task_endpoint(api, monkeypatch):
    client, _, _ = api
    monkeypatch.setattr(planning_job_service, "query_planning_job", lambda *_args: None)
    monkeypatch.setattr("app.main.query_task", lambda task_id: {
        "code": 0, "message": "success", "data": {"task_id": task_id, "legacy": True},
    })
    response = client.get("/api/v1/plan/tasks/tsk-legacy00001")
    assert response.status_code == 200
    assert response.json()["data"] == {
        "task_id": "tsk-legacy00001", "legacy": True,
    }


def test_save_list_and_get_versioned_plan(api, monkeypatch):
    client, _, _ = api
    saved = _save_task(client)
    job = client.post("/api/v1/plan/sr4-flow", json=_submit(saved)).json()["data"]
    _complete_job(monkeypatch, saved)

    response = client.post("/api/v1/plans", json=_save_plan(saved, job))
    assert response.status_code == 200
    body = response.json()
    assert body["code"] == 0
    plan = body["data"]
    assert plan["task_id"] == saved["task_id"]
    assert plan["job_id"] == job["job_id"]
    assert plan["revision"] == saved["revision"]
    assert plan["version"] == "V1"
    assert plan["state"] == "SAVED"
    assert plan["result"]["snapshot_id"] == "SNP-PL-TEST0001"
    assert plan["outcome"]["code"] == 0
    assert plan["current"] is False

    job_response = client.get(f"/api/v1/plan/tasks/{job['job_id']}", params={
        "context": "planning", "task_id": saved["task_id"],
    }).json()
    assert job_response["data"]["saved_plan_id"] == plan["plan_id"]

    listing = client.get("/api/v1/plans", params={
        "context": "planning", "task_id": saved["task_id"],
        "page": 1, "page_size": 10, "search": "太行山区",
    })
    assert listing.status_code == 200
    assert listing.json()["data"] == {
        "items": [plan], "total": 1, "page": 1, "page_size": 10,
    }
    detail = client.get(f"/api/v1/plans/{plan['plan_id']}", params={
        "context": "planning", "task_id": saved["task_id"],
    })
    assert detail.status_code == 200
    assert detail.json()["data"] == plan


def test_save_infeasible_plan_as_reference(api, monkeypatch):
    client, _, _ = api
    saved = _save_task(client)
    job = client.post("/api/v1/plan/sr4-flow", json=_submit(saved)).json()["data"]
    _complete_job(monkeypatch, saved, conflict=True)
    response = client.post("/api/v1/plans", json=_save_plan(saved, job))
    assert response.status_code == 200
    assert response.json()["data"]["state"] == "REFERENCE"
    assert response.json()["data"]["outcome"]["code"] == 3001


def test_save_plan_freezes_raw_algorithm_snapshot(api, monkeypatch):
    client, factory, _ = api
    saved = _save_task(client)
    job = client.post("/api/v1/plan/sr4-flow", json=_submit(saved)).json()["data"]
    raw = _raw_plan(saved)
    raw["detail"]["devices"] = [{"device_id": "DV-TEST", "node_id": saved["input"]["stationIds"][0]}]
    monkeypatch.setattr(planning_job_service.planning_w3, "get_task_record", lambda _: {
        "status": "SUCCESS", "progress": 1.0, "stage": "完成", "error": None,
        "result": {"solution_id": raw["solution_id"]},
    })
    monkeypatch.setattr(planning_job_service.planning_w4, "get_plan", lambda _: {
        "code": 0, "data": raw,
    })
    plan = client.post("/api/v1/plans", json=_save_plan(saved, job)).json()["data"]
    assert "raw_snapshot" not in plan
    with factory() as session:
        stored = session.get(planning_plan_service.PlanningPlan, plan["plan_id"])
        assert stored.raw_snapshot["detail"]["devices"][0]["device_id"] == "DV-TEST"
        assert stored.raw_snapshot["solution_id"] == raw["solution_id"]
        assert len(stored.raw_snapshot["terrain_fingerprint"]) == 64


def test_adjusted_submit_uses_frozen_physical_bindings(api, monkeypatch):
    client, _, calls = api
    saved = _save_task(client)
    job = client.post("/api/v1/plan/sr4-flow", json=_submit(saved)).json()["data"]
    raw = _raw_plan(saved)
    source, target = saved["input"]["stationIds"][:2]
    models = {row["model_id"]: row for row in
              planning_job_service.planning_w4._load("device_model.csv")}
    devices = planning_job_service.planning_w4._load("device.csv")
    bound = [next(row for row in devices if row["node_id"] == node_id
                  and models[row["model_id"]]["device_class"] == "HF")
             for node_id in (source, target)]
    raw["detail"]["devices"] = bound
    raw["detail"]["links"][0].update(device_a_id=bound[0]["device_id"],
                                        device_b_id=bound[1]["device_id"])
    monkeypatch.setattr(planning_job_service.planning_w3, "get_task_record", lambda _: {
        "status": "SUCCESS", "progress": 1.0, "stage": "完成", "error": None,
        "result": {"solution_id": raw["solution_id"]},
    })
    monkeypatch.setattr(planning_job_service.planning_w4, "get_plan", lambda _: {
        "code": 0, "data": raw,
    })
    plan_response = client.post("/api/v1/plans", json=_save_plan(saved, job))
    assert plan_response.status_code == 200
    parent = plan_response.json()["data"]
    request = _submit(saved, "submit-adjusted")
    request["parent_plan_id"] = parent["plan_id"]
    request["adjustment"] = {
        "schema_version": 1, "source_plan_id": parent["plan_id"],
        "source_snapshot_id": parent["result"]["snapshot_id"],
        "manual_override": [{"device_id": bound[0]["device_id"],
                             "tx_power_dbm": 40.0}],
        "manual_paths": [{"demand_id": saved["input"]["demands"][0]["id"],
                          "node_path": [source, target]}],
    }
    response = client.post("/api/v1/plan/sr4-flow", json=request)
    assert response.status_code == 202, response.text
    payload = calls[-1][0]
    assert payload["parent_snapshot"]["detail"]["links"][0]["device_a_id"] == bound[0]["device_id"]
    assert payload["manual_device_refs"] == [{"device_id": bound[0]["device_id"],
                                              "node_id": source, "band": "HF"}]
    assert response.json()["data"]["adjustment"]["manual_override"] == request[
        "adjustment"]["manual_override"]


def test_save_plan_is_idempotent_and_rejects_duplicate_job(api, monkeypatch):
    client, _, _ = api
    saved = _save_task(client)
    job = client.post("/api/v1/plan/sr4-flow", json=_submit(saved)).json()["data"]
    _complete_job(monkeypatch, saved)
    request = _save_plan(saved, job)
    first = client.post("/api/v1/plans", json=request)
    replay = client.post("/api/v1/plans", json=request)
    assert replay.status_code == 200
    assert replay.json()["data"]["plan_id"] == first.json()["data"]["plan_id"]

    reused = dict(request, note="不同内容")
    conflict = client.post("/api/v1/plans", json=reused)
    assert conflict.status_code == 409
    assert conflict.json()["code"] == 1204

    duplicate = client.post(
        "/api/v1/plans", json=_save_plan(saved, job, "save-plan-2")
    )
    assert duplicate.status_code == 409
    assert duplicate.json()["code"] == 1204
    assert duplicate.json()["data"]["plan_id"] == first.json()["data"]["plan_id"]


def test_save_plan_checks_revision_owner_and_request_shape(api, monkeypatch):
    client, _, _ = api
    saved = _save_task(client)
    job = client.post("/api/v1/plan/sr4-flow", json=_submit(saved)).json()["data"]
    _complete_job(monkeypatch, saved)

    stale = _save_plan(saved, job)
    stale["expected_revision"] = "TREV-stale"
    response = client.post("/api/v1/plans", json=stale)
    assert response.status_code == 409
    assert response.json()["code"] == 1203

    wrong_owner = _save_plan(saved, job, "save-wrong-owner")
    wrong_owner["task_id"] = "another-task"
    response = client.post("/api/v1/plans", json=wrong_owner)
    assert response.status_code == 404
    assert response.json()["code"] == 2001

    incomplete_solution = _save_plan(saved, job, "save-solution")
    incomplete_solution["solution_id"] = "SOLUTION-1"
    response = client.post("/api/v1/plans", json=incomplete_solution)
    assert response.status_code == 422
    assert response.json()["code"] == 1001

    missing = client.get("/api/v1/plans/PLAN-missing", params={
        "context": "planning", "task_id": saved["task_id"],
    })
    assert missing.status_code == 404
    assert missing.json()["code"] == 2001


def test_confirm_saved_plan_and_replay_request(api, monkeypatch):
    client, _, _ = api
    saved = _save_task(client)
    job = client.post("/api/v1/plan/sr4-flow", json=_submit(saved)).json()["data"]
    _complete_job(monkeypatch, saved)
    plan = client.post("/api/v1/plans", json=_save_plan(saved, job)).json()["data"]
    request = _confirm_plan(saved, plan)

    response = client.post(f"/api/v1/plans/{plan['plan_id']}/confirm", json=request)
    assert response.status_code == 200
    confirmed = response.json()["data"]
    assert confirmed["state"] == "CONFIRMED"
    assert confirmed["current"] is True
    assert confirmed["confirmed_at"] is not None

    replay = client.post(f"/api/v1/plans/{plan['plan_id']}/confirm", json=request)
    assert replay.status_code == 200
    assert replay.json()["data"] == confirmed

    repeated = client.post(
        f"/api/v1/plans/{plan['plan_id']}/confirm",
        json=_confirm_plan(saved, plan, "confirm-plan-2"),
    )
    assert repeated.status_code == 200
    assert repeated.json()["data"]["confirmed_at"] == confirmed["confirmed_at"]

    reused = dict(request, expected_revision="TREV-other")
    conflict = client.post(f"/api/v1/plans/{plan['plan_id']}/confirm", json=reused)
    assert conflict.status_code == 409
    assert conflict.json()["code"] == 1204


def test_confirm_rejects_reference_and_changed_resource(api, monkeypatch):
    client, _, _ = api
    saved = _save_task(client)
    job = client.post("/api/v1/plan/sr4-flow", json=_submit(saved)).json()["data"]
    _complete_job(monkeypatch, saved, conflict=True)
    reference = client.post(
        "/api/v1/plans", json=_save_plan(saved, job)
    ).json()["data"]
    response = client.post(
        f"/api/v1/plans/{reference['plan_id']}/confirm",
        json=_confirm_plan(saved, reference),
    )
    assert response.status_code == 409
    assert response.json()["code"] == 3001
    assert response.json()["data"]["gap"]

    other_saved = _save_task(client, task_id="frontend-task-2")
    other_job = client.post(
        "/api/v1/plan/sr4-flow", json=_submit(other_saved, "submit-other")
    ).json()["data"]
    _complete_job(monkeypatch, other_saved)
    plan = client.post(
        "/api/v1/plans", json=_save_plan(other_saved, other_job, "save-other")
    ).json()["data"]
    monkeypatch.setattr(frontend_integration, "resource_revision",
                        lambda: "resource-new-version")
    changed = client.post(
        f"/api/v1/plans/{plan['plan_id']}/confirm",
        json=_confirm_plan(other_saved, plan, "confirm-other"),
    )
    assert changed.status_code == 409
    assert changed.json()["code"] == 1205


def test_confirm_checks_revision_owner_and_switches_current_plan(api, monkeypatch):
    client, _, _ = api
    saved = _save_task(client)
    first_job = client.post(
        "/api/v1/plan/sr4-flow", json=_submit(saved)
    ).json()["data"]
    _complete_job(monkeypatch, saved)
    first = client.post(
        "/api/v1/plans", json=_save_plan(saved, first_job)
    ).json()["data"]

    stale = _confirm_plan(saved, first)
    stale["expected_revision"] = "TREV-stale"
    response = client.post(f"/api/v1/plans/{first['plan_id']}/confirm", json=stale)
    assert response.status_code == 409
    assert response.json()["code"] == 1203

    wrong_owner = _confirm_plan(saved, first, "confirm-wrong-owner")
    wrong_owner["task_id"] = "another-task"
    response = client.post(
        f"/api/v1/plans/{first['plan_id']}/confirm", json=wrong_owner
    )
    assert response.status_code == 404
    assert response.json()["code"] == 2001

    client.post(
        f"/api/v1/plans/{first['plan_id']}/confirm",
        json=_confirm_plan(saved, first, "confirm-first"),
    )
    second_job = client.post(
        "/api/v1/plan/sr4-flow", json=_submit(saved, "submit-2")
    ).json()["data"]
    second = client.post(
        "/api/v1/plans", json=_save_plan(saved, second_job, "save-plan-2")
    ).json()["data"]
    response = client.post(
        f"/api/v1/plans/{second['plan_id']}/confirm",
        json=_confirm_plan(saved, second, "confirm-second"),
    )
    assert response.status_code == 200
    assert response.json()["data"]["current"] is True

    first_detail = client.get(f"/api/v1/plans/{first['plan_id']}", params={
        "context": "planning", "task_id": saved["task_id"],
    }).json()["data"]
    second_detail = client.get(f"/api/v1/plans/{second['plan_id']}", params={
        "context": "planning", "task_id": saved["task_id"],
    }).json()["data"]
    assert first_detail["state"] == "CONFIRMED"
    assert first_detail["current"] is False
    assert second_detail["state"] == "CONFIRMED"
    assert second_detail["current"] is True


def test_compare_versioned_plans_returns_complete_rows_in_request_order(
        api, monkeypatch):
    client, _, _ = api
    saved = _save_task(client)
    first_job = client.post(
        "/api/v1/plan/sr4-flow", json=_submit(saved)
    ).json()["data"]
    _complete_job(monkeypatch, saved)
    first = client.post(
        "/api/v1/plans", json=_save_plan(saved, first_job)
    ).json()["data"]
    client.post(
        f"/api/v1/plans/{first['plan_id']}/confirm",
        json=_confirm_plan(saved, first),
    )

    second_job = client.post(
        "/api/v1/plan/sr4-flow", json=_submit(saved, "submit-compare-2")
    ).json()["data"]
    second = client.post(
        "/api/v1/plans", json=_save_plan(saved, second_job, "save-compare-2")
    ).json()["data"]
    response = client.post("/api/v1/plans/compare", json={
        "context": "planning",
        "task_id": saved["task_id"],
        "solution_ids": [second["plan_id"], first["plan_id"]],
    })

    assert response.status_code == 200
    compared = response.json()["data"]
    assert [row["plan_id"] for row in compared] == [
        second["plan_id"], first["plan_id"],
    ]
    assert compared[0]["result"] == second["result"]
    assert compared[1]["state"] == "CONFIRMED"
    assert compared[1]["current"] is True

    extra = []
    for index in (3, 4):
        job = client.post(
            "/api/v1/plan/sr4-flow",
            json=_submit(saved, f"submit-compare-{index}"),
        ).json()["data"]
        extra.append(client.post(
            "/api/v1/plans",
            json=_save_plan(saved, job, f"save-compare-{index}"),
        ).json()["data"])
    four_ids = [first["plan_id"], second["plan_id"],
                extra[0]["plan_id"], extra[1]["plan_id"]]
    four = client.post("/api/v1/plans/compare", json={
        "context": "planning",
        "task_id": saved["task_id"],
        "solution_ids": four_ids,
    })
    assert four.status_code == 200
    assert [row["plan_id"] for row in four.json()["data"]] == four_ids


def test_compare_versioned_plans_validates_selection_and_owner(api, monkeypatch):
    client, _, _ = api
    saved = _save_task(client)
    job = client.post("/api/v1/plan/sr4-flow", json=_submit(saved)).json()["data"]
    _complete_job(monkeypatch, saved)
    plan = client.post(
        "/api/v1/plans", json=_save_plan(saved, job)
    ).json()["data"]
    base = {"context": "planning", "task_id": saved["task_id"]}

    for solution_ids in ([plan["plan_id"]], [plan["plan_id"], plan["plan_id"]],
                         ["P-1", "P-2", "P-3", "P-4", "P-5"]):
        response = client.post(
            "/api/v1/plans/compare", json={**base, "solution_ids": solution_ids}
        )
        assert response.status_code == 422
        assert response.json()["code"] == 1001

    missing = client.post("/api/v1/plans/compare", json={
        **base, "solution_ids": [plan["plan_id"], "PLAN-missing"],
    })
    assert missing.status_code == 404
    assert missing.json()["code"] == 2001

    other_saved = _save_task(client, task_id="frontend-task-compare-other")
    other_job = client.post(
        "/api/v1/plan/sr4-flow", json=_submit(other_saved, "submit-compare-other")
    ).json()["data"]
    _complete_job(monkeypatch, other_saved)
    other_plan = client.post(
        "/api/v1/plans", json=_save_plan(
            other_saved, other_job, "save-compare-other"
        )
    ).json()["data"]
    cross_task = client.post("/api/v1/plans/compare", json={
        **base, "solution_ids": [plan["plan_id"], other_plan["plan_id"]],
    })
    assert cross_task.status_code == 404
    assert cross_task.json()["code"] == 2001


def test_compare_rejects_different_frozen_conditions_and_keeps_legacy(
        api, monkeypatch):
    client, _, _ = api
    saved = _save_task(client)
    first_job = client.post(
        "/api/v1/plan/sr4-flow", json=_submit(saved)
    ).json()["data"]
    _complete_job(monkeypatch, saved)
    first = client.post(
        "/api/v1/plans", json=_save_plan(saved, first_job)
    ).json()["data"]

    changed_input = deepcopy(saved["input"])
    changed_input["demands"][0]["maxDelayMs"] += 1
    updated = client.put(f"/api/v1/task-inputs/{saved['task_id']}", json={
        "context": "planning",
        "task_id": saved["task_id"],
        "expected_revision": saved["revision"],
        "resource_revision": saved["resource_revision"],
        "request_id": "update-compare-condition",
        "input": changed_input,
    }).json()["data"]
    second_job = client.post(
        "/api/v1/plan/sr4-flow", json=_submit(updated, "submit-after-update")
    ).json()["data"]
    _complete_job(monkeypatch, updated)
    second = client.post(
        "/api/v1/plans", json=_save_plan(updated, second_job, "save-after-update")
    ).json()["data"]

    response = client.post("/api/v1/plans/compare", json={
        "context": "planning",
        "task_id": saved["task_id"],
        "solution_ids": [first["plan_id"], second["plan_id"]],
    })
    assert response.status_code == 409
    assert response.json()["code"] == 1001
    assert response.json()["data"]["suggestions"]

    legacy = client.post(
        "/api/v1/plans/compare", json={"solution_ids": ["PL-missing"]}
    )
    assert legacy.status_code == 200
    assert legacy.json()["code"] == 2001


def test_legacy_compare_rejects_empty_or_invalid_selection(api):
    client, _, _ = api
    for payload in ({}, {"solution_ids": []}, {"solution_ids": None},
                    {"solution_ids": [""]}, {"solution_ids": "PL-1"}):
        response = client.post("/api/v1/plans/compare", json=payload)
        assert response.status_code == 422
        assert response.json()["code"] == 1001
        assert "solution_ids" in response.json()["message"]


def test_export_versioned_plan_as_json_and_xlsx(api, monkeypatch):
    client, _, _ = api
    saved = _save_task(client)
    job = client.post("/api/v1/plan/sr4-flow", json=_submit(saved)).json()["data"]
    _complete_job(monkeypatch, saved)
    plan = client.post(
        "/api/v1/plans", json=_save_plan(saved, job)
    ).json()["data"]
    path = f"/api/v1/plans/{plan['plan_id']}/export"
    params = {"context": "planning", "task_id": saved["task_id"]}

    exported_json = client.get(path, params={**params, "format": "json"})
    assert exported_json.status_code == 200
    assert exported_json.json()["data"] == plan

    exported_xlsx = client.get(path, params={**params, "format": "xlsx"})
    assert exported_xlsx.status_code == 200
    assert exported_xlsx.content.startswith(b"PK")
    assert exported_xlsx.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert f"plan-{plan['plan_id']}.xlsx" in exported_xlsx.headers[
        "content-disposition"
    ]

    book = load_workbook(BytesIO(exported_xlsx.content), read_only=True)
    try:
        assert {"方案概要", "任务输入", "通联需求", "节点", "设备", "链路",
                "部署结果", "设备参数", "路由", "频率", "指标", "检查项",
                "完整快照"}.issubset(book.sheetnames)
        summary = {row[0]: row[1] for row in book["方案概要"].iter_rows(
            min_row=2, values_only=True
        )}
        assert summary["plan_id"] == plan["plan_id"]
        chunks = [row[1] for row in book["完整快照"].iter_rows(
            min_row=2, values_only=True
        )]
        assert json.loads("".join(chunks))["result"] == plan["result"]
    finally:
        book.close()


def test_export_versioned_plan_validates_owner_format_and_legacy(api, monkeypatch):
    client, _, _ = api
    saved = _save_task(client)
    job = client.post("/api/v1/plan/sr4-flow", json=_submit(saved)).json()["data"]
    _complete_job(monkeypatch, saved)
    plan = client.post(
        "/api/v1/plans", json=_save_plan(saved, job)
    ).json()["data"]
    path = f"/api/v1/plans/{plan['plan_id']}/export"

    wrong_owner = client.get(path, params={
        "context": "planning", "task_id": "another-task", "format": "json",
    })
    assert wrong_owner.status_code == 404
    assert wrong_owner.json()["code"] == 2001

    invalid_format = client.get(path, params={
        "context": "planning", "task_id": saved["task_id"], "format": "csv",
    })
    assert invalid_format.status_code == 422
    assert invalid_format.json()["code"] == 1001

    incomplete = client.get(path, params={"task_id": saved["task_id"]})
    assert incomplete.status_code == 422
    assert incomplete.json()["code"] == 1001

    legacy = client.get(
        "/api/v1/plans/PL-missing/export", params={"format": "csv"}
    )
    assert legacy.status_code == 200
    assert legacy.json()["code"] == 2001


def test_cancel_pending_job_is_persistent_and_idempotent(api, monkeypatch):
    client, factory, _ = api
    saved = _save_task(client)
    job = client.post("/api/v1/plan/sr4-flow", json=_submit(saved)).json()["data"]
    path = f"/api/v1/tasks/{job['job_id']}"
    request = _cancel(saved)

    first = client.request("DELETE", path, json=request)
    replay = client.request("DELETE", path, json=request)
    assert first.status_code == replay.status_code == 200
    assert first.json() == replay.json()
    assert first.json()["request_id"] == request["request_id"]
    assert first.json()["data"]["status"] == "CANCELLED"
    assert first.json()["data"]["input"] == saved["input"]

    polled = client.get(f"/api/v1/plan/tasks/{job['job_id']}", params={
        "context": "planning", "task_id": saved["task_id"],
    })
    assert polled.json()["data"]["status"] == "CANCELLED"
    monkeypatch.setattr(planning_job_service.planning_w3, "get_task_record",
                        lambda _: {"status": "SUCCESS", "progress": 1.0,
                                   "stage": "晚到结果", "result": {"solution_id": "STALE"}})
    stale = client.get(f"/api/v1/plan/tasks/{job['job_id']}", params={
        "context": "planning", "task_id": saved["task_id"],
    })
    assert stale.json()["data"]["status"] == "CANCELLED"
    with factory() as session:
        assert session.get(planning_job_models.PlanningJob, job["job_id"]).status == "CANCELLED"
        receipts = session.scalars(select(planning_job_models.PlanningJobCancellation)).all()
        assert len(receipts) == 1
        assert receipts[0].outcome_status == "CANCELLED"


def test_cancel_checks_owner_frozen_revision_and_request_digest(api):
    client, _, _ = api
    saved = _save_task(client)
    job = client.post("/api/v1/plan/sr4-flow", json=_submit(saved)).json()["data"]
    path = f"/api/v1/tasks/{job['job_id']}"

    wrong_owner = client.request("DELETE", path, json={
        **_cancel(saved), "task_id": "different-task",
    })
    assert wrong_owner.status_code == 404
    assert wrong_owner.json()["code"] == 2001

    wrong_revision = client.request("DELETE", path, json={
        **_cancel(saved), "expected_revision": "TREV-stale",
    })
    assert wrong_revision.status_code == 409
    assert wrong_revision.json()["code"] == 1203

    malformed = client.request("DELETE", path, json={**_cancel(saved), "extra": 1})
    assert malformed.status_code == 422
    assert malformed.json()["code"] == 1001

    missing = client.request("DELETE", "/api/v1/tasks/tsk-missing0000", json=_cancel(saved))
    assert missing.status_code == 404
    assert missing.json()["code"] == 2001

    assert client.request("DELETE", path, json=_cancel(saved)).status_code == 200
    reused = client.request("DELETE", path, json={
        **_cancel(saved), "expected_revision": "TREV-stale",
    })
    assert reused.status_code == 409
    assert reused.json()["code"] == 1204


def test_cancel_uses_job_revision_after_task_input_changes(api):
    client, _, _ = api
    saved = _save_task(client)
    job = client.post("/api/v1/plan/sr4-flow", json=_submit(saved)).json()["data"]
    changed = deepcopy(saved["input"])
    changed["name"] = "已修改的任务"
    latest = client.put(f"/api/v1/task-inputs/{saved['task_id']}", json={
        "context": "planning", "task_id": saved["task_id"],
        "expected_revision": saved["revision"],
        "resource_revision": saved["resource_revision"],
        "request_id": "save-task-again", "input": changed,
    })
    assert latest.status_code == 200
    assert latest.json()["data"]["revision"] != saved["revision"]
    response = client.request("DELETE", f"/api/v1/tasks/{job['job_id']}",
                              json=_cancel(saved))
    assert response.status_code == 200
    assert response.json()["data"]["revision"] == saved["revision"]
    assert response.json()["data"]["status"] == "CANCELLED"


def test_cancel_does_not_rewrite_completed_job_or_lose_result(api, monkeypatch):
    client, factory, _ = api
    saved = _save_task(client)
    job = client.post("/api/v1/plan/sr4-flow", json=_submit(saved)).json()["data"]
    _complete_job(monkeypatch, saved, conflict=True)
    path = f"/api/v1/tasks/{job['job_id']}"

    response = client.request("DELETE", path, json=_cancel(saved))
    assert response.status_code == 200
    assert response.json()["code"] == 0
    assert response.json()["data"]["status"] == "SUCCESS"
    assert response.json()["data"]["result"] is not None
    assert response.json()["data"]["outcome"]["code"] == 3001
    with factory() as session:
        assert session.get(planning_job_models.PlanningJob, job["job_id"]).status == "SUCCESS"
        receipt = session.get(planning_job_models.PlanningJobCancellation, "cancel-1")
        assert receipt.outcome_status == "SUCCESS"


def test_cancel_failed_job_preserves_error(api):
    client, factory, _ = api
    saved = _save_task(client)
    job = client.post("/api/v1/plan/sr4-flow", json=_submit(saved)).json()["data"]
    planning_job_service._update_job(factory, job["job_id"], {
        "status": "FAILED", "progress": 1.0, "stage": "失败", "error": "算法错误",
    })
    response = client.request("DELETE", f"/api/v1/tasks/{job['job_id']}",
                              json=_cancel(saved))
    assert response.status_code == 200
    assert response.json()["data"]["status"] == "FAILED"
    assert response.json()["data"]["error"] == "算法错误"


def test_cancel_stops_worker_at_checkpoint_and_rejects_late_callback(api, monkeypatch):
    client, factory, _ = api
    started = threading.Event()
    release = threading.Event()
    finished = threading.Event()

    def controlled(_payload, tick):
        try:
            started.set()
            assert release.wait(5)
            tick(0.8, "取消后的检查点")
            return {"solution_id": "SHOULD-NOT-SAVE"}
        finally:
            finished.set()

    monkeypatch.setattr(planning_job_service.planning_w4, "submit_sr4",
                        lambda payload, *, job_id, on_update:
                        planning_job_service.planning_w3._spawn(
                            "SR4_FLOW", payload, controlled, task_id=job_id,
                            on_update=on_update))
    saved = _save_task(client)
    job = client.post("/api/v1/plan/sr4-flow", json=_submit(saved)).json()["data"]
    try:
        assert started.wait(5)
        response = client.request("DELETE", f"/api/v1/tasks/{job['job_id']}",
                                  json=_cancel(saved))
        assert response.status_code == 200
        assert response.json()["data"]["status"] == "CANCELLED"
    finally:
        release.set()
        assert finished.wait(5)
    planning_job_service._update_job(factory, job["job_id"], {
        "status": "SUCCESS", "progress": 1.0, "stage": "晚到回调", "error": None,
    })
    polled = client.get(f"/api/v1/plan/tasks/{job['job_id']}", params={
        "context": "planning", "task_id": saved["task_id"],
    })
    assert polled.json()["data"]["status"] == "CANCELLED"
    assert planning_job_service.planning_w3.get_task_record(job["job_id"])["status"] == "CANCELLED"
    with factory() as session:
        assert session.get(planning_job_models.PlanningJob, job["job_id"]).status == "CANCELLED"


def test_sr4_checks_cancellation_before_saving_plan(monkeypatch):
    import sr4_flow
    from app import planning_w3, planning_w4

    result = {
        "task_id": "T1", "total_s": 1, "within_limit": True,
        "optimization": {"best": {"objectives": {}}, "feasible": True},
        "frequency": {"conflicts": []},
        "check": {"passed": True, "summary": {}},
        "interference": {"analysis": {"summary": {"affected_links": 0}}},
        "steps": [], "routes": {}, "params": {}, "links": [], "nodes": [],
    }
    monkeypatch.setattr(sr4_flow, "run", lambda **_kwargs: result)
    monkeypatch.setattr(planning_w4, "_save_plan",
                        lambda *_args: pytest.fail("已取消的方案不应保存"))
    stages = []

    def checkpoint(progress, stage):
        stages.append((progress, stage))
        if progress == 0.99:
            raise planning_w3.PlanningTaskCancelled()

    with pytest.raises(planning_w3.PlanningTaskCancelled):
        planning_w4._do_sr4({"task_id": "T1"}, checkpoint)
    assert stages[-1] == (0.99, "正在保存方案")


@pytest.fixture
def local_owner_api(monkeypatch, tmp_path):
    """R10仅验证数据库/受理语义，算法运行记录为显式桩，不启动线程。"""
    from app import database, resource_dataset_service
    assert tmp_path.resolve().is_relative_to((ROOT / "data/runtime/tmp").resolve())
    path = tmp_path / "local-owner.sqlite"
    assert not path.exists(), "每轮使用新GUID基目录，不重建已有库"
    engine = create_engine("sqlite:///" + str(path), connect_args={"check_same_thread": False})
    Base.metadata.create_all(engine)
    factory = sessionmaker(bind=engine, expire_on_commit=False)
    for module in (database, resource_dataset_service, task_input_service,
                   planning_job_service, planning_plan_service):
        monkeypatch.setattr(module, "get_session_factory", lambda: factory)
    monkeypatch.setenv("FRONTEND_INTEGRATION_DEMO", "1")
    monkeypatch.setenv("RESOURCE_DATASET_MODE", "database")
    resource_dataset_service.initialize(factory)
    runtime_root = ROOT / "data/runtime/r2"
    monkeypatch.setattr(planning_job_service.planning_w3, "TASK_DIR", runtime_root / "owner-tests/tasks")
    monkeypatch.setattr(planning_job_service.planning_w4, "PLAN_DIR", runtime_root / "owner-tests/plans")
    submitted, cancelled, records = [], [], {}
    def fake_submit(payload, *, job_id, on_update):
        submitted.append(job_id)
        records[job_id] = {"status": "PENDING", "progress": 0, "stage": "显式R10测试桩"}
        on_update(deepcopy(records[job_id]))
        return {"code": 0, "data": {"task_id": job_id}}
    def fake_cancel(job_id):
        cancelled.append(job_id)
        records[job_id].update(status="CANCELLED", progress=1, stage="桩取消请求，不是CPU确认")
        return deepcopy(records[job_id])
    monkeypatch.setattr(planning_job_service.planning_w4, "submit_sr4", fake_submit)
    monkeypatch.setattr(planning_job_service.planning_w3, "get_task_record", lambda job_id: deepcopy(records.get(job_id)))
    monkeypatch.setattr(planning_job_service.planning_w3, "cancel_task", fake_cancel)
    try:
        yield TestClient(app), factory, submitted, cancelled
    finally:
        engine.dispose()


def _local_owner_case(local_owner_api, kind):
    import hashlib
    from app.planning_job_schemas import PlanningJobCancelRequest, PlanningJobSubmitRequest
    client, factory, _, _ = local_owner_api
    saved = _save_task(client)
    submission = PlanningJobSubmitRequest.model_validate(_submit(saved, "R10-SUBMIT"))
    if kind == "submit":
        def invoke(*, actor_id=None, local_owner=False, factory_override=None):
            return planning_job_service.submit_planning_job(submission, factory_override or factory,
                actor_id=actor_id, local_owner=local_owner)
        return invoke, planning_job_models.PlanningJob, planning_job_service._digest(submission), "planning.submit"
    job = planning_job_service.submit_planning_job(submission, factory)
    payload = PlanningJobCancelRequest.model_validate(_cancel(saved, "R10-CANCEL"))
    if kind == "child_cancel":
        with factory.begin() as session:
            session.get(planning_job_models.PlanningJob, job["job_id"]).adjustment = {
                "interference_origin": {"adoption_id": "UNIT-OWNER", "actor_id": "ORIGIN-OWNER"}}
    def invoke(*, actor_id=None, local_owner=False, factory_override=None):
        return planning_job_service.cancel_planning_job(job["job_id"], payload, factory_override or factory,
            actor_id=actor_id, local_owner=local_owner)
    digest = hashlib.sha256(f"{job['job_id']}:{planning_job_service._digest(payload)}".encode("utf-8")).hexdigest()
    return invoke, planning_job_models.PlanningJobCancellation, digest, "planning.cancel"


@pytest.mark.parametrize("kind", ["submit", "cancel", "child_cancel"])
def test_local_owner_same_account_replays_and_other_account_conflicts(local_owner_api, kind):
    from app.local_operation_digest import bind_local_operation_digest
    invoke, model, body_digest, operation = _local_owner_case(local_owner_api, kind)
    _, factory, submitted, cancelled = local_owner_api
    first = invoke(actor_id="ACCOUNT-A", local_owner=True)
    assert invoke(actor_id="ACCOUNT-A", local_owner=True) == first
    side_effects = (len(submitted), len(cancelled))
    with pytest.raises(planning_job_service.PlanningJobError) as rejected:
        invoke(actor_id="ACCOUNT-B", local_owner=True)
    assert (rejected.value.code, rejected.value.status) == (1204, 409)
    assert (len(submitted), len(cancelled)) == side_effects
    with factory() as session:
        rows = session.scalars(select(model)).all()
        assert len(rows) == 1
        assert rows[0].request_digest == bind_local_operation_digest(
            body_digest, operation=operation, account_id="ACCOUNT-A")
    assert len(submitted) == 1


@pytest.mark.parametrize("kind", ["submit", "cancel", "child_cancel"])
def test_local_owner_never_replays_legacy_unbound_receipt(local_owner_api, kind):
    invoke, model, _, _ = _local_owner_case(local_owner_api, kind)
    _, factory, submitted, cancelled = local_owner_api
    invoke()
    with factory() as session:
        legacy = session.scalar(select(model)).request_digest
    side_effects = (len(submitted), len(cancelled))
    with pytest.raises(planning_job_service.PlanningJobError) as rejected:
        invoke(actor_id="ACCOUNT-A", local_owner=True)
    assert (rejected.value.code, rejected.value.status) == (1204, 409)
    assert (len(submitted), len(cancelled)) == side_effects
    with factory() as session:
        assert session.scalar(select(model)).request_digest == legacy


@pytest.mark.parametrize("kind", ["submit", "cancel", "child_cancel"])
def test_local_owner_off_keeps_legacy_digest_and_ignores_public_actor(local_owner_api, kind):
    from app.interference_analysis_engine import digest
    invoke, model, body_digest, _ = _local_owner_case(local_owner_api, kind)
    _, factory, _, _ = local_owner_api
    first = invoke()
    assert invoke(actor_id="IGNORED-OFF-ACCOUNT", local_owner=False) == first
    with factory() as session:
        row = session.scalar(select(model))
        if kind == "child_cancel":
            from app.planning_job_schemas import PlanningJobCancelRequest
            job = session.get(planning_job_models.PlanningJob, row.job_id)
            payload = PlanningJobCancelRequest(context="planning", request_id="R10-CANCEL",
                task_id=job.task_id, expected_revision=job.revision_id)
            expected = digest({"job_id": job.job_id, "body": payload.model_dump(), "actor_id": None})
        else:
            expected = body_digest
        assert row.request_digest == expected


@pytest.mark.parametrize("kind", ["submit", "cancel", "child_cancel"])
@pytest.mark.parametrize("actor_id", [None, "", "   "])
def test_local_owner_missing_actor_fails_closed_before_database(local_owner_api, kind, actor_id):
    invoke, model, _, _ = _local_owner_case(local_owner_api, kind)
    _, factory, submitted, cancelled = local_owner_api
    def forbidden_factory():
        pytest.fail("LOCAL缺可信actor必须在访问数据库前503")
    side_effects = (len(submitted), len(cancelled))
    with pytest.raises(planning_job_service.PlanningJobError) as rejected:
        invoke(actor_id=actor_id, local_owner=True, factory_override=forbidden_factory)
    assert (rejected.value.code, rejected.value.status) == (5001, 503)
    assert (len(submitted), len(cancelled)) == side_effects
    with factory() as session:
        assert session.scalar(select(func.count()).select_from(model)) == 0


@pytest.mark.parametrize("bound_actor", [None, "ACCOUNT-A"])
@pytest.mark.parametrize("replay_actor", ["ACCOUNT-A", "ACCOUNT-B"])
def test_local_owner_submit_integrity_race_rechecks_same_owner_digest(local_owner_api, monkeypatch, bound_actor, replay_actor):
    from sqlalchemy.exc import IntegrityError
    invoke, model, _, _ = _local_owner_case(local_owner_api, "submit")
    _, factory, submitted, _ = local_owner_api
    first = invoke(actor_id=bound_actor, local_owner=bound_actor is not None)
    def race(*args):
        raise IntegrityError("显式模拟唯一号竞争回读", {}, RuntimeError("UNIT-RACE"))
    monkeypatch.setattr(planning_job_service, "_reserve", race)
    if replay_actor == bound_actor:
        assert invoke(actor_id=replay_actor, local_owner=True) == first
    else:
        with pytest.raises(planning_job_service.PlanningJobError) as rejected:
            invoke(actor_id=replay_actor, local_owner=True)
        assert (rejected.value.code, rejected.value.status) == (1204, 409)
    assert len(submitted) == 1
    with factory() as session:
        assert session.scalar(select(func.count()).select_from(model)) == 1


@pytest.mark.parametrize("replay_actor", ["ACCOUNT-A", "ACCOUNT-B"])
def test_local_owner_cancel_second_receipt_read_rechecks_owner(local_owner_api, monkeypatch, replay_actor):
    invoke, model, _, _ = _local_owner_case(local_owner_api, "cancel")
    _, factory, _, cancelled = local_owner_api
    first = invoke(actor_id="ACCOUNT-A", local_owner=True)
    real_get = factory.class_.get
    reads = []
    def raced_get(session, entity, ident, *args, **kwargs):
        if entity is model and ident == "R10-CANCEL":
            reads.append(ident)
            if len(reads) == 1:
                return None
        return real_get(session, entity, ident, *args, **kwargs)
    monkeypatch.setattr(factory.class_, "get", raced_get)
    if replay_actor == "ACCOUNT-A":
        assert invoke(actor_id=replay_actor, local_owner=True) == first
    else:
        with pytest.raises(planning_job_service.PlanningJobError) as rejected:
            invoke(actor_id=replay_actor, local_owner=True)
        assert (rejected.value.code, rejected.value.status) == (1204, 409)
    assert len(reads) == 2 and len(cancelled) == 1


def test_local_owner_internal_adoption_and_shutdown_keep_origin_digest(local_owner_api, monkeypatch):
    from app.interference_analysis_engine import digest
    from app.planning_job_schemas import PlanningJobCancelRequest
    _, model, _, _ = _local_owner_case(local_owner_api, "child_cancel")
    _, factory, _, _ = local_owner_api
    with factory() as session:
        job = session.scalar(select(planning_job_models.PlanningJob))
    monkeypatch.setenv("AUTH_MODE", "LOCAL_PASSWORD")
    payload = PlanningJobCancelRequest(context="planning", request_id="UNIT-INTERNAL-CANCEL",
        task_id=job.task_id, expected_revision=job.revision_id)
    planning_job_service.request_interference_cancellation(job.job_id, payload, factory, actor_id="ORIGIN-OWNER")
    planning_job_service.request_interference_cancellation(job.job_id, payload, factory, actor_id="ORIGIN-OWNER")
    assert planning_job_service.stop_frozen_planning_jobs(factory, timeout=0) is False
    with factory() as session:
        rows = session.scalars(select(model)).all()
        assert len(rows) == 2
        for row in rows:
            body = payload.model_copy(update={"request_id": row.request_id})
            assert row.request_digest == digest({"job_id": job.job_id, "body": body.model_dump(),
                "actor_id": "ORIGIN-OWNER"})
            assert row.outcome_status == "CANCEL_REQUESTED"


def test_local_owner_invalid_scope_flag_is_wiring_failure(local_owner_api):
    invoke, _, _, _ = _local_owner_case(local_owner_api, "submit")
    with pytest.raises(planning_job_service.PlanningJobError) as rejected:
        invoke(actor_id="ACCOUNT-A", local_owner="false")
    assert (rejected.value.code, rejected.value.status) == (5001, 503)
