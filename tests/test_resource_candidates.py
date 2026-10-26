"""候选服务纯内存验证，不初始化数据库、不启动服务、不读取业务原件。"""
from copy import deepcopy
import csv
from datetime import datetime, time, timedelta
from io import BytesIO, StringIO
import json
from pathlib import Path
import sys
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "algorithm-service"))
from app import resource_candidate_service as service
from app import resource_import_parser as parser
from app.resource_candidate_schemas import ASSESSMENT_FIELDS, CONFIG_FIELDS
from app.resource_dataset_service import ResourceDatasetError


@pytest.fixture
def snapshot():
    return {"revision": "resource-test", "source": "SYNTHETIC_INTEGRATION", "tables": {
        "node.csv": [{"node_id": f"N{i}"} for i in range(1, 4)],
        "device.csv": [
            {"device_id": "D1", "node_id": "N1", "model_id": "M1"},
            {"device_id": "D2", "node_id": "N2", "model_id": "M2"},
            {"device_id": "D3", "node_id": "N3", "model_id": "M1"},
            {"device_id": "V1", "node_id": "N1", "model_id": "M3"},
            {"device_id": "V2", "node_id": "N2", "model_id": "M3"},
        ],
        "device_model.csv": [
            {"model_id": "M1", "device_class": "HF", "freq_min_khz": "1000", "freq_max_khz": "30000"},
            {"model_id": "M2", "device_class": "HF", "freq_min_khz": "2000", "freq_max_khz": "20000"},
            {"model_id": "M3", "device_class": "VUHF", "freq_min_khz": "30000", "freq_max_khz": "300000"},
        ],
        "link_metric.csv": [{"record_id": "historic", "path_loss_db": "80"}],
    }}


def manual(**changes):
    return {"id": "C1", "nodeAId": "N1", "nodeBId": "N2", "deviceAId": "D1", "deviceBId": "D2",
            "deviceClass": "HF", "freqKhz": None, "bandwidthKhz": None, "isBackup": None,
            **dict.fromkeys(ASSESSMENT_FIELDS), "assessmentState": "PENDING",
            "source": {"kind": "MANUAL", "file": None, "row": None}, "reference": None, **changes}


def file_row(**changes):
    return {"link_id": "C1", "node_a_id": "N1", "node_b_id": "N2", "device_a_id": "D1", "device_b_id": "D2",
            "device_class": "HF", "freq_khz": "", "bandwidth_khz": "", "distance_m": "", "path_loss_db": "",
            "rx_power_dbm": "", "link_margin_db": "", "is_available": "", "link_state": "", "is_backup": "", **changes}


def csv_bytes(rows, head=None, bom=False):
    head = service.HEADERS if head is None else head
    output = StringIO(newline="")
    writer = csv.writer(output)
    writer.writerow(head)
    for row in rows:
        writer.writerow([row.get(field, "") for field in head])
    return (b"\xef\xbb\xbf" if bom else b"") + output.getvalue().encode("utf-8")


def parse(snapshot, rows, filename="input.csv", **kwargs):
    return service.parse_candidate_import(snapshot, filename, csv_bytes(rows, **kwargs))


def workbook_bytes(rows=None, mutate=None):
    book = Workbook()
    try:
        sheet = book.active
        sheet.title = service.SHEET
        sheet.append(service.HEADERS)
        for row in [file_row()] if rows is None else rows:
            sheet.append([row.get(field, "") for field in service.HEADERS])
            for cell in sheet[sheet.max_row]:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
        if mutate:
            mutate(book)
        return parser._save_memory(book)
    finally:
        book.close()


def rewrite_zip(data, mutate):
    output = BytesIO()
    with ZipFile(BytesIO(data)) as original, ZipFile(output, "w", ZIP_DEFLATED) as target:
        members = {item.filename: original.read(item) for item in original.infolist()}
        mutate(members)
        for name, value in members.items():
            target.writestr(name, value)
    return output.getvalue()


def assert_failure(snapshot, result):
    assert result["issues"]
    assert result["links"] == service.get_links(snapshot)
    assert all(set(issue) == {"sheet", "row", "objectId", "field", "message"} for issue in result["issues"])
    assert all(issue["sheet"] == "topology_links" for issue in result["issues"])


def assert_rejected(snapshot, changes, field=None):
    before = deepcopy(snapshot)
    original_changes = deepcopy(changes)
    with pytest.raises(ResourceDatasetError) as caught:
        service.apply_candidate_changes(snapshot, changes)
    assert caught.value.code == 1001 and caught.value.status == 422
    assert caught.value.issues
    if field:
        assert any(field in issue["field"] for issue in caught.value.issues)
    assert snapshot == before and changes == original_changes


def with_reference(snapshot, **changes):
    result = parse(snapshot, [file_row(path_loss_db="0", rx_power_dbm="-100", is_available="false", **changes)])
    assert not result["issues"]
    snapshot[service.SNAPSHOT_KEY] = result["links"]
    return result["links"][0]


def test_default_is_empty_and_historic_metrics_never_become_current(snapshot):
    before = deepcopy(snapshot)
    assert service.get_links(snapshot) == []
    assert service.apply_candidate_changes(snapshot, {"stations": [{"anything": "ignored"}], "devices": None}) == []
    assert snapshot == before


def test_get_returns_deeply_detached_canonical_links(snapshot):
    with_reference(snapshot)
    snapshot[service.SNAPSHOT_KEY][0]["notices"] = ["原提示"]
    before = deepcopy(snapshot)
    links = service.get_links(snapshot)
    links[0]["reference"]["data"]["pathLossDb"] = 999
    links[0]["reference"]["source"]["file"] = "another.csv"
    links[0]["notices"].append("修改")
    assert snapshot == before


@pytest.mark.parametrize("bad", [None, {}, "links", [manual(), manual()], [{"id": "incomplete"}]])
def test_invalid_snapshot_candidates_are_rejected(snapshot, bad):
    snapshot[service.SNAPSHOT_KEY] = bad
    with pytest.raises(ResourceDatasetError) as caught:
        service.get_links(snapshot)
    assert (caught.value.code, caught.value.status) == (1001, 422)


def test_manual_add_retains_nulls_and_inputs(snapshot):
    before, row = deepcopy(snapshot), manual()
    changes = {"newCandidateLinks": [row], "stations": [{"id": "unrelated"}], "interferences": []}
    result = service.apply_candidate_changes(snapshot, changes)
    assert result == [row]
    assert snapshot == before and changes["newCandidateLinks"] == [manual()]
    result[0]["source"]["kind"] = "IMPORT"
    assert row == manual()


@pytest.mark.parametrize("field,value", [
    ("freqKhz", 0), ("freqKhz", -1), ("freqKhz", "2000"), ("freqKhz", True),
    ("freqKhz", float("nan")), ("freqKhz", float("inf")),
    ("bandwidthKhz", 0), ("bandwidthKhz", -1), ("bandwidthKhz", False),
    ("isBackup", 0), ("isBackup", "false"), ("id", " C1"), ("id", "C 1"),
    ("id", ""), ("id", "a" * 121), ("nodeAId", "N1\n"), ("deviceClass", "UHF"),
    ("distanceM", 0), ("pathLossDb", -20), ("isAvailable", False), ("linkState", "GOOD"),
    ("extra", 1), ("reference", {}), ("notices", [1]),
])
def test_manual_strict_fields_and_no_forged_assessments(snapshot, field, value):
    assert_rejected(snapshot, {"newCandidateLinks": [manual(**{field: value})]})


@pytest.mark.parametrize("source", [
    {"kind": "IMPORT", "file": "forged.csv", "row": 2},
    {"kind": "SOURCE", "file": "forged.csv", "row": 2},
    {"kind": "MANUAL", "file": "forged.csv", "row": None},
    {"kind": "MANUAL", "file": None, "row": 2},
    {"kind": "MANUAL", "file": None, "row": None, "extra": True},
])
def test_manual_cannot_forge_origin(snapshot, source):
    assert_rejected(snapshot, {"newCandidateLinks": [manual(source=source)]}, "source")


@pytest.mark.parametrize("key,value", [("candidateLinks", None), ("newCandidateLinks", {}),
                                       ("importedCandidateLinks", ""), ("importedCandidateLinks", [manual()])])
def test_change_collection_structure_and_import_bypass_are_rejected(snapshot, key, value):
    assert_rejected(snapshot, {key: value})


def test_full_collections_are_never_replaced_or_deleted(snapshot):
    snapshot[service.SNAPSHOT_KEY] = [manual(), manual(id="C2")]
    assert service.apply_candidate_changes(snapshot, {"candidateLinks": []}) == snapshot[service.SNAPSHOT_KEY]
    result = service.apply_candidate_changes(snapshot, {"candidateLinks": [{"id": "C1", "isBackup": True}]})
    assert result[1] == manual(id="C2")
    assert_rejected(snapshot, {"candidateLinks": [{"id": "C1", "deleted": True}]}, "deleted")
    assert_rejected(snapshot, {"candidateLinks": [{"id": "C3", "isBackup": False}]}, "id")
    assert_rejected(snapshot, {"newCandidateLinks": [manual()]}, "id")
    assert_rejected(snapshot, {"candidateLinks": [manual()]})


@pytest.mark.parametrize("changes", [
    {"newCandidateLinks": [manual(), manual()]},
    {"newCandidateLinks": [manual()], "candidateLinks": [{"id": "C1", "freqKhz": 2500}]},
])
def test_duplicate_changes_are_atomic(snapshot, changes):
    assert_rejected(snapshot, changes, "id")


@pytest.mark.parametrize("field,value", [("distanceM", 1), ("pathLossDb", 100), ("rxPowerDbm", -100),
                                        ("linkMarginDb", 0), ("isAvailable", False), ("linkState", "GOOD"),
                                        ("source", {}), ("reference", None), ("assessmentState", "PENDING"),
                                        ("notices", [])])
def test_patch_only_accepts_eight_config_fields(snapshot, field, value):
    snapshot[service.SNAPSHOT_KEY] = [manual()]
    assert_rejected(snapshot, {"candidateLinks": [{"id": "C1", field: value}]}, field)


@pytest.mark.parametrize("patch", [
    {"freqKhz": 3000}, {"bandwidthKhz": 9999999},
    {"nodeBId": "N3", "deviceBId": "D3"},
    {"nodeAId": "N2", "nodeBId": "N1", "deviceAId": "D2", "deviceBId": "D1"},
    {"deviceClass": "VUHF", "deviceAId": "V1", "deviceBId": "V2", "freqKhz": 50000},
])
def test_config_edits_clear_six_assessments_and_freeze_reference(snapshot, patch):
    old = deepcopy(with_reference(snapshot))
    before = deepcopy(snapshot)
    links = service.apply_candidate_changes(snapshot, {"candidateLinks": [{"id": "C1", **patch}]})
    link = links[0]
    assert all(link[field] is None for field in ASSESSMENT_FIELDS)
    assert link["assessmentState"] == "PENDING" and link["reference"] == old["reference"]
    assert link["source"] == old["source"] and snapshot == before
    assert link["reference"]["data"]["deviceBId"] == "D2"
    links[0]["reference"]["data"]["pathLossDb"] = 200
    assert snapshot == before


@pytest.mark.parametrize("backup", [True, False, None])
def test_backup_only_changes_do_not_invalidate_reference(snapshot, backup):
    old = deepcopy(with_reference(snapshot))
    links = service.apply_candidate_changes(snapshot, {"candidateLinks": [{"id": "C1", "isBackup": backup}]})
    assert links[0] == {**old, "isBackup": backup}


def test_unchanged_config_and_returning_to_reference_do_not_recertify(snapshot):
    old = deepcopy(with_reference(snapshot))
    no_op = {"id": "C1", **{field: old[field] for field in CONFIG_FIELDS}}
    assert service.apply_candidate_changes(snapshot, {"candidateLinks": [no_op]}) == [old]
    changed = service.apply_candidate_changes(snapshot, {"candidateLinks": [{"id": "C1", "freqKhz": 2500}]})
    snapshot[service.SNAPSHOT_KEY] = changed
    restored = service.apply_candidate_changes(snapshot, {"candidateLinks": [{"id": "C1", "freqKhz": None}]})[0]
    assert restored["assessmentState"] == "PENDING" and restored["pathLossDb"] is None
    assert restored["reference"] == old["reference"]


@pytest.mark.parametrize("changes", [
    {"nodeAId": "missing"}, {"nodeBId": "N1"}, {"deviceBId": "D1"}, {"deviceAId": "missing"},
    {"deviceAId": "D2"}, {"deviceClass": "VUHF"}, {"deviceBId": "V2"},
    {"freqKhz": 1500}, {"freqKhz": 20001}, {"nodeAId": None}, {"deviceClass": None},
])
def test_endpoint_and_both_model_frequency_checks(snapshot, changes):
    assert_rejected(snapshot, {"newCandidateLinks": [manual(**changes)]})


@pytest.mark.parametrize("frequency", [2000, 20000])
def test_frequency_boundaries_are_inclusive_and_bandwidth_limit_not_invented(snapshot, frequency):
    row = manual(freqKhz=frequency, bandwidthKhz=1e9)
    assert service.apply_candidate_changes(snapshot, {"newCandidateLinks": [row]}) == [row]


@pytest.mark.parametrize("value", ["", "NaN", "inf", "0", "-1", "2000 ", True])
def test_malformed_model_ranges_block_frequency_validation(snapshot, value):
    snapshot["tables"]["device_model.csv"][1]["freq_min_khz"] = value
    assert_rejected(snapshot, {"newCandidateLinks": [manual(freqKhz=2500)]}, "freqKhz")


def test_updated_snapshot_tables_are_used_and_old_reference_is_not_revalidated(snapshot):
    old = deepcopy(with_reference(snapshot))
    snapshot["tables"]["node.csv"] = [{"node_id": "N3"}, {"node_id": "N2"}]
    snapshot["tables"]["device.csv"] = [row for row in snapshot["tables"]["device.csv"] if row["device_id"] != "D1"]
    link = service.apply_candidate_changes(snapshot, {"candidateLinks": [{"id": "C1", "nodeAId": "N3", "deviceAId": "D3"}]})[0]
    assert link["reference"] == old["reference"] and link["assessmentState"] == "PENDING"


@pytest.mark.parametrize("bom", [False, True])
def test_csv_import_nulls_false_zero_negative_and_frozen_reference(snapshot, bom):
    before = deepcopy(snapshot)
    result = parse(snapshot, [file_row(path_loss_db="-5", rx_power_dbm="0", is_available="false", is_backup="false")], bom=bom)
    assert not result["issues"]
    link = result["links"][0]
    assert link["assessmentState"] == "REFERENCE" and link["pathLossDb"] == -5
    assert link["rxPowerDbm"] == 0 and link["isAvailable"] is False and link["isBackup"] is False
    assert link["freqKhz"] is None and link["bandwidthKhz"] is None and link["linkState"] is None
    assert link["source"] == {"kind": "IMPORT", "file": "input.csv", "row": 2}
    assert link["reference"]["data"]["pathLossDb"] == -5
    link["pathLossDb"] = 500
    link["source"]["file"] = "changed.csv"
    assert link["reference"]["data"]["pathLossDb"] == -5 and link["reference"]["source"]["file"] == "input.csv"
    assert snapshot == before


@pytest.mark.parametrize("field,value", [("distance_m", "0"), ("path_loss_db", "0"), ("rx_power_dbm", "0"),
                                        ("link_margin_db", "0"), ("is_available", "false"), ("link_state", "FAULT")])
def test_any_single_assessment_is_reference(snapshot, field, value):
    result = parse(snapshot, [file_row(**{field: value})])
    assert not result["issues"] and result["links"][0]["assessmentState"] == "REFERENCE"


def test_empty_assessments_stay_pending_and_warning_values_are_preserved(snapshot):
    result = parse(snapshot, [file_row()])
    assert not result["issues"]
    assert result["links"][0]["assessmentState"] == "PENDING" and result["links"][0]["reference"] is None
    warned = parse(snapshot, [file_row(path_loss_db="0", is_available="true", link_state="FAULT")])
    assert not warned["issues"]
    assert {warning["field"] for warning in warned["warnings"]} == {"reference", "path_loss_db", "link_state"}
    assert warned["links"][0]["pathLossDb"] == 0 and warned["links"][0]["isAvailable"] is True


def test_repeated_original_file_is_noop_with_unchanged_reference_and_notices(snapshot):
    row = file_row(path_loss_db="100")
    first = parse(snapshot, [row])
    snapshot[service.SNAPSHOT_KEY] = first["links"] + [manual(id="C2")]
    snapshot[service.SNAPSHOT_KEY][0]["notices"] = ["保留服务提示"]
    before = deepcopy(snapshot)
    second = parse(snapshot, [row])
    assert not second["issues"] and second["links"] == snapshot[service.SNAPSHOT_KEY]
    assert second["rows"] == [{"sheet": "topology_links", "row": 2, "id": "C1", "action": "UNCHANGED", "differences": []}]
    second["links"][0]["reference"]["data"]["freqKhz"] = 999
    assert snapshot == before


@pytest.mark.parametrize("change", ["config", "assessment", "source", "row"])
def test_config_reference_and_origin_changes_produce_updated_links(snapshot, change):
    with_reference(snapshot)
    row, filename = file_row(path_loss_db="0", rx_power_dbm="-100", is_available="false"), "input.csv"
    if change == "config":
        row["freq_khz"] = "2500"
    elif change == "assessment":
        row["path_loss_db"] = "80"
    elif change == "source":
        filename = "another.csv"
    data = csv_bytes([row])
    if change == "row":
        head, rest = data.split(b"\r\n", 1)
        data = head + b"\r\n\r\n" + rest
    result = service.parse_candidate_import(snapshot, filename, data)
    assert not result["issues"] and result["links"] != snapshot[service.SNAPSHOT_KEY]
    assert result["rows"][0]["action"] == "UPDATE" and result["rows"][0]["differences"]
    assert result["links"][0]["reference"]["source"] == result["links"][0]["source"]


def test_edit_then_reimport_original_explicitly_replaces_pending_with_reference(snapshot):
    with_reference(snapshot)
    snapshot[service.SNAPSHOT_KEY] = service.apply_candidate_changes(snapshot, {"candidateLinks": [{"id": "C1", "freqKhz": 2500}]})
    result = parse(snapshot, [file_row(path_loss_db="0", rx_power_dbm="-100", is_available="false")])
    assert not result["issues"] and result["rows"][0]["action"] == "UPDATE"
    assert result["links"][0]["freqKhz"] is None and result["links"][0]["assessmentState"] == "REFERENCE"


@pytest.mark.parametrize("field,value", [
    ("link_id", ""), ("link_id", "bad id"), ("link_id", "x" * 121), ("link_id", "A\nB"),
    ("node_a_id", "absent"), ("node_b_id", "N1"), ("device_a_id", "D2"), ("device_b_id", "D1"),
    ("device_b_id", "absent"), ("device_class", "VUHF"), ("freq_khz", "1500"), ("freq_khz", "20001"),
    ("freq_khz", "0"), ("bandwidth_khz", "-1"), ("freq_khz", "NaN"), ("freq_khz", "inf"),
    ("freq_khz", "1e999"), ("freq_khz", "true"), ("path_loss_db", "=1+2"),
    ("rx_power_dbm", "NaN"), ("distance_m", "1_000"), ("link_margin_db", "0x10"),
    ("freq_khz", " 2500"), ("is_available", "TRUE"), ("is_backup", "0"), ("link_state", "NORMAL"),
])
def test_import_issues_located_and_whole_batch_rolled_back(snapshot, field, value):
    snapshot[service.SNAPSHOT_KEY] = [manual(id="existing")]
    before = deepcopy(snapshot)
    result = parse(snapshot, [file_row(link_id="valid"), file_row(**{field: value})])
    assert_failure(snapshot, result)
    assert any(issue["row"] == 3 for issue in result["issues"])
    assert snapshot == before


def test_duplicate_import_ids_and_unknown_delete_columns_are_rejected(snapshot):
    duplicate = parse(snapshot, [file_row(), file_row(freq_khz="2500")])
    assert_failure(snapshot, duplicate)
    assert duplicate["issues"][0]["row"] == 3 and duplicate["issues"][0]["field"] == "link_id"
    for header in ("source", "reference", "delete", "unknown"):
        assert_failure(snapshot, parse(snapshot, [file_row()], head=[*service.HEADERS, header]))


@pytest.mark.parametrize("mutation", ["missing", "duplicate", "short", "long", "quote"])
def test_csv_headers_and_records_are_strict(snapshot, mutation):
    head = list(service.HEADERS)
    if mutation == "missing":
        head.pop()
    elif mutation == "duplicate":
        head[-1] = head[0]
    data = csv_bytes([file_row()], head)
    if mutation == "short":
        data = data[:-3] + b"\r\n"
    elif mutation == "long":
        data = data[:-2] + b",extra\r\n"
    elif mutation == "quote":
        data = data.split(b"\r\n")[0] + b'\r\n"unclosed'
    assert_failure(snapshot, service.parse_candidate_import(snapshot, "input.csv", data))


@pytest.mark.parametrize("value", ["~rp~invalid", "~rp~null", "~rp~false", "~rp~[]", '~rp~{"x":1}'])
def test_csv_invalid_text_escape(snapshot, value):
    result = parse(snapshot, [file_row(link_id=value)])
    assert_failure(snapshot, result)
    assert result["issues"][0]["field"] == "link_id"


@pytest.mark.parametrize("prefix", ["=N", "+N", "-N", "@N", "~rp~N"])
@pytest.mark.parametrize("format", ["csv", "xlsx"])
def test_template_escapes_real_reference_ids_and_roundtrips(snapshot, prefix, format):
    snapshot["tables"]["node.csv"][0]["node_id"] = prefix
    snapshot["tables"]["device.csv"][0]["node_id"] = prefix
    data, mime, filename = service.make_candidate_template(snapshot, format)
    result = service.parse_candidate_import(snapshot, filename, data)
    assert not result["issues"] and result["links"][0]["nodeAId"] == prefix
    if format == "csv":
        assert b"~rp~" in data


@pytest.mark.parametrize("filename,data", [
    ("input.csv", b""), ("input.csv", b"\xff"), ("input.csv", b"\x00"),
    ("input.csv", b"x" * (service.MAX_FILE_BYTES + 1)), ("input.xls", b"x"),
    ("input.xlsm", b"x"), ("input.xlsx", b"not zip"), (None, b"x"),
    (" input.csv", b"x"), ("input\n.csv", b"x"), ("input.csv", "text"),
], ids=["empty", "invalid-utf8", "nul", "oversize", "xls", "xlsm", "bad-xlsx", "null-name", "space-name", "control-name", "not-bytes"])
def test_file_type_encoding_size_and_name_gates(snapshot, filename, data):
    assert_failure(snapshot, service.parse_candidate_import(snapshot, filename, data))


@pytest.mark.parametrize("count", [0, 5000, 5001])
def test_csv_business_row_limits(snapshot, count):
    data = csv_bytes([file_row(link_id=f"C{i}") for i in range(count)])
    result = service.parse_candidate_import(snapshot, "input.csv", data)
    if count == 5000:
        assert not result["issues"] and len(result["links"]) == 5000
    else:
        assert_failure(snapshot, result)


def test_import_preview_shape_and_reordered_columns(snapshot):
    result = parse(snapshot, [file_row(freq_khz="2.5e3", bandwidth_khz="+1")], head=list(reversed(service.HEADERS)))
    assert set(result) == {"links", "rows", "issues", "warnings"} and not result["issues"]
    assert set(result["rows"][0]) == {"sheet", "row", "id", "action", "differences"}
    assert all(set(diff) == {"field", "before", "after"} for diff in result["rows"][0]["differences"])
    assert result["links"][0]["freqKhz"] == 2500


@pytest.mark.parametrize("mutation", [
    lambda book: setattr(book.active["I2"], "value", "=1+1"),
    lambda book: setattr(book.active["I2"], "value", "#DIV/0!"),
    lambda book: setattr(book.active["I2"], "value", datetime(2026, 1, 1)),
    lambda book: setattr(book.active["I2"], "value", time(1, 2)),
    lambda book: setattr(book.active["I2"], "value", timedelta(days=1)),
    lambda book: setattr(book.active.row_dimensions[2], "hidden", True),
    lambda book: setattr(book.active.column_dimensions["B"], "hidden", True),
    lambda book: setattr(book.create_sheet("说明"), "sheet_state", "hidden"),
    lambda book: book.create_sheet("nodes"),
    lambda book: setattr(book.active, "title", "other"),
    lambda book: book.active.merge_cells("I2:J2"),
    lambda book: setattr(book.active["B2"], "hyperlink", "https://example.com/"),
    lambda book: setattr(book.active["A2"], "value", 1),
    lambda book: setattr(book.active["G2"], "value", True),
    lambda book: setattr(book.active["A1"], "value", "unknown"),
    lambda book: setattr(book.active["P2"], "value", "extra"),
    lambda book: setattr(book.active["A5002"], "value", "C5002"),
    lambda book: setattr(book.create_sheet("说明")["A1"], "value", "=1+1"),
])
def test_xlsx_unsafe_cells_and_structure_rejected(snapshot, mutation):
    data = workbook_bytes(mutate=mutation)
    assert_failure(snapshot, service.parse_candidate_import(snapshot, "input.xlsx", data))


def test_xlsx_native_booleans_numbers_and_optional_notes(snapshot):
    row = file_row(freq_khz=2500, path_loss_db=0, rx_power_dbm=-20, is_available=False, is_backup=True)
    data = workbook_bytes([row], mutate=lambda book: book.create_sheet("说明").append(["说明", "非算法评估"]))
    result = service.parse_candidate_import(snapshot, "input.XLSX", data)
    assert not result["issues"]
    assert result["links"][0]["isAvailable"] is False and result["links"][0]["isBackup"] is True
    assert result["links"][0]["pathLossDb"] == 0


@pytest.mark.parametrize("member,content", [
    ("xl/vbaProject.bin", b"macro"), ("xl/externalLinks/externalLink1.xml", b"<a/>"),
    ("xl/embeddings/data", b"embedded"), ("xl/activeX/x.xml", b"<a/>"),
    ("../unsafe", b"x"), ("/unsafe", b"x"),
    ("unrelated.xml", b'<!DOCTYPE a [<!ENTITY x "attack">]><a>&x;</a>'),
    ("unrelated.xml", b"<unclosed>"), ("unrelated.xml", b"<a>\x00</a>"),
    ("unrelated.rels", b'<Relationships><Relationship TargetMode="External" Target="https://example.com"/></Relationships>'),
    ("huge.txt", b"x" * (parser.MAX_ZIP_MEMBER_BYTES + 1)),
], ids=["macro", "external-link", "embedded", "activex", "traversal", "absolute", "dtd", "bad-xml", "nul-xml", "external-relationship", "oversize-member"])
def test_zip_and_xml_safety_reuses_shared_gate(snapshot, member, content):
    data = rewrite_zip(workbook_bytes(), lambda members: members.update({member: content}))
    assert_failure(snapshot, service.parse_candidate_import(snapshot, "input.xlsx", data))


@pytest.mark.parametrize("mutation", ["duplicate", "many", "ratio", "macro_content", "hidden_sheet", "repeated_cell", "repeated_row", "bad_dimension"])
def test_raw_zip_and_xml_gate_is_not_bypassed(snapshot, mutation):
    def mutate(members):
        if mutation == "duplicate":
            members["XL/WORKBOOK.XML"] = members["xl/workbook.xml"]
        elif mutation == "many":
            members.update({f"extra/{i}.txt": b"x" for i in range(parser.MAX_ZIP_MEMBERS)})
        elif mutation == "ratio":
            members["compressed.txt"] = b"0" * 2000000
        elif mutation == "macro_content":
            members["[Content_Types].xml"] = members["[Content_Types].xml"].replace(b"spreadsheetml.sheet.main+xml", b"ms-excel.sheet.macroEnabled.main+xml")
        elif mutation == "hidden_sheet":
            members["xl/workbook.xml"] = members["xl/workbook.xml"].replace(b'state="visible"', b'state="hidden"')
        else:
            name = "xl/worksheets/sheet1.xml"
            data = members[name]
            if mutation == "repeated_cell":
                data = data.replace(b'<c r="B2"', b'<c r="A2"')
            elif mutation == "repeated_row":
                data = data.replace(b'<row r="2"', b'<row r="1"')
            else:
                data = data.replace(b'ref="A1:O2"', b'ref="A1:A1"')
                data = data.replace(b'<c r="O2"', b'<c r="Q2"')
            members[name] = data
    data = rewrite_zip(workbook_bytes(), mutate)
    assert_failure(snapshot, service.parse_candidate_import(snapshot, "input.xlsx", data))


@pytest.mark.parametrize("format", ["csv", "xlsx"])
def test_template_real_endpoints_unknowns_ascii_filename_and_stable_reimport(snapshot, format):
    before = deepcopy(snapshot)
    data, mime, filename = service.make_candidate_template(snapshot, format)
    assert filename.isascii() and filename.endswith("." + format)
    assert mime == (service.XLSX_MIME if format == "xlsx" else "text/csv; charset=utf-8")
    assert not data.startswith(b"\xef\xbb\xbf") and len(data) < service.MAX_FILE_BYTES
    result = service.parse_candidate_import(snapshot, filename, data)
    assert not result["issues"] and len(result["links"]) == 1
    link = result["links"][0]
    assert link["deviceAId"] == "D1" and link["deviceBId"] == "D2"
    assert link["nodeAId"] == "N1" and link["nodeBId"] == "N2" and link["deviceClass"] == "HF"
    assert all(link[field] is None for field in (*ASSESSMENT_FIELDS, "freqKhz", "bandwidthKhz", "isBackup"))
    assert link["assessmentState"] == "PENDING" and link["reference"] is None
    assert snapshot == before
    snapshot[service.SNAPSHOT_KEY] = result["links"]
    repeated = service.parse_candidate_import(snapshot, filename, data)
    assert not repeated["issues"] and repeated["links"] == result["links"]
    assert repeated["rows"][0]["action"] == "UNCHANGED"
    if format == "xlsx":
        book = load_workbook(BytesIO(data), read_only=True)
        try:
            assert book.sheetnames == ["topology_links", "说明"]
        finally:
            book.close()


def test_template_skips_same_node_and_class_mismatch_without_fabricating_devices(snapshot):
    snapshot["tables"]["device.csv"][1]["node_id"] = "N1"
    data, _, filename = service.make_candidate_template(snapshot, "csv")
    result = service.parse_candidate_import(snapshot, filename, data)
    assert not result["issues"] and result["links"][0]["deviceBId"] == "D3"
    snapshot["tables"]["device.csv"] = snapshot["tables"]["device.csv"][:1]
    with pytest.raises(ResourceDatasetError) as caught:
        service.make_candidate_template(snapshot, "csv")
    assert caught.value.code == 1001 and caught.value.status == 422


def test_template_does_not_reuse_existing_id(snapshot):
    snapshot[service.SNAPSHOT_KEY] = [manual(id="NEW-CANDIDATE-01")]
    data, _, filename = service.make_candidate_template(snapshot, "csv")
    result = service.parse_candidate_import(snapshot, filename, data)
    assert not result["issues"] and result["links"][1]["id"] == "NEW-CANDIDATE-02"


def test_all_operations_avoid_database_http_algorithm_and_runtime_files(snapshot, monkeypatch):
    import tempfile
    from app import resource_dataset_service as datasets
    from openpyxl.worksheet import _writer

    def forbidden(*args, **kwargs):
        raise AssertionError("unexpected external operation")

    monkeypatch.setattr(datasets, "current", forbidden)
    monkeypatch.setattr(datasets, "get_session_factory", forbidden)
    monkeypatch.setattr(tempfile, "NamedTemporaryFile", forbidden)
    monkeypatch.setattr(_writer, "create_temporary_file", forbidden)
    monkeypatch.setattr(Path, "open", forbidden)
    assert service.get_links(snapshot) == []
    assert service.apply_candidate_changes(snapshot, {"newCandidateLinks": [manual()]})
    for format in ("csv", "xlsx"):
        data, _, filename = service.make_candidate_template(snapshot, format)
        result = service.parse_candidate_import(snapshot, filename, data)
        assert not result["issues"]
