"""GIS 导出：真实隔离 SQLite、全库范围及纯内存文件，不接路由或监听端口。"""
from copy import deepcopy
import csv
from datetime import datetime, timezone
import hashlib
from io import BytesIO, StringIO
import json
import os
from pathlib import Path
import sys
from types import SimpleNamespace
from uuid import uuid4

import pytest
from openpyxl import load_workbook
from pydantic import ValidationError
from sqlalchemy import create_engine, event, func, select
from sqlalchemy.orm import sessionmaker

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "algorithm-service"))
import app

# 独立工作树显式复用主树真实共享源码；不复制依赖、不提供查询替身。
if os.environ.get("GIS_EXPORT_SHARED_APP"):
    app.__path__.append(str(Path(os.environ["GIS_EXPORT_SHARED_APP"]).resolve()))

from app import resource_candidate_service as candidates, resource_dataset_service as datasets
from app import resource_gis_export_service as service, resource_import_parser as parser
from app import resource_read_projection as projection, resource_read_query_schemas as read_schemas
from app.database import Base
from app.resource_candidate_schemas import ASSESSMENT_FIELDS
from app.resource_gis_export_schemas import CandidateExportFilter, ResourceGisExportRequest
from app.resource_dataset_models import ResourceDatasetCurrent, ResourceDatasetOperation, ResourceDatasetVersion
from app.resource_export_service import HEADERS


def node(identifier, name, parent="", status="NORMAL", mobility="FIXED", relay="true"):
    return dict(zip(HEADERS["nodes"], [identifier, name, "II", "II_FIXED", parent, "TASK", mobility,
                "114.5", "37.5", "", "测试区域", "HF;VUHF", relay, status, "false", "原备注"]))


def device(identifier, node_id, model, status="NORMAL", frequency="", availability="ACTIVE"):
    return {"device_id": identifier, "node_id": node_id, "model_id": model, "antenna_id": "A1",
            "antenna_height_m": "10", "tilt_deg": "", "azimuth_deg": "", "tx_power_dbm": "40",
            "work_freq_khz": frequency, "status": status, "install_time": "", "availability": availability}


def candidate(identifier="C1", imported=False, pending_reference=False):
    data = {"id": identifier, "nodeAId": "N1", "nodeBId": "N2", "deviceAId": "D1", "deviceBId": "D2",
            "deviceClass": "HF", "freqKhz": 2500 if imported else None, "bandwidthKhz": 2.5 if imported else None,
            "isBackup": None, **dict.fromkeys(ASSESSMENT_FIELDS)}
    source = {"kind": "IMPORT" if imported else "MANUAL", "file": "external.csv" if imported else None,
              "row": 2 if imported else None}
    if imported:
        data.update(distanceM=0, pathLossDb=-5, rxPowerDbm=0, linkMarginDb=-1, isAvailable=False, linkState="FAULT")
    result = {**data, "source": source, "assessmentState": "REFERENCE" if imported else "PENDING",
              "reference": {"data": deepcopy(data), "source": deepcopy(source)} if imported else None}
    if pending_reference:
        result.update(freqKhz=3000, assessmentState="PENDING", **dict.fromkeys(ASSESSMENT_FIELDS))
    return result


def source_snapshot():
    tables = {
        "node.csv": [node("N0", "上级节点"), node("N1", "目标节点", "N0", "FAULT", "VEHICLE"),
                     node("N2", "第二节点", "N1", "UNKNOWN", "MANPACK", "false")],
        "device.csv": [device("D1", "N1", "MH", "FAULT", "2500"),
                       device("V1", "N1", "MV", "NORMAL", "50000", "OFF_SHELF"),
                       device("D2", "N2", "MH", "UNKNOWN"), device("V2", "N2", "MV", "OFFLINE", "40000")],
        "device_model.csv": [
            {"model_id": "MH", "model_name": "短波型号", "device_class": "HF", "freq_min_khz": "1000", "freq_max_khz": "30000", "tx_power_max_dbm": "50"},
            {"model_id": "MV", "model_name": "超短波型号", "device_class": "VUHF", "freq_min_khz": "30000", "freq_max_khz": "300000", "tx_power_max_dbm": "50"}],
        "antenna_model.csv": [{"antenna_id": "A1", "antenna_name": "构造天线", "device_class": "HF;VUHF"}],
        "frequency_resource.csv": [{"freq_id": "F1"}],
        "interference_source.csv": [dict(zip(HEADERS["interference_sources"], [
            "I1", "目标干扰", "114.5", "37.5", "", "0", "2500", "1", "-10", "", "NOISE", "", "", "", "false"]))],
        "task_scenario.csv": [{"task_id": "T1"}],
        "comm_demand.csv": [{"demand_id": "DEM1", "task_id": "T1", "src_node_id": "N1", "dst_node_id": "N2"}],
        "link_metric.csv": [{"record_id": "HISTORY", "link_id": "OLD-LINK", "path_loss_db": "999"}],
        "candidate_site.csv": [{"site_id": "S1"}],
    }
    return {"schema_version": 1, "source_updated_at": "2026-10-23T00:00:00Z", "tables": tables,
            "source_files": {name: {"path": "data/constructed-test/" + name, "sha256": "0" * 64} for name in tables},
            candidates.SNAPSHOT_KEY: [candidate(), candidate("C2", True), candidate("C3", True, True)]}


@pytest.fixture
def gis_db(tmp_path, monkeypatch, request):
    engines = []

    def create(mutate=None):
        raw = source_snapshot()
        if mutate:
            mutate(raw)
        directory = tmp_path / uuid4().hex
        directory.mkdir()
        path = directory / "gis-export.sqlite"
        engine = create_engine("sqlite:///" + str(path), connect_args={"check_same_thread": False})
        engines.append(engine)
        factory = sessionmaker(bind=engine, expire_on_commit=False)
        Base.metadata.create_all(engine, tables=[model.__table__ for model in (
            ResourceDatasetVersion, ResourceDatasetCurrent, ResourceDatasetOperation)])
        digest = datasets.content_hash(raw)
        revision = "resource-" + digest[:16]
        with factory.begin() as session:
            session.add(ResourceDatasetVersion(revision=revision, content_hash=digest, source=datasets.SOURCE,
                data_origin=datasets.ORIGIN, snapshot=deepcopy(raw), created_at=datetime.now(timezone.utc)))
            session.flush()
            session.add(ResourceDatasetCurrent(scope=datasets.SCOPE, revision=revision))
        monkeypatch.setenv("RESOURCE_DATASET_MODE", "database")
        monkeypatch.setattr(datasets, "get_session_factory", lambda: factory)
        statements = []

        @event.listens_for(engine, "before_cursor_execute")
        def record(connection, cursor, statement, parameters, context, executemany):
            statements.append(statement)

        request.node.user_properties.append(("database", str(path)))
        for label, module in (("projection", projection), ("read_schemas", read_schemas)):
            module_path = Path(module.__file__).resolve()
            request.node.user_properties.extend([(label + "_path", str(module_path)),
                (label + "_sha256", hashlib.sha256(module_path.read_bytes()).hexdigest())])
        return SimpleNamespace(factory=factory, raw=deepcopy(raw), revision=revision, path=path, statements=statements)

    yield create
    for engine in engines:
        engine.dispose()


def query(case=None, **changes):
    return {"context": "gis", "revision": case.revision if case else "resource-" + "0" * 16,
            "format": "csv", "kinds": ["devices"], "range": "all", "browse": "devices", **changes}


def unchanged(case):
    with case.factory() as session:
        assert session.scalar(select(func.count()).select_from(ResourceDatasetVersion)) == 1
        assert session.scalar(select(func.count()).select_from(ResourceDatasetOperation)) == 0
        assert session.get(ResourceDatasetCurrent, datasets.SCOPE).revision == case.revision
        row = session.get(ResourceDatasetVersion, case.revision)
        assert row.snapshot == case.raw
        assert row.content_hash == datasets.content_hash(case.raw)
    assert case.statements and all(statement.lstrip().upper().startswith("SELECT") for statement in case.statements)


def exported_csv(case, **options):
    data, mime, filename = service.export_file(query(case, **options))
    assert mime == "text/csv; charset=utf-8" and filename.isascii() and filename.endswith(".csv")
    assert not data.startswith(b"\xef\xbb\xbf")
    return list(csv.DictReader(StringIO(data.decode("utf-8")))), data, filename


def workbook(case, **options):
    data, mime, filename = service.export_file(query(case, format="xlsx", **options))
    assert mime == service.XLSX_MIME and filename.isascii() and filename.endswith(".xlsx")
    return load_workbook(BytesIO(data), read_only=True, data_only=False)


def test_preview_canonical_all_and_real_sql_reads_only(gis_db):
    case = gis_db()
    result = service.preview(query(case, format="xlsx", kinds=["devices", "nodes", "interference_sources"]))
    assert result == {"revision": case.revision, "data_origin": "SYNTHETIC_INTEGRATION", "source": datasets.SOURCE,
                      "tables": [{"kind": "devices", "count": 4}, {"kind": "nodes", "count": 3},
                                 {"kind": "interference_sources", "count": 1}], "added": 0,
                      "warnings": [service.BASELINE_NOTICE]}
    unchanged(case)


@pytest.mark.parametrize("browse", ["devices", "stations"])
def test_filtered_classification_has_explicit_browse_semantics(gis_db, browse):
    case = gis_db()
    options = {"range": "filtered", "filter": {"device_class": "HF"}, "browse": browse}
    rows, _, _ = exported_csv(case, **options)
    expected = {"D1", "D2"} if browse == "devices" else {"D1", "D2", "V1", "V2"}
    assert {row["device_id"] for row in rows} == expected
    preview = service.preview(query(case, **options))
    assert preview["tables"] == [{"kind": "devices", "count": len(expected)}]
    assert (service.STATION_NOTICE in preview["warnings"]) == (browse == "stations")
    unchanged(case)


def test_station_search_exports_all_its_devices_and_xlsx_explains(gis_db):
    case = gis_db()
    options = {"range": "filtered", "filter": {"search": "目标节点", "device_class": "HF"}, "browse": "stations"}
    book = workbook(case, kinds=["devices", "nodes"], **options)
    try:
        device_rows = list(book["devices"].values)
        assert {row[0] for row in device_rows[1:]} == {"D1", "V1"}
        assert {row[0] for row in list(book["nodes"].values)[1:]} == {"N0", "N1"}
        assert any(service.STATION_NOTICE in row for row in book["说明"].values)
        assert any(service.BASELINE_NOTICE in row for row in book["说明"].values)
    finally:
        book.close()
    unchanged(case)


@pytest.mark.parametrize("condition,expected", [
    ({"status": "FAULT"}, {"D1"}), ({"status": "UNKNOWN"}, {"D2"}),
    ({"alert": "UNKNOWN"}, {"D2"}), ({"alert": "NO"}, {"V1"}),
    ({"availability": "OFF_SHELF"}, {"V1"}), ({"mobility": "MANPACK"}, {"D2", "V2"}),
    ({"node_role": "TASK", "relay_capable": True}, {"D1", "V1"}),
    ({"search": "第二节点"}, {"D2", "V2"}),
    ({"frequency": {"mode": "MODEL", "min_khz": 25000, "max_khz": 35000}}, {"D1", "D2", "V1", "V2"}),
    ({"frequency": {"mode": "DEVICE", "presence": "MISSING"}}, {"D2"}),
    ({"frequency": {"mode": "DEVICE", "min_khz": 3000}}, {"V1", "V2"}),
])
def test_device_filters_apply_to_full_baseline_projection(gis_db, condition, expected):
    case = gis_db()
    rows, _, _ = exported_csv(case, range="filtered", filter=condition)
    assert {row["device_id"] for row in rows} == expected
    unchanged(case)


@pytest.mark.parametrize("condition,expected", [
    ({"availability": "OFF_SHELF"}, {"D1", "V1"}),
    ({"frequency": {"mode": "DEVICE", "presence": "MISSING"}}, {"D2", "V2"}),
    ({"status": "UNKNOWN"}, {"D2", "V2"}),
    ({"mobility": "VEHICLE", "relay_capable": True}, {"D1", "V1"}),
])
def test_node_filters_use_association_existence_not_second_device_filter(gis_db, condition, expected):
    case = gis_db()
    rows, _, _ = exported_csv(case, range="filtered", browse="stations", filter=condition)
    assert {row["device_id"] for row in rows} == expected
    unchanged(case)


def test_filtered_matches_beyond_fifty_and_does_not_use_directory_pages(gis_db, monkeypatch):
    def extend(raw):
        raw["tables"]["device.csv"].extend(device(f"TAIL-{number:03d}", "N2", "MH") for number in range(75))

    case = gis_db(extend)
    def forbidden(*args, **kwargs):
        raise AssertionError("A paged GET cannot define export scope")
    monkeypatch.setattr(projection, "query_devices", forbidden)
    monkeypatch.setattr(projection, "query_nodes", forbidden)
    rows, _, _ = exported_csv(case, range="filtered", filter={"device_class": "HF"})
    assert len(rows) == 77 and rows[-1]["device_id"] == "TAIL-074"
    unchanged(case)


def test_interference_search_is_independent_of_device_class_and_node_filters(gis_db):
    case = gis_db()
    rows, _, _ = exported_csv(case, kinds=["interference_sources"], range="filtered",
                              filter={"search": "目标", "device_class": "VUHF", "mobility": "MANPACK"})
    assert [row["interference_id"] for row in rows] == ["I1"]
    unchanged(case)


def test_selected_multi_type_ids_validate_even_unexported_buckets_without_expanding(gis_db):
    case = gis_db()
    rows, _, _ = exported_csv(case, range="selected", selected_ids={"devices": ["D1"], "nodes": ["N2"], "interference_sources": []})
    assert [row["device_id"] for row in rows] == ["D1"]
    with pytest.raises(datasets.ResourceDatasetError) as caught:
        service.preview(query(case, range="selected", selected_ids={"devices": ["D1"], "nodes": ["MISSING"]}))
    assert (caught.value.code, caught.value.status) == (1002, 404)
    unchanged(case)


def test_xlsx_closure_only_when_nodes_requested_and_counts_unique_parents(gis_db):
    case = gis_db()
    options = {"range": "selected", "selected_ids": {"devices": ["D2"], "nodes": []}}
    preview = service.preview(query(case, format="xlsx", kinds=["devices", "nodes"], **options))
    assert preview["added"] == 3 and preview["tables"] == [{"kind": "devices", "count": 1}, {"kind": "nodes", "count": 3}]
    book = workbook(case, kinds=["devices", "nodes"], **options)
    try:
        assert book.sheetnames == ["devices", "nodes", "说明"]
        assert len(list(book["devices"].values)) == 2
        assert {row[0] for row in list(book["nodes"].values)[1:]} == {"N0", "N1", "N2"}
    finally:
        book.close()
    book = workbook(case, **options)
    try:
        assert book.sheetnames == ["devices", "说明"]
    finally:
        book.close()
    assert service.preview(query(case, format="xlsx", **options))["added"] == 0
    unchanged(case)


def test_csv_node_selection_has_no_implicit_parent_closure(gis_db):
    case = gis_db()
    rows, _, _ = exported_csv(case, kinds=["nodes"], range="selected", selected_ids={"nodes": ["N2"]})
    assert [row["node_id"] for row in rows] == ["N2"]
    unchanged(case)


@pytest.mark.parametrize("format", ["csv", "xlsx"])
def test_candidates_use_saved_sidecar_only_with_original_fifteen_columns(gis_db, format):
    case = gis_db()
    request = query(case, format=format, kinds=["topology_links"])
    preview = service.preview(request)
    assert preview["tables"] == [{"kind": "topology_links", "count": 3}] and preview["added"] == 0
    assert service.CANDIDATE_NOTICE in preview["warnings"]
    data, _, filename = service.export_file(request)
    assert "candidate-config-external-reference" in filename
    if format == "csv":
        rows = list(csv.reader(StringIO(data.decode("utf-8"))))
    else:
        book = load_workbook(BytesIO(data), read_only=True)
        try:
            assert book.sheetnames == ["topology_links", "说明"]
            rows = list(book["topology_links"].values)
        finally:
            book.close()
    assert list(rows[0]) == list(candidates.HEADERS)
    assert all(len(row) == 15 for row in rows)
    assert {row[0] for row in rows[1:]} == {"C1", "C2", "C3"}
    unchanged(case)


def test_candidate_csv_roundtrips_and_cannot_relabel_old_reference_as_current(gis_db):
    case = gis_db()
    rows, data, filename = exported_csv(case, kinds=["topology_links"])
    by_id = {row["link_id"]: row for row in rows}
    for identifier in ("C1", "C3"):
        assert all(by_id[identifier][header] == "" for header in candidates.HEADERS[8:14])
    assert by_id["C2"]["path_loss_db"] == "-5.0" and by_id["C2"]["rx_power_dbm"] == "0.0"
    assert by_id["C2"]["is_available"] == "false"
    parsed = candidates.parse_candidate_import(datasets.current(), filename, data)
    assert not parsed["issues"]
    parsed_links = {row["id"]: row for row in parsed["links"]}
    assert parsed_links["C3"]["assessmentState"] == "PENDING" and parsed_links["C3"]["reference"] is None
    assert parsed_links["C2"]["source"]["file"] == filename
    assert parsed_links["C2"]["reference"]["data"]["pathLossDb"] == -5
    assert service.CSV_REFERENCE_NOTICE in service.preview(query(case, kinds=["topology_links"]))["warnings"]
    unchanged(case)


def test_xlsx_notes_freeze_full_original_source_and_reference_without_overwriting_config(gis_db):
    case = gis_db()
    book = workbook(case, kinds=["topology_links"])
    try:
        notes = list(book["说明"].values)
        by_id = {row[0]: row for row in notes if row[0] in {"C1", "C2", "C3"}}
        assert by_id["C3"][1] == "PENDING"
        assert json.loads(by_id["C3"][2]) == case.raw[candidates.SNAPSHOT_KEY][2]["source"]
        reference = json.loads(by_id["C3"][3])
        assert reference == case.raw[candidates.SNAPSHOT_KEY][2]["reference"]
        assert reference["data"]["freqKhz"] == 2500 and reference["data"]["pathLossDb"] == -5
        values = {row[0]: row for row in list(book["topology_links"].values)[1:]}
        assert values["C3"][6] == 3000 and all(value is None for value in values["C3"][8:14])
        assert values["C2"][12] is False and values["C2"][10] == 0
    finally:
        book.close()
    unchanged(case)


@pytest.mark.parametrize("condition,ids", [({"search": "C2"}, {"C2"}), ({"search": "目标节点"}, {"C1", "C2", "C3"}),
                                         ({"device_class": "VUHF"}, set()), ({"alert": "YES"}, set())])
def test_candidate_filters_are_independent_and_no_quality_as_operational_alert(gis_db, condition, ids):
    case = gis_db()
    rows, _, _ = exported_csv(case, kinds=["topology_links"], range="filtered", filter=condition)
    assert {row["link_id"] for row in rows} == ids
    unchanged(case)


def test_selected_candidate_has_no_node_device_or_formal_link_closure(gis_db):
    case = gis_db()
    request = query(case, format="xlsx", kinds=["topology_links"], range="selected", selected_ids={"topology_links": ["C2"]})
    assert service.preview(request)["tables"] == [{"kind": "topology_links", "count": 1}]
    book = workbook(case, kinds=["topology_links"], range="selected", selected_ids={"topology_links": ["C2"]})
    try:
        assert book.sheetnames == ["topology_links", "说明"] and len(list(book["topology_links"].values)) == 2
    finally:
        book.close()
    unchanged(case)


@pytest.mark.parametrize("mode", ["missing", "empty"])
def test_missing_or_empty_sidecar_does_not_export_historic_link_metrics(gis_db, mode):
    def mutate(raw):
        if mode == "missing":
            raw.pop(candidates.SNAPSHOT_KEY)
        else:
            raw[candidates.SNAPSHOT_KEY] = []
    case = gis_db(mutate)
    rows, _, _ = exported_csv(case, kinds=["topology_links"])
    assert rows == [] and service.preview(query(case, kinds=["topology_links"]))["tables"][0]["count"] == 0
    unchanged(case)


@pytest.mark.parametrize("field,value", [
    ("context", "ledger"), ("revision", "bad id"), ("format", "xlsm"), ("kinds", []),
    ("kinds", ["devices", "devices"]), ("kinds", ["unknown"]), ("kinds", ["nodes", "devices"]),
    ("kinds", ["topology_links", "devices"]), ("kinds", ("devices",)),
    ("browse", "map"), ("range", "page"), ("filter", None), ("filter", {}),
    ("selected_ids", None), ("selected_ids", {}), ("page", 1), ("ids", {"devices": ["D1"]}),
    ("status_mode", "observed"), ("observation_snapshot_id", "O1"), ("draft", {}),
])
def test_strict_scope_and_no_hidden_all_conditions(field, value):
    with pytest.raises(ValidationError):
        ResourceGisExportRequest.model_validate(query(**{field: value}))


@pytest.mark.parametrize("payload", [
    {"range": "filtered"}, {"range": "filtered", "filter": None},
    {"range": "filtered", "filter": {}, "selected_ids": {}},
    {"range": "selected"}, {"range": "selected", "selected_ids": {}},
    {"range": "selected", "selected_ids": {"devices": []}},
    {"range": "selected", "selected_ids": {"nodes": ["N1"]}},
    {"range": "selected", "selected_ids": {"devices": ["D1", "D1"]}},
    {"range": "selected", "selected_ids": {"devices": [True]}},
    {"range": "selected", "selected_ids": {"devices": ["D 1"]}},
    {"range": "selected", "selected_ids": {"devices": [f"D{i}" for i in range(51)]}},
    {"range": "selected", "selected_ids": {"devices": ["D1"], "unknown": []}},
    {"range": "selected", "selected_ids": {"devices": ["D1"], "topology_links": []}},
    {"range": "selected", "selected_ids": {"devices": ["D1"]}, "filter": None},
])
def test_filtered_selected_cross_fields_and_ids_are_strict(payload):
    with pytest.raises(ValidationError):
        ResourceGisExportRequest.model_validate(query(**payload))


@pytest.mark.parametrize("field,value", [
    ("mobility", "FIXED"), ("node_role", "TASK"), ("relay_capable", False), ("mobility", None),
    ("status", "FAULT"), ("availability", "ACTIVE"), ("frequency", {"mode": "ALL"}),
    ("frequency", None), ("alert", "NO"), ("alert", "UNKNOWN"), ("status_mode", "baseline"),
    ("observation_snapshot_id", None), ("page", 1), ("search", "BAD\x00"),
])
def test_candidates_reject_undefined_or_hidden_filter_conditions(field, value):
    with pytest.raises(ValidationError):
        ResourceGisExportRequest.model_validate(query(kinds=["topology_links"], range="filtered", filter={field: value}))


def test_candidate_filter_model_is_explicitly_different_from_device_filter():
    request = ResourceGisExportRequest.model_validate(query(kinds=["topology_links"], range="filtered", filter={"search": "C"}))
    assert isinstance(request.filter, CandidateExportFilter)
    assert set(CandidateExportFilter.model_fields) == {"search", "device_class", "alert"}


@pytest.mark.parametrize("condition", [
    {"relay_capable": "true"}, {"search": "BAD\n"}, {"status": "GOOD"},
    {"frequency": {"mode": "ALL", "min_khz": 1}}, {"frequency": {"mode": "MODEL", "presence": "MISSING"}},
    {"frequency": {"mode": "DEVICE", "min_khz": True}},
    {"frequency": {"mode": "DEVICE", "min_khz": float("nan")}},
    {"frequency": {"mode": "MODEL", "min_khz": 20, "max_khz": 10}},
    {"status_mode": "observed"}, {"observation_snapshot_id": "OBS"},
])
def test_device_filter_reuses_strict_finite_and_combination_validation(condition):
    with pytest.raises(ValidationError):
        ResourceGisExportRequest.model_validate(query(range="filtered", filter=condition))


@pytest.mark.parametrize("mutation", [
    "missing-parent", "cycle", "duplicate-node", "missing-model", "missing-antenna", "missing-device-node",
    "nonfinite", "bad-status", "bad-boolean", "bad-availability", "bad-candidate-node", "bad-candidate-device",
    "bad-candidate-class", "bad-candidate-frequency", "duplicate-candidate", "candidate-metadata", "control-text",
])
def test_source_damage_fails_closed_before_preview_or_file_without_sql_writes(gis_db, mutation):
    def mutate(raw):
        tables = raw["tables"]
        if mutation == "missing-parent":
            tables["node.csv"][1]["parent_node_id"] = "MISSING"
        elif mutation == "cycle":
            tables["node.csv"][0]["parent_node_id"] = "N2"
        elif mutation == "duplicate-node":
            tables["node.csv"].append(deepcopy(tables["node.csv"][0]))
        elif mutation in {"missing-model", "missing-antenna", "missing-device-node"}:
            field = {"missing-model": "model_id", "missing-antenna": "antenna_id", "missing-device-node": "node_id"}[mutation]
            tables["device.csv"][0][field] = "MISSING"
        elif mutation == "nonfinite":
            tables["device.csv"][0]["tx_power_dbm"] = "NaN"
        elif mutation == "bad-status":
            tables["node.csv"][0]["status"] = "GOOD"
        elif mutation == "bad-boolean":
            tables["node.csv"][0]["relay_capable"] = "1"
        elif mutation == "bad-availability":
            tables["device.csv"][0]["availability"] = "INVALID"
        elif mutation == "control-text":
            tables["node.csv"][0]["remark"] = "bad\x00text"
        elif mutation == "duplicate-candidate":
            raw[candidates.SNAPSHOT_KEY].append(deepcopy(raw[candidates.SNAPSHOT_KEY][0]))
        elif mutation == "candidate-metadata":
            raw[candidates.SNAPSHOT_KEY][0]["assessmentState"] = "REFERENCE"
        else:
            field, value = {"bad-candidate-node": ("nodeAId", "MISSING"), "bad-candidate-device": ("deviceAId", "V1"),
                            "bad-candidate-class": ("deviceClass", "VUHF"), "bad-candidate-frequency": ("freqKhz", 999999)}[mutation]
            raw[candidates.SNAPSHOT_KEY][0][field] = value
    case = gis_db(mutate)
    for operation in (service.preview, service.export_file):
        with pytest.raises(datasets.ResourceDatasetError) as caught:
            operation(query(case))
        assert (caught.value.code, caught.value.status) == (5002, 503)
    unchanged(case)


@pytest.mark.parametrize("kind,ids", [("devices", ["MISSING"]), ("nodes", ["D1"]), ("interference_sources", ["N1"]), ("topology_links", ["OLD-LINK"])])
def test_unknown_or_wrong_type_ids_are_not_silently_dropped(gis_db, kind, ids):
    case = gis_db()
    with pytest.raises(datasets.ResourceDatasetError) as caught:
        service.preview(query(case, kinds=[kind], range="selected", selected_ids={kind: ids}))
    assert (caught.value.code, caught.value.status) == (1002, 404)
    unchanged(case)


def test_wrong_revision_file_mode_and_schema_errors_never_read_raw_files(gis_db, monkeypatch):
    case = gis_db()
    def forbidden(*args, **kwargs):
        raise AssertionError("Raw files cannot be an export fallback")
    monkeypatch.setattr(datasets, "capture_files", forbidden)
    with pytest.raises(datasets.ResourceDatasetError) as caught:
        service.preview(query(case, revision="resource-" + "f" * 16))
    assert (caught.value.code, caught.value.status) == (1205, 409)
    monkeypatch.setenv("RESOURCE_DATASET_MODE", "file")
    with pytest.raises(datasets.ResourceDatasetError) as caught:
        service.export_file(query(case))
    assert (caught.value.code, caught.value.status) == (5003, 503)
    with pytest.raises(datasets.ResourceDatasetError) as caught:
        service.preview(query(case, filter={}))
    assert (caught.value.code, caught.value.status) == (1001, 422)
    unchanged(case)


@pytest.mark.parametrize("prefix", ["=ID", "+ID", "-ID", "@ID", "~rp~ID"])
def test_candidate_csv_formula_prefix_text_roundtrips_without_metadata_columns(gis_db, prefix):
    def mutate(raw):
        raw[candidates.SNAPSHOT_KEY][0]["id"] = prefix
    case = gis_db(mutate)
    _, data, filename = exported_csv(case, kinds=["topology_links"])
    assert b"~rp~" in data
    parsed = candidates.parse_candidate_import(datasets.current(), filename, data)
    assert not parsed["issues"] and prefix in {row["id"] for row in parsed["links"]}
    unchanged(case)


def test_xlsx_text_is_never_a_formula_and_export_never_creates_runtime_files(gis_db, monkeypatch):
    case = gis_db(lambda raw: raw["tables"]["node.csv"][1].update(node_name="=1+1"))
    import tempfile
    from openpyxl.worksheet import _writer
    def forbidden(*args, **kwargs):
        raise AssertionError("File export must stay in memory")
    monkeypatch.setattr(tempfile, "NamedTemporaryFile", forbidden)
    monkeypatch.setattr(_writer, "create_temporary_file", forbidden)
    monkeypatch.setattr(Path, "open", forbidden)
    book = workbook(case, kinds=["nodes", "devices", "interference_sources"])
    try:
        assert book["nodes"]["B3"].value == "=1+1" and book["nodes"]["B3"].data_type == "s"
        assert not any(cell.data_type == "f" for sheet in book for row in sheet for cell in row)
    finally:
        book.close()
    unchanged(case)


@pytest.mark.parametrize("format", ["csv", "xlsx"])
def test_empty_filtered_results_produce_headers_not_fallback_to_all(gis_db, format):
    case = gis_db()
    request = query(case, format=format, range="filtered", filter={"search": "NEVER-MATCH"})
    assert service.preview(request)["tables"] == [{"kind": "devices", "count": 0}]
    data, _, _ = service.export_file(request)
    if format == "csv":
        assert len(list(csv.reader(StringIO(data.decode("utf-8"))))) == 1
    else:
        book = load_workbook(BytesIO(data), read_only=True)
        try:
            assert len(list(book["devices"].values)) == 1
        finally:
            book.close()
    unchanged(case)


@pytest.mark.parametrize("format", ["csv", "xlsx"])
def test_size_limit_rejects_whole_file_without_truncation_or_sql_writes(gis_db, monkeypatch, format):
    case = gis_db()
    monkeypatch.setattr(service, "MAX_FILE_BYTES", 1)
    with pytest.raises(datasets.ResourceDatasetError) as caught:
        service.export_file(query(case, format=format))
    assert (caught.value.code, caught.value.status) == (1001, 413)
    unchanged(case)


def test_shared_parent_dag_closure_is_not_a_cycle_and_counts_each_node_once(gis_db):
    case = gis_db(lambda raw: raw["tables"]["node.csv"][2].update(parent_node_id="N1;N0"))
    result = service.preview(query(case, format="xlsx", kinds=["nodes", "devices"], range="selected",
                                   selected_ids={"nodes": ["N1"], "devices": ["D2"]}))
    assert result["tables"] == [{"kind": "nodes", "count": 3}, {"kind": "devices", "count": 1}]
    assert result["added"] == 2
    unchanged(case)


def test_pending_historical_reference_does_not_require_removed_devices_to_exist_now(gis_db):
    def mutate(raw):
        reference = raw[candidates.SNAPSHOT_KEY][2]["reference"]["data"]
        reference.update(nodeAId="HISTORICAL-NODE", deviceAId="HISTORICAL-DEVICE")
    case = gis_db(mutate)
    book = workbook(case, kinds=["topology_links"])
    try:
        row = next(row for row in book["说明"].values if row[0] == "C3")
        assert json.loads(row[3])["data"]["deviceAId"] == "HISTORICAL-DEVICE"
        current = next(row for row in list(book["topology_links"].values)[1:] if row[0] == "C3")
        assert current[1] == "N1" and current[3] == "D1" and current[9] is None
    finally:
        book.close()
    unchanged(case)


def test_xlsx_overlong_utf16_cell_is_rejected_while_csv_retains_complete_text(gis_db):
    text = "\U00020000" * 16384
    case = gis_db(lambda raw: raw["tables"]["node.csv"][0].update(remark=text))
    for operation in (service.preview, service.export_file):
        with pytest.raises(datasets.ResourceDatasetError) as caught:
            operation(query(case, format="xlsx", kinds=["nodes"]))
        assert (caught.value.code, caught.value.status) == (1001, 413)
    rows, _, _ = exported_csv(case, kinds=["nodes"])
    assert rows[0]["remark"] == text
    unchanged(case)


def test_existing_project_constructed_ten_tables_export_from_isolated_database(gis_db):
    _, captured = datasets.capture_files()
    original = deepcopy(captured)
    def seed(raw):
        raw.clear()
        raw.update(deepcopy(captured))
    case = gis_db(seed)
    request = query(case, format="xlsx", kinds=["nodes", "devices", "interference_sources"])
    result = service.preview(request)
    assert result["tables"] == [{"kind": kind, "count": len(captured["tables"][service.TABLES[kind]])}
                                for kind in request["kinds"]]
    book = workbook(case, kinds=request["kinds"])
    try:
        assert book.sheetnames == ["nodes", "devices", "interference_sources", "说明"]
        assert len(list(book["devices"].values)) == len(captured["tables"]["device.csv"]) + 1
    finally:
        book.close()
    assert captured == original
    unchanged(case)
